"""Independent check of T24 HORSE / JOCKEY (SKY STARS) and SKY STARS SUMMARY (own code; families.py for scores only).
usage: check_ss_indep.py <after.xlsx> [<before.xlsx>]"""
import sys, re, ast
sys.path.insert(0, '/Users/eddieplace/astro-harmonics/tools')
import openpyxl
import families as F

src = open('/Users/eddieplace/astro-harmonics/tools/jupiter_tables.py').read()
m = re.search(r"^T1_T2_T24_STARS = (\[.*?\])", src, re.S | re.M)
STARS67 = ast.literal_eval(m.group(1))
assert len(STARS67) == 67
FILL = {'green': 'C6EFCE', 'yellow': 'FFEB9C', 'light': 'FFF8DC', 'red': 'FF8080', '': None}
FAMNAME = {'whole': 'whole', 'half': 'half', 'golden': 'golden', 'phi power': 'phi power', 'ninth': 'ninth', 'sqrt2': 'sqrt2',
           'silver': 'silver', 'conj': 'conjunction'}
MM = []
def mm(part, where, exp, got):
    MM.append((part, where, exp, got))
    if len(MM) < 80:
        print('MISMATCH', part, where, 'expected', repr(exp)[:200], 'got', repr(got)[:200])

def rgb(c):
    f = c.fill
    if f is None or f.fill_type is None:
        return None
    return str(f.start_color.rgb)[-6:]

def bord(c):
    b = c.border
    return tuple((getattr(b, s).style, str(getattr(b, s).color.rgb) if getattr(b, s).color is not None else None)
                 for s in ('left', 'right', 'top', 'bottom'))

def com(c):
    return c.comment.text if c.comment else None

def dist(a, b, plane):
    if plane == 'RA':
        d = abs(a[0] - b[0]) % 360
        return 360 - d if d > 180 else d
    return abs(a[1] - b[1])

def expect(v, plane):
    b = F.best(v, plane)
    if not b:
        return None, None
    conj = v < F.CONJ[plane]
    return b, ('red' if conj else F.colour(b['score']))

after = sys.argv[1]
wb = openpyxl.load_workbook(after)
# ---------- POSITIONS
pos, bodies = {}, {}
for r in wb['POSITIONS'].iter_rows(values_only=True):
    if r[0] and r[2] in ('S', 'H', 'J') and r[4] in ('body', 'star', 'equator') and isinstance(r[5], (int, float)):
        key = (r[0], r[2], r[3])
        pos[key] = (r[5], r[6])
        if r[4] == 'body':
            bodies.setdefault((r[0], r[2]), []).append(r[3])
runners = sorted({k[0] for k in pos if k[0] != 'SKY'})
n = {'a': 0, 'b_t24': 0, 'b_recalc': 0, 'b_scores': 0, 'b_hdr': 0, 'c': 0}
# ---------- (a) selection
sj = pos[('SKY', 'S', 'Jupiter')]
sel, hit = [], {}
sky = wb['SKY JUPITER']
skyrow = {}
for r in range(1, sky.max_row + 1):
    if sky.cell(r, 1).value == 'SKY' and str(sky.cell(r, 2).value).startswith('S Jupiter → S '):
        skyrow.setdefault(sky.cell(r, 2).value[len('S Jupiter → S '):], r)
for s in STARS67:
    if ('SKY', 'S', s) not in pos:
        continue
    parts = []
    for plane, scol in (('RA', 8), ('Dec', 17)):
        v = dist(sj, pos[('SKY', 'S', s)], plane)
        b, col = expect(v, plane)
        got = sky.cell(skyrow[s], scol).value if s in skyrow else 'NO ROW'
        n['a'] += 1
        if b:
            if got is None or abs(got - round(b['score'], 2)) > 0.0051:
                mm('a', f'SKY block {s} {plane} score', round(b['score'], 2), got)
            tgt = '0' if b['family'] == 'conj' else b['target']
            parts.append(f"{plane} {tgt} {FAMNAME[b['family']]} ({b['score']:.2f})")
        elif got not in (None, ''):
            mm('a', f'SKY block {s} {plane} score', None, got)
    if parts:
        sel.append(s); hit[s] = '; '.join(parts)
print('selection:', len(sel), sel)

def find_blocks(ws, aid):
    rows = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == aid]
    return rows

def table_rows(ws, title_id):
    """title row, header rows (2), data rows of block with col A id == title_id (title row first)"""
    rs = find_blocks(ws, title_id)
    t = rs[0]
    data = [r for r in rs if r > t]
    return t, data

def t24_blocks(ws):
    ts = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == 'T24' and ws.cell(r + 1, 1).value == 'table']
    out = []
    for t in ts[:2]:
        rows = {}
        r = t + 3
        while ws.cell(r, 1).value == 'T24':
            rows[ws.cell(r, 2).value] = r; r += 1
        out.append((t, rows))
    return out

def merged_in(ws, r0, r1):
    return sorted((m.min_row - r0, m.min_col, m.max_row - r0, m.max_col) for m in ws.merged_cells.ranges if r0 <= m.min_row <= r1)

runner_tables = {}
for p in runners:
    ws = wb[p]
    (th, hrows), (tj, jrows) = t24_blocks(ws)
    nb = len(bodies[(p, 'H')])
    t24w = 2 + nb + 1 + nb + 1 + 1   # A,B, RA bodies, RA mid, Dec bodies, Dec mid, repeats -> last col
    for ch, tid, (t24t, t24rows) in (('H', 'T24 HORSE (SKY STARS)', (th, hrows)), ('J', 'T24 JOCKEY (SKY STARS)', (tj, jrows))):
        bl = [ws.cell(t24t + 2, c).value for c in range(3, 3 + nb)]   # T24 header order (checked = POSITIONS body set, both halves)
        if sorted(bl) != sorted(bodies[(p, ch)]) or [ws.cell(t24t + 2, c).value for c in range(4 + nb, 4 + 2 * nb)] != bl:
            mm('b', f'{p} {ch} T24 header body set', sorted(bodies[(p, ch)]), bl)
        rs = find_blocks(ws, tid)
        if not rs:
            mm('b', f'{p} {tid}', 'present', 'missing'); continue
        t = rs[0]
        if ws.cell(t + 1, 1).value != 'table':
            mm('b', f'{p} {tid} header', 'table', ws.cell(t + 1, 1).value)
        data = [r for r in rs if r > t]
        if data != list(range(t + 3, t + 3 + len(data))):
            mm('b', f'{p} {tid} rows contiguous', 'contiguous', data)
        runner_tables[(p, ch)] = (t, t + 2 + len(data))
        title = ws.cell(t, 2).value or ''
        exp_sel = [s for s in sel if (p, ch, s) in pos]
        if f'Selected: {len(exp_sel)} stars' not in title:
            mm('b', f'{p} {tid} title', f'Selected: {len(exp_sel)} stars', title[-60:])
        # header rows
        lastcol = 1 + t24w + 2 * nb
        n['b_hdr'] += 1
        h1 = {c: ws.cell(t + 1, c).value for c in range(1, lastcol + 2) if ws.cell(t + 1, c).value is not None}
        exp_h1 = {1: 'table', 2: 'star', 3: 'sky hit', 4: 'RA', 4 + nb: 'RA midpoints', 5 + nb: 'Dec', 5 + 2 * nb: 'Dec midpoints',
                  6 + 2 * nb: 'repeats', 7 + 2 * nb: 'SCORES RA', 7 + 3 * nb: 'SCORES Dec'}
        if h1 != exp_h1:
            mm('b', f'{p} {tid} header row 1', exp_h1, h1)
        h2 = [ws.cell(t + 2, c).value for c in range(4, lastcol + 1)]
        t24h2 = [ws.cell(t24t + 2, c).value for c in range(3, t24w + 1)]
        exp_h2 = t24h2 + bl + bl
        if h2 != exp_h2:
            mm('b', f'{p} {tid} header row 2', exp_h2, h2)
        if t24h2[:nb] != bl or t24h2[nb + 1:2 * nb + 1] != bl:
            mm('b', f'{p} T24 header bodies', bl, t24h2)
        mg = merged_in(ws, t, t + 2)
        exp_mg = sorted([(1, 4, 1, 3 + nb), (1, 5 + nb, 1, 4 + 2 * nb), (1, 7 + 2 * nb, 1, 6 + 3 * nb), (1, 7 + 3 * nb, 1, 6 + 4 * nb)])
        if mg != exp_mg:
            mm('b', f'{p} {tid} merged', exp_mg, mg)
        # rows
        got_labels = [ws.cell(r, 2).value for r in data]
        exp_labels = [f'{ch} {s}' for s in exp_sel]
        if got_labels != exp_labels:
            mm('b', f'{p} {tid} rows', exp_labels, got_labels)
        jup = pos[(p, ch, 'Jupiter')]
        for r in data:
            lab = ws.cell(r, 2).value
            s = lab[2:]
            if ws.cell(r, 3).value != hit.get(s):
                mm('b', f'{p} {tid} {lab} C sky hit', hit.get(s), ws.cell(r, 3).value)
            tr = t24rows.get(lab)
            if tr is None:
                mm('b', f'{p} {lab}', 'T24 row', None); continue
            # copy check: T24 cols 3..t24w -> new cols 4..t24w+1
            for c in range(3, t24w + 1):
                a, b2 = ws.cell(tr, c), ws.cell(r, c + 1)
                n['b_t24'] += 1
                for what, f in (('value', lambda x: x.value), ('fill', rgb), ('comment', com), ('border', bord),
                                ('numfmt', lambda x: x.number_format), ('font', lambda x: (x.font.b, x.font.i, str(x.font.color.rgb) if x.font.color else None))):
                    if f(a) != f(b2):
                        mm('b', f'{p} {b2.coordinate} {what} vs T24 {a.coordinate}', f(a), f(b2))
            # recompute
            sp = pos[(p, ch, s)]
            for k, plane in enumerate(('RA', 'Dec')):
                for i, body in enumerate(bl):
                    vc = 4 + i + k * (nb + 1)
                    scc = 7 + 2 * nb + i + k * nb
                    v = dist(sp, pos[(p, ch, body)], plane)
                    cell = ws.cell(r, vc); scell = ws.cell(r, scc)
                    n['b_recalc'] += 1; n['b_scores'] += 1
                    if not isinstance(cell.value, (int, float)) or abs(cell.value - round(v, 6)) > 6e-7:
                        mm('b', f'{p} {cell.coordinate} {lab} {plane} {body} value', round(v, 6), cell.value)
                    b, col = expect(v, plane)
                    if rgb(cell) != FILL[col or '']:
                        mm('b', f'{p} {cell.coordinate} {lab} {plane} {body} fill', FILL[col or ''], rgb(cell))
                    cm = (com(cell) or '').split('\n')[0]
                    if b:
                        if cm.startswith('nearest') or abs(float(cm.split(' · ')[-1]) - round(b['score'], 2)) > 0.0051:
                            mm('b', f'{p} {cell.coordinate} comment score', round(b['score'], 2), cm)
                        if scell.value is None or abs(scell.value - round(b['score'], 2)) > 1e-9:
                            mm('b', f'{p} {scell.coordinate} {lab} SCORES {plane} {body}', round(b['score'], 2), scell.value)
                        if rgb(scell) != rgb(cell):
                            mm('b', f'{p} {scell.coordinate} SCORES fill', rgb(cell), rgb(scell))
                    else:
                        if not cm.startswith('nearest'):
                            mm('b', f'{p} {cell.coordinate} comment', 'nearest: ...', cm)
                        if scell.value not in (None, ''):
                            mm('b', f'{p} {scell.coordinate} SCORES blank', None, scell.value)
                        if rgb(scell) is not None:
                            mm('b', f'{p} {scell.coordinate} SCORES fill', None, rgb(scell))
            for c in range(lastcol + 1, lastcol + 6):
                if ws.cell(r, c).value is not None:
                    mm('b', f'{p} {ws.cell(r, c).coordinate}', None, ws.cell(r, c).value)
# ---------- (c) summary
names = wb.sheetnames
if 'SKY STARS SUMMARY' not in names:
    mm('c', 'sheet', 'SKY STARS SUMMARY', names)
else:
    sm = wb['SKY STARS SUMMARY']
    r = 1
    def cmp_block(ws, r0, r1, sr):
        for i in range(r1 - r0 + 1):
            for c in range(1, max(ws.max_column, sm.max_column) + 1):
                a, b2 = ws.cell(r0 + i, c), sm.cell(sr + i, c)
                n['c'] += 1
                for what, f in (('value', lambda x: x.value), ('fill', rgb), ('comment', com), ('border', bord),
                                ('numfmt', lambda x: x.number_format)):
                    if f(a) != f(b2):
                        mm('c', f'{ws.title} {a.coordinate} -> SUMMARY {b2.coordinate} {what}', f(a), f(b2))
        mws = merged_in(ws, r0, r1); msm = merged_in(sm, sr, sr + r1 - r0)
        if mws != msm:
            mm('c', f'merged {ws.title} {r0} -> {sr}', mws, msm)
    def blank(rr):
        for c in range(1, sm.max_column + 1):
            n['c'] += 1
            if sm.cell(rr, c).value is not None:
                mm('c', f'SUMMARY {sm.cell(rr, c).coordinate} blank', None, sm.cell(rr, c).value)
    for p in runners:
        ws = wb[p]
        n['c'] += 1
        if sm.cell(r, 1).value != ws.cell(1, 1).value:
            mm('c', f'SUMMARY A{r} runner title', ws.cell(1, 1).value, sm.cell(r, 1).value)
        r += 1
        t0, t1 = runner_tables[(p, 'H')]; cmp_block(ws, t0, t1, r); r += t1 - t0 + 1
        blank(r); r += 1
        t0, t1 = runner_tables[(p, 'J')]; cmp_block(ws, t0, t1, r); r += t1 - t0 + 1
        blank(r); blank(r + 1); r += 2
    if sm.max_row >= r:
        extra = [(rr, sm.cell(rr, 1).value) for rr in range(r, sm.max_row + 1) if any(sm.cell(rr, c).value is not None for c in range(1, sm.max_column + 1))]
        if extra:
            mm('c', 'SUMMARY rows after the last runner', [], extra[:5])
print('counts', n)
print('mismatches', len(MM))
