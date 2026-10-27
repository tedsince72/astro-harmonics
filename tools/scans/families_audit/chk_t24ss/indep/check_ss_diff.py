"""(d) before vs after, every sheet cell by cell; (e) default vs --full new tables. usage: check_ss_diff.py before after [other_after]"""
import sys
import openpyxl
IDS = ('T24 HORSE (SKY STARS)', 'T24 JOCKEY (SKY STARS)')
def rgb(c):
    f = c.fill
    return None if f is None or f.fill_type is None else (f.fill_type, str(f.start_color.rgb))
def bord(c):
    b = c.border
    return tuple((getattr(b, s).style, str(getattr(b, s).color.rgb) if getattr(b, s).color is not None else None) for s in ('left', 'right', 'top', 'bottom'))
def sig(c):
    return (c.value, rgb(c), c.comment.text if c.comment else None, bord(c), c.number_format,
            (c.font.b, c.font.i, str(c.font.color.rgb) if c.font.color else None), c.alignment.wrap_text, c.alignment.horizontal)
MM = []
WD = []
def mm(*a):
    MM.append(a)
    if len(MM) < 60: print('MISMATCH', *a)
old = openpyxl.load_workbook(sys.argv[1]); new = openpyxl.load_workbook(sys.argv[2])
print('old sheets', old.sheetnames); print('new sheets', new.sheetnames)
on, nn = old.sheetnames, new.sheetnames
if [s for s in nn if s != 'SKY STARS SUMMARY'] != on or nn.count('SKY STARS SUMMARY') != 1:
    mm('sheet list', on, nn)
runners = [s for s in on if len(s) == 3 and s[0] == 'P' and s[1:].isdigit()]
i = nn.index('SKY STARS SUMMARY')
last_rel = max(j for j, s in enumerate(nn) if s[:3] in runners)   # last P## / P## RATIOS / P## NUMBERS
if i != last_rel + 1:
    mm('summary position', last_rel + 1, i)
print('summary at index', i, 'after', nn[i - 1], 'before', nn[i + 1] if i + 1 < len(nn) else None)
cells = 0
for name in on:
    a, b = old[name], new[name]
    if name in runners:
        rows_new = [r for r in range(1, b.max_row + 1) if b.cell(r, 1).value in IDS]
        start = min(rows_new)
        # old row where the first block after T24 JOCKEY RATIOS starts = first non-empty col A row after the last 'T24 JOCKEY RATIOS' row
        lr = max(r for r in range(1, a.max_row + 1) if a.cell(r, 1).value == 'T24 JOCKEY RATIOS')
        nxt = next(r for r in range(lr + 1, a.max_row + 1) if a.cell(r, 1).value is not None)
        nxt_new = next(r for r in range(max(rows_new) + 1, b.max_row + 1) if b.cell(r, 1).value is not None)
        shift = nxt_new - nxt
        # inserted rows: new rows [nxt, nxt_new) hold only the new tables and blanks
        for r in range(nxt, nxt_new):
            v = b.cell(r, 1).value
            if v not in (None, 'table') + IDS:
                mm(name, 'inserted region row', r, v)
            if v is None and b.cell(r - 1, 1).value != 'table' and any(b.cell(r, c).value is not None for c in range(1, b.max_column + 1)):
                mm(name, 'inserted blank row not blank', r)
        mapr = lambda r: r if r < nxt else r + shift
        print(name, 'old next block row', nxt, 'new', nxt_new, 'shift', shift, 'new tables start', start)
        if b.max_row != a.max_row + shift: mm(name, 'max_row', a.max_row + shift, b.max_row)
    else:
        mapr = lambda r: r
        if b.max_row != a.max_row: mm(name, 'max_row', a.max_row, b.max_row)
    mc = max(a.max_column, b.max_column) if name not in runners else a.max_column
    for r in range(1, a.max_row + 1):
        R = mapr(r)
        for c in range(1, mc + 1):
            cells += 1
            x, y = sig(a.cell(r, c)), sig(b.cell(R, c))
            if x != y:
                mm(name, a.cell(r, c).coordinate, '->', b.cell(R, c).coordinate, x, y)
    # cells right of old max column in new (outside inserted region) must be empty
    if name in runners:
        for r in range(1, a.max_row + 1):
            for c in range(a.max_column + 1, b.max_column + 1):
                if b.cell(mapr(r), c).value is not None:
                    mm(name, 'extra cell', b.cell(mapr(r), c).coordinate, b.cell(mapr(r), c).value)
    om = sorted((mapr(m.min_row), m.min_col, mapr(m.max_row), m.max_col) for m in a.merged_cells.ranges)
    nm = sorted((m.min_row, m.min_col, m.max_row, m.max_col) for m in b.merged_cells.ranges
                if name not in runners or not (nxt <= m.min_row < nxt_new))
    if om != nm: mm(name, 'merged ranges', len(om), len(nm))
    ow = {k: v.width for k, v in a.column_dimensions.items()}
    nw = {k: v.width for k, v in b.column_dimensions.items()}
    for k in set(ow) | set(nw):
        if ow.get(k) != nw.get(k):
            WD.append((name, k, ow.get(k), nw.get(k)))
    if a.freeze_panes != b.freeze_panes: mm(name, 'freeze', a.freeze_panes, b.freeze_panes)
print('cells compared before vs after', cells)
if len(sys.argv) > 3:
    o2 = openpyxl.load_workbook(sys.argv[3]); ec = 0
    def region(wb, p):
        ws = wb[p]; rs = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value in IDS]
        return ws, min(rs), max(rs)
    for p in runners:
        (w1, s1, e1), (w2, s2, e2) = region(new, p), region(o2, p)
        if e1 - s1 != e2 - s2: mm('e', p, 'region size', e1 - s1, e2 - s2); continue
        for k in range(e1 - s1 + 1):
            for c in range(1, max(w1.max_column, w2.max_column) + 1):
                ec += 1
                if sig(w1.cell(s1 + k, c)) != sig(w2.cell(s2 + k, c)):
                    mm('e', p, w1.cell(s1 + k, c).coordinate, w2.cell(s2 + k, c).coordinate)
    a1, a2 = new['SKY STARS SUMMARY'], o2['SKY STARS SUMMARY']
    if a1.max_row != a2.max_row: mm('e', 'summary rows', a1.max_row, a2.max_row)
    for r in range(1, max(a1.max_row, a2.max_row) + 1):
        for c in range(1, max(a1.max_column, a2.max_column) + 1):
            ec += 1
            if sig(a1.cell(r, c)) != sig(a2.cell(r, c)): mm('e', 'summary', a1.cell(r, c).coordinate)
    m1 = sorted(str(m) for m in a1.merged_cells.ranges); m2 = sorted(str(m) for m in a2.merged_cells.ranges)
    if m1 != m2: mm('e', 'summary merged')
    print('cells compared default vs full', ec)
import collections
print('width changes', len(WD), collections.Counter((w[2], w[3]) for w in WD))
print('P01 width changes', sorted((w[1], w[2], w[3]) for w in WD if w[0]=='P01' and w[2] is not None))
print('mismatches', len(MM))
