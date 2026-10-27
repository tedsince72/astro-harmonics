"""Scoring sweep of every distance cell in a jupiter_tables / pairs workbook, own family code (fam_own.py).
usage: python3 sweep.py <workbook.xlsx> <out.json>"""
import sys, re, json
from collections import Counter, defaultdict
import openpyxl
from openpyxl.utils import get_column_letter as L
sys.path.insert(0, __file__.rsplit('/', 1)[0])
from fam_own import evaluate, best, parse_label, same_target, FAMNORM, ninth_ambiguous

FAMS = r'(whole|half|golden|phi power|ninth|sqrt2|silver|conj)'
RE_GROUP = re.compile(r'^(?P<lab>\S+) \(?' + FAMS + r'\)?(?P<rest>.*)$')
RE_LEG = re.compile(r'^→ (?P<lab>\S+) \(' + FAMS + r', (?P<sc>-?[\d.]+)\)')
RE_CHIT = re.compile(r'^(?P<lab>\S+) · (?P<fam>whole|half|golden|phi power|ninth|sqrt2|silver|conjunction \((?:RA|Dec) < [\d.]+\)) · (?P<off>\S+) · (?P<sc>-?[\d.]+)$')
RE_CNO = re.compile(r'^nearest: (?P<lab>\S+) (?:· )?(?P<fam>whole|half|golden|phi power|ninth|sqrt2|silver) — ')
PLANES = ('RA', 'Dec', 'Equator')

path, out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(path)
counts = Counter()
problems = []
boundary = []
ties = []
unk = Counter()
unk_ex = {}
together = Counter()
together_ex = []
ninth_amb = []
samples = {}
rounding_ok = [0]


def shown_num(x):
    if isinstance(x, (int, float)):
        return float(x)
    try:
        return float(x)
    except Exception:
        return None


def check(kind, sheet, coord, tab, v, plane, shown):
    """shown: None (no hit) or (fam, labval, score or None)."""
    counts[(tab, kind)] += 1
    if counts[(tab, kind)] <= 3:
        samples.setdefault(f'{tab}|{kind}', []).append([sheet, coord, v, plane, str(shown)])
    b, tie = best(v, plane)
    # boundary: does the answer change within +-5e-7 (6-dp rounding)?
    bl, _ = best(v - 5e-7, plane)
    bh, _ = best(v + 5e-7, plane)
    key = lambda x: None if x is None else (x[0], round(x[1], 6))
    bnd = key(bl) != key(b) or key(bh) != key(b)
    if tie:
        ties.append((sheet, coord, tab, v, plane, evaluate(v, plane)))
    if v is not None and ninth_ambiguous(v):
        ninth_amb.append((sheet, coord, tab, v, plane, b, shown))
    err = None
    if b is None and shown is not None:
        err = 'sheet shows a hit, rule gives none'
    elif b is not None and shown is None:
        err = 'rule gives a hit, sheet shows none'
    elif b is not None:
        fam, lv, sc = shown
        if fam != b[0] or not same_target(b[0], b[1], lv):
            err = 'different target'
        elif sc is not None and abs(sc - b[2]) > 0.0051:
            # the shown value is rounded to 6 dp: accept a score the unrounded value (within +-5e-7) can give
            rng = []
            for dv in (-5e-7, 5e-7):
                for hh in evaluate(v + dv, plane):
                    if hh[0] == b[0] and abs(hh[1] - b[1]) < 1e-9:
                        rng.append(hh[2])
            if rng and min(rng) - 0.0051 <= sc <= max(rng) + 0.0051:
                rounding_ok[0] += 1
            else:
                err = 'different score'
    if err:
        rec = dict(sheet=sheet, cell=coord, table=tab, kind=kind, value=v, plane=plane, err=err,
                   expected=None if b is None else [b[0], round(b[1], 6), round(b[2], 4)], shown=shown)
        (boundary if bnd else problems).append(rec)


for ws in wb.worksheets:
    sheet = ws.title
    rows = list(ws.iter_rows())
    nrow = len(rows)
    grid = [[c.value for c in r] for r in rows]
    tab = None
    # row -> table id (column A)
    tabs = []
    for r in grid:
        a = r[0] if r else None
        if isinstance(a, str) and len(a) <= 40:
            tab = a
        tabs.append(tab)
    # header rows with both 'RA' and 'Dec' (for T24-like grids)
    radec_hdr = {}
    for i, r in enumerate(grid):
        cols = {v: j for j, v in enumerate(r) if v in ('RA', 'Dec', 'RA midpoints', 'Dec midpoints')}
        if 'RA' in cols and 'Dec' in cols:
            radec_hdr[i] = cols['Dec']
    hdr_rows = sorted(radec_hdr)
    # timeline header rows
    tl_hdr = {}
    cur = None
    for i, r in enumerate(grid):
        if 'kind' in r and 'value at the off' in r and 'score at the off' in r:
            cur = {v: j for j, v in enumerate(r) if isinstance(v, str)}
        elif r and isinstance(r[0], str) and r[0] == 'table':
            cur = None
        tl_hdr[i] = cur

    def plane_from_col(i, j):
        # scan upward for a string in column j naming a plane (group tables)
        for k in range(i - 1, max(-1, i - 400), -1):
            x = grid[k][j] if j < len(grid[k]) else None
            if isinstance(x, str):
                if 'Equator' in x or x.startswith('equator'):
                    return 'Equator'
                if 'Dec' in x:
                    return 'Dec'
                if 'RA' in x:
                    return 'RA'
                return None
        return None

    def plane_from_row(r):
        ps = [x for x in r if isinstance(x, str) and x in PLANES]
        return ps[0] if ps else None

    def plane_t24(i, j):
        for k in reversed(hdr_rows):
            if k < i and i - k < 400:
                return 'Dec' if j >= radec_hdr[k] else 'RA'
        return None

    for i, row in enumerate(rows):
        r = grid[i]
        t = tabs[i]
        rowtxt = ' '.join(str(x) for x in r if isinstance(x, str))
        # ---- timelines
        h = tl_hdr.get(i)
        if h and r and r[0] != 'table' and h.get('kind') is not None and len(r) > h['kind'] and r[h['kind']] == 'distance':
            v = shown_num(r[h['value at the off']])
            tg = r[h['target']]
            sc = shown_num(r[h['score at the off']])
            pl = r[h['plane']]
            m = RE_GROUP.match(tg or '') if isinstance(tg, str) else None
            if v is not None and m:
                fam = FAMNORM[m.group(2)]
                shown = (fam, parse_label(m.group('lab'), fam), sc) if (sc is not None and m.group('rest') == '') else None
                if sc is not None and m.group('rest') != '':
                    unk[('tl-score-with-outside', sheet)] += 1
                check('timeline', sheet, row[h['value at the off']].coordinate, t, v, pl, shown)
            else:
                unk[('tl-unparsed', str(tg)[:40])] += 1
            # sky equivalent
            if 'sky value at the off' in h:
                sv = shown_num(r[h['sky value at the off']])
                stg = r[h['sky target']]
                ssc = shown_num(r[h['sky score']])
                if sv is not None and isinstance(stg, str):
                    m = RE_GROUP.match(stg)
                    if m:
                        fam = FAMNORM[m.group(2)]
                        shown = (fam, parse_label(m.group('lab'), fam), ssc) if (ssc is not None and m.group('rest') == '') else None
                        check('timeline sky equiv', sheet, row[h['sky value at the off']].coordinate, t, sv, pl, shown)
                    else:
                        unk[('tl-sky-unparsed', stg[:40])] += 1
            continue
        if h and r and r[0] != 'table':
            continue  # other timeline rows (ratios, midpoints)
        for j, c in enumerate(row):
            x = c.value
            # ---- comment cells (T24 kinds)
            if c.comment is not None and isinstance(x, (int, float)):
                first = c.comment.text.split('\n')[0]
                m = RE_CHIT.match(first)
                if m:
                    fam = m.group('fam')
                    if fam.startswith('conjunction'):
                        pl_c = 'RA' if '(RA' in fam else 'Dec'
                        fam = 'conj'
                    else:
                        pl_c = None
                    pl = plane_t24(i, j)
                    if pl_c and pl and pl != pl_c:
                        unk[('conj-plane-mismatch', sheet)] += 1
                    pl = pl or pl_c
                    check('comment', sheet, c.coordinate, t, float(x), pl,
                          (fam, parse_label(m.group('lab'), fam), float(m.group('sc'))))
                    continue
                m = RE_CNO.match(first)
                if m:
                    check('comment', sheet, c.coordinate, t, float(x), plane_t24(i, j), None)
                    continue
                if not first.startswith('nearest:') and ' · ' not in first:
                    unk[('comment-unparsed', re.sub(r'[-+]?\d+(\.\d+)?', '#', first)[:60])] += 1
                continue
            if not isinstance(x, str):
                continue
            # ---- leg cells
            if x == 'no family' or x.startswith('→ '):
                v = shown_num(r[j - 1]) if j >= 1 else None
                if v is None:
                    unk[('leg-no-value', sheet)] += 1
                    continue
                pl = plane_from_row(r)
                if x == 'no family':
                    shown = None
                else:
                    m = RE_LEG.match(x)
                    if not m:
                        unk[('leg-unparsed', re.sub(r'[-+]?\d+(\.\d+)?', '#', x)[:50])] += 1
                        continue
                    fam = FAMNORM[m.group(2)]
                    shown = (fam, parse_label(m.group('lab'), fam), float(m.group('sc')))
                check('leg', sheet, r and row[j - 1].coordinate, t, v, pl, shown)
                continue
            # ---- group target cells
            m = RE_GROUP.match(x)
            if m and j >= 2 and isinstance(r[j - 2], (int, float)) and not isinstance(r[j - 2], bool):
                rest = m.group('rest')
                if rest and not (rest.startswith(' — outside') or rest.startswith(' — not a 4-digit match')):
                    unk[('group-rest', re.sub(r'[-+]?\d+(\.\d+)?', '#', rest)[:40])] += 1
                    continue
                v = float(r[j - 2])
                pl = plane_from_col(i, j - 2)
                sc = shown_num(r[j + 3]) if j + 3 < len(r) else None
                fam = FAMNORM[m.group(2)]
                try:
                    parse_label(m.group('lab'), fam)
                except ValueError:
                    unk[('group-label-unparsed', x[:60], sheet[:8], tabs[i])] += 1
                    continue
                if rest == '':
                    shown = (fam, parse_label(m.group('lab'), fam), sc)
                    if sc is None:
                        unk[('group-hit-without-score', sheet)] += 1
                else:
                    shown = None
                    if sc is not None:
                        unk[('group-outside-with-score', sheet)] += 1
                check('group', sheet, row[j - 2].coordinate, t, v, pl, shown)
            elif m:
                unk[('group-like-no-value', re.sub(r'[-+]?\d+(\.\d+)?', '#', x)[:50], sheet[:5])] += 1
                unk_ex.setdefault(('group-like-no-value', re.sub(r'[-+]?\d+(\.\d+)?', '#', x)[:50], sheet[:5]), (sheet, c.coordinate, x[:80], tabs[i]))

res = dict(file=path, counts={f'{k[0]}|{k[1]}': n for k, n in counts.items()}, problems=problems, boundary=boundary,
           ties=[list(map(str, t)) for t in ties], unknown={str(k): n for k, n in unk.items()},
           unknown_ex={str(k): v for k, v in unk_ex.items()}, ninth_amb=[list(map(str, t)) for t in ninth_amb], samples=samples, rounding_ok=rounding_ok[0])
json.dump(res, open(out, 'w'), indent=1, ensure_ascii=False)
print(path, 'cells', sum(counts.values()), 'problems', len(problems), 'boundary', len(boundary), 'ties', len(ties),
      'ninth-amb', len(ninth_amb), 'score diffs explained by 6-dp rounding', rounding_ok[0])
for k, n in sorted(unk.items(), key=lambda x: -x[1])[:40]:
    print('  unk', n, k, unk_ex.get(k, ''))
for p in problems[:30]:
    print('  P', p)
