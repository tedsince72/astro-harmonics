"""Independent checker for MATCH HORSE | JOCKEY (own code; does not import jupiter_tables or chk_matchhj).

usage: check_hj.py before.xlsx after.xlsx [--no-before] [--hand-ripon]
Task 1: every before cell identical in after (value, fill, bold, number format, border, comment); widths, freeze,
merges (only the two new per runner sheet); nothing else added.
Task 2: the new table rebuilt from MATCH HORSE / MATCH JOCKEY on the after sheet, compared row by row.
"""
import sys, re, math
from collections import defaultdict, OrderedDict
import openpyxl
from openpyxl.utils import get_column_letter

PHI = (1 + 5 ** 0.5) / 2
R2 = 2 ** 0.5
NEW = 'MATCH HORSE | JOCKEY'
GREY = 'D9D9D9'
NC = 29  # A..AC

errors = []


def err(msg):
    errors.append(msg)


def fill_of(c):
    f = c.fill
    if f is None or f.fill_type is None:
        return None
    rgb = f.fgColor.rgb if f.fgColor is not None else None
    return (f.fill_type, str(rgb)[-6:] if rgb else None)


def border_of(c):
    b = c.border
    return tuple((getattr(b, s).style if getattr(b, s) is not None else None,
                  str(getattr(b, s).color.rgb) if getattr(b, s) is not None and getattr(b, s).color is not None else None)
                 for s in ('left', 'right', 'top', 'bottom'))


def sig(c):
    return (c.value, fill_of(c), bool(c.font.b) if c.font else False, c.number_format, border_of(c),
            c.comment.text if c.comment else None)


def sig_nocomment(c):
    return sig(c)[:5]


# ---------------- task 1 ----------------
def compare_before_after(wb0, wb1):
    n = 0
    if wb0.sheetnames != wb1.sheetnames:
        err(f'sheet names differ {wb0.sheetnames} vs {wb1.sheetnames}')
    new_ranges = {}
    for name in wb0.sheetnames:
        a, b = wb0[name], wb1[name]
        runner = re.fullmatch(r'P\d\d', name) is not None
        mr0, mc0 = a.max_row, a.max_column
        for r in range(1, mr0 + 1):
            for c in range(1, max(mc0, b.max_column) + 1):
                n += 1
                if sig(a.cell(r, c)) != sig(b.cell(r, c)):
                    err(f'{name}!{get_column_letter(c)}{r}: before {sig(a.cell(r, c))} after {sig(b.cell(r, c))}')
        # rows after before.max_row
        if not runner:
            if b.max_row != mr0:
                err(f'{name}: max_row {mr0} -> {b.max_row}')
        else:
            # two blank rows then the new table to the end
            for r in (mr0 + 1, mr0 + 2):
                for c in range(1, b.max_column + 1):
                    if b.cell(r, c).value is not None or fill_of(b.cell(r, c)) is not None:
                        err(f'{name}!{get_column_letter(c)}{r}: not blank between MATCH JOCKEY and the new table')
            if b.cell(mr0 + 3, 1).value != NEW:
                err(f'{name}: row {mr0 + 3} col A = {b.cell(mr0 + 3, 1).value!r}, expected the new table')
            for r in range(mr0 + 3, b.max_row + 1):
                if b.cell(r, 1).value != (('table') if r == mr0 + 5 else NEW):
                    err(f'{name}!A{r}: {b.cell(r, 1).value!r} inside / after the new table')
                for c in range(NC + 1, b.max_column + 1):
                    if b.cell(r, c).value is not None or fill_of(b.cell(r, c)) is not None:
                        err(f'{name}!{get_column_letter(c)}{r}: content beyond AC')
            new_ranges[name] = (mr0 + 3, b.max_row)
        # widths
        keys = set(a.column_dimensions.keys()) | set(b.column_dimensions.keys())
        for k in keys:
            wa = a.column_dimensions[k].width if k in a.column_dimensions else None
            wb_ = b.column_dimensions[k].width if k in b.column_dimensions else None
            if wa != wb_:
                err(f'{name}: width {k} {wa} -> {wb_}')
        if a.freeze_panes != b.freeze_panes:
            err(f'{name}: freeze {a.freeze_panes} -> {b.freeze_panes}')
        m0 = set(str(m) for m in a.merged_cells.ranges)
        m1 = set(str(m) for m in b.merged_cells.ranges)
        if not m0 <= m1:
            err(f'{name}: merges lost {sorted(m0 - m1)}')
        added = m1 - m0
        if runner:
            band = mr0 + 4
            exp = {f'N{band}:U{band}', f'V{band}:AC{band}'}
            if added != exp:
                err(f'{name}: merges added {sorted(added)} expected {sorted(exp)}')
        elif added:
            err(f'{name}: merges added {sorted(added)}')
    return n, new_ranges


# ---------------- task 2 ----------------
def tval(fam, t):
    t = str(t)
    if fam == 'whole':
        return float(int(t))
    if fam == 'ninth':
        return float(t)
    if fam == 'sqrt2':
        return {'0.707': R2 / 2, '7.07': 10 * R2 / 2, '70.7': 100 * R2 / 2,
                '1.414': R2, '14.14': 10 * R2, '141.4': 100 * R2}[t]
    if fam == 'silver':
        return {'0.414': R2 - 1, '4.14': 10 * (R2 - 1), '41.4': 100 * (R2 - 1),
                '2.414': R2 + 1, '24.14': 10 * (R2 + 1)}[t]
    m = re.fullmatch(r'(?:(\d+)×)?φ\^(-?\d+)', t)
    if m:
        return (int(m.group(1)) if m.group(1) else 1) * PHI ** int(m.group(2))
    m = re.fullmatch(r'(\d+)×φ', t)
    if m:
        return int(m.group(1)) * PHI
    if t == 'φ':
        return PHI
    raise ValueError(f'unknown target {fam} {t}')


def pq(s):
    s = str(s)
    if '/' in s:
        p, q = s.split('/')
        return int(p), int(q)
    return int(s), 1


def read_source(ws, table):
    """groups: OrderedDict sky -> {'head': prefix, 'rows': [row cell lists]}"""
    groups = OrderedDict()
    cur = None
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value != table:
            continue
        b = ws.cell(r, 2).value
        if b in ('EXACT', 'RATIO'):
            row = [ws.cell(r, c) for c in range(1, 22)]
            sky = (row[7].value, row[8].value)
            if cur is None or cur != sky:
                err(f'{ws.title} {table} row {r}: data row outside its group header')
            groups[sky]['rows'].append((r, row))
        elif isinstance(b, str) and ' — EXACT ' in b and fill_of(ws.cell(r, 2)) and fill_of(ws.cell(r, 2))[1] == GREY:
            prefix = b.split(' — ')[0]
            m = re.fullmatch(r'(.+) (RA|Dec) (-?[\d.]+) (\S+) (.+?) \(([\d.]+)\)', prefix)
            sky = (m.group(1), m.group(2))
            groups[sky] = {'head': prefix, 'score': float(m.group(6)), 'rows': [], 'hdr': b}
            cur = sky
    # own counts vs header
    for sky, g in groups.items():
        ne = sum(1 for _, row in g['rows'] if row[1].value == 'EXACT')
        nr = len(g['rows']) - ne
        if g['hdr'] != f"{g['head']} — EXACT {ne}, RATIO {nr}":
            err(f'{ws.title} {table}: source header {g["hdr"]!r} vs counts {ne}/{nr}')
    return groups


def merge_order(oh, oj, scores):
    """union of the two group orders, respecting both and non-increasing 2-dp score."""
    i = j = 0
    out = []
    ties = []
    while i < len(oh) or j < len(oj):
        if i < len(oh) and oh[i] in out:
            i += 1; continue
        if j < len(oj) and oj[j] in out:
            j += 1; continue
        if i >= len(oh):
            out.append(oj[j]); j += 1; continue
        if j >= len(oj):
            out.append(oh[i]); i += 1; continue
        a, b = oh[i], oj[j]
        if a == b:
            out.append(a); i += 1; j += 1; continue
        sa, sb = scores[a], scores[b]
        if sa > sb:
            out.append(a); i += 1
        elif sb > sa:
            out.append(b); j += 1
        else:
            # 2-dp tie between different sky cells: if one appears later in the other list, the order is fixed by it
            ties.append((a, b, sa))
            if b in oh[i:]:
                out.append(a); i += 1
            elif a in oj[j:]:
                out.append(b); j += 1
            else:
                out.append(a); i += 1  # ambiguous: reported
    return out, ties


def expected_table(ws):
    H = read_source(ws, 'MATCH HORSE')
    J = read_source(ws, 'MATCH JOCKEY')
    scores = {}
    heads = {}
    for G in (H, J):
        for sky, g in G.items():
            if sky in scores and (scores[sky] != g['score'] or heads[sky] != g['head']):
                err(f'{ws.title}: sky header differs H/J {heads[sky]!r} vs {g["head"]!r}')
            scores[sky] = g['score']
            heads[sky] = g['head']
    order, ties = merge_order(list(H), list(J), scores)
    for k in range(len(order) - 1):
        if scores[order[k]] < scores[order[k + 1]]:
            err(f'{ws.title}: group order not by score')
    lines = []  # list of ('hdr', text) or ('row', shared_src, hrow, jrow, key)
    stats = dict(lines=0, both=0, h=0, j=0, sky=len(order),
                 nh=sum(len(g['rows']) for g in H.values()), nj=sum(len(g['rows']) for g in J.values()))
    for sky in order:
        hr = H[sky]['rows'] if sky in H else []
        jr = J[sky]['rows'] if sky in J else []
        cnt = lambda rows, s: sum(1 for _, row in rows if row[1].value == s)
        lines.append(('hdr', f"{heads[sky]} — EXACT {cnt(hr, 'EXACT')} | {cnt(jr, 'EXACT')}, "
                             f"RATIO {cnt(hr, 'RATIO')} | {cnt(jr, 'RATIO')}"))
        keyed = defaultdict(lambda: [[], []])
        for side, rows in ((0, hr), (1, jr)):
            for idx, (r, row) in enumerate(rows):
                key = (sky, row[1].value, str(row[2].value), row[17].value, str(row[18].value))
                keyed[key][side].append((-float(row[19].value), idx, row))

        def sk(key):
            s, sec, ratio, fam, tgt = key
            p, q = pq(ratio)
            return (0 if sec == 'EXACT' else 1, q, p, tval(fam, tgt), fam, tgt)
        keys = sorted(keyed, key=sk)
        for a, b in zip(keys, keys[1:]):
            if sk(a)[:4] == sk(b)[:4]:
                err(f'{ws.title}: two keys tie on numeric target {a} {b}')
        for key in keys:
            hs = [x[2] for x in sorted(keyed[key][0], key=lambda x: (x[0], x[1]))]
            js = [x[2] for x in sorted(keyed[key][1], key=lambda x: (x[0], x[1]))]
            if hs and js:
                stats['both'] += 1
            elif hs:
                stats['h'] += 1
            else:
                stats['j'] += 1
            for n in range(max(len(hs), len(js))):
                h = hs[n] if n < len(hs) else None
                j = js[n] if n < len(js) else None
                if h is not None and j is not None:
                    for c in range(1, 13):
                        if sig_nocomment(h[c]) != sig_nocomment(j[c]):
                            err(f'{ws.title}: key {key} shared col {get_column_letter(c + 1)} differs H/J '
                                f'{h[c].value!r} / {j[c].value!r}')
                lines.append(('row', h if h is not None else j, h, j, key))
                stats['lines'] += 1
    return lines, stats, ties


EMPTY = (None, None, False, 'General', ((None, None),) * 4, None)


def check_table(ws, lo, hi, lines, stats):
    rows = list(range(lo, hi + 1))
    # title
    t = ws.cell(lo, 2).value or ''
    m = re.search(r'(\d+) lines, (\d+) keys: (\d+) with both sides, (\d+) horse only, (\d+) jockey only; (\d+) sky cells '
                  r'\(MATCH HORSE (\d+) rows, MATCH JOCKEY (\d+) rows\)', t)
    if not m:
        err(f'{ws.title}: title totals not found')
    else:
        got = tuple(int(x) for x in m.groups())
        exp = (stats['lines'], stats['both'] + stats['h'] + stats['j'], stats['both'], stats['h'], stats['j'],
               stats['sky'], stats['nh'], stats['nj'])
        if got != exp:
            err(f'{ws.title}: title totals {got} expected {exp}')
    band = lo + 1
    for c in range(2, NC + 1):
        v = ws.cell(band, c).value
        e = 'HORSE' if c == 14 else 'JOCKEY' if c == 22 else None
        if v != e:
            err(f'{ws.title}!{get_column_letter(c)}{band}: band {v!r} expected {e!r}')
    names = ['table', 'section', 'ratio p/q', 'ratio', 'shared factor', 'circle factor', 'slice of 360',
             'sky: relationship', 'plane', 'value', 'family', 'target', 'score']
    side = ['T24: row', 'body', 'plane', 'value', 'family', 'target', 'score', 'lower score']
    expn = names + side + side
    gotn = [ws.cell(lo + 2, c).value for c in range(1, NC + 1)]
    if gotn != expn:
        err(f'{ws.title}: column names {gotn}')
    data = rows[3:]
    if len(data) != len(lines):
        err(f'{ws.title}: {len(data)} sheet rows after the headers, expected {len(lines)}')
    for r, ln in zip(data, lines):
        if ln[0] == 'hdr':
            if ws.cell(r, 2).value != ln[1]:
                err(f'{ws.title}!B{r}: header {ws.cell(r, 2).value!r} expected {ln[1]!r}')
            for c in range(1, NC + 1):
                f = fill_of(ws.cell(r, c))
                if f != ('solid', GREY):
                    err(f'{ws.title}!{get_column_letter(c)}{r}: header fill {f}')
                if c > 2 and ws.cell(r, c).value is not None:
                    err(f'{ws.title}!{get_column_letter(c)}{r}: header row value {ws.cell(r, c).value!r}')
            if fill_of(ws.cell(r, NC + 1)) is not None:
                err(f'{ws.title}!AD{r}: grey beyond AC')
            continue
        _, sh, h, j, key = ln
        if ws.cell(r, 1).value != NEW:
            err(f'{ws.title}!A{r}: {ws.cell(r, 1).value!r}')
        if fill_of(ws.cell(r, 1)) != fill_of(sh[0]):
            err(f'{ws.title}!A{r}: fill {fill_of(ws.cell(r, 1))}')
        for c in range(2, 14):
            if sig(ws.cell(r, c)) != sig(sh[c - 1]):
                err(f'{ws.title}!{get_column_letter(c)}{r}: {sig(ws.cell(r, c))} expected '
                    f'{sig(sh[c - 1])} key {key}')
        for base, src in ((14, h), (22, j)):
            for k in range(8):
                c = base + k
                got = sig(ws.cell(r, c))
                exp = sig(src[13 + k]) if src is not None else EMPTY
                if got != exp:
                    err(f'{ws.title}!{get_column_letter(c)}{r}: {got} expected {exp} key {key}')


def main():
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    before, after = args
    wb1 = openpyxl.load_workbook(after)
    if '--no-before' not in sys.argv:
        wb0 = openpyxl.load_workbook(before)
        n, ranges = compare_before_after(wb0, wb1)
        print(f'task 1: cells compared {n}')
    else:
        ranges = {}
        for name in wb1.sheetnames:
            if re.fullmatch(r'P\d\d', name):
                ws = wb1[name]
                rs = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == NEW]
                ranges[name] = (rs[0], rs[-1])
    for name in wb1.sheetnames:
        if not re.fullmatch(r'P\d\d', name):
            continue
        ws = wb1[name]
        lo, hi = ranges[name]
        lines, st, ties = expected_table(ws)
        check_table(ws, lo, hi, lines, st)
        print(f"{name}: lines {st['lines']}, keys {st['both'] + st['h'] + st['j']}: both {st['both']} / horse only "
              f"{st['h']} / jockey only {st['j']}, sky cells {st['sky']}, rows H {st['nh']} J {st['nj']}, "
              f"2-dp sky score ties {len(ties)} {ties if ties else ''}")
    if '--hand-ripon' in sys.argv:
        ws = wb1['P01']
        lo, hi = ranges['P01']
        for r in range(lo, hi + 1):
            if str(ws.cell(r, 2).value).startswith('S Jupiter → S Wezen'):
                print('hand:', ws.cell(r, 2).value)
                print('hand:', [ws.cell(r + 1, c).value for c in range(1, NC + 1)])
                break
    print(f'mismatches: {len(errors)}')
    for e in errors[:30]:
        print('  ', e)


if __name__ == '__main__':
    main()
