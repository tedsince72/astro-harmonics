import sys, openpyxl
from openpyxl.utils import get_column_letter as L
for new, old in zip(sys.argv[1::2], sys.argv[2::2]):
    a, b = openpyxl.load_workbook(new), openpyxl.load_workbook(old)
    bad, narrow, diff, n = [], [], 0, 0
    for s in a.sheetnames:
        wa, wb_ = a[s], b[s]
        for c in range(1, max(wa.max_column, wb_.max_column) + 1):        # never narrower
            wx = wa.column_dimensions[L(c)].width if L(c) in wa.column_dimensions else 8.43
            wy = wb_.column_dimensions[L(c)].width if L(c) in wb_.column_dimensions else 8.43
            if wx < wy: narrow.append((s, L(c), wx, wy))
        for ra, rb in zip(wa.iter_rows(), wb_.iter_rows()):              # values, fills, comments unchanged
            for p, q in zip(ra, rb):
                fa = p.fill.start_color.rgb if p.fill.fill_type else None
                fb = q.fill.start_color.rgb if q.fill.fill_type else None
                if (p.value, fa, p.comment and p.comment.text) != (q.value, fb, q.comment and q.comment.text): diff += 1
        if wa.max_row != wb_.max_row: diff += 1
        blk = False
        for r in wa.iter_rows():                                          # every T24 cell fits its column
            if r[0].value == 'T24' and str(r[1].value).startswith('T24 '): blk = True; continue
            if blk and r[0].value == 'T24':
                for c in r[1:]:
                    if c.value is None: continue
                    txt = f'{c.value:.6f}' if isinstance(c.value, float) else str(c.value)
                    w = wa.column_dimensions[c.column_letter].width
                    n += 1
                    if len(txt) + 1 > w: bad.append((s, c.coordinate, txt, w))
            elif r[0].value != 'table' and r[0].value is None and blk and r[1].value is None and r[2].value is None: blk = blk
    print(new.split('/')[-1], 'T24 cells', n, 'too narrow', bad[:3], len(bad), '| columns narrower than before', narrow[:3], len(narrow), '| cells changed', diff)
