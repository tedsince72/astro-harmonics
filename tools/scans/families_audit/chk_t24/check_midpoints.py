import sys, openpyxl
from openpyxl.utils import get_column_letter as L
BODIES=['Sun','Moon','Mercury','Venus','Mars','Jupiter','Saturn','Uranus','Neptune','Pluto','Chiron','Ceres','Pallas','Juno','Vesta','Eris','Sedna','Haumea','Makemake','Gonggong','Quaoar','Orcus','Transpluto','Rahu','Ketu']
def arc(a,b):
    d=abs(a-b)%360; return min(d,360-d)
def calc(plane,s,a,b):
    if plane=='RA':
        mid=(a+(((b-a+180)%360)-180)/2)%360
        near=((s-mid+180)%360)-180; far=((s-mid)%360)-180
        if abs(far)<abs(near): off,side=far,'far'
        else: off,side=near,'near'
        d1,d2,dab=arc(s,a),arc(s,b),arc(a,b); tol=0.5
    else:
        off=s-(a+b)/2; side=None; d1,d2,dab=abs(s-a),abs(s-b),abs(a-b); tol=0.15
    clo=100*min(d1,d2)/max(d1,d2) if max(d1,d2)>0 else 100
    rin=100-clo<=10; din=abs(off)<=tol
    rh=0.5*(10*(10-(100-clo))/10)**2 if rin else None
    dh=0.5*(10*(tol-abs(off))/tol)**2 if din else None
    btw=abs(dab-(d1+d2))<=0.01*(d1+d2)
    return dict(off=off,side=side,d1=d1,d2=d2,dab=dab,clo=clo,hit=rin and din,tot=(rh+dh) if rin and din else None,btw=btw)
def fillc(t): return 'C6EFCE' if t>=80 else ('FFEB9C' if t>=50 else 'FFF8DC')
def run(path,sheet,runner,chart,blockname):
    wb=openpyxl.load_workbook(path); P=wb['POSITIONS']
    pos={}; stars=[]
    for r in P.iter_rows(min_row=5,values_only=True):
        if r[0]==runner and r[2]==chart:
            pos[r[3]]=(r[5],r[6])
            if r[4]=='star': stars.append(r[3])
    ws=wb[sheet]
    T=None
    for r in range(1,ws.max_row+1):
        if ws.cell(r,1).value=='T24' and str(ws.cell(r,2).value or '').startswith('T24 '+blockname): T=r;break
    h1,h2=T+1,T+2
    cols={}
    for c in range(3,ws.max_column+1):
        v=ws.cell(h1,c).value
        if v=='RA midpoints': raM=c
        if v=='Dec midpoints': deM=c;break
    bl=[b for b in BODIES if b!='Moon' or chart=='S']
    raC=[ws.cell(h2,c).value for c in range(3,raM)]; deC=[ws.cell(h2,c).value for c in range(raM+1,deM)]
    mism=[]
    if raC!=bl or deC!=bl: mism.append(('header',raC,deC))
    bl=[b for b in bl if b in pos]
    rowstars=[]; r=T+3
    while ws.cell(r,1).value=='T24' and ws.cell(r,2).value: rowstars.append((r,ws.cell(r,2).value)); r+=1
    exp=[chart+' '+s for s in stars]
    if [s for _,s in rowstars]!=exp: mism.append(('stars',[s for _,s in rowstars],exp))
    cells=0; nh={'RA':0,'Dec':0}; dropped=[]
    for r,lab in rowstars:
        st=lab[2:]
        for plane,base,mc,idx in (('RA',3,raM,0),('Dec',raM+1,deM,1)):
            cells+=1
            s=pos[st][idx]; hits=[]; res={}
            for i,a in enumerate(bl):
                for b in bl[i+1:]:
                    res[(a,b)]=calc(plane,s,pos[a][idx],pos[b][idx])
            for (a,b),x in res.items():
                if not x['hit']: continue
                why=None
                if {a,b}=={'Rahu','Ketu'}: why='Rahu–Ketu'
                elif x['dab']<(1.5 if plane=='RA' else 0.3): why='together'
                elif 'Ketu' in (a,b):
                    o=a if b=='Ketu' else b
                    k=tuple(sorted([o,'Rahu'],key=BODIES.index)); y=res.get(k)
                    if y and y['hit'] and {round(y['d1'],6),round(y['d2'],6)}=={round(x['d1'],6),round(x['d2'],6)}: why='Ketu repeat'
                if why: dropped.append((lab,plane,a,b,why,round(x['tot'],2))); continue
                hits.append((a,b,x))
            hits.sort(key=lambda h:-h[2]['tot'])
            nh[plane]+=len(hits)
            def txt(a,b,x):
                pl=(x['side']+', ' if plane=='RA' else '')+('between' if x['btw'] else 'far side')
                return f"{a} = {b} {x['tot']:.2f} ({pl})"
            et='; '.join(txt(*h) for h in hits)
            got=ws.cell(r,mc).value or ''
            if got!=et: mism.append((lab,plane,'text',got,et))
            c=ws.cell(r,mc); f=c.fill.fgColor.rgb[-6:] if c.fill.fill_type else None
            ef=fillc(hits[0][2]['tot']) if hits else None
            if f!=ef: mism.append((lab,plane,'fill',f,ef))
            lines={b:[] for b in bl}
            for a,b,x in hits:
                lines[a].append(f"midpoint with {b}: total {x['tot']:.2f}")
                lines[b].append(f"midpoint with {a}: total {x['tot']:.2f}")
            for j,bn in enumerate(bl):
                cc=ws.cell(r,base+([b for b in BODIES if b!='Moon' or chart=='S'].index(bn)))
                bd=cc.border; thick=all(getattr(bd,k).style=='thick' for k in ('left','right','top','bottom'))
                anyth=any(getattr(bd,k).style=='thick' for k in ('left','right','top','bottom'))
                if bool(lines[bn])!=thick or (not lines[bn] and anyth): mism.append((lab,plane,bn,'border',thick,bool(lines[bn])))
                cl=[l for l in (cc.comment.text.split('\n') if cc.comment else []) if l.startswith('midpoint')]
                if cl!=lines[bn]:
                    kind='comment-order' if sorted(cl)==sorted(lines[bn]) else 'comment'
                    mism.append((lab,plane,bn,kind,cl,lines[bn]))
    print(f'== {path.split("/")[-1]} {sheet} T24 {blockname}: star-plane cells {cells}, hits RA {nh["RA"]} Dec {nh["Dec"]}, mismatches {len(mism)}')
    for m in mism[:10]: print('  MISMATCH',m)
    for d in dropped: print('  dropped',d)
D='/private/tmp/claude-501/-Users-eddieplace-astro-harmonics/3d86481b-e19f-4be4-91e4-da248144d5be/scratchpad/chk29/'
run(D+'ascot_gc.xlsx','SKY JUPITER','SKY','S','SKY')
run(D+'ascot_gc.xlsx','P01','P01','H','HORSE')
run(D+'ascot_gc.xlsx','P01','P01','J','JOCKEY')
run(D+'lingfield.xlsx','SKY JUPITER','SKY','S','SKY')
run(D+'lingfield.xlsx','P07','P07','H','HORSE')
run(D+'lingfield.xlsx','P07','P07','J','JOCKEY')
