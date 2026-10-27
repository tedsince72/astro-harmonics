"""Independent checker (27 Oct 2026) for the display change after a2d20618: SKY STARS SUMMARY two header rows on top + freeze C3,
grey 'stars' separator rows, column widths. Own code (openpyxl only).

usage: check_display_indep.py <a2d20618 tables> <after tables> <35255b16 tables> <variant tables (lay(m) only)> [pairs_a pairs_b]
"""
import sys
import openpyxl
from openpyxl.utils import get_column_letter as L

SUM = 'SKY STARS SUMMARY'


def col(c):
    if c is None:
        return None
    return (c.type, c.rgb if c.type == 'rgb' else c.theme if c.type == 'theme' else c.indexed, round(c.tint or 0, 6))


def side(s):
    return (s.style, col(s.color)) if s is not None else None


def sig(cell, with_fill=True):
    f, fi, b, a = cell.font, cell.fill, cell.border, cell.alignment
    return (cell.value, cell.number_format,
            (f.name, f.sz, f.b, f.i, f.u, f.strike, col(f.color)),
            (fi.patternType, col(fi.fgColor), col(fi.bgColor)) if with_fill else None,
            tuple(side(getattr(b, k)) for k in ('left', 'right', 'top', 'bottom', 'diagonal')),
            (a.horizontal, a.vertical, a.wrap_text, a.indent, a.text_rotation),
            cell.comment.text if cell.comment else None)


def cells(ws):
    return {k for k in ws._cells}


def widths(ws):
    w = {}
    for d in ws.column_dimensions.values():
        if d.min is None:
            continue
        for c in range(d.min, d.max + 1):
            if d.customWidth and d.width:
                w[c] = d.width
    return w


def merges(ws, shift=0):
    return {(m.min_row + shift, m.min_col, m.max_row + shift, m.max_col) for m in ws.merged_cells.ranges}


def sep_rows(ws):
    """('stars' separator row, table width) for every T24 (SKY STARS) table: width = last column used by its header rows"""
    out = []
    hdr = None
    for r in range(1, ws.max_row + 1):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        if a == 'table' and b == 'star' and ws.cell(r, 3).value == 'sky hit':
            last = max(c for (rr, c) in ws._cells if rr in (r, r + 1) and ws._cells[(rr, c)].value is not None)
            for m in ws.merged_cells.ranges:
                if m.min_row in (r, r + 1):
                    last = max(last, m.max_col)
            hdr = last
        if isinstance(a, str) and a.endswith('(SKY STARS)') and b == 'stars':
            out.append((r, hdr))
    return out


def t24_headers(ws):
    """[(row, {col: need})] for every T24 / T24 (SKY STARS) header pair (row with 'RA midpoints')"""
    res = []
    for (r, c), cell in list(ws._cells.items()):
        if cell.value == 'RA midpoints':
            need = {2: 18}
            mr = [m for m in ws.merged_cells.ranges if m.min_row == r]
            for m in mr:
                lab = ws.cell(r, m.min_col).value
                if lab in ('RA', 'Dec', 'SCORES RA', 'SCORES Dec'):
                    for k in range(m.min_col, m.max_col + 1):
                        need[k] = 11
            for k in range(1, ws.max_column + 1):
                v = ws.cell(r, k).value
                if v in ('RA midpoints', 'Dec midpoints'):
                    need[k] = 40
                elif v == 'repeats':
                    need[k] = 60
                elif v == 'sky hit':
                    need[k] = 36
            res.append((r, need))
    return res


def main():
    fa, fn, fh, fv = sys.argv[1:5]
    A = openpyxl.load_workbook(fa)
    N = openpyxl.load_workbook(fn)
    H = openpyxl.load_workbook(fh)
    V = openpyxl.load_workbook(fv)
    bad = []
    n_cells = 0
    print('sheets', A.sheetnames == N.sheetnames, N.sheetnames)
    if A.sheetnames != N.sheetnames:
        bad.append(('sheetnames', A.sheetnames, N.sheetnames))
    wchg = []
    for sn in N.sheetnames:
        a, n = A[sn], N[sn]
        sh = 2 if sn == SUM else 0
        seps = {r: w for r, w in sep_rows(n)} if (sn.startswith('P') and sn != 'POSITIONS') or sn == SUM else {}
        # separators found the same way in the before build (shifted)
        seps_a = {r + sh: w for r, w in sep_rows(a)} if seps else {}
        if seps and seps != seps_a:
            bad.append((sn, 'separator rows / widths differ', sorted(seps.items()), sorted(seps_a.items())))
        # cells
        keys = cells(n) | {(r + sh, c) for (r, c) in cells(a)}
        if sh:
            keys |= {(r, c) for (r, c) in cells(a) if r in (3, 4)} | {(r - 2, c) for (r, c) in cells(a) if r in (3, 4)}
        nsep = 0
        for (r, c) in sorted(keys):
            ncell = n.cell(r, c)
            if sh and r <= 2:
                acell = a.cell(r + 2, c)
            else:
                if r - sh < 1:
                    continue
                acell = a.cell(r - sh, c)
            n_cells += 1
            if r in seps and c <= seps[r]:
                nsep += 1
                fi = ncell.fill
                if not (fi.patternType == 'solid' and fi.fgColor.type == 'rgb' and fi.fgColor.rgb[-6:] == 'D9D9D9'):
                    bad.append((sn, ncell.coordinate, 'separator not grey', fi.patternType, col(fi.fgColor)))
                if sig(ncell, False) != sig(acell, False):
                    bad.append((sn, ncell.coordinate, 'separator cell changed', sig(acell, False), sig(ncell, False)))
                continue
            if sig(ncell) != sig(acell):
                bad.append((sn, ncell.coordinate, sig(acell), sig(ncell)))
        if seps:
            print(sn, 'separators', len(seps), 'grey cells', nsep, 'widths', sorted(set(seps.values())))
        # merges
        ma = merges(a, sh)
        if sh:
            ma |= {(r1 - 4, c1, r2 - 4, c2) for (r1, c1, r2, c2) in ma if 5 <= r1 and r2 <= 6}
        if merges(n) != ma:
            bad.append((sn, 'merges', sorted(merges(n) - ma), sorted(ma - merges(n))))
        fz = n.freeze_panes
        exp = 'C3' if sn == SUM else a.freeze_panes
        if fz != exp:
            bad.append((sn, 'freeze', a.freeze_panes, fz))
        if sn == SUM:
            print(SUM, 'freeze before', a.freeze_panes, 'after', fz, 'rows', a.max_row, '->', n.max_row)
        # widths
        wa, wn, wv = widths(a), widths(n), widths(V[sn])
        wh = widths(H[sn]) if sn in H.sheetnames else {}
        for c in sorted(set(wa) | set(wn) | set(wh)):
            x, y, z = wa.get(c), wn.get(c), wh.get(c)
            if y is None and (x is not None or z is not None) or (y is not None and ((x or 0) > y + 1e-9 or (z or 0) > y + 1e-9)):
                bad.append((sn, 'narrower', L(c), 'a2', x, '35255b16', z, 'after', y))
            if x != y:
                wchg.append((sn, L(c), x, y, z, wv.get(c)))
        for r, need in t24_headers(n):
            for c, w in need.items():
                y = wn.get(c)
                if y is None or y + 1e-9 < w:
                    bad.append((sn, 'header col too narrow', r, L(c), w, y))
    print('cells compared', n_cells)
    print('widths changed vs a2d20618 (sheet, col, a2, after, 35255b16, variant lay(m) only):')
    for t in wchg:
        print('  ', t)
    # columns wider than the current layout needs: after > variant
    extra = []
    for sn in N.sheetnames:
        wn, wv = widths(N[sn]), widths(V[sn])
        for c in sorted(set(wn) | set(wv)):
            if wn.get(c) != wv.get(c):
                extra.append((sn, L(c), wv.get(c), wn.get(c)))
    print('columns where after != variant (kept wide only by the old layout):', len(extra))
    for t in extra:
        print('  ', t)
    if len(sys.argv) > 6:
        pa, pb = openpyxl.load_workbook(sys.argv[5]), openpyxl.load_workbook(sys.argv[6])
        np = 0
        same = pa.sheetnames == pb.sheetnames
        for sn in pa.sheetnames:
            x, y = pa[sn], pb[sn]
            for k in cells(x) | cells(y):
                np += 1
                if sig(x.cell(*k)) != sig(y.cell(*k)):
                    same = False
                    bad.append(('pairs', sn, k))
            if widths(x) != widths(y) or merges(x) != merges(y) or x.freeze_panes != y.freeze_panes:
                same = False
                bad.append(('pairs layout', sn))
        print('pairs cells', np, 'identical', same)
    print('MISMATCHES', len(bad))
    for t in bad[:200]:
        print('  ', t)


if __name__ == '__main__':
    main()
