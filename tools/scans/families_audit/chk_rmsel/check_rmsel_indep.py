"""Independent check (22 Oct 2026): T24 STAR SELECT removed; everything else unchanged. Usage: check_rmsel_indep.py <before_dir> <after_dir>"""
import sys, os, openpyxl
from openpyxl.utils import range_boundaries
B, A = sys.argv[1], sys.argv[2]
TID = 'T24 STAR SELECT'
mism = []

def cellsig(c):
    f = c.fill
    fg = (f.fill_type, f.fgColor.rgb if f.fgColor is not None else None) if f is not None else None
    b = c.border
    bd = tuple((getattr(b, s).style if getattr(b, s) is not None else None) for s in ('left', 'right', 'top', 'bottom'))
    return (c.value, fg, c.comment.text if c.comment else None, bd, c.number_format, bool(c.font and c.font.b))

def rows(ws):
    out = []
    for r in ws.iter_rows(min_row=1, max_row=ws.max_row):
        sig = [cellsig(c) for c in r]
        while sig and sig[-1][0] is None and sig[-1][1] in (None, (None, '00000000')) and sig[-1][2] is None and sig[-1][3] == (None,)*4:
            sig.pop()
        out.append(sig)
    return out

def find_select_region(rws):
    """indices (0-based) of rows in the SELECT region: from the first row with col A == TID to the start of the next block
    (next row whose col A is a non-empty value other than TID / 'table'), including blank rows"""
    idx = [i for i, r in enumerate(rws) if r and r[0][0] == TID]
    if not idx:
        return set(), 0
    s = idx[0]
    e = s
    i = s
    while i < len(rws):
        a = rws[i][0][0] if rws[i] else None
        if a not in (None, TID, 'table'):
            break
        i += 1
    e = i
    # trailing blank rows belong to the region (write_t24sel adds one per table); keep them in
    titles = sum(1 for r in rws[s:e] if r and r[0][0] == TID and len(r) > 1 and isinstance(r[1][0], str) and r[1][0].startswith(TID + ' —'))
    return set(range(s, e)), titles

def comp(fb, fa, label, select_expected):
    wb = openpyxl.load_workbook(fb); wa = openpyxl.load_workbook(fa)
    res = {}
    if wb.sheetnames != wa.sheetnames:
        mism.append(f'{label}: sheet names differ {wb.sheetnames} vs {wa.sheetnames}')
    for sn in wa.sheetnames:
        sb, sa = wb[sn], wa[sn]
        rb, ra = rows(sb), rows(sa)
        for i, r in enumerate(ra):
            for c in r:
                if isinstance(c[0], str) and TID in c[0]:
                    mism.append(f'{label} {sn} row {i+1}: contains {TID}')
        reg, ntab = find_select_region(rb) if select_expected else (set(), 0)
        kept = [i for i in range(len(rb)) if i not in reg]
        # after may have trailing empty rows trimmed; compare
        ra2 = ra[:]
        while ra2 and not ra2[-1]: ra2.pop()
        kb = [rb[i] for i in kept]
        while kb and not kb[-1]: kb.pop()
        if len(kb) != len(ra2):
            mism.append(f'{label} {sn}: kept rows {len(kb)} vs after {len(ra2)}')
        for j, (x, y) in enumerate(zip(kb, ra2)):
            if x != y:
                k = next((q for q in range(max(len(x), len(y))) if (x[q] if q < len(x) else None) != (y[q] if q < len(y) else None)), None)
                mism.append(f'{label} {sn}: after row {j+1} col {k+1 if k is not None else "?"} differs: {x[k] if k is not None and k < len(x) else None} | {y[k] if k is not None and k < len(y) else None}')
                if len(mism) > 50: return res
        # region contents: only SELECT rows / headers / blanks
        for i in reg:
            a = rb[i][0][0] if rb[i] else None
            assert a in (None, TID, 'table')
        # merged ranges
        if reg:
            s0, e0 = min(reg) + 1, max(reg) + 1; n = e0 - s0 + 1
        mb = []
        for m in sb.merged_cells.ranges:
            c1, r1, c2, r2 = range_boundaries(str(m))
            if reg and s0 <= r1 <= e0:
                continue
            if reg and r1 > e0:
                r1 -= n; r2 -= n
            mb.append((c1, r1, c2, r2))
        ma = [range_boundaries(str(m)) for m in sa.merged_cells.ranges]
        if sorted(mb) != sorted(ma):
            mism.append(f'{label} {sn}: merged ranges differ ({len(mb)} expected, {len(ma)} after)')
        nmerge_sel = len(sb.merged_cells.ranges) - len(mb)
        # column widths
        wdiff = []
        for L in set(sb.column_dimensions) | set(sa.column_dimensions):
            w1 = sb.column_dimensions[L].width if L in sb.column_dimensions else None
            w2 = sa.column_dimensions[L].width if L in sa.column_dimensions else None
            if w1 != w2: wdiff.append((L, w1, w2))
        res[sn] = (len(reg), ntab, nmerge_sel, len(sa.merged_cells.ranges), len(wdiff), sorted(wdiff)[:3])
    return res, wa

out = {}
for race in ('lingfield', 'ascot'):
    for mode in ('def', 'full'):
        lab = f'{race}_{mode}'
        r, wa = comp(f'{B}/{lab}_jupiter_tables.xlsx', f'{A}/{lab}_jupiter_tables.xlsx', lab, True)
        out[lab] = r
        print(lab, 'size MB', round(os.path.getsize(f'{A}/{lab}_jupiter_tables.xlsx') / 1e6, 2))
        for sn, v in r.items():
            if v[0] or v[2] or v[4]:
                print('  ', sn, 'removed rows', v[0], 'tables', v[1], 'select merges', v[2], 'merges after', v[3], 'width diffs', v[4], v[5])
        # 4: sky star timelines
        ws = wa['SKY JUPITER']; rr = list(ws.iter_rows(values_only=True))
        for blk in ('STAR TIMELINE', 'LONG STAR TIMELINE'):
            st = [i for i, x in enumerate(rr) if x[0] == blk]
            heads = set(); bad = 0; n = 0
            for i in st:
                for v in rr[i]:
                    if v in ('> 95 since', '100', '< 95 from'): heads.add(v)
                rel = rr[i][4]
                if isinstance(rel, str) and rel not in ('relationship',) and not rel.startswith('S Jupiter') and rel not in ('OFF', 'OFF + 5 MIN', 'SEPARATING', 'APPLYING'):
                    bad += 1; print('    non-SJ', blk, i + 1, rel)
                n += 1
            print(f'   {blk}: rows {n}, headers {sorted(heads)}, non-S Jupiter rows {bad}')
            if heads != {'> 95 since', '100', '< 95 from'} or bad:
                mism.append(f'{lab} {blk}: headers {heads} bad {bad}')
    # 6 pairs
    for mode in ('def', 'full'):
        lab = f'{race}_{mode}'
        r, _ = comp(f'{B}/{lab}_pairs.xlsx', f'{A}/{lab}_pairs.xlsx', lab + '_pairs', False)
        print(lab, 'pairs sheets', len(r), 'size MB', round(os.path.getsize(f'{A}/{lab}_pairs.xlsx') / 1e6, 2))
    # 7 default vs full after
    wd = openpyxl.load_workbook(f'{A}/{race}_def_jupiter_tables.xlsx'); wf = openpyxl.load_workbook(f'{A}/{race}_full_jupiter_tables.xlsx')
    for sn in wd.sheetnames:
        if sn not in wf.sheetnames:
            mism.append(f'{race} def sheet {sn} missing in full'); continue
        if rows(wd[sn]) != rows(wf[sn]):
            mism.append(f'{race} def vs full {sn} differ')
    print(race, 'default sheets', len(wd.sheetnames), 'checked = --full')
print('MISMATCHES', len(mism))
for m in mism[:60]: print(m)
