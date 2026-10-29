"""Check 4: plant errors in copies of an after workbook, one per copy, and run the checkers on each.
usage: plant.py <new.xlsx> <old.xlsx> <outdir>"""
import os, sys, subprocess, copy
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

src, old, out = sys.argv[1:4]
HERE = os.path.dirname(os.path.abspath(__file__))
PY = sys.executable


def find(ws, pred):
    for row in ws.iter_rows():
        for c in row:
            if pred(c):
                return c


def p1(wb):  # half text back into a SKY block target cell
    ws = wb['SKY JUPITER']
    c = find(ws, lambda c: c.column == 5 and isinstance(c.value, str) and c.value.endswith('(whole) — outside ±0.02'))
    c.value = c.value.replace('(whole) — outside ±0.02', '(half) — outside ±0.01').replace(c.value.split(' ')[0], str(int(c.value.split(' ')[0]) - 0.5)) if False else '14.5 (half)'
    return f'SKY JUPITER {c.coordinate} -> 14.5 (half)'


def p2(wb):  # change a T1 score
    ws = wb['P01']
    c = find(ws, lambda c: c.column == 8 and ws.cell(c.row, 1).value == 'T1' and isinstance(c.value, (int, float)))
    c.value = round(c.value + 1.0, 2)
    return f'P01 {c.coordinate} score +1'


def p3(wb):  # change a T24 fill (a light-yellow hit cell to green)
    ws = wb['P03']
    c = find(ws, lambda c: ws.cell(c.row, 1).value == 'T24' and c.fill.fgColor.rgb == '00FFF8DC' and c.column > 2)
    c.fill = PatternFill('solid', fgColor='C6EFCE')
    return f'P03 {c.coordinate} fill -> green'


def p4(wb):  # wrong nearest target in a T24 comment
    ws = wb['P05']
    c = find(ws, lambda c: ws.cell(c.row, 1).value == 'T24' and c.comment and c.comment.text.startswith('nearest: ') and '· whole — outside' in c.comment.text)
    t = c.comment.text
    lab = t.split('nearest: ')[1].split(' ·')[0]
    c.comment = Comment(t.replace(f'nearest: {lab} · whole', f'nearest: {float(lab) + 0.5:g} · half', 1).replace('±0.02', '±0.01', 1), 'x')
    return f'P05 {c.coordinate} comment nearest -> half'


def p5(wb):  # T5 offset changed
    ws = wb['P07']
    c = find(ws, lambda c: c.column == 7 and ws.cell(c.row, 1).value == 'T5' and isinstance(c.value, str) and c.value[:1] in '+-')
    c.value = f'{float(c.value) + 0.001:+.5f}'
    return f'P07 {c.coordinate} offset +0.001'


def p6(wb):  # a half repeat added to a T24 repeats cell
    ws = wb['P09']
    c = find(ws, lambda c: ws.cell(c.row, 1).value == 'T24' and c.column == 55 and ws.cell(c.row, 2).value not in ('stars', None)
             and '→' not in str(ws.cell(c.row, 2).value) and not str(ws.cell(c.row, 2).value).startswith('T24'))
    c.value = '27.5 half ×2 (Sun RA (91.23), Mars Dec (50.00))'
    return f'P09 {c.coordinate} repeats -> half repeat'


def p7(wb):  # pair-block leg: a 'no family' leg given a half hit text
    ws = wb['SKY JUPITER']
    c = find(ws, lambda c: c.column == 6 and c.value == 'no family' and ws.cell(c.row, 1).value == 'SKY STARS')
    c.value = '→ 14.5 (half, 85.64) · separating, exact 14:05 (24.0 min ago)'
    return f'SKY JUPITER {c.coordinate} leg -> half'


def p8(wb):  # T24 value nudged
    ws = wb['P11']
    c = find(ws, lambda c: ws.cell(c.row, 1).value == 'T24' and isinstance(c.value, float) and c.column == 10)
    c.value = round(c.value + 0.00002, 6)
    return f'P11 {c.coordinate} value +0.00002'


def p9(wb):  # T2 SKY cell: nearest target swapped to another family
    ws = wb['P01']
    c = find(ws, lambda c: c.column == 21 and ws.cell(c.row, 1).value == 'T2' and isinstance(c.value, str) and '(ninth) — outside' in c.value)
    c.value = '24.14 (silver) — outside ±0.121'
    return f'P01 {c.coordinate} SKY nearest -> silver'


def p10(wb):  # dotted outline removed from a repeated cell
    ws = wb['P01']
    c = find(ws, lambda c: ws.cell(c.row, 1).value == 'T24' and c.border.left.style == 'dotted')
    from openpyxl.styles import Border
    c.border = Border()
    return f'P01 {c.coordinate} dotted outline removed'


def p0(wb):  # control: saved unchanged
    return 'no change (control)'


PLANTS = [p0, p1, p2, p3, p4, p5, p6, p7, p8, p9, p10]
os.makedirs(out, exist_ok=True)
caught = 0
for p in PLANTS:
    wb = openpyxl.load_workbook(src)
    what = p(wb)
    f = os.path.join(out, f'plant_{p.__name__}.xlsx')
    wb.save(f)
    res = []
    for script, args in (('check_cells.py', [f]), ('check_repeats.py', [f]), ('check_pairs.py', [old, f]), ('scan_half_text.py', [f])):
        o = subprocess.run([PY, os.path.join(HERE, script)] + args, capture_output=True, text=True).stdout
        flag = ('errors 0' not in o.splitlines()[0]) if script in ('check_cells.py', 'check_repeats.py') else \
               ('leg errors 0' not in o.splitlines()[0]) if script == 'check_pairs.py' else ('half-family texts: 0' not in o)
        if flag:
            res.append(script)
    caught += bool(res) and p is not p0
    print(f'{p.__name__}: {what}: caught by {res or "NONE"}')
print(f'caught {caught} of {len(PLANTS) - 1} planted errors')
