"""Check 2: recompute every T1-T8 cell (natal + SKY cells), SKY block cell and T24 SKY / HORSE / JOCKEY cell from POSITIONS.
usage: check_cells.py <workbook.xlsx> [--half]   (--half = evaluate on the before rules, with the half family)"""
import re, sys, collections
import openpyxl
sys.path.insert(0, __import__('os').path.dirname(__file__))
import ifam

FAMS = r'(whole|half|golden|phi power|ninth|sqrt2|silver)'
RE_HIT = re.compile(r'^(.*) \(' + FAMS + r'\)$')
RE_OUT = re.compile(r'^(.*) \(' + FAMS + r'\) — outside ±(\S+)$')
RE_DIG = re.compile(r'^(.*) ' + FAMS + r' — not a 4-digit match \((\S+)\)$')
RE_CHIT = re.compile(r'^(.*) · ' + FAMS + r' · (\S+) · (\S+)$')
RE_CCONJ = re.compile(r'^0 · conjunction \((RA|Dec) < (\S+)\) · (\S+) · (\S+)$')
RE_CNOUT = re.compile(r'^nearest: (.*) · ' + FAMS + r' — outside ±(\S+) · (\S+)$')
RE_CNDIG = re.compile(r'^nearest: (.*) ' + FAMS + r' — not a 4-digit match \((\S+)\)$')
RED = '00FF8080'


def wrap(a, b):
    d = abs(a - b) % 360
    return 360 - d if d > 180 else d


class Pos:
    def __init__(self, wb):
        self.p = {}
        for r in wb['POSITIONS'].iter_rows(min_row=5, values_only=True):
            if r[0] and r[2] in ('S', 'H', 'J') and isinstance(r[5], (int, float)):
                key = ('SKY' if r[2] == 'S' else r[0], r[2], r[3])
                self.p[key] = (r[5], r[6], r[7] or 0.0, r[8] or 0.0)

    def get(self, runner, chart, point, t=0):
        ra, dec, rr, rd = self.p[('SKY' if chart == 'S' else runner, chart, point)]
        if chart == 'S':
            return ra + rr * t, dec + rd * t
        return ra, dec

    def dist(self, runner, a, b, plane, t=0):
        if b[1] == 'equator':
            ra1, d1 = self.get(runner, *a, t)
            return abs(d1)
        if a[1] == 'equator':
            ra1, d1 = self.get(runner, *b, t)
            return abs(d1)
        ra1, d1 = self.get(runner, *a, t)
        ra2, d2 = self.get(runner, *b, t)
        return wrap(ra1, ra2) if plane == 'RA' else abs(d1 - d2)


def pt(s):
    c, n = s.split(' ', 1)
    return c, n


class Checker:
    def __init__(self, path, half=False):
        self.wb = openpyxl.load_workbook(path)
        self.pos = Pos(self.wb)
        self.half = half
        self.n = collections.Counter()
        self.bad = []
        self.fams = collections.Counter()

    def err(self, where, msg):
        self.bad.append(f'{where}: {msg}')

    # ---- one T-table / SKY block value group --------------------------------------------------------------
    def tcell(self, where, plane, v, v5, cells, timing_present=None):
        """cells: dict value, plus5, target, offset, score, arrow, also, fill"""
        self.n['tcells'] += 1
        e = ifam.evaluate(v, plane, self.half)
        self.fams[('hit:' + e['fam']) if e['hit'] else ('near:' + e['fam'] + ':' + e['kind'])] += 1
        if cells['value'] is None or abs(cells['value'] - round(v, 6)) > 1.5e-6:
            self.err(where, f'value {cells["value"]} vs {v:.6f}')
        if v5 is not None and (cells['plus5'] is None or abs(cells['plus5'] - round(v5, 6)) > 1.5e-6):
            self.err(where, f'+5 {cells["plus5"]} vs {v5:.6f}')
        txt = cells['target'] or ''
        if 'half' in txt and not self.half:
            self.err(where, f'half text {txt!r}')
        exp_off = f'{e["off"]:+.5f}'
        if e['hit']:
            if e['fam'] == 'conj':
                ok = txt == 'conj (conj)'
                exp_off = f'{v:+.5f}'
            else:
                m = RE_HIT.match(txt)
                ok = bool(m) and m.group(2) == e['fam'] and ifam.label_eq(e['fam'], ifam.parse_label(e['fam'], m.group(1)), e['lab'])
            if not ok:
                self.err(where, f'target {txt!r} vs hit {e["fam"]} {e["lab"]} ({e["score"]:.2f})')
            sc = cells['score']
            if sc is None or abs(float(sc) - round(e['score'], 2)) > 0.0101:
                self.err(where, f'score {sc} vs {e["score"]:.4f}')
            if cells['fill'] != ifam.fill_for(e['score']):
                self.err(where, f'fill {cells["fill"]} vs {ifam.fill_for(e["score"])}')
            if v5 is not None:
                exp_arrow = '-> toward' if abs(v5 - e['t']) < abs(v - e['t']) else '<- away'
                if cells.get('arrow') != exp_arrow:
                    self.err(where, f'arrow {cells.get("arrow")!r} vs {exp_arrow}')
            # also: the other family hits
            others = sorted((h[1], str(h[2])) for h in e['allhits'] if h is not max(e['allhits'], key=lambda x: x[0]) )
            got = []
            for part in (cells.get('also') or '').split(';'):
                part = part.strip()
                mm = re.match(r'^(.*) \((whole|half|golden|phi power|ninth|sqrt2|silver|conj)\) (\S+) \[(\S+)\]$', part)
                if mm:
                    got.append((mm.group(2), str(ifam.parse_label(mm.group(2), mm.group(1)))))
                elif part:
                    got.append(('?', part))
            if [g[0] for g in sorted(got)] != [o[0] for o in others]:
                self.err(where, f'also {cells.get("also")!r} vs {others}')
        else:
            if e['kind'] == 'digit':
                m = RE_DIG.match(txt)
                ok = bool(m) and m.group(2) == e['fam'] and ifam.label_eq(e['fam'], ifam.parse_label(e['fam'], m.group(1)), e['lab']) \
                    and abs(float(m.group(3)) - float(exp_off)) < 1.5e-5
            else:
                m = RE_OUT.match(txt)
                ok = bool(m) and m.group(2) == e['fam'] and ifam.label_eq(e['fam'], ifam.parse_label(e['fam'], m.group(1)), e['lab'])
            if not ok:
                self.err(where, f'nearest {txt!r} vs {e["fam"]} {e["lab"]} {e["kind"]} {exp_off}')
            if cells['score'] not in (None, ''):
                self.err(where, f'score {cells["score"]} on a no-hit')
            if cells['fill'] != '00000000':
                self.err(where, f'fill {cells["fill"]} on a no-hit')
            if cells.get('arrow'):
                self.err(where, f'arrow on a no-hit')
        off = cells['offset']
        if off is None or abs(float(off) - float(exp_off)) > 1.5e-5:
            self.err(where, f'offset {off} vs {exp_off}')
        if timing_present is not None:
            want = e['hit'] and v5 is not None and v5 != v
            if bool(timing_present) != want and not (e['hit'] and e['fam'] == 'conj'):
                self.err(where, f'timing present {bool(timing_present)} vs hit {e["hit"]}')

    # ---- T24 comment cell --------------------------------------------------------------------------------------
    def t24cell(self, where, plane, v, c):
        self.n['t24cells'] += 1
        e = ifam.evaluate(v, plane, self.half)
        self.fams[('hit:' + e['fam']) if e['hit'] else ('near:' + e['fam'] + ':' + e['kind'])] += 1
        if not isinstance(c.value, (int, float)) or abs(c.value - round(v, 6)) > 1.5e-6:
            self.err(where, f'value {c.value} vs {v:.6f}')
        line = c.comment.text.split('\n')[0] if c.comment else ''
        if 'half' in (c.comment.text if c.comment else '') and not self.half:
            self.err(where, f'half comment {line!r}')
        exp_off = f'{e["off"]:+.5f}'
        fill = c.fill.fgColor.rgb
        conjrange = v < (1.5 if plane == 'RA' else 0.3)
        if e['hit']:
            if e['fam'] == 'conj':
                m = RE_CCONJ.match(line)
                ok = bool(m) and m.group(1) == plane and abs(float(m.group(3)) - v) < 1.5e-5 and abs(float(m.group(4)) - round(e['score'], 2)) <= 0.0101
            else:
                m = RE_CHIT.match(line)
                ok = bool(m) and m.group(2) == e['fam'] and ifam.label_eq(e['fam'], ifam.parse_label(e['fam'], m.group(1)), e['lab']) \
                    and abs(float(m.group(3)) - float(exp_off)) < 1.5e-5 and abs(float(m.group(4)) - round(e['score'], 2)) <= 0.0101
            if not ok:
                self.err(where, f'comment {line!r} vs hit {e["fam"]} {e["lab"]} {exp_off} {e["score"]:.2f}')
            want = RED if conjrange else ifam.fill_for(e['score'])
        else:
            if e['kind'] == 'digit':
                m = RE_CNDIG.match(line)
                ok = bool(m) and m.group(2) == e['fam'] and ifam.label_eq(e['fam'], ifam.parse_label(e['fam'], m.group(1)), e['lab']) \
                    and abs(float(m.group(3)) - float(exp_off)) < 1.5e-5
            else:
                m = RE_CNOUT.match(line)
                ok = bool(m) and m.group(2) == e['fam'] and ifam.label_eq(e['fam'], ifam.parse_label(e['fam'], m.group(1)), e['lab']) \
                    and abs(float(m.group(4)) - float(exp_off)) < 1.5e-5
            if not ok:
                self.err(where, f'comment {line!r} vs nearest {e["fam"]} {e["lab"]} {e["kind"]} {exp_off}')
            want = RED if conjrange else '00000000'
        if fill != want:
            self.err(where, f'fill {fill} vs {want}')

    # ---- sheets -------------------------------------------------------------------------------------------------
    def run(self):
        for ws in self.wb:
            if ws.title == 'POSITIONS':
                continue
            runner = ws.title if ws.title.startswith('P') else 'SKY'
            rows = list(ws.iter_rows())
            hdr = None
            t24 = None
            for i, row in enumerate(rows):
                a = row[0].value
                if a == 'table':
                    hdr = {c.column: c.value for c in row}
                    if t24 is not None:
                        t24 = None
                    if row[1].value == 'star':
                        names = {c.column: c.value for c in rows[i + 1]}
                        planes = {}
                        cur = None
                        for c in row:
                            if c.value in ('RA', 'Dec'):
                                cur = c.value
                            elif c.value is not None and c.column > 2:
                                cur = None
                            planes[c.column] = cur
                        t24 = {col: (planes.get(col), nm) for col, nm in names.items()
                               if nm and planes.get(col) and not str(nm).startswith('star =')}
                    continue
                if a in ('T1', 'T2', 'T3', 'T4', 'T5', 'T6', 'T7', 'T8', 'SKY') and hdr and isinstance(row[1].value, str) and row[1].value.count('→') == 1 and row[1].value.split(' → ')[0].count(' ') == 1 and 'bodies' not in row[1].value and 'sky points' not in row[1].value:
                    self.trow(ws.title, runner, a, row, hdr)
                elif a == 'T24' and t24 and isinstance(row[1].value, str) and row[1].value != 'stars':
                    self.t24row(ws.title, runner, row, t24)
        return self

    def trow(self, sheet, runner, table, row, hdr):
        L, R = [s.strip() for s in row[1].value.split('→')]
        a, b = pt(L), pt(R)
        cols = {}
        for c in row:
            cols[c.column] = c
        # value columns = headers 'RA','Dec','SKY RA','SKY Dec'
        for col, h in hdr.items():
            if h not in ('RA', 'Dec', 'SKY RA', 'SKY Dec'):
                continue
            plane = h.split()[-1]
            sky = h.startswith('SKY') or table == 'SKY'
            pa, pb = (('S', a[1]), ('S', b[1])) if (h.startswith('SKY')) else (a, b)
            vc = cols.get(col)
            if vc is None or vc.value is None:
                # RA of an equator row is blank
                if b[1] == 'equator' and plane == 'RA':
                    continue
                self.err(f'{sheet} {row[0].row} {h}', 'missing value')
                continue
            moving = 'S' in (pa[0], pb[0])
            v = self.pos.dist(runner, pa, pb, plane)
            v5 = self.pos.dist(runner, pa, pb, plane, 5) if moving else None
            g = lambda k: cols[col + k].value if col + k in cols else None
            names = [hdr.get(col + k) for k in range(1, 9)]
            idx = {n: k for k, n in zip(range(1, 9), names) if n}
            cells = dict(value=vc.value, plus5=g(idx['+5 min']), target=g(idx['nearest target']), offset=g(idx['offset']),
                         score=g(idx['score']), arrow=g(idx['off->+5']), also=g(idx['also (score)']), fill=vc.fill.fgColor.rgb)
            tp = g(idx['timing']) if 'timing' in idx else None
            self.tcell(f'{sheet} r{row[0].row} {row[1].value} {h}', plane, v, v5, cells, tp if 'timing' in idx else None)

    def t24row(self, sheet, runner, row, t24):
        lab = row[1].value
        cols = {c.column: c for c in row}
        for col, (plane, body) in t24.items():
            c = cols.get(col)
            if '→' in lab:
                L, R = [s.strip() for s in lab.split('→')]
                a, b = pt(L), pt(R)
                if body == 'S Jupiter':
                    if c is not None and c.value is not None:
                        self.err(f'{sheet} r{row[0].row} {lab} {plane} S Jupiter', 'not blank')
                    continue
                if body == 'Jupiter':
                    if a[0] == b[0]:
                        if c is not None and c.value is not None:
                            self.err(f'{sheet} r{row[0].row} {lab} {plane} Jupiter', 'not blank')
                        continue
                    natal = b[0] if b[0] != 'S' else a[0]
                    pa, pb = ('S', 'Jupiter'), (natal, 'Jupiter')
                else:
                    pa = (a[0], 'Jupiter' if a[1] == 'Jupiter' else body)
                    pb = (b[0], 'Jupiter' if b[1] == 'Jupiter' else body)
            else:
                ch, star = pt(lab)
                if body == 'S Jupiter':
                    pa, pb = ('S', 'Jupiter'), ('S', star)
                else:
                    pa, pb = (ch, star), (ch, body)
            try:
                v = self.pos.dist(runner, pa, pb, plane)
            except KeyError:
                if c is not None and c.value is not None:
                    self.err(f'{sheet} r{row[0].row} {lab} {plane} {body}', 'value but no position')
                continue
            if c is None or c.value is None:
                self.err(f'{sheet} r{row[0].row} {lab} {plane} {body}', 'blank')
                continue
            self.t24cell(f'{sheet} r{row[0].row} {lab} {plane} {body}', plane, v, c)


if __name__ == '__main__':
    ch = Checker(sys.argv[1], '--half' in sys.argv).run()
    print(sys.argv[1], dict(ch.n), 'errors', len(ch.bad))
    print(' families:', dict(sorted(ch.fams.items())))
    for b in ch.bad[:40]:
        print('  ', b)
