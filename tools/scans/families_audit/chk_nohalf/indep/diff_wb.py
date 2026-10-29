"""Check 3: before (with half) vs after (no half), every cell of every sheet. Rows aligned per sheet by difflib on
row content; every changed cell / dropped / added row is classified, anything not explained by half is listed.
usage: diff_wb.py <old.xlsx> <new.xlsx>"""
import re, sys, collections, difflib
import openpyxl
sys.path.insert(0, __import__('os').path.dirname(__file__))
import ifam
from check_cells import Pos, pt

HALFTXT = re.compile(r'\(half\)|· half|half —| half \(|\bhalf,|half ×', re.I)


def rowsig(row):
    return tuple((c.value, c.comment.text if c.comment else None, c.fill.fgColor.rgb, c.border.left.style) for c in row)


def key(row):
    # identity of a row: table id + label cells (A-D) (values that do not depend on families)
    return tuple(str(c.value) for c in row[:4])


class Diff:
    def __init__(self, fo, fn):
        self.wo = openpyxl.load_workbook(fo)
        self.wn = openpyxl.load_workbook(fn)
        self.pos = Pos(self.wn)
        self.cnt = collections.Counter()
        self.unexplained = []
        self.notes = collections.defaultdict(list)

    def un(self, msg):
        self.unexplained.append(msg)

    def run(self):
        assert self.wo.sheetnames == self.wn.sheetnames, (self.wo.sheetnames, self.wn.sheetnames)
        for name in self.wo.sheetnames:
            so, sn = self.wo[name], self.wn[name]
            ro, rn = list(so.iter_rows(max_col=60)), list(sn.iter_rows(max_col=60))
            ko, kn = [key(r) for r in ro], [key(r) for r in rn]
            sm = difflib.SequenceMatcher(None, ko, kn, autojunk=False)
            ctx = {'hdr': None, 'block': None}
            for op, i1, i2, j1, j2 in sm.get_opcodes():
                if op == 'equal' or (op == 'replace' and i2 - i1 == j2 - j1):
                    for a, b in zip(ro[i1:i2], rn[j1:j2]):
                        self.track(ctx, b)
                        if rowsig(a) != rowsig(b):
                            self.rowdiff(name, a, b, ctx)
                else:
                    for a in ro[i1:i2]:
                        self.dropped(name, a)
                    for b in rn[j1:j2]:
                        self.track(ctx, b)
                        self.added(name, b)
        return self

    def track(self, ctx, row):
        a = row[0].value
        if a == 'table':
            ctx['hdr'] = {c.column: c.value for c in row}
            ctx['hdrrow'] = row
            ctx['names'] = None
        elif ctx.get('hdr') and ctx['hdr'].get(2) == 'star' and ctx.get('names') is None and a is None:
            ctx['names'] = {c.column: c.value for c in row}

    # ------------------------------------------------------------------------------------------------------
    def rowdiff(self, sheet, a, b, ctx):
        tab = b[0].value
        hdr = ctx.get('hdr') or {}
        lab = b[1].value
        for ca, cb in zip(a, b):
            if (ca.value, ca.comment.text if ca.comment else None, ca.fill.fgColor.rgb, ca.border.left.style) == \
               (cb.value, cb.comment.text if cb.comment else None, cb.fill.fgColor.rgb, cb.border.left.style):
                continue
            col = cb.column
            h = hdr.get(col)
            where = f'{sheet} r{cb.row} {cb.coordinate} [{tab}] {lab}'
            if isinstance(tab, str) and tab in ('T1', 'T2', 'T3', 'T4', 'T5', 'T6', 'T7', 'T8', 'SKY') and isinstance(lab, str) and '→' in lab:
                self.tdiff(where, a, b, col, hdr, ca, cb)
            elif tab == 'T24' and isinstance(lab, str) and ctx.get('names') and not lab.startswith('T24 '):
                self.t24diff(where, sheet, a, b, col, hdr, ctx, ca, cb)
            elif tab == 'T24':
                # title rows: repeat counts
                o, n = str(ca.value), str(cb.value)
                if 'REPEATS' in o and re.sub(r'\d+ of \d+ rows with a repeat, \d+ repeated targets', '', o) == \
                        re.sub(r'\d+ of \d+ rows with a repeat, \d+ repeated targets', '', n):
                    self.cnt['T24 title repeat count'] += 1
                    self.notes['T24 title'].append(f'{sheet}: {re.search(r"\d+ of \d+ rows with a repeat, \d+ repeated targets", o).group()} -> {re.search(r"\d+ of \d+ rows with a repeat, \d+ repeated targets", n).group()}')
                else:
                    self.un(f'{where}: {ca.value!r} -> {cb.value!r}')
            elif isinstance(tab, str) and tab.startswith(('SKY PAIRS', 'SKY STARS', 'SKY ROYAL', 'SKY EQ')):
                self.pairdiff(where, ca, cb, hdr)
            elif isinstance(tab, str) and 'TIMELINE' in tab:
                self.tldiff(where, a, b, ca, cb, hdr)
            else:
                o, n = str(ca.value), str(cb.value)
                if re.sub(r'\d+', 'N', o) == re.sub(r'\d+', 'N', n) and isinstance(ca.value, str) and len(o) > 30:
                    self.cnt['title count'] += 1
                    self.notes['titles'].append(f'{sheet} {cb.coordinate}: ' + ' | '.join(
                        f'{x}->{y}' for x, y in zip(re.findall(r'\d+', o), re.findall(r'\d+', n)) if x != y) + f'  [{o[:60]}]')
                else:
                    self.un(f'{where}: {ca.value!r} -> {cb.value!r}')

    # value group in T tables / SKY block: changed cells must belong to a value whose before eval involved half
    def tdiff(self, where, a, b, col, hdr, ca, cb):
        # find the plane block: nearest value column at or left of col
        vcol = max(c for c, h in hdr.items() if h in ('RA', 'Dec', 'SKY RA', 'SKY Dec') and c <= col)
        h = hdr[vcol]
        plane = h.split()[-1]
        v = b[vcol - 1].value
        tgt_o = a[vcol + 1].value if hdr.get(vcol + 2) == 'nearest target' else None
        tgt_o = next(a[c - 1].value for c in range(vcol + 1, vcol + 9) if hdr.get(c) == 'nearest target')
        tgt_n = next(b[c - 1].value for c in range(vcol + 1, vcol + 9) if hdr.get(c) == 'nearest target')
        if a[vcol - 1].value != v:
            self.un(f'{where}: value changed {a[vcol-1].value} -> {v}')
            return
        eo = ifam.evaluate(v, plane, True)
        en = ifam.evaluate(v, plane, False)
        kind = self.classify(eo, en, str(tgt_o))
        if kind is None:
            self.un(f'{where} ({h}): {ca.value!r} -> {cb.value!r}; before {tgt_o!r} after {tgt_n!r}')
            return
        if col == vcol or hdr.get(col) == 'nearest target' or (kind, vcol) not in getattr(self, '_seen', set()):
            pass
        k = (where.split(' [')[0].rsplit(' ', 1)[0], vcol, b[0].row, a[0].parent.title if hasattr(a[0], 'parent') else '')
        if k not in self.__dict__.setdefault('_vg', set()):
            self._vg.add(k)
            self.cnt[f'T-table value group: {kind}'] += 1
            if len(self.notes[kind]) < 6:
                self.notes[kind].append(f'{where} {h} {v}: {tgt_o!r} -> {tgt_n!r}')
        self.cnt['T-table cells changed'] += 1

    def classify(self, eo, en, txt_o):
        if eo['fam'] == 'half' and eo['hit']:
            if en['hit']:
                return 'half hit -> conj' if en['fam'] == 'conj' else f'half hit -> {en["fam"]}'
            return 'half hit -> no score'
        if not eo['hit'] and eo['fam'] == 'half':
            return f'nearest half -> nearest {en["fam"]}' + (' (4-digit)' if en.get('kind') == 'digit' else '')
        if eo['hit'] and any(hh[1] == 'half' for hh in eo['allhits']):
            return 'half in also only'
        return None

    def t24diff(self, where, sheet, a, b, col, hdr, ctx, ca, cb):
        names = ctx['names']
        body = names.get(col)
        hcol = hdr.get(col)
        if hcol == 'repeats' or (body is None and hdr.get(col) == 'repeats'):
            self.repdiff(where, ca, cb)
            return
        if body is None or str(body).startswith('star ='):
            self.un(f'{where} col {cb.coordinate} (midpoints / other): {ca.value!r} -> {cb.value!r}')
            return
        # plane
        plane = None
        for c in sorted(hdr):
            if c <= col and hdr[c] in ('RA', 'Dec'):
                plane = hdr[c]
        v = cb.value
        if ca.value != v:
            self.un(f'{where} {body}: value {ca.value} -> {v}')
            return
        co = ca.comment.text if ca.comment else ''
        cn = cb.comment.text if cb.comment else ''
        lo, ln = co.split('\n'), cn.split('\n')
        eo = ifam.evaluate(v, plane, True)
        en = ifam.evaluate(v, plane, False)
        kind = None
        if lo[0] != ln[0] or ca.fill.fgColor.rgb != cb.fill.fgColor.rgb:
            kind = self.classify(eo, en, lo[0])
            if kind is None:
                self.un(f'{where} {body} {plane} {v}: {lo[0]!r} -> {ln[0]!r} fill {ca.fill.fgColor.rgb}->{cb.fill.fgColor.rgb}')
                return
            self.cnt[f'T24 cell: {kind}'] += 1
            if len(self.notes['T24 ' + kind]) < 4:
                self.notes['T24 ' + kind].append(f'{where} {body} {plane} {v}: {lo[0]!r} -> {ln[0]!r}')
        # remaining comment lines / border: repeat lines only
        ro = [x for x in lo[1:] if not x.startswith('repeat:')]
        rn = [x for x in ln[1:] if not x.startswith('repeat:')]
        if ro != rn:
            self.un(f'{where} {body}: non-repeat comment lines changed {ro} -> {rn}')
        rpo = [x for x in lo[1:] if x.startswith('repeat:')]
        rpn = [x for x in ln[1:] if x.startswith('repeat:')]
        if rpo != rpn or ca.border.left.style != cb.border.left.style:
            if all('whole' not in x or True for x in rpo) and any(' half ×' in x for x in rpo) or kind:
                self.cnt['T24 repeat line / outline change on a half cell'] += 1
            else:
                self.un(f'{where} {body}: repeat lines {rpo} -> {rpn}, border {ca.border.left.style}->{cb.border.left.style}')

    def repdiff(self, where, ca, cb):
        o = [x.strip() for x in str(ca.value or '').split(';') if x.strip()]
        n = [x.strip() for x in str(cb.value or '').split(';') if x.strip()]
        lost = [x for x in o if x not in n]
        gained = [x for x in n if x not in o]
        if all(' half ×' in x for x in lost) and not gained:
            self.cnt['T24 repeats: half target lost'] += len(lost)
            self.notes['repeats lost'].extend(f'{where}: {x[:90]}' for x in lost)
        else:
            # a non-half repeat may change because one of its cells moved off half to the same target
            self.un(f'{where} repeats: lost {lost} gained {gained}')

    def pairdiff(self, where, ca, cb, hdr):
        o, n = str(ca.value), str(cb.value)
        if HALFTXT.search(o) and not HALFTXT.search(n):
            self.cnt['pair-block leg cell: half -> ' + ('no family' if n == 'no family' else re.sub(r'.*\((\w[\w ]*?),.*', r'\1', n))] += 1
        elif re.sub(r'\d+', 'N', o) == re.sub(r'\d+', 'N', n) and len(o) > 30:
            self.cnt['pair-block title count'] += 1
            self.notes['titles'].append(f'{where[:40]}: ' + ' | '.join(f'{x}->{y}' for x, y in zip(re.findall(r'\d+', o), re.findall(r'\d+', n)) if x != y) + f' [{o[:50]}]')
        else:
            self.un(f'{where}: {o[:120]!r} -> {n[:120]!r}')

    def tldiff(self, where, a, b, ca, cb, hdr):
        o, n = str(ca.value), str(cb.value)
        h = hdr.get(cb.column)
        if h == 'note' and ('also' in o or 'also' in n):
            self.cnt['timeline note ("also") changed'] += 1
            self.notes['tl note'].append(f'{where}: {o!r} -> {n!r}')
        elif re.sub(r'\d+', 'N', o) == re.sub(r'\d+', 'N', n) and (len(o) > 30 or 'events' in o):
            self.cnt['timeline title / section count'] += 1
            self.notes['titles'].append(f'{where[:40]}: ' + ' | '.join(f'{x}->{y}' for x, y in zip(re.findall(r'\d+', o), re.findall(r'\d+', n)) if x != y) + f' [{o[:50]}]')
        else:
            self.un(f'{where} [{h}]: {o[:100]!r} -> {n[:100]!r}')

    def dropped(self, sheet, row):
        tab = str(row[0].value)
        txt = ' | '.join(str(c.value) for c in row if c.value is not None)
        if row[0].value is None and all(c.value is None for c in row):
            self.cnt[f'{sheet} blank row dropped'] += 1
            return
        if 'TIMELINE' in tab:
            if HALFTXT.search(txt):
                self.cnt[f'{tab}: half event dropped'] += 1
                self.notes['tl dropped'].append(f'{sheet} {tab}: ' + ' | '.join(str(c.value) for c in row[1:9]))
            else:
                self.cnt[f'{tab}: non-half event dropped'] += 1
                self.notes['tl dropped (non-half)'].append(f'{sheet} {tab}: ' + ' | '.join(str(c.value) for c in row[1:13]))
            return
        if tab.startswith(('SKY PAIRS', 'SKY STARS', 'SKY ROYAL')):
            if HALFTXT.search(txt):
                self.cnt[f'{tab} row dropped (half leg)'] += 1
                self.notes['pair rows dropped'].append(f'{sheet}: ' + ' | '.join(str(c.value)[:40] for c in row[:9]))
            else:
                self.un(f'{sheet} dropped row without half: {txt[:200]}')
            return
        self.un(f'{sheet} dropped row r{row[0].row}: {txt[:200]}')

    def added(self, sheet, row):
        tab = str(row[0].value)
        txt = ' | '.join(str(c.value) for c in row if c.value is not None)
        if row[0].value is None and all(c.value is None for c in row):
            self.cnt[f'{sheet} blank row added'] += 1
            return
        if 'TIMELINE' in tab:
            self.cnt[f'{tab}: event added'] += 1
            self.notes['tl added'].append(f'{sheet} {tab}: ' + ' | '.join(str(c.value) for c in row[1:13]))
            return
        self.un(f'{sheet} added row r{row[0].row}: {txt[:200]}')


if __name__ == '__main__':
    d = Diff(sys.argv[1], sys.argv[2]).run()
    print('==', sys.argv[2])
    for k, v in sorted(d.cnt.items()):
        print(f'  {v:6d}  {k}')
    for k, v in d.notes.items():
        print(' --', k, len(v))
        for x in v[:int(sys.argv[3]) if len(sys.argv) > 3 else 8]:
            print('     ', x)
    print('  UNEXPLAINED', len(d.unexplained))
    for u in d.unexplained[:40]:
        print('    ', u)
