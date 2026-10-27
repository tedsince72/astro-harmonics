#!/usr/bin/env python3
"""Independent checker (27 Oct 2026): S Jupiter column in T24 HORSE / JOCKEY, the four rows and SCORES of T24 HORSE / JOCKEY
(SKY STARS), SKY STARS SUMMARY. Own code from POSITIONS; families.py used only for targets / scores (evaluate, outside, CONJ).
usage: check_cells.py <tables.xlsx> [...]"""
import os, sys, math, collections
import openpyxl
from openpyxl.utils import get_column_letter as L

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..', '..', '..'))
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..', '..'))
import families as F  # noqa: E402

BODIES = ['Sun', 'Mercury', 'Venus', 'Mars', 'Jupiter', 'Saturn', 'Uranus', 'Neptune', 'Pluto', 'Chiron', 'Ceres', 'Pallas', 'Juno',
          'Vesta', 'Eris', 'Sedna', 'Haumea', 'Makemake', 'Gonggong', 'Quaoar', 'Orcus', 'Transpluto', 'Rahu', 'Ketu']
FILL = {'green': 'FFC6EFCE', 'yellow': 'FFFFEB9C', 'light': 'FFFFF8DC', 'red': 'FFFF8080', '': None}
GREEN, YELLOW = 80, 50

def dist(a, b, plane):
    if plane == 'RA':
        d = abs(a[0] - b[0]) % 360.0
        return min(d, 360.0 - d)
    return abs(a[1] - b[1])

def expect(v, plane):
    """-> (value 6dp, comment first line, fill key, score or None)"""
    v6 = round(v, 6)
    hits = F.evaluate(v, plane, 'deg')       # scored on the unrounded distance (shown value 6 dp)
    if hits:
        b = hits[0]
        s = b['score']
        if b['family'] == 'conj':
            txt = f"0 · conjunction ({plane} < {F.CONJ[plane]:g}) · {v:+.5f} · {s:.2f}"
        else:
            txt = f"{b['target']} · {b['family']} · {b['off']:+.5f} · {s:.2f}"
        fill = 'red' if v6 < F.CONJ[plane] else ('green' if s >= GREEN else 'yellow' if s >= YELLOW else 'light')
        return v6, txt, fill, s
    o = F.outside(v, plane, 'deg')
    if o['digit']:
        txt = f"nearest: {o['target']} {o['family']} — not a 4-digit match ({o['off']:+.5f})"
    else:
        txt = f"nearest: {o['target']} · {o['family']} — outside ±{o['tol']:.3g} · {o['off']:+.5f}"
    return v6, txt, ('red' if v6 < F.CONJ[plane] else ''), None

def rgb(c):
    f = c.fill
    if f is None or f.fill_type in (None, 'none'):
        return None
    x = f.fgColor.rgb
    return None if x in ('00000000', None) else 'FF' + x[2:]

def com(c):
    return c.comment.text if c.comment else None

class R:
    def __init__(s):
        s.n = collections.Counter(); s.bad = []
    def ok(s, key, cond, msg):
        s.n[key] += 1
        if not cond:
            s.bad.append(f'[{key}] {msg}')

def load_pos(wb):
    pos = {}
    for r in wb['POSITIONS'].iter_rows(min_row=5, values_only=True):
        if r[0] and r[2] in ('S', 'H', 'J') and isinstance(r[5], (int, float)):
            pos[(r[0], r[2], r[3])] = (r[5], r[6])
    return pos

def check_cell(rep, key, c, v, plane, where):
    v6, txt, fk, s = expect(v, plane)
    rep.ok(key + ' value', isinstance(c.value, (int, float)) and abs(c.value - v6) < 5e-7, f'{where} value {c.value} vs {v6}')
    cm = com(c)
    first = cm.split('\n')[0] if cm else None
    rep.ok(key + ' comment', first == txt, f'{where} comment {first!r} vs {txt!r}')
    rep.ok(key + ' fill', rgb(c) == FILL[fk], f'{where} fill {rgb(c)} vs {FILL[fk]} ({v6})')
    rep.ok(key + ' format', c.number_format == '0.000000', f'{where} number format {c.number_format}')
    return v6, s, fk

def header_cols(ws, hr2, first_col):
    """body-name header row -> dict name -> [cols]"""
    out = collections.defaultdict(list)
    for c in range(first_col, ws.max_column + 1):
        x = ws.cell(hr2, c).value
        if x is not None:
            out[x].append(c)
    return out

def main(path):
    rep = R()
    wb = openpyxl.load_workbook(path)
    pos = load_pos(wb)
    sky = wb['SKY JUPITER']
    skyblock = {}
    for r in range(1, sky.max_row + 1):
        if sky.cell(r, 1).value == 'SKY' and str(sky.cell(r, 2).value or '').startswith('S Jupiter → S '):
            skyblock[sky.cell(r, 2).value[len('S Jupiter → S '):]] = (sky.cell(r, 3), sky.cell(r, 8), sky.cell(r, 12), sky.cell(r, 17))
    sj = pos[('SKY', 'S', 'Jupiter')]
    runners = [n for n in wb.sheetnames if n.startswith('P') and n[1:3].isdigit() and len(n) == 3]
    sstables = {}
    for p in runners:
        ws = wb[p]
        rows = {}
        for r in range(1, ws.max_row + 1):
            a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
            rows.setdefault((a, (b or '')[:30]), []).append(r)
        lab = {}
        for r in range(1, ws.max_row + 1):
            a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
            if isinstance(b, str):
                lab.setdefault((a, b), r)
        for ch, name in (('H', 'HORSE'), ('J', 'JOCKEY')):
            # ---------------- (a) T24 HORSE / JOCKEY ----------------
            t = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == 'T24' and str(ws.cell(r, 2).value or '').startswith(f'T24 {name} (natal)')]
            rep.ok('a block', len(t) == 1, f'{p} T24 {name} title rows {t}')
            tr = t[0]; h1, h2 = tr + 1, tr + 2
            hc = header_cols(ws, h2, 3)
            cols_ra = list(range(3, 3 + 24)); cols_dec = list(range(29, 29 + 24))
            for i, bname in enumerate(BODIES):
                rep.ok('a header', ws.cell(h2, cols_ra[i]).value == bname and ws.cell(h2, cols_dec[i]).value == bname, f'{p} {name} header {bname}')
            sjc = {'RA': 27, 'Dec': 53}
            rep.ok('a header', ws.cell(h2, 27).value == 'S Jupiter' and ws.cell(h2, 53).value == 'S Jupiter', f'{p} {name} S Jupiter header {ws.cell(h2,27).value} {ws.cell(h2,53).value}')
            rep.ok('a header', ws.cell(h1, 28).value == 'RA midpoints' and ws.cell(h1, 54).value == 'Dec midpoints' and ws.cell(h1, 55).value == 'repeats',
                   f'{p} {name} header row1 {ws.cell(h1,28).value} {ws.cell(h1,54).value} {ws.cell(h1,55).value}')
            merged = [str(m) for m in ws.merged_cells.ranges if m.min_row == h1]
            rep.ok('a header', 'C%d:AA%d' % (h1, h1) in merged and 'AC%d:BA%d' % (h1, h1) in merged, f'{p} {name} merged {merged}')
            for cl in ('AA', 'BA'):
                w = ws.column_dimensions[cl].width
                rep.ok('a width', w is not None and w >= 11, f'{p} column {cl} width {w}')
            r = h2 + 1
            nstar = 0
            while ws.cell(r, 1).value == 'T24' and str(ws.cell(r, 2).value or '').startswith(ch + ' '):
                star = ws.cell(r, 2).value[2:]
                sp = pos[('SKY', 'S', star)]
                for plane in ('RA', 'Dec'):
                    c = ws.cell(r, sjc[plane])
                    v6, s, fk = check_cell(rep, 'a S Jupiter', c, dist(sj, sp, plane), plane, f'{p} {name} {star} {plane}')
                    sb = skyblock.get(star)
                    rep.ok('a = SKY block', sb is not None, f'{star} not in SKY block')
                    if sb:
                        vc, sc = (sb[0], sb[1]) if plane == 'RA' else (sb[2], sb[3])
                        rep.ok('a = SKY block', vc.value == c.value and (sc.value if sc.value != '' else None) == (round(s, 2) if s is not None else None),
                               f'{p} {name} {star} {plane}: T24 {c.value} / {s} vs SKY {vc.value} / {sc.value}')
                    # not used in midpoints / repeats
                    cm = com(c) or ''
                    rep.ok('a no mid/rep line', '\n' not in cm, f'{p} {name} {star} {plane} extra comment lines {cm!r}')
                    rep.ok('a border', not (c.border.left.style or c.border.right.style or c.border.top.style or c.border.bottom.style),
                           f'{p} {name} {star} {plane} border {c.border.left.style}')
                # calibration: every T24 body cell of the row by the same formatter (first comment line)
                for plane, c0 in (('RA', 3), ('Dec', 29)):
                    for i, x in enumerate(BODIES):
                        if (p, ch, star) in pos and (p, ch, x) in pos:
                            check_cell(rep, 'z calibration T24 body', ws.cell(r, c0 + i), dist(pos[(p, ch, star)], pos[(p, ch, x)], plane), plane,
                                       f'{p} {name} {star} {x} {plane}')
                for mc in (28, 54, 55):
                    rep.ok('a no mid/rep line', 'S Jupiter' not in str(ws.cell(r, mc).value or ''), f'{p} {name} {star} col {mc} names S Jupiter')
                nstar += 1; r += 1
            rep.n[f'a star rows'] += nstar
            # ---------------- (b), (c) (SKY STARS) ----------------
            tt = f'T24 {name} (SKY STARS)'
            t = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == tt and str(ws.cell(r, 2).value or '').startswith(tt)]
            rep.ok('b block', len(t) == 1, f'{p} {tt} title rows {t}')
            tr = t[0]; h1, h2 = tr + 1, tr + 2
            hr = [ws.cell(h2, c).value for c in range(1, 83)]
            # layout: C sky hit, D..AA bodies RA, AB S Jupiter, AC mids, AD..BA Dec, BB S Jupiter, BC mids, BD repeats, BE..CC SCORES RA, CD..DB SCORES Dec
            RA0, DEC0, SRA0, SDEC0 = 4, 30, 57, 82
            for i, bname in enumerate(BODIES + ['S Jupiter']):
                for c0 in (RA0, DEC0, SRA0, SDEC0):
                    rep.ok('b header', ws.cell(h2, c0 + i).value == bname, f'{p} {tt} header col {L(c0+i)} {ws.cell(h2, c0+i).value} vs {bname}')
            rep.ok('b header', ws.cell(h1, 55).value == 'SCORES Dec' if False else True, '')
            h1v = {L(c): ws.cell(h1, c).value for c in range(1, 110) if ws.cell(h1, c).value is not None}
            rep.ok('b header', h1v == {'A': 'table', 'B': 'star', 'C': 'sky hit', 'D': 'RA', 'AC': 'RA midpoints', 'AD': 'Dec', 'BC': 'Dec midpoints',
                                      'BD': 'repeats', 'BE': 'SCORES RA', 'CD': 'SCORES Dec'}, f'{p} {tt} header row1 {h1v}')
            rep.ok('b header', ws.cell(h2, 107).value is None and ws.cell(h2, 105).value == 'Ketu' and ws.cell(h2, 106).value == 'S Jupiter',
                   f'{p} {tt} SCORES Dec end {ws.cell(h2,106).value} {ws.cell(h2,107).value} {ws.cell(h2,108).value}')
            merged = sorted(str(m) for m in ws.merged_cells.ranges if m.min_row == h1)
            exp_m = sorted(['D%d:AB%d' % (h1, h1), 'AD%d:BB%d' % (h1, h1), 'BE%d:CC%d' % (h1, h1), 'CD%d:DB%d' % (h1, h1)])
            rep.ok('b header', all(m in merged for m in exp_m), f'{p} {tt} merged {merged} expected {exp_m}')
            pj = (p, ch, 'Jupiter')
            nj = pos[pj]
            labels = [f'S Jupiter → S body', f'{ch} Jupiter → {ch} body', f'S Jupiter → {ch} body', f'S body → {ch} Jupiter']
            # reference tables
            def find(tab, lbl):
                rr = lab.get((tab, lbl))
                return rr
            for k, lb in enumerate(labels):
                r = h2 + 1 + k
                rep.ok('b labels', ws.cell(r, 1).value == tt and ws.cell(r, 2).value == lb, f'{p} {tt} row {r} label {ws.cell(r,2).value!r} vs {lb!r}')
                for col in (3, 28, 29, 54, 55, 56):
                    rep.ok('b blanks', ws.cell(r, col).value is None and rgb(ws.cell(r, col)) is None and com(ws.cell(r, col)) is None,
                           f'{p} {tt} {lb} col {L(col)} not blank: {ws.cell(r,col).value!r}')
                for plane, c0, s0 in (('RA', RA0, SRA0), ('Dec', DEC0, SDEC0)):
                    for i, x in enumerate(BODIES):
                        c = ws.cell(r, c0 + i); sc = ws.cell(r, s0 + i)
                        if k == 0:
                            a, b = sj, pos[('SKY', 'S', x)]; blank = x == 'Jupiter'
                        elif k == 1:
                            a, b = nj, pos[(p, ch, x)]; blank = x == 'Jupiter'
                        elif k == 2:
                            a, b = sj, pos[(p, ch, x)]; blank = False
                        else:
                            a, b = pos[('SKY', 'S', x)], nj; blank = False
                        where = f'{p} {tt} {lb} {x} {plane}'
                        if blank:
                            rep.ok('b blanks', c.value is None and com(c) is None and rgb(c) is None and sc.value is None and rgb(sc) is None, f'{where} not blank {c.value}')
                            continue
                        v6, s, fk = check_cell(rep, 'b four rows', c, dist(a, b, plane), plane, where)
                        cm = com(c) or ''
                        rep.ok('b four rows', '\n' not in cm, f'{where} extra comment lines {cm!r}')
                        # SCORES
                        exp_s = round(s, 2) if s is not None else None
                        rep.ok('c SCORES four rows', (sc.value == exp_s) if exp_s is not None else sc.value is None,
                               f'{where} SCORES {sc.value} vs {exp_s}')
                        rep.ok('c SCORES four rows', rgb(sc) == rgb(c), f'{where} SCORES fill {rgb(sc)} vs {rgb(c)}')
                        # vs other tables
                        if k == 0:
                            rr = find('T1' if ch == 'H' else 'T2', f'{ch} Jupiter → {ch} {x}')
                            ref = ws.cell(rr, 19 if plane == 'RA' else 28) if rr else None; rs = ws.cell(rr, 24 if plane == 'RA' else 33) if rr else None
                            tag = 'b = T1/T2 SKY'
                        elif k == 1:
                            rr = find('T1' if ch == 'H' else 'T2', f'{ch} Jupiter → {ch} {x}')
                            ref = ws.cell(rr, 3 if plane == 'RA' else 11) if rr else None; rs = ws.cell(rr, 8 if plane == 'RA' else 16) if rr else None
                            tag = 'b = T1/T2'
                        elif k == 2:
                            if x == 'Jupiter':
                                rr = find('T9', f'S Jupiter – {ch} Jupiter'); tag = 'b = T9'
                            else:
                                rr = find('T5' if ch == 'H' else 'T6', f'S Jupiter → {ch} {x}'); tag = 'b = T5/T6'
                            ref = ws.cell(rr, 3 if plane == 'RA' else 11) if rr else None; rs = ws.cell(rr, 8 if plane == 'RA' else 16) if rr else None
                        else:
                            if x == 'Jupiter':
                                rr = find('T9', f'S Jupiter – {ch} Jupiter'); tag = 'b = T9'
                            else:
                                rr = find('T7' if ch == 'H' else 'T8', f'S {x} → {ch} Jupiter'); tag = 'b = T7/T8'
                            ref = ws.cell(rr, 3 if plane == 'RA' else 11) if rr else None; rs = ws.cell(rr, 8 if plane == 'RA' else 16) if rr else None
                        if ref is None:
                            rep.n[tag + ' (no row)'] += 1
                            rep.bad.append(f'[{tag}] {where}: no reference row')
                            continue
                        rsv = rs.value if rs.value not in ('', None) else None
                        rep.ok(tag, ref.value == c.value and rsv == exp_s, f'{where}: {c.value} / {exp_s} vs {tag[4:]} row {rr} {ref.value} / {rsv}')
            sep = h2 + 5
            rep.ok('b separator', ws.cell(sep, 1).value == tt and ws.cell(sep, 2).value == 'stars' and all(ws.cell(sep, c).value is None for c in range(3, 110)),
                   f'{p} {tt} separator row {sep}: {ws.cell(sep,2).value!r}')
            # star rows: copy of T24 row + S Jupiter + SCORES
            t24 = {ws.cell(r, 2).value: r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == 'T24' and str(ws.cell(r, 2).value or '').startswith(ch + ' ')}
            r = sep + 1; nst = 0
            while ws.cell(r, 1).value == tt:
                lbl = ws.cell(r, 2).value; src = t24.get(lbl)
                rep.ok('c star rows', src is not None, f'{p} {tt} {lbl} not in T24')
                for c in range(3, 56):          # T24 cols C..BC -> D..BD
                    a, b = ws.cell(src, c), ws.cell(r, c + 1)
                    rep.ok('c copied cells', a.value == b.value and rgb(a) == rgb(b) and com(a) == com(b) and a.border.left.style == b.border.left.style
                           and a.border.top.style == b.border.top.style, f'{p} {tt} {lbl} col {L(c+1)}: {b.value} vs T24 {a.value}')
                for plane, c0, s0 in (('RA', RA0, SRA0), ('Dec', DEC0, SDEC0)):
                    for i in range(25):
                        c = ws.cell(r, c0 + i); sc = ws.cell(r, s0 + i)
                        cm = com(c)
                        exp_s = None
                        if cm and not cm.startswith('nearest:'):
                            exp_s = round(float(cm.split('\n')[0].split(' · ')[-1]), 2)
                        if i == 24:   # S Jupiter: recompute
                            _, _, _, s = expect(dist(sj, pos[('SKY', 'S', lbl[2:])], plane), plane)
                            rep.ok('c SCORES S Jupiter', (round(s, 2) if s is not None else None) == exp_s, f'{p} {tt} {lbl} {plane} SJ score {s} vs comment {exp_s}')
                        if c.value is None:
                            rep.ok('c SCORES star rows', sc.value is None, f'{p} {tt} {lbl} {plane} {L(s0+i)} SCORES without value')
                            continue
                        rep.ok('c SCORES star rows', sc.value == exp_s, f'{p} {tt} {lbl} {plane} {L(s0+i)} SCORES {sc.value} vs {exp_s}')
                        rep.ok('c SCORES star rows', rgb(sc) == (rgb(c) if exp_s is not None else None), f'{p} {tt} {lbl} {plane} {L(s0+i)} fill {rgb(sc)} vs {rgb(c)}')
                nst += 1; r += 1
            rep.n['c star rows count'] += nst
            sstables[(p, ch)] = (tr, r - 1)
    # ---------------- (d) summary ----------------
    sm = wb['SKY STARS SUMMARY']
    srow = 1
    rows_sm = [r for r in range(1, sm.max_row + 1)]
    # locate each runner's blocks by matching first title row
    for p in runners:
        ws = wb[p]
        for ch in ('H', 'J'):
            tr, last = sstables[(p, ch)]
            title = ws.cell(tr, 2).value
            cand = [r for r in rows_sm if sm.cell(r, 2).value == title and sm.cell(r + 3, 2).value is not None]
            # pick the one followed by this runner's rows: find by A1 line preceding
            found = None
            for r0 in cand:
                ok = all(sm.cell(r0 + k, c).value == ws.cell(tr + k, c).value for k in range(last - tr + 1) for c in range(1, 110))
                if ok:
                    found = r0; break
            rep.ok('d summary block', found is not None, f'{p} {ch} block not found in summary')
            if found is None:
                continue
            for k in range(last - tr + 1):
                for c in range(1, 110):
                    a, b = ws.cell(tr + k, c), sm.cell(found + k, c)
                    rep.ok('d summary cells', a.value == b.value and rgb(a) == rgb(b) and com(a) == com(b) and a.border.left.style == b.border.left.style
                           and a.border.top.style == b.border.top.style and a.number_format == b.number_format,
                           f'summary {p} {ch} r{found+k} {L(c)}: {b.value} vs {a.value}')
    return rep

if __name__ == '__main__':
    for pth in sys.argv[1:]:
        rep = main(pth)
        print('==', pth)
        for k, v in sorted(rep.n.items()):
            print(f'  {k}: {v}')
        print('  mismatches:', len(rep.bad))
        for b in rep.bad[:60]:
            print('   ', b)
