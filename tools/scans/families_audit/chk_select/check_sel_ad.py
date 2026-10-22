"""check_sel_ad.py: (a) sky STAR / LONG STAR TIMELINE after = before minus the star-to-star distance events (old columns identical);
(d) every other row of every sheet identical to before (values, fills, comments), default sheets = --full sheets.
usage: check_sel_ad.py before.xlsx after.xlsx [default_after.xlsx]"""
import difflib, re, sys
import openpyxl

STARB = {'STAR TIMELINE': 13, 'LONG STAR TIMELINE': 15}


def sig(row):
    out = []
    for c in row:
        v = c.value
        f = c.fill.fgColor.rgb if c.fill is not None and c.fill.fill_type else None
        cm = c.comment.text if c.comment else None
        out.append((v, f, cm))
    while out and out[-1] == (None, None, None):
        out.pop()
    return tuple(out)


def sheet_rows(ws):
    return [sig(r) for r in ws.iter_rows()]


def key(s):
    a = s[0][0] if s else None
    if a in STARB:
        t = list(s[:STARB[a]])
        while t and t[-1] == (None, None, None):
            t.pop()
        return tuple(t)
    return s


def short(s):
    return [x[0] for x in s][:8]


before = openpyxl.load_workbook(sys.argv[1])
after = openpyxl.load_workbook(sys.argv[2])
print('sheets before == after:', before.sheetnames == after.sheetnames)
dump = {}
for name in after.sheetnames:
    A = sheet_rows(after[name])
    dump[name] = A
    if name not in before.sheetnames:
        print('NEW SHEET', name)
        continue
    B = sheet_rows(before[name])
    kb, ka = [key(s) for s in B], [key(s) for s in A]
    sm = difflib.SequenceMatcher(None, kb, ka, autojunk=False)
    nd = 0
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op == 'equal':
            continue
        nd += 1
        print(f'[{name}] {op} before {i1 + 1}-{i2} after {j1 + 1}-{j2}')
        for s in B[i1:i2][:60]:
            print('   -', short(s))
        if op != 'delete':
            ins = A[j1:j2]
            ids = {}
            for s in ins:
                ids[s[0][0] if s else None] = ids.get(s[0][0] if s else None, 0) + 1
            print('   + rows by col A:', {str(k)[:60]: v for k, v in ids.items()})
            for s in ins[:6]:
                print('   +', [str(x)[:100] for x in short(s)])
    print(f'[{name}] rows before {len(B)} after {len(A)}, diff blocks {nd}')

if len(sys.argv) > 3:
    d = openpyxl.load_workbook(sys.argv[3])
    print('default sheets:', d.sheetnames)
    for name in d.sheetnames:
        D = sheet_rows(d[name])
        print(f'default {name}: identical to --full = {D == dump[name]} ({len(D)} rows)')
