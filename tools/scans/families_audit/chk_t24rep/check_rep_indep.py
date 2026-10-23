"""Independent checker for T24 REPEATS (22 Oct 2026). Own code; families.py for targets/scores only."""
import sys, collections
sys.path.insert(0, 'tools')
import families, openpyxl
from openpyxl.utils import get_column_letter

S = sys.argv[1]
mism = []
def bad(*a):
    mism.append(' '.join(str(x) for x in a))

def positions(wb):
    ws = wb['POSITIONS']; P = {}
    for r in range(5, ws.max_row + 1):
        run, ch, nm, ra, dec = (ws.cell(r, c).value for c in (1, 3, 4, 6, 7))
        if ch in ('S', 'H', 'J') and isinstance(ra, (int, float)):
            P[(ch, nm)] = (ra, dec)
    return P

def blocks(ws):
    out = []
    for r in range(1, ws.max_row + 1):
        b = ws.cell(r, 2).value
        if ws.cell(r, 1).value == 'T24' and isinstance(b, str) and b.split(' (')[0].split(':')[0] in ('T24 SKY', 'T24 HORSE', 'T24 JOCKEY'):
            out.append((r, b.split(' (')[0].split(':')[0]))
    return out

def layout(ws, r):
    h1, h2 = r + 1, r + 2
    cols = {}
    plane = None; ramid = decmid = rep = None
    for c in range(3, ws.max_column + 1):
        v1 = ws.cell(h1, c).value
        if v1 == 'RA': plane = 'RA'
        elif v1 == 'RA midpoints': ramid = c; plane = None; continue
        elif v1 == 'Dec': plane = 'Dec'
        elif v1 == 'Dec midpoints': decmid = c; plane = None; continue
        elif v1 == 'repeats': rep = c; continue
        if plane and ws.cell(h2, c).value:
            cols[c] = (plane, ws.cell(h2, c).value)
    return cols, ramid, decmid, rep

def fmt(s): return f'{s:.2f}'

def expected(ch, star, cols, P):
    sra, sdec = P[(ch, star)]
    cells = []  # (col, plane, body, value, fam, target, score)
    vals = {}
    for c, (pl, body) in cols.items():
        bra, bdec = P[(ch, body)]
        if pl == 'RA':
            d = abs(sra - bra) % 360; d = 360 - d if d > 180 else d
        else:
            d = abs(sdec - bdec)
        vals[c] = d
        if (pl == 'RA' and d < 1.5) or (pl == 'Dec' and d < 0.3):
            continue
        ev = families.evaluate(d, pl, 'deg')
        if not ev: continue
        b = ev[0]
        if b['family'] == 'conj': bad('conj best outside range?', ch, star, body, pl, d); continue
        cells.append((c, pl, body, d, b['family'], b['target'], b['score']))
    g = collections.defaultdict(list)
    for x in cells: g[(x[4], x[5])].append(x)
    reps = []
    for (fam, tgt), xs in g.items():
        if len(xs) < 2: continue
        xs.sort(key=lambda x: (-x[6], 0 if x[1] == 'RA' else 1, x[0]))
        reps.append((fam, tgt, xs))
    reps.sort(key=lambda t: (-len(t[2]), -max(x[6] for x in t[2])))
    # tie check
    for i in range(len(reps) - 1):
        a, b = reps[i], reps[i + 1]
        if len(a[2]) == len(b[2]) and fmt(max(x[6] for x in a[2])) == fmt(max(x[6] for x in b[2])):
            print('  NOTE tie in target order', ch, star, a[1], b[1])
    text = '; '.join(f"{tgt} {fam} ×{len(xs)} (" + ', '.join(f"{x[2]} {x[1]} ({fmt(x[6])})" for x in xs) + ')' for fam, tgt, xs in reps)
    lines = {}
    for fam, tgt, xs in reps:
        for x in xs:
            others = ', '.join(f"{y[2]} {y[1]} ({fmt(y[6])})" for y in xs if y is not x)
            lines.setdefault(x[0], []).append(f"repeat: {tgt} {fam} ×{len(xs)} (with {others})")
    return text, lines, vals, reps

def bstyle(b):
    return tuple((s.style, s.color.rgb if s.color is not None else None) for s in (b.left, b.right, b.top, b.bottom))

def cmt(c): return c.comment.text if c.comment else None
def fill(c): return (c.fill.fill_type, c.fill.fgColor.rgb)

def run(tag):
    bwb = openpyxl.load_workbook(f'{S}/before/ascot_{tag}_jupiter_tables.xlsx')
    awb = openpyxl.load_workbook(f'{S}/after/ascot_{tag}_jupiter_tables.xlsx')
    P = positions(awb)
    if bwb.sheetnames != awb.sheetnames: bad(tag, 'sheetnames differ')
    handled = collections.defaultdict(set)  # sheet -> {(r,c)} cells checked by block logic
    stats = []
    scorecases = []
    for sn in ('SKY JUPITER', 'P01'):
        ws, wb_ = awb[sn], bwb[sn]
        for r0, name in blocks(ws):
            ch = {'T24 SKY': 'S', 'T24 HORSE': 'H', 'T24 JOCKEY': 'J'}[name]
            cols, ramid, decmid, rep = layout(ws, r0)
            if rep is None or rep != decmid + 1: bad(tag, sn, name, 'repeats column not after Dec midpoints', rep, decmid)
            if ws.cell(r0 + 2, rep).value is None: bad(tag, sn, name, 'no row-2 header on repeats')
            print(f'  {tag} {sn} {name} header row2 repeats: {ws.cell(r0+2, rep).value!r}; width {ws.column_dimensions[get_column_letter(rep)].width}')
            if ws.column_dimensions[get_column_letter(rep)].width != 60: bad(tag, sn, name, 'width')
            # header cells
            for hr in (r0 + 1, r0 + 2):
                handled[sn].add((hr, rep))
            r = r0 + 3; nrows = nrep = ntg = 0; maxn = 0
            while ws.cell(r, 1).value == 'T24' and isinstance(ws.cell(r, 2).value, str) and ws.cell(r, 2).value.startswith(ch + ' '):
                star = ws.cell(r, 2).value[2:]
                text, lines, vals, reps = expected(ch, star, cols, P)
                nrows += 1
                if reps:
                    nrep += 1; ntg += len(reps); maxn = max(maxn, max(len(x[2]) for x in reps))
                    for fam, tgt, xs in reps:
                        for x in xs: scorecases.append((sn, name, star, x))
                got = ws.cell(r, rep).value
                if (got or '') != text: bad(tag, sn, name, star, 'repeats text', repr(got), '!=', repr(text))
                if text and not ws.cell(r, rep).alignment.wrap_text: bad(tag, sn, name, star, 'no wrap')
                handled[sn].add((r, rep))
                for c in cols:
                    a, b = ws.cell(r, c), wb_.cell(r, c)
                    handled[sn].add((r, c))
                    if a.value is None or abs(a.value - round(vals[c], 6)) > 1.5e-6: bad(tag, sn, name, star, c, 'value', a.value, vals[c])
                    if a.value != b.value or fill(a) != fill(b) or a.number_format != b.number_format: bad(tag, sn, name, star, c, 'value/fill/fmt changed')
                    bc, ac = cmt(b), cmt(a)
                    if c in lines:
                        exp = (bc + '\n' if bc else '') + '\n'.join(lines[c])
                        if ac != exp: bad(tag, sn, name, star, cols[c], 'comment', repr(ac), repr(exp))
                        bb, ab = bstyle(b.border), bstyle(a.border)
                        thick = any(s[0] in ('thick', 'medium') for s in bb)
                        if thick:
                            if ab != bb: bad(tag, sn, name, star, cols[c], 'thick border changed', bb, ab)
                        else:
                            if any(s[0] != 'dotted' for s in ab): bad(tag, sn, name, star, cols[c], 'not dotted', ab)
                    else:
                        if ac != bc: bad(tag, sn, name, star, cols[c], 'comment changed')
                        if bstyle(a.border) != bstyle(b.border): bad(tag, sn, name, star, cols[c], 'border changed')
                r += 1
            title = ws.cell(r0, 2).value
            suf = f'. REPEATS (22 Oct 2026): {nrep} of {nrows} rows with a repeat, {ntg} repeated targets'
            if not title.endswith(suf): bad(tag, sn, name, 'title suffix', title[-90:], 'expected', suf)
            handled[sn].add((r0, 2))
            stats.append((sn, name, nrep, nrows, ntg, maxn))
            # title before vs after: after desc part may also have changed (note)
            bt = wb_.cell(r0, 2).value
            if title[:-len(suf)] != bt:
                print(f'  NOTE {tag} {name}: title text before the suffix also changed (len before {len(bt)}, after-minus-suffix {len(title)-len(suf)})')
    # every other cell of every sheet
    ncell = 0
    for sn in awb.sheetnames:
        a, b = awb[sn], bwb[sn]
        if a.max_row != b.max_row: bad(tag, sn, 'max_row', b.max_row, a.max_row)
        mc = max(a.max_column, b.max_column)
        for r in range(1, max(a.max_row, b.max_row) + 1):
            for c in range(1, mc + 1):
                if (r, c) in handled[sn]: continue
                x, y = a.cell(r, c), b.cell(r, c)
                ncell += 1
                if (x.value != y.value or fill(x) != fill(y) or cmt(x) != cmt(y) or bstyle(x.border) != bstyle(y.border)
                        or x.number_format != y.number_format):
                    bad(tag, sn, get_column_letter(c) + str(r), 'differs', repr(y.value)[:60], '->', repr(x.value)[:60])
        for k, d in b.column_dimensions.items():
            if k in a.column_dimensions and a.column_dimensions[k].width != d.width:
                bad(tag, sn, 'col width', k, d.width, a.column_dimensions[k].width)
        if sorted(map(str, a.merged_cells.ranges)) != sorted(map(str, b.merged_cells.ranges)):
            print('  NOTE merged ranges differ on', sn, set(map(str, a.merged_cells.ranges)) ^ set(map(str, b.merged_cells.ranges)))
    print(f'{tag}: other cells compared {ncell}')
    for s in stats: print(f'  {tag} {s[0]} {s[1]}: {s[2]} of {s[3]} rows with a repeat, {s[4]} repeated targets, largest x{s[5]}')
    return scorecases

sc = None
for tag in ('def', 'full'):
    x = run(tag)
    sc = sc or x
import json
json.dump([(a, b, c, list(x)) for a, b, c, x in sc], open(f'{S}/indep_scorecases.json', 'w'))
print('mismatches:', len(mism))
for m in mism[:60]: print(' ', m)
