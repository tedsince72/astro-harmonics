import sys, math, re
sys.path.insert(0,'/Users/eddieplace/astro-harmonics/rebuild_kit'); sys.path.insert(0,'.')
import openpyxl, eph3
from active_vibrations import score_whole, score_phi_interval, score_phi_powers, score_ninth, _NINTH_RANGES, _PHI_POWER_TARGETS
INP=eph3.INP; OUT='k1915.xlsx'
BODIES=['Sun','Mercury','Venus','Mars','Jupiter','Saturn','Uranus','Neptune','Pluto','Chiron','Ceres','Pallas','Juno','Vesta','Rahu','Ketu','Eris','Sedna','Haumea','Makemake','Orcus','Quaoar','Gonggong','Transpluto']
ORD=[('S','H'),('S','J'),('H','S'),('H','J'),('J','S'),('J','H')]
wi=openpyxl.load_workbook(INP,data_only=True)
def load(tab): return {r[1]:(r[4],r[8]) for r in wi[tab].iter_rows(min_row=2,values_only=True) if r[3]=='transit_self'}
S=load('TRANS'); s0=eph3.sky(0); s5=eph3.sky(5)
S5={}
for n,(ra,dec) in S.items():
    if n in s0:
        dra=((s5[n][0]-s0[n][0]+180)%360)-180; S5[n]=((ra+dra)%360,dec+s5[n][1]-s0[n][1])
def addketu(d):
    if 'Rahu' in d and 'Ketu' not in d: ra,dec=d['Rahu']; d['Ketu']=(None,-dec)
    elif 'Ketu' in d: d['Ketu']=(None,d['Ketu'][1])
for d in (S,S5): addketu(d)
def ra_sep(a,b): x=abs(a-b)%360; return 360-x if x>180 else x
import importlib.util
# reuse checker's degree families (independent of the tool)
src=open('chk16.py').read(); i0=src.index('def phi_label'); i1=src.index('# ---------- intervals')
PHI=(1+5**0.5)/2; R2=2**0.5; exec(src[i0:i1].replace("'--exactlabels' in sys.argv","False"))
def colour(s):
    if s>=80: return 'C6EFCE'
    if s>=50: return 'FFEB9C'
    if s>0: return 'FFF8DC'
def fill(c):
    if c.fill and c.fill.fill_type:
        rgb=c.fill.fgColor.rgb
        if rgb and rgb!='00000000': return rgb[-6:]
wo=openpyxl.load_workbook(OUT)
runners=[n for n in wo.sheetnames if re.match(r'^P\d\d$',n)]
total=0; mism=[]; hits=[]; notes=[]
for P in runners:
    H=load(P); J=load('P%02d'%(int(P[1:])+1)); addketu(H); addketu(J)
    C={'S':S,'H':H,'J':J}
    ws=wo[P]; rows={}
    for r in range(1,ws.max_row+1):
        if ws.cell(r,1).value=='T16' and ws.cell(r,2).value and ws.cell(r,2).value.startswith('Jupiter–'): rows[ws.cell(r,2).value]=r
        if ws.cell(r,1).value=='table' and str(ws.cell(r,3).value).startswith('S RA = H Dec'):
            hdr=[ws.cell(r,c).value for c in range(3,3+8*6)]
    exp_rows=[x for x in BODIES if x!='Jupiter' and x in S and x in H and x in J]
    if list(rows)!=['Jupiter–'+x for x in exp_rows]: mism.append((P,'row set',list(rows),exp_rows))
    for x in exp_rows:
        r=rows.get('Jupiter–'+x)
        if r is None: continue
        for k,(a,b) in enumerate(ORD):
            c0=3+8*k; got=[ws.cell(r,c0+j).value for j in range(8)]; dcol=fill(ws.cell(r,c0+2)); fcol=fill(ws.cell(r,c0+6))
            total+=8
            def M(w,g,e): mism.append((P,x,'%s RA = %s Dec'%(a,b),w,g,e))
            if x=='Ketu':
                if any(v is not None for v in got): M('Ketu cell present',got,None)
                continue
            ra=ra_sep(C[a]['Jupiter'][0],C[a][x][0]); de=abs(C[b]['Jupiter'][1]-C[b][x][1]); d=ra-de
            exp5=None
            if 'S' in (a,b):
                A5=S5 if a=='S' else C[a]; B5=S5 if b=='S' else C[b]
                if all(n in A5 for n in ('Jupiter',x)) and all(n in B5 for n in ('Jupiter',x)):
                    d5=ra_sep(A5['Jupiter'][0],A5[x][0])-abs(B5['Jupiter'][1]-B5[x][1])
                    exp5=(d5,'-> toward' if abs(d5)<abs(d) else '<- away')
                else: exp5='n/a'
            hit=abs(d)<0.02; sc=100*(0.02-abs(d))/0.02 if hit else None
            if got[0] is None or abs(got[0]-ra)>2e-6: M('a RA',got[0],ra)
            if got[1] is None or abs(got[1]-de)>2e-6: M('b Dec',got[1],de)
            if got[2] is None or abs(got[2]-d)>2e-6: M('RA-Dec',got[2],d)
            if hit:
                if got[3] is None or abs(got[3]-sc)>0.011: M('score',got[3],sc)
                if dcol!=colour(sc): M('diff colour',dcol,colour(sc))
                hits.append((P,x,a,b,ra,de,d,sc,exp5))
                mean=(ra+de)/2; f=[z for z in fams_deg(mean,None)]
                ef=None if not f else '%s (%s) %+.5f [%.1f]'%(f[0][1],f[0][0],f[0][2],f[0][3])
                notes.append((P,x,a,b,'mean',mean,'families',f))
                if got[6]!=ef: M('family of the value',got[6],ef)
                if ef and fcol!=colour(f[0][3]): M('family colour',fcol,colour(f[0][3]))
            else:
                if got[3] is not None: M('score on non-hit',got[3],None)
                if dcol is not None: M('colour on non-hit',dcol,None)
                if got[6] is not None: M('family on non-hit',got[6],None)
            if exp5 is None:
                if got[4] is not None or got[5] is not None: M('+5 on natal-only',got[4:6],None)
            elif exp5=='n/a':
                if got[5]!='n/a': M('+5 n/a',got[4:6],'n/a')
            else:
                if got[4] is None or abs(got[4]-exp5[0])>2e-5: M('+5',got[4],exp5[0])
                if got[5]!=exp5[1]: M('arrow',got[5],exp5[1],)
            if got[7] is not None: M('spacer col not empty',got[7],None)
print('runners',runners,'T16 cells compared',total)
print('mismatches',len(mism))
for m in mism[:60]: print(' ',m)
print('hits',hits); print('notes',notes)
# FIELD 16
wf=wo['FIELD 16']; fgot=[]; ford=None; fc=0; fm=[]
names={}
for r in range(1,wf.max_row+1):
    a=wf.cell(r,1).value
    if a and ' RA = ' in a:
        ford=a; cols={wf.cell(r,c).value[:3]:c for c in range(2,wf.max_column+1) if isinstance(wf.cell(r,c).value,str) and wf.cell(r,c).value.startswith('P')}; continue
    if a and a.startswith('Jupiter–'):
        for P,c in cols.items():
            v=wf.cell(r,c).value; s=wf.cell(r,c+1).value; fc+=2
            if v is not None or s is not None: fgot.append((P,a[8:],ford,v,s,fill(wf.cell(r,c))))
exp=[(P,x,'%s RA = %s Dec'%(a,b),d,sc) for P,x,a,b,ra,de,d,sc,_ in hits]
print('FIELD16 non-empty',fgot,'expected',exp,'cells scanned',fc)
for e in exp:
    g=[z for z in fgot if z[:3]==e[:3]]
    if not g: fm.append(('missing',e)); continue
    g=g[0]
    if abs(g[3]-e[3])>2e-6: fm.append(('value',g,e))
    es='crossing %+.5f [%.1f]'%(e[3],e[4])
    if g[4]!=es: fm.append(('string',g[4],es))
    if g[5]!=colour(e[4]): fm.append(('colour',g[5],colour(e[4])))
for g in fgot:
    if not any(g[:3]==e[:3] for e in exp): fm.append(('extra',g))
print('FIELD16 mismatches',fm)
# near misses for info
