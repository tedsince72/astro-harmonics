"""check67_t1_t2_t24.py - independent checker (21 Oct 2026, adapted from chk_stars46) for the 67-star list in T1, T2 and T24
(T24 SKY / HORSE / JOCKEY incl. midpoints and the red conjunction fill FF8080). Own code from POSITIONS; families.py for family
targets / scores only; no import of tools/jupiter_tables.py.
usage: check67_t1_t2_t24.py <tables.xlsx>"""
import sys, os, math, datetime as dt
import openpyxl
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..', '..'))
sys.path.insert(0, os.path.dirname(__file__))
import families as F
from check67_common import STARS67, NEW21, BODIES as _B
STARS35 = STARS67
NEW = NEW21
BODIES = ['Sun', 'Moon', 'Mercury', 'Venus', 'Mars', 'Jupiter', 'Saturn', 'Uranus', 'Neptune', 'Pluto', 'Chiron', 'Ceres', 'Pallas',
          'Juno', 'Vesta', 'Eris', 'Sedna', 'Haumea', 'Makemake', 'Gonggong', 'Quaoar', 'Orcus', 'Transpluto', 'Rahu', 'Ketu']
FILL = {'green': 'C6EFCE', 'yellow': 'FFEB9C', 'light': 'FFF8DC'}
RED = 'FF8080'
FAMWORD = {'conj': 'conjunction'}
OFF = dt.datetime(2012, 6, 19, 14, 30)
mism, counts = [], {}


def bump(k, n=1):
    counts[k] = counts.get(k, 0) + n


def cmp(tag, coord, got, want, tolnum=1e-6):
    bump(tag)
    if isinstance(want, float) and isinstance(got, (int, float)):
        if abs(got - want) <= tolnum:
            return
    elif got == want or (got in (None, '') and want in (None, '')):
        return
    mism.append((tag, coord, got, want))


def wrap(d):
    d = (d + 180.0) % 360.0 - 180.0
    return d


def dist(a, b, plane):
    return abs(wrap(a[0] - b[0])) if plane == 'RA' else abs(a[1] - b[1])


def load_pos(wb):
    P = {}
    for r in wb['POSITIONS'].iter_rows(values_only=True):
        if r[0] in (None, 'runner') or r[2] not in ('S', 'H', 'J') or r[5] is None:
            continue
        P[(r[2], r[3])] = (float(r[5]), float(r[6]), float(r[7] or 0.0), float(r[8] or 0.0))
    for ch in 'SHJ':                                         # Ketu = Rahu + 180, -Dec, if not listed
        if (ch, 'Rahu') in P and (ch, 'Ketu') not in P:
            ra, de, rr, rd = P[(ch, 'Rahu')]
            P[(ch, 'Ketu')] = ((ra + 180) % 360, -de, rr, -rd)
    return P


def plus5(p):
    return (p[0] + 5 * p[2], p[1] + 5 * p[3], p[2], p[3])


def tol_txt(t):
    return f'{t:.3g}'


def fam_hits(v, plane):
    return F.evaluate(v, plane, 'deg')


def target_txt(h):
    return f"{h['target']} ({h['family']})"


def fill_of(c):
    return c.fill.fgColor.rgb[-6:] if c.fill is not None and c.fill.fill_type == 'solid' else None


def timing(v, v5, tv):
    rate = (v5 - v) / 5.0
    if rate == 0:
        return None
    t = (tv - v) / rate
    when = OFF + dt.timedelta(minutes=t)
    when = when.replace(second=0, microsecond=0)
    clock = when.strftime('%H:%M') if when.date() == OFF.date() else when.strftime('%d %b %Y %H:%M')
    return (f'applying, exact {clock} (in {t:.1f} min)' if t > 0 else f'separating, exact {clock} ({-t:.1f} min ago)'), t


# ---------------------------------------------------------------- T1 / T2
def check_cells(ws, r, col, v, v5, plane, tag, sky):
    """8 natal cells (value, +5, target, group, offset, score, arrow, also) or 9 sky cells (+ timing)"""
    cell = lambda k: ws.cell(r, col + k)
    hits = fam_hits(v, plane)
    cmp(tag + ' value', cell(0).coordinate, cell(0).value, round(v, 6))
    cmp(tag + ' +5', cell(1).coordinate, cell(1).value, round(v5, 6) if sky else None)
    cmp(tag + ' group', cell(3).coordinate, cell(3).value, None)
    if hits:
        b = hits[0]
        cmp(tag + ' target', cell(2).coordinate, cell(2).value, target_txt(b))
        cmp(tag + ' offset', cell(4).coordinate, cell(4).value, f"{b['off']:+.5f}")
        cmp(tag + ' score', cell(5).coordinate, cell(5).value, round(b['score'], 2), 0.0051)
        also = '; '.join(f"{target_txt(h)} {h['off']:+.5f} [{h['score']:.2f}]" for h in hits[1:])
        cmp(tag + ' also', cell(7).coordinate, cell(7).value, also or None)
        cmp(tag + ' fill', cell(0).coordinate, fill_of(cell(0)), FILL[F.colour(b['score'])])
        if sky:
            arrow = '-> toward' if abs(v5 - b['tv']) < abs(v - b['tv']) else '<- away'
            cmp(tag + ' arrow', cell(6).coordinate, cell(6).value, arrow)
            tm = timing(v, v5, b['tv'])
            cmp(tag + ' timing', cell(8).coordinate, cell(8).value, tm[0] if tm else None)
            if tm:
                bump('timed sky hits ' + tag.split()[0])
        else:
            cmp(tag + ' arrow', cell(6).coordinate, cell(6).value, None)
    else:
        o = F.outside(v, plane, 'deg')
        cmp(tag + ' target', cell(2).coordinate, cell(2).value, f"{o['target']} ({o['family']}) — outside ±{tol_txt(o['tol'])}")
        cmp(tag + ' offset', cell(4).coordinate, cell(4).value, f"{o['off']:+.5f}")
        cmp(tag + ' score', cell(5).coordinate, cell(5).value, None)
        cmp(tag + ' arrow', cell(6).coordinate, cell(6).value, None)
        cmp(tag + ' also', cell(7).coordinate, cell(7).value, None)
        cmp(tag + ' fill', cell(0).coordinate, fill_of(cell(0)), None)
        if sky:
            cmp(tag + ' timing', cell(8).coordinate, cell(8).value, None)


def check_t12(wb, P, tab, ch):
    ws = wb['P01']
    rows = [(r[0].row, r[1].value) for r in ws.iter_rows(max_col=2) if r[0].value == tab and r[1].value
            and str(r[1].value).startswith(f'{ch} Jupiter → {ch} ') and ',' not in str(r[1].value)]
    names = [lab.split(' → ')[1][2:] for _, lab in rows]
    stars = [n for n in names if n not in BODIES and n != 'equator']
    cmp(f'{tab} star rows (list)', tab, stars, STARS35)
    cmp(f'{tab} no Polaris', tab, 'Polaris' in names, False)
    cmp(f'{tab} order tail', tab, names[-2:], ['Ketu', 'equator'])
    body_rows = names[:names.index(STARS35[0])]
    cmp(f'{tab} bodies before stars', tab, all(n in BODIES for n in body_rows), True)
    jn, js = P[(ch, 'Jupiter')], P[('S', 'Jupiter')]
    for r, lab in rows:
        x = lab.split(' → ')[1][2:]
        if x not in STARS35:
            continue
        bump(f'{tab} star rows')
        hx, sx = P[(ch, x)], P[('S', x)]
        for plane, nc, sc in (('RA', 3, 19), ('Dec', 11, 28)):
            check_cells(ws, r, nc, dist(jn, hx, plane), None, plane, f'{tab} natal{" NEW" if x in NEW else ""}', False)
            check_cells(ws, r, sc, dist(js, sx, plane), dist(plus5(js), plus5(sx), plane), plane,
                        f'{tab} SKY{" NEW" if x in NEW else ""}', True)


# ---------------------------------------------------------------- T24
def comment_first(v, plane):
    hits = fam_hits(v, plane)
    if hits:
        b = hits[0]
        if b['family'] == 'conj':
            return f"0 · conjunction ({plane} < {F.CONJ[plane]}) · {b['off']:+.5f} · {b['score']:.2f}", b
        return f"{b['target']} · {b['family']} · {b['off']:+.5f} · {b['score']:.2f}", b
    o = F.outside(v, plane, 'deg')
    return f"nearest: {o['target']} · {o['family']} — outside ±{tol_txt(o['tol'])} · {o['off']:+.5f}", None


def mid_hits(s, pts, bodies, plane):
    out = []
    tol = 0.5 if plane == 'RA' else 0.15
    for i, a in enumerate(bodies):
        for b in bodies[i + 1:]:
            if {a, b} == {'Rahu', 'Ketu'}:
                continue
            A, B = pts[a], pts[b]
            dab = dist(A, B, plane)
            if dab < F.CONJ[plane]:
                continue
            d1, d2 = dist(s, A, plane), dist(s, B, plane)
            if max(d1, d2) == 0:
                continue
            clo = 100.0 * min(d1, d2) / max(d1, d2)
            if clo < 90:
                continue
            if plane == 'RA':
                m = A[0] + wrap(B[0] - A[0]) / 2.0
                on, of = wrap(s[0] - m), wrap(s[0] - (m + 180.0))
                off, side = (on, 'near') if abs(on) <= abs(of) else (of, 'far')
            else:
                off, side = s[1] - (A[1] + B[1]) / 2.0, None
            if abs(off) >= tol:
                continue
            rh = 0.5 * (10 * (10 - (100 - clo)) / 10) ** 2
            dh = 0.5 * (10 * (tol - abs(off)) / tol) ** 2
            btw = 'between' if abs(dab - (d1 + d2)) <= 0.01 * (d1 + d2) else 'far side'
            out.append(dict(a=a, b=b, total=rh + dh, side=side, btw=btw, d=(round(d1, 6), round(d2, 6))))
    # Ketu repeats of a Rahu pair (same rounded distances)
    rahu = {(h['a'] if h['b'] == 'Rahu' else h['b'], tuple(sorted(h['d']))) for h in out if 'Rahu' in (h['a'], h['b'])}
    out = [h for h in out if not ('Ketu' in (h['a'], h['b']) and
                                  (h['a'] if h['b'] == 'Ketu' else h['b'], tuple(sorted(h['d']))) in rahu)]
    return sorted(out, key=lambda h: -h['total'])


def check_t24(ws, P, ch, title_prefix, tag):
    bodies = [b for b in BODIES if (ch, b) in P and not (ch != 'S' and b == 'Moon')]
    hdr = next(r[0].row for r in ws.iter_rows(max_col=2) if r[0].value == 'T24' and str(r[1].value or '').startswith(title_prefix))
    h1, h2 = hdr + 1, hdr + 2
    nb = len(bodies)
    ra0, ram, de0, dem = 3, 3 + nb, 4 + nb, 4 + 2 * nb
    cmp(tag + ' header', f'row {h1}', [ws.cell(h1, ra0).value, ws.cell(h1, ram).value, ws.cell(h1, de0).value, ws.cell(h1, dem).value],
        ['RA', 'RA midpoints', 'Dec', 'Dec midpoints'])
    cmp(tag + ' header bodies', f'row {h2}', [ws.cell(h2, ra0 + k).value for k in range(nb)] + [ws.cell(h2, de0 + k).value for k in range(nb)],
        bodies + bodies)
    merged = {str(m) for m in ws.merged_cells.ranges}
    from openpyxl.utils import get_column_letter as L
    for c0 in (ra0, de0):
        cmp(tag + ' merge', f'row {h1}', f'{L(c0)}{h1}:{L(c0 + nb - 1)}{h1}' in merged, True)
    names, r = [], h2 + 1
    while ws.cell(r, 1).value == 'T24' and ws.cell(r, 2).value:
        names.append(ws.cell(r, 2).value); r += 1
    cmp(tag + ' star rows', tag, names, [f'{ch} {s}' for s in STARS35])
    pts = {b: P[(ch, b)] for b in bodies}
    for i, star in enumerate(STARS35):
        r = h2 + 1 + i
        s = P[(ch, star)]
        isnew = star in NEW
        for plane, c0, cm in (('RA', ra0, ram), ('Dec', de0, dem)):
            mh = mid_hits(s, pts, bodies, plane)
            lines = {}
            for h in mh:
                lines.setdefault(h['a'], []).append(f"midpoint with {h['b']}: total {h['total']:.2f}")
                lines.setdefault(h['b'], []).append(f"midpoint with {h['a']}: total {h['total']:.2f}")
            for k, b in enumerate(bodies):
                c = ws.cell(r, c0 + k)
                v = dist(s, pts[b], plane)
                bump(tag + ' cells')
                cmp(tag + ' value', c.coordinate, c.value, round(v, 6))
                first, bh = comment_first(v, plane)
                want = '\n'.join([first] + lines.get(b, []))
                cmp(tag + ' comment', c.coordinate, c.comment.text if c.comment else None, want)
                red = v < F.CONJ[plane]                        # red = inside the conjunction tolerance (any best family)
                cmp(tag + ' fill', c.coordinate, fill_of(c), RED if red else (FILL[F.colour(bh['score'])] if bh else None))
                if red:
                    bump(tag + ' red cells' + (' NEW' if isnew else ''))
                    if bh['family'] != 'conj':
                        print(f'   note: red cell whose best family is not conjunction: {tag} {c.coordinate} {star} {b} {plane} {v:.6f} -> {first}')
                thick = c.border.left.style == 'thick' and c.border.top.style == 'thick'
                cmp(tag + ' outline', c.coordinate, thick, b in lines)
            c = ws.cell(r, cm)
            if isnew:
                bump(tag + ' NEW midpoint cells')
                bump(tag + ' NEW midpoint hits', len(mh))
            if plane == 'RA':
                txt = '; '.join(f"{h['a']} = {h['b']} {h['total']:.2f} ({h['side']}, {h['btw']})" for h in mh)
            else:
                txt = '; '.join(f"{h['a']} = {h['b']} {h['total']:.2f} ({h['btw']})" for h in mh)
            cmp(tag + ' midpoint text' + (' NEW' if isnew else ''), c.coordinate, c.value, txt or None)
            cmp(tag + ' midpoint fill', c.coordinate, fill_of(c), FILL[F.colour(mh[0]['total'])] if mh else None)
        bump(tag + ' rows')


def main():
    wb = openpyxl.load_workbook(sys.argv[1])
    P = load_pos(wb)
    for st in NEW:
        for ch in 'SHJ':
            assert (ch, st) in P, (ch, st)
    check_t12(wb, P, 'T1', 'H')
    check_t12(wb, P, 'T2', 'J')
    check_t24(wb['SKY JUPITER'], P, 'S', 'T24 SKY', 'T24 SKY')
    check_t24(wb['P01'], P, 'H', 'T24 HORSE', 'T24 HORSE')
    check_t24(wb['P01'], P, 'J', 'T24 JOCKEY', 'T24 JOCKEY')
    print('checks per kind:')
    for k in sorted(counts):
        print(f'   {k}: {counts[k]}')
    print('MISMATCHES:', len(mism))
    for m in mism[:60]:
        print('   ', m)


if __name__ == '__main__':
    main()
