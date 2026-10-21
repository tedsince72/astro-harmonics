"""check67_engine.py (independent checker, 21 Oct 2026): the 21 new stars in the engine (J2000.0 vs catalogue), the 67 in engine RA order,
Mirfak |Dec|, the rebuilt batch vs the 46-star batch, POSITIONS rows of the 21 = batch values. Own code.
usage: check67_engine.py <old batch.xlsx> <new batch.xlsx> <tables.xlsx>"""
import os, sys
import openpyxl
REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..', '..', '..'))
sys.path[:0] = [os.path.join(REPO, 'rebuild_kit'), os.path.dirname(__file__)]
import swisseph as swe
import celestial_bodies_swisseph as E
from check67_common import CAT, NEW21, STARS67

swe.set_ephe_path(os.path.expanduser('~/astronomy-project/ephe')); swe.set_topo(-0.6738, 51.4106, 0)
bad = []
B = dict(E.BEHENIAN_STARS)
print('engine names:', {x: B.get(x) for x in NEW21})
for x, (ra, dec) in CAT.items():
    r, nm, _ = swe.fixstar2_ut(B[x], 2451545.0, E._FLAG)
    ok = abs(r[0] - ra) < 0.02 and abs(r[1] - dec) < 0.02
    print(f'  {x:21s} {nm:30s} {r[0]:11.6f} {r[1]:+10.6f}  dRA {r[0] - ra:+.4f} dDec {r[1] - dec:+.4f} {"ok" if ok else "OFF"}')
    if not ok:
        bad.append(('catalogue', x, r[:2]))
era = {x: swe.fixstar2_ut(B[x], 2451545.0, E._FLAG)[0] for x in STARS67}
order = sorted(STARS67, key=lambda x: era[x][0])
print('67 in engine RA order = list:', order == STARS67, '| max |Dec|', max(abs(v[1]) for v in era.values()), '| Mirfak Dec', era['Mirfak'][1])
if order != STARS67:
    bad.append(('order', [(a, b) for a, b in zip(order, STARS67) if a != b]))
if abs(era['Mirfak'][1]) > 50:
    bad.append(('Mirfak', era['Mirfak'][1]))
for x in ('Bellatrix', 'El Nath', 'Meissa', 'Orion Nebula Cluster', 'Alnilam', 'Phact', 'Facies', 'Vega'):
    print(f'   {x:21s} RA {era[x][0]:.6f} Dec {era[x][1]:+.6f}')
# Polaris present in engine but not in the 67
print('Polaris in the 67:', 'Polaris' in STARS67)

old, new = openpyxl.load_workbook(sys.argv[1], read_only=True), openpyxl.load_workbook(sys.argv[2], read_only=True)
added_exp = {f'{x}_B' for x in NEW21}
batch = {}
print('sheets old', old.sheetnames, 'new', new.sheetnames)
for sh in new.sheetnames:
    ro = [tuple(r) for r in old[sh].iter_rows(values_only=True)] if sh in old.sheetnames else []
    rn = [tuple(r) for r in new[sh].iter_rows(values_only=True)]
    isnew = lambda r: bool(r) and any(v in added_exp for v in r[1:3])
    rest = [r for r in rn if not isnew(r)]
    added = sorted({v for r in rn if isnew(r) for v in r[1:3] if v in added_exp})
    nadd = sum(1 for r in rn if isnew(r))
    same = rest == ro
    nd = 0
    if not same:
        diffs = [(i, a, b) for i, (a, b) in enumerate(zip(ro, rest)) if a != b]
        nd = len(diffs) + abs(len(ro) - len(rest))
    print(f'batch {sh}: old rows {len(ro)}, new {len(rn)}, added names {len(added)} (rows {nadd}), rest identical {same} (differing {nd})')
    if not same:
        for i, a, b in diffs[:6]:
            print('     ', i, str(a)[:160], '|', str(b)[:160])
        if sh != 'META':
            bad.append(('batch', sh, nd))
    if sh in ('TRANS', 'P01', 'P02') and set(added) != added_exp:
        bad.append(('batch added', sh, sorted(added_exp - set(added))))
    batch[sh] = {r[1]: r for r in rn if r and len(r) > 1 and r[1]}
wb = openpyxl.load_workbook(sys.argv[3], read_only=True)
SHMAP = {'S': 'TRANS', 'H': 'P01', 'J': 'P02'}
n = 0
for r in wb['POSITIONS'].iter_rows(values_only=True):
    if r[2] in SHMAP and r[3] in NEW21:
        b = batch[SHMAP[r[2]]][f'{r[3]}_B']
        n += 1
        if abs(float(r[5]) - float(b[4])) > 1e-9 or abs(float(r[6]) - float(b[8])) > 1e-9:
            bad.append(('POSITIONS', r[2], r[3], r[5:7], (b[4], b[8])))
print('POSITIONS rows of the 21 checked against the batch:', n)
print('MISMATCHES:', len(bad))
for x in bad:
    print('  ', x)
