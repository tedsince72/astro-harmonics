#!/usr/bin/env python3
"""check67_sky_block.py - independent checker (21 Oct 2026, adapted from chk_stars46) for the 67-star SKY block in jupiter_tables.
Own code: positions read from POSITIONS, distances, +5, targets (families.evaluate / families.outside), text formats, timing, fills.
No import of jupiter_tables (constants in check67_common.py).
usage: check67_sky_block.py default_tables.xlsx full_tables.xlsx"""
import datetime as dt
import math
import os
import sys

import openpyxl

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))))
sys.path[:0] = [os.path.join(REPO, 'tools'), os.path.join(REPO, 'tools', 'jup2'), os.path.join(REPO, 'rebuild_kit')]
import families                                                     # noqa: E402
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from check67_common import STARS67 as T1_T2_T24_STARS, OLD_STARS as STARS, NEW21  # noqa: E402
NEW_STARS = [x for x in T1_T2_T24_STARS if x not in STARS]

OFF = dt.datetime(2012, 6, 19, 14, 30)
mism = []
counts = {}


def note(key, n=1):
    counts[key] = counts.get(key, 0) + n


def cmp(where, got, want, key):
    note(key)
    if isinstance(got, float) and isinstance(want, float):
        ok = abs(got - want) < 1e-9
    else:
        ok = got == want
    if not ok:
        mism.append((where, got, want))


# ---------- 1a. engine RA order ----------
def engine_order():
    import swisseph as swe
    import celestial_bodies_swisseph as cb
    swe.set_ephe_path(os.path.expanduser('~/astronomy-project/ephe'))
    swe.set_topo(-0.6738, 51.4106, 0.0)   # Ascot (topocentric flag needs a site; stars: no effect at this precision)
    out = []
    for name, sname in cb.BEHENIAN_STARS:
        if name in T1_T2_T24_STARS:
            out.append((swe.fixstar2_ut(sname, 2451545.0, cb._FLAG)[0][0], name))
    out.sort()
    return out


# ---------- positions ----------
def positions(wb):
    ws = wb['POSITIONS']
    pos = {}
    for r in ws.iter_rows(values_only=True):
        if r[0] in ('SKY', 'P01') and r[2] in ('S', 'H', 'J') and isinstance(r[5], (int, float)):
            pos[(r[2], r[3])] = (r[5], r[6], r[7] or 0.0, r[8] or 0.0)
    return pos


def ra_d(a, b):
    d = abs(a - b) % 360.0
    return 360.0 - d if d > 180 else d


def tolfmt(t):
    return f'±{t:.3g}'


def clock(t):
    when = OFF + dt.timedelta(minutes=t)
    when = when.replace(second=0, microsecond=0)
    if when.date() == OFF.date():
        return when.strftime('%H:%M')
    return when.strftime('%d %b %Y %H:%M')


def timing(v0, v5, target):
    rate = (v5 - v0) / 5.0
    if rate == 0:
        return None
    t = (target - v0) / rate
    if t >= 0:
        return f'applying, exact {clock(t)} (in {t:.1f} min)'
    return f'separating, exact {clock(t)} ({-t:.1f} min ago)'


def expected_cells(v0, v5, plane):
    """9 cells: value, +5, target text, group, offset, score, arrow, also, timing; + expected fill of the value cell"""
    ev = families.evaluate(v0, plane, 'deg')
    if ev:
        b = ev[0]
        tv = b['tv']
        arrow = '-> toward' if abs(v5 - tv) < abs(v0 - tv) else '<- away'
        also = '; '.join(f"{d['target']} ({d['family']}) {d['off']:+.5f} [{d['score']:.2f}]" for d in ev[1:]) or None
        s = b['score']
        fill = 'C6EFCE' if s >= 80 else 'FFEB9C' if s >= 50 else 'FFF8DC'
        return [round(v0, 6), round(v5, 6), f"{b['target']} ({b['family']})", None, f"{b['off']:+.5f}", round(s, 2), arrow, also,
                timing(v0, v5, tv)], fill, ev
    o = families.outside(v0, plane, 'deg')
    return [round(v0, 6), round(v5, 6), f"{o['target']} ({o['family']}) — outside {tolfmt(o['tol'])}", None, f"{o['off']:+.5f}",
            None, None, None, None], None, []


def sky_expect(pos, x):
    """-> {plane: (cells, fill, ev)} for S Jupiter -> S X"""
    J = pos[('S', 'Jupiter')]
    out = {}
    if x == 'equator':
        out['Dec'] = expected_cells(abs(J[1]), abs(J[1] + 5 * J[3]), 'Dec')
        return out
    X = pos[('S', x)]
    out['RA'] = expected_cells(ra_d(J[0], X[0]), ra_d(J[0] + 5 * J[2], X[0] + 5 * X[2]), 'RA')
    out['Dec'] = expected_cells(abs(J[1] - X[1]), abs(J[1] + 5 * J[3] - X[1] - 5 * X[3]), 'Dec')
    return out


def fill_rgb(c):
    if c.fill is None or c.fill.fill_type is None:
        return None
    return str(c.fill.start_color.rgb)[-6:]


def check_sky(wb, pos):
    ws = wb['SKY JUPITER']
    sky_rows, order = {}, []
    for row in ws.iter_rows():
        if row[0].value == 'SKY' and row[1].value and str(row[1].value).startswith('S Jupiter → S '):
            x = str(row[1].value)[len('S Jupiter → S '):]
            sky_rows[x] = row
            order.append(x)
    stars_in = [x for x in order if x in set(STARS) | set(NEW_STARS)]
    cmp('SKY block star order', stars_in, T1_T2_T24_STARS, 'order')
    cmp('SKY block Polaris absent', 'Polaris' in order, False, 'order')
    hand = {}
    for x in order:
        exp = sky_expect(pos, x)
        row = sky_rows[x]
        for plane, c0 in (('RA', 2), ('Dec', 11)):
            cells = [row[c0 + k] for k in range(9)]
            if plane not in exp:
                for k, c in enumerate(cells):
                    cmp(f'SKY {x} {plane} col {c.column_letter}', c.value, None, 'sky_cells')
                continue
            want, fill, ev = exp[plane]
            for k, c in enumerate(cells):
                got = c.value
                w = want[k]
                if k in (0, 1) and isinstance(got, (int, float)):
                    note('sky_cells')
                    if abs(got - w) > 1.5e-6:
                        mism.append((f'SKY {x} {plane} {c.coordinate}', got, w))
                    continue
                cmp(f'SKY {x} {plane} {c.coordinate}', got, w, 'sky_cells')
            cmp(f'SKY {x} {plane} {cells[0].coordinate} fill', fill_rgb(cells[0]), fill, 'sky_fills')
            # other cells: no fill expected except possibly the score cell - record any
            for c in cells[1:]:
                f = fill_rgb(c)
                if f and f not in ('000000', None):
                    note(f'other fill {plane} col {c.column - c0} {f}')
            hand[(x, plane)] = (want, ev)
    return sky_rows, order, hand


def check_t1_sky(wb, sky_rows):
    ws = wb['P01']
    for blk in ('T1', 'T2'):
        for row in ws.iter_rows():
            if row[0].value == blk and row[1].value and ' → ' in str(row[1].value) and not str(row[1].value).endswith('(natal); SKY cells S Jupiter → S X beside each row'):
                lab = str(row[1].value)
                x = lab.split(' → ')[1].split(' ', 1)[1]
                if x == 'bodies, stars, equator (natal); SKY cells S Jupiter':
                    continue
                if x not in sky_rows:
                    note(f'{blk} rows without SKY row: {x}')
                    continue
                s = sky_rows[x]
                for k in range(9):
                    a, b = row[18 + k], s[2 + k]
                    cmp(f'{blk} {x} RA SKY {a.coordinate}', a.value, b.value, 't1_sky_cells')
                    a, b = row[27 + k], s[11 + k]
                    cmp(f'{blk} {x} Dec SKY {a.coordinate}', a.value, b.value, 't1_sky_cells')
                cmp(f'{blk} {x} RA SKY fill', fill_rgb(row[18]), fill_rgb(s[2]), 't1_sky_cells')
                cmp(f'{blk} {x} Dec SKY fill', fill_rgb(row[27]), fill_rgb(s[11]), 't1_sky_cells')


def block_orders(wb):
    res = {}
    for sh in ('P01', 'SKY JUPITER'):
        ws = wb[sh]
        for row in ws.iter_rows(max_col=2, values_only=True):
            a, b = row
            if a in ('T1', 'T2') and b and ' → ' in str(b):
                x = str(b).split(' → ')[1].split(' ', 1)[1]
                if x in T1_T2_T24_STARS or x == 'Polaris':
                    res.setdefault(f'{sh} {a}', []).append(x)
            if a == 'T24' and b and not str(b).startswith('T24') and str(b)[:2] in ('S ', 'H ', 'J '):
                res.setdefault(f'{sh} T24 {str(b)[0]}', []).append(str(b)[2:])
    for k, v in res.items():
        cmp(f'order {k}', v, T1_T2_T24_STARS, 'order')
    return res


ALLSTARS = set(STARS) | set(NEW_STARS)


def main():
    tables, full = sys.argv[1:3]
    for path in (tables, full):
        wb = openpyxl.load_workbook(path)
        pos = positions(wb)
        sky_rows, order, hand = check_sky(wb, pos)
        print(path.split('/')[-1], 'SKY block rows:', len(order), 'stars:', len([x for x in order if x in ALLSTARS]),
              'new-star rows:', len([x for x in order if x in NEW21]))
        check_t1_sky(wb, sky_rows)
        bo = block_orders(wb)
        print('  block orders checked:', {k: len(v) for k, v in bo.items()})
        if path == tables:
            for key in (('Meissa', 'RA'), ('Al Rescha', 'RA'), ('Orion Nebula Cluster', 'RA'), ('El Nath', 'Dec')):
                want, ev = hand[key]
                print('  HAND', key, want, [(d['family'], d['target'], round(d['tv'], 6), round(d['score'], 4)) for d in ev])
            # SKY block rows of the 21 with a hit and timing
            nt = sum(1 for (x, pl), (want, ev) in hand.items() if x in NEW21 and want[8])
            print('  new-star SKY plane cells with timing:', nt)
    print('COUNTS', sorted(counts.items()))
    print('MISMATCHES', len(mism))
    for m in mism[:80]:
        print('  ', m)


if __name__ == '__main__':
    main()
