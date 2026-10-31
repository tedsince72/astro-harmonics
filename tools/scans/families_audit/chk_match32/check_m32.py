"""MATCH groups in blocks (FIXED STARS, CONJUNCTIONS, OTHER BODIES) and SUMMARY on unrounded scores (Eddie, 31 Oct 2026): checks one build
by the new code against the build before (3da4ae61), same race.
usage: check_m32.py <before.xlsx> <after.xlsx>
1. Sheets, POSITIONS, SKY JUPITER, every runner-sheet cell before the MATCH tables, widths and freeze identical.
2. Each MATCH table (MATCH HORSE, MATCH JOCKEY, MATCH HORSE | JOCKEY): the after table minus its band rows = the before table with its groups
   (grey header + rows) moved only - every group the same cell for cell, nothing else changed but the title (column B).
3. Order: blocks FIXED STARS (star or equator sky cell, not conjunction), CONJUNCTIONS (star / equator conjunction), OTHER BODIES (body
   cells, not conjunction, then body conjunctions); inside each by the sky score unrounded, recomputed from POSITIONS (families via
   jupiter_tables.nearest), highest first, ties SKY block order (RA before Dec). A band row (column B the block name, A-last column darker grey
   A6A6A6) exactly before the first group of each block present, no other; title counts 'groups: FIXED STARS a, CONJUNCTIONS b, OTHER BODIES
   c (d of them body conjunctions, at the end)'.
4. SUMMARY rebuilt from MATCH HORSE | JOCKEY with the scores recomputed from POSITIONS (unrounded) - lines, group headers and counts, bands
   only for blocks with a shown group, cells and styles, merges, blank rows; lines the 2-dp filter would add or drop listed.
Prints per runner the groups per block (H, J, H | J) and 'mismatches: n' last."""
import sys
import openpyxl
sys.path.insert(0, 'tools')
import jupiter_tables as JT

SEC = ('EXACT', 'RATIO', 'BODY', 'CONJ STAR')
BANDS = ('FIXED STARS', 'CONJUNCTIONS', 'OTHER BODIES')
bad = []


def sig(c):
    return (c.value, repr(c.font), repr(c.fill), repr(c.border), c.number_format, repr(c.alignment), c.comment.text if c.comment else None)


def last_row(ws):
    m = ws.max_row
    while m > 0 and all(c.value is None for c in ws[m]):
        m -= 1
    return m


B, A = openpyxl.load_workbook(sys.argv[1]), openpyxl.load_workbook(sys.argv[2])
if A.sheetnames != B.sheetnames:
    bad.append('sheets')
for sh in ('POSITIONS', 'SKY JUPITER'):
    for r in range(1, max(B[sh].max_row, A[sh].max_row) + 1):
        for k in range(1, max(B[sh].max_column, A[sh].max_column) + 1):
            if sig(B[sh].cell(row=r, column=k)) != sig(A[sh].cell(row=r, column=k)):
                bad.append(f'{sh} {r},{k}')
POS = {}
for x in A['POSITIONS'].iter_rows(min_row=5, values_only=True):
    if x and x[0] and x[3] and isinstance(x[5], (int, float)):
        POS[(x[0], x[2], x[3])] = (x[5], x[6])
SKYORD = {}
for r in A['SKY JUPITER'].iter_rows(values_only=True):
    if r and r[0] == 'SKY' and str(r[1]).startswith('S Jupiter → '):
        for pl in ('RA', 'Dec'):
            SKYORD[(r[1], pl)] = len(SKYORD)


def dist(p, q, pl):
    if pl == 'RA':
        x = abs(p[0] - q[0]) % 360
        return min(x, 360 - x)
    return abs(p[1] - q[1])


def score(v, pl):
    ev = JT.nearest(v, 'dist', pl)[0]
    return ev['score'] if ev else None


def sky_score(lab, pl):
    x = lab[len('S Jupiter → S '):]
    sj = POS[('SKY', 'S', 'Jupiter')]
    if x == 'equator':
        return score(abs(sj[1]), 'Dec')
    return score(dist(sj, POS[('SKY', 'S', x)], pl), pl)


def t24_score(P, lab, body, pl):
    C = lab[0] if not lab.startswith('S ') else lab.split(' → ')[1][0]
    if lab.startswith(f'{C} Jupiter → '):
        a, b = POS[(P, C, 'Jupiter')], POS[(P, C, body)]
    elif lab.startswith('S Jupiter → '):
        a, b = POS[('SKY', 'S', 'Jupiter')], POS[(P, C, body)]
    elif lab.startswith('S body → '):
        a, b = POS[('SKY', 'S', body)], POS[(P, C, 'Jupiter')]
    else:
        a, b = POS[(P, C, lab.split(' ', 1)[1])], POS[(P, C, body)]
    return score(dist(a, b, pl), pl)


def block(lab, fam):
    x = lab[len('S Jupiter → S '):]
    star = x == 'equator' or (x not in JT.SKY_BODIES and x not in JT.T24_BODIES)
    return (0 if star else 2) + (fam == 'conj')


def table(ws, tid):
    return [r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value == tid]


def parse(ws, rows, n):
    """-> (title row, header rows before the first group, [band rows], [(header row, [data rows])])"""
    groups, bands, pre = [], [], []
    for r in rows[1:]:
        b = ws.cell(row=r, column=2).value
        if b in SEC:
            groups[-1][1].append(r)
        elif b in BANDS:
            bands.append(r)
        elif isinstance(b, str) and ' — EXACT ' in b:
            groups.append((r, []))
        else:
            pre.append(r)
    return rows[0], pre, bands, groups


summ = {}
for P in [s for s in B.sheetnames if s.startswith('P') and s[1:].isdigit()]:
    b, a = B[P], A[P]
    if a.freeze_panes != b.freeze_panes:
        bad.append(f'{P} freeze')
    for L, d in b.column_dimensions.items():
        if a.column_dimensions[L].width != d.width:
            bad.append(f'{P} width {L}')
    s0 = table(b, 'MATCH HORSE')[0]
    if s0 != table(a, 'MATCH HORSE')[0]:
        bad.append(f'{P} MATCH start')
    for r in range(1, s0):
        for k in range(1, b.max_column + 1):
            if sig(b.cell(row=r, column=k)) != sig(a.cell(row=r, column=k)):
                bad.append(f'{P} {r},{k}')
    line = []
    for tid, n in (('MATCH HORSE', 21), ('MATCH JOCKEY', 21), ('MATCH HORSE | JOCKEY', 29)):
        tb, pb, bb, gb = parse(b, table(b, tid), n)
        ta, pa, ba, ga = parse(a, table(a, tid), n)
        if bb:
            bad.append(f'{P} {tid} band rows before?')
        for x, y in zip(pb, pa):                           # band row of MATCH | (HORSE / JOCKEY), column names
            if [sig(b.cell(row=x, column=k)) for k in range(1, n + 1)] != [sig(a.cell(row=y, column=k)) for k in range(1, n + 1)]:
                bad.append(f'{P} {tid} header row {x}')
        if len(pb) != len(pa):
            bad.append(f'{P} {tid} header rows')
        if b.cell(row=tb, column=1).value != a.cell(row=ta, column=1).value:
            bad.append(f'{P} {tid} title')
        sigs = lambda ws, g: tuple(tuple(sig(ws.cell(row=r, column=k)) for k in range(1, n + 1)) for r in [g[0]] + g[1])
        before = sorted(sigs(b, g) for g in gb)
        after = sorted(sigs(a, g) for g in ga)
        if before != after:
            bad.append(f'{P} {tid} groups differ ({len(gb)} / {len(ga)})')
        # order and bands
        keys = []
        for hr, rs in ga:
            v = [a.cell(row=rs[0], column=k).value for k in (8, 9, 11)]
            blk = block(v[0], v[2])
            keys.append((blk, -sky_score(v[0], v[1]), SKYORD[(v[0], v[1])], hr))
        if keys != sorted(keys):
            bad.append(f'{P} {tid} group order')
        want = []
        for i, k in enumerate(keys):
            if i == 0 or min(k[0], 2) != min(keys[i - 1][0], 2):
                want.append((min(k[0], 2), k[3] - 1))
        got = [(BANDS.index(a.cell(row=r, column=2).value), r) for r in ba]
        if got != want:
            bad.append(f'{P} {tid} bands {got} want {want}')
        for r in ba:
            if a.cell(row=r, column=1).value != tid or any(a.cell(row=r, column=k).fill.start_color.rgb != '00A6A6A6' for k in range(1, n + 1)) \
                    or any(a.cell(row=r, column=k).value is not None for k in range(3, n + 1)) or not a.cell(row=r, column=2).font.b:
                bad.append(f'{P} {tid} band row {r} format')
        c = [sum(1 for k in keys if k[0] == z) for z in range(4)]
        txt = f'groups: FIXED STARS {c[0]}, CONJUNCTIONS {c[1]}, OTHER BODIES {c[2] + c[3]} ({c[3]} of them body conjunctions, at the end)'
        if txt not in str(a.cell(row=ta, column=2).value) or 'ORDER (31 Oct 2026)' not in str(a.cell(row=ta, column=2).value):
            bad.append(f'{P} {tid} title counts / ORDER')
        if (ta, pa[-1] if pa else None) and (ga and ga[0][0] - 1 not in ba):
            bad.append(f'{P} {tid} first group not under a band')
        line.append(f"{tid.replace('MATCH ', '')} {c[0]}/{c[1]}/{c[2]}+{c[3]}")
        if tid == 'MATCH HORSE | JOCKEY':
            hj = (ta, pa, ba, ga)
    if last_row(a) != table(a, 'MATCH HORSE | JOCKEY')[-1]:
        bad.append(f'{P} rows after the tables')
    # SUMMARY expected (unrounded scores from POSITIONS)
    ta, pa, ba, ga = hj
    big = lambda x: x is not None and x > 50
    items = sorted([('band', r) for r in ba] + [('group', g) for g in ga], key=lambda it: it[1] if it[0] == 'band' else it[1][0])
    exp, edge = [('a1', 1)], []
    shown = {}
    for hr, rs in ga:
        sel = []
        for r in rs:
            v = [a.cell(row=r, column=k).value for k in range(1, 30)]
            sk = sky_score(v[7], v[8])
            h = t24_score(P, v[13], v[14], v[15]) if v[13] else None
            j = t24_score(P, v[21], v[22], v[23]) if v[21] else None
            for x, cell in ((sk, v[12]), (h, v[19]), (j, v[27])):
                if (x is None) != (cell is None) or (x is not None and round(x, 2) != cell):
                    bad.append(f'{P} recomputed score {x} vs cell {cell} row {r}')
            ok = big(sk) and (big(h) or big(j))
            ok2 = big(v[12]) and (big(v[19]) or big(v[27]))
            if ok != ok2:
                edge.append((P, r, ok))
            if ok:
                sel.append(r)
        shown[hr] = sel
    if any(shown.values()):
        exp += [('row', pa[0]), ('row', pa[0] + 1)]
        for i, it in enumerate(items):
            if it[0] == 'band':
                nxt = next((k for k in range(i + 1, len(items)) if items[k][0] == 'band'), len(items))
                if any(shown[items[k][1][0]] for k in range(i + 1, nxt)):
                    exp.append(('row', it[1]))
                continue
            hr = it[1][0]
            sel = shown[hr]
            if not sel:
                continue
            cnt = [(sum(1 for r in sel if a.cell(row=r, column=2).value == s and a.cell(row=r, column=14).value is not None),
                    sum(1 for r in sel if a.cell(row=r, column=2).value == s and a.cell(row=r, column=22).value is not None)) for s in SEC]
            exp.append(('head', hr, a.cell(row=hr, column=2).value.split(' — EXACT ')[0] + ' — ' + ', '.join(f'{s} {h} | {j}' for s, (h, j) in zip(SEC, cnt))))
            exp += [('row', r) for r in sel]
    else:
        exp.append(('none',))
    summ[P] = exp
    print(P, 'groups per block (FIXED STARS / CONJUNCTIONS / OTHER BODIES + body conjunctions):', '; '.join(line),
          f"| summary {sum(1 for e in exp if e[0] == 'row' and a.cell(row=e[1], column=2).value in SEC)}", f'| 2-dp edge lines {edge}')

S = A['SUMMARY']
r = 1
merges = set()
for P, exp in summ.items():
    a = A[P]
    for e in exp:
        if e[0] == 'a1':
            if sig(S.cell(row=r, column=1)) != sig(a.cell(row=1, column=1)):
                bad.append(f'SUMMARY {r} A1 of {P}')
        elif e[0] == 'none':
            if S.cell(row=r, column=1).value != 'no lines':
                bad.append(f'SUMMARY {r} no lines')
        else:
            for k in range(1, 30):
                s, t = a.cell(row=e[1], column=k), S.cell(row=r, column=k)
                if e[0] == 'head' and k == 2:
                    if t.value != e[2] or (repr(s.font), repr(s.fill)) != (repr(t.font), repr(t.fill)):
                        bad.append(f'SUMMARY {r} head {t.value!r} want {e[2]!r}')
                elif sig(s) != sig(t):
                    bad.append(f'SUMMARY {r},{k} ({P} row {e[1]}): {t.value!r} want {s.value!r}')
            if a.cell(row=e[1], column=14).value == 'HORSE':
                merges |= {f'N{r}:U{r}', f'V{r}:AC{r}'}
        r += 1
    for k in range(2):
        if any(S.cell(row=r + k, column=c).value is not None for c in range(1, 30)):
            bad.append(f'SUMMARY {r + k} not blank')
    r += 2
if last_row(S) != r - 3:
    bad.append(f'SUMMARY last row {last_row(S)} want {r - 3}')
if {str(m) for m in S.merged_cells.ranges} != merges:
    bad.append('SUMMARY merges')
for x in bad[:40]:
    print('  ', x)
print('mismatches:', len(bad))
