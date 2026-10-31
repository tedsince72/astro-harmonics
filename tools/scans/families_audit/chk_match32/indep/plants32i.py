"""Planted errors for check32i.py. Writes copies to <outdir>, never touches the c31 / c32 builds.
Usage: python3 plants32i.py <after_dir> <outdir>
"""
import sys, os, copy, math
import openpyxl
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import check32i as C

A, OUT = sys.argv[1], sys.argv[2]
os.makedirs(OUT, exist_ok=True)


def grab(ws, r, nc=30):
    return [(ws.cell(r, k).value, copy.copy(ws.cell(r, k)._style),
             copy.copy(ws.cell(r, k).comment)) for k in range(1, nc + 1)]


def put(ws, r, cells):
    for k, (v, st, cm) in enumerate(cells, 1):
        c = ws.cell(r, k)
        c.value = v
        c._style = copy.copy(st)
        c.comment = copy.copy(cm) if cm is not None else None


def table(ws, tid):
    t = C.find_titles(ws)[tid]
    T = C.parse_table(ws, tid, t)
    return T


def body_rows(T):
    return [r for _, r in T['items']]


def groups(ws, T, tid):
    g, b = C.groups_of(ws, T, C.NCOL[tid])
    return g, b


def rewrite(ws, rows, new_order_rows):
    snap = {r: grab(ws, r) for r in rows}
    for r, src in zip(rows, new_order_rows):
        put(ws, r, snap[src])


def move_group(ws, tid, key, before_index):
    """move group with sky key to before the item at position before_index (in the items list); returns nothing"""
    T = table(ws, tid)
    rows = body_rows(T)
    g, _ = groups(ws, T, tid)
    gg = [x for x in g if x['key'] == key][0]
    seg = [gg['hrow']] + gg['rows']
    rest = [r for r in rows if r not in seg]
    tgt = rows[before_index] if before_index < len(rows) else None
    i = rest.index(tgt) if tgt in rest else len(rest)
    rewrite(ws, rows, rest[:i] + seg + rest[i:])


def item_index(ws, tid, pred):
    T = table(ws, tid)
    for i, (kind, r) in enumerate(T['items']):
        if pred(kind, r):
            return i
    raise KeyError


def save(wb, name):
    p = os.path.join(OUT, name)
    wb.save(p)
    print('wrote', p)


def load(race):
    return openpyxl.load_workbook(os.path.join(A, race + '_jupiter_tables.xlsx'))


plant = sys.argv[3] if len(sys.argv) > 3 else 'all'


def want(n):
    return plant in ('all', n)


# 0 control: round trip, no change
if want('p00'):
    wb = load('goodwood'); save(wb, 'p00_control_goodwood.xlsx')

# 1 equator group into OTHER BODIES (goodwood P01 MATCH HORSE): move it just after the OTHER BODIES band
if want('p01'):
    wb = load('goodwood'); ws = wb['P01']
    i = item_index(ws, 'MATCH HORSE', lambda k, r: k == 'band' and ws.cell(r, 2).value == 'OTHER BODIES')
    move_group(ws, 'MATCH HORSE', ('S Jupiter → S equator', 'Dec'), i + 1)
    # keep bands where they are: FIXED STARS band is now above the Capella group (fine)
    save(wb, 'p01_equator_in_other_goodwood.xlsx')

# 2 a star group in OTHER BODIES (goodwood P03 MATCH JOCKEY): last FIXED STARS group moved to the end
if want('p02'):
    wb = load('goodwood'); ws = wb['P03']
    T = table(ws, 'MATCH JOCKEY'); g, b = groups(ws, T, 'MATCH JOCKEY')
    fixed = [x for x in g if x['key'][0].split('→ S ')[1] not in C.BODIES]
    move_group(ws, 'MATCH JOCKEY', fixed[-1]['key'], len(T['items']))
    save(wb, 'p02_star_in_other_goodwood.xlsx')

# 3 body conjunction before a non-conjunction body group (ripon P01 MATCH HORSE, S Juno Dec conj)
if want('p03'):
    wb = load('ripon'); ws = wb['P01']
    i = item_index(ws, 'MATCH HORSE', lambda k, r: k == 'band' and ws.cell(r, 2).value == 'OTHER BODIES')
    move_group(ws, 'MATCH HORSE', ('S Jupiter → S Juno', 'Dec'), i + 1)
    save(wb, 'p03_bodyconj_first_ripon.xlsx')

# 4 missing band (goodwood P05 HJ: OTHER BODIES band row deleted, rows below shifted up)
if want('p04'):
    wb = load('goodwood'); ws = wb['P05']
    T = table(ws, 'MATCH HORSE | JOCKEY')
    r = [r for k, r in T['items'] if k == 'band' and ws.cell(r, 2).value == 'OTHER BODIES'][0]
    rows = list(range(r, ws.max_row + 1))
    rewrite(ws, rows, rows[1:] + [rows[0]])
    put(ws, ws.max_row, [(None, copy.copy(ws.cell(1, 200)._style), None)] * 30)
    save(wb, 'p04_missing_band_goodwood.xlsx')

# 5 wrong order inside FIXED STARS (goodwood P07 MATCH HORSE: swap the first two groups)
if want('p05'):
    wb = load('goodwood'); ws = wb['P07']
    T = table(ws, 'MATCH HORSE'); g, b = groups(ws, T, 'MATCH HORSE')
    move_group(ws, 'MATCH HORSE', g[1]['key'], [i for i, (k, r) in enumerate(T['items']) if r == g[0]['hrow']][0])
    save(wb, 'p05_order_in_block_goodwood.xlsx')

# 6 band fill short (A..U on HJ) and band text wrong (newcastle P01 HJ CONJUNCTIONS -> 'CONJUNCTION')
if want('p06'):
    wb = load('newcastle'); ws = wb['P01']
    T = table(ws, 'MATCH HORSE | JOCKEY')
    bands = [r for k, r in T['items'] if k == 'band']
    r = bands[0]
    for k in range(22, 30):
        ws.cell(r, k).fill = openpyxl.styles.PatternFill()
    rc = [r for r in bands if ws.cell(r, 2).value == 'CONJUNCTIONS']
    if rc:
        ws.cell(rc[0], 2).value = 'CONJUNCTION'
    save(wb, 'p06_band_fill_text_newcastle.xlsx')

# 7 title group count wrong (lingfield P03 MATCH JOCKEY)
if want('p07'):
    wb = load('lingfield'); ws = wb['P03']
    t = C.find_titles(ws)['MATCH JOCKEY']
    s = ws.cell(t, 2).value
    import re
    s2 = re.sub(r'groups: FIXED STARS (\d+)', lambda m: 'groups: FIXED STARS %d' % (int(m.group(1)) + 1), s)
    ws.cell(t, 2).value = s2
    save(wb, 'p07_title_count_lingfield.xlsx')

# 8 conjunction star group in FIXED STARS (newcastle: Sadalmelik conj group moved to the top of its MATCH JOCKEY)
if want('p08'):
    wb = load('newcastle'); ws = wb['P01']
    T = table(ws, 'MATCH JOCKEY'); g, b = groups(ws, T, 'MATCH JOCKEY')
    key = ('S Jupiter → S Sadalmelik', 'RA')
    move_group(ws, 'MATCH JOCKEY', key, 1)
    save(wb, 'p08_conjstar_in_fixed_newcastle.xlsx')


# --- SUMMARY rounding edge plants: change POSITIONS so a line's T24 whole score is 49.996 / 50.004 (shown 50.00)
def edge(race, target_keep):
    wb = load(race)
    P = C.load_positions(wb)
    order, sky = C.load_sky(wb, P)
    S = wb['SUMMARY']
    shown = set()
    for r in range(1, S.max_row + 1):
        if S.cell(r, 1).value == 'MATCH HORSE | JOCKEY' and S.cell(r, 2).value in ('EXACT', 'RATIO'):
            shown.add(tuple(S.cell(r, k).value for k in range(1, 30)))
    runners = [s for s in wb.sheetnames if s not in ('POSITIONS', 'SKY JUPITER', 'SUMMARY')]
    for sn in runners:
        ws = wb[sn]
        T = table(ws, 'MATCH HORSE | JOCKEY'); g, b = groups(ws, T, 'MATCH HORSE | JOCKEY')
        for gg in g:
            if sky[gg['key']]['score'] <= 50.5:
                continue
            for r in gg['rows']:
                v = tuple(ws.cell(r, k).value for k in range(1, 30))
                lab = v[13]
                if lab is None or '→' in lab or v[17] != 'whole':
                    continue
                jsc = C.fnum(v[27]) if v[21] is not None else None
                if jsc is not None and jsc > 49:
                    continue
                hsc = C.fnum(v[19])
                if target_keep and v in shown:
                    continue
                if not target_keep and v not in shown:
                    continue
                # horse star row: move the H star so d -> t + dev with score 50.004 (keep) or 49.996 (drop)
                star, body, plane = lab.split(' ', 1)[1], v[14], v[15]
                tval = float(v[18])
                d_now = C.t24_dist(P, sn, lab, body, plane)
                sc = 50.004 if target_keep else 49.996
                dev = 0.02 * (1 - math.sqrt(sc) / 10)
                newd = tval + (dev if d_now >= tval else -dev)
                # find the star row in POSITIONS
                Pw = wb['POSITIONS']
                for pr in range(5, Pw.max_row + 1):
                    if Pw.cell(pr, 1).value == sn and Pw.cell(pr, 3).value == 'H' and Pw.cell(pr, 4).value == star:
                        break
                col = 6 if plane == 'RA' else 7
                sp = list(P[(sn, 'H', star)]); bp = P[(sn, 'H', body)]
                idx = 0 if plane == 'RA' else 1
                # shift star coordinate by (newd - d_now) in the direction that grows distance
                sgn = 1 if (sp[idx] - bp[idx]) % 360 < 180 or plane == 'Dec' and sp[idx] >= bp[idx] else -1
                if plane == 'Dec':
                    sgn = 1 if sp[idx] >= bp[idx] else -1
                Pw.cell(pr, col).value = sp[idx] + sgn * (newd - d_now)
                P2 = C.load_positions(wb)
                d2 = C.t24_dist(P2, sn, lab, body, plane)
                txt = '%.6f' % d2
                s2 = (10 * (0.02 - abs(d2 - tval)) / 0.02) ** 2
                print(race, sn, 'plant edge', 'keep' if target_keep else 'drop', lab, body, plane, 'd', d2, 'score', s2, 'was', hsc)
                # update every copy of this cell value / score (MATCH tables, SUMMARY): value col Q (17), score T (20), lower U (21)
                for wsx in [ws, S]:
                    for rr in range(1, wsx.max_row + 1):
                        if wsx.cell(rr, 14).value == lab and wsx.cell(rr, 15).value == body and wsx.cell(rr, 16).value == plane:
                            wsx.cell(rr, 17).value = txt
                            wsx.cell(rr, 20).value = '50'
                            low = min(50.0, C.fnum(wsx.cell(rr, 13).value))
                            wsx.cell(rr, 21).value = ('%g' % round(low, 2))
                return wb
    return None


if want('p09'):
    wb = edge('kelso', True)
    if wb: save(wb, 'p09_edge_keep_kelso.xlsx')
if want('p10'):
    wb = edge('kelso', False)
    if wb: save(wb, 'p10_edge_drop_kelso.xlsx')

# 11 SUMMARY group header count wrong (ripon)
if want('p11'):
    wb = load('ripon'); S = wb['SUMMARY']
    for r in range(1, S.max_row + 1):
        c = S.cell(r, 2)
        if C.fillrgb(S.cell(r, 1)) == C.GREY and isinstance(c.value, str) and 'EXACT ' in c.value:
            c.value = c.value.replace('EXACT ', 'EXACT 9', 1)
            break
    save(wb, 'p11_summary_header_ripon.xlsx')

# 12 SUMMARY band missing (ripon: first band row in SUMMARY blanked to a plain row of grey header? -> delete its text/fill)
if want('p12'):
    wb = load('ripon'); S = wb['SUMMARY']
    for r in range(1, S.max_row + 1):
        if C.fillrgb(S.cell(r, 1)) == C.BAND_FILL and r > 10:
            for k in range(1, 31):
                S.cell(r, k).fill = openpyxl.styles.PatternFill()
            S.cell(r, 2).value = None
            break
    save(wb, 'p12_summary_band_missing_ripon.xlsx')

# 13 SUMMARY line with sky score <= 50 inserted? -> instead: SUMMARY groups order swapped (two groups of P01)
if want('p13'):
    wb = load('kelso'); S = wb['SUMMARY']
    hdr = [r for r in range(1, S.max_row + 1) if C.fillrgb(S.cell(r, 1)) == C.GREY]
    a, b = hdr[0], hdr[1]
    c = b + 1
    while C.fillrgb(S.cell(c, 1)) not in (C.GREY, C.BAND_FILL) and S.cell(c, 1).value is not None:
        c += 1
    rows = list(range(a, c))
    rewrite(S, rows, list(range(b, c)) + list(range(a, b)))
    save(wb, 'p13_summary_group_order_kelso.xlsx')
