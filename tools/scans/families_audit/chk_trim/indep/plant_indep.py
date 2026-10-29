#!/usr/bin/env python3
"""Plant errors in copies of the AFTER default workbook and confirm check_trim_indep.py finds each.

Usage: plant_indep.py <old_def> <old_full> <new_def> <outdir>
"""
import copy
import os
import subprocess
import sys

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill, Border, Side

HERE = os.path.dirname(os.path.abspath(__file__))
old_def, old_full, new_def, outdir = sys.argv[1:5]
os.makedirs(outdir, exist_ok=True)


def find(ws, a, b):
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == a and ws.cell(r, 2).value == b:
            return r
    raise KeyError((ws.title, a, b))


def t24_rows(ws, which):
    """first row of the five rows of T24 HORSE / JOCKEY."""
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 2).value
        if ws.cell(r, 1).value == 'T24' and isinstance(v, str) and v.startswith(which):
            return r, r + 3
    raise KeyError(which)


def p1(wb):  # five-row value
    ws = wb['P01']; _, r = t24_rows(ws, 'T24 HORSE')
    ws.cell(r + 2, 3).value = ws.cell(r + 2, 3).value + 0.00001
    return f'P01 C{r + 2} value +0.00001'


def p2(wb):  # drop a fill on a five-row cell
    ws = wb['P03']; _, r = t24_rows(ws, 'T24 JOCKEY')
    for rr in range(r, r + 4):
        for c in range(3, 55):
            if ws.cell(rr, c).fill.fill_type:
                ws.cell(rr, c).fill = PatternFill(fill_type=None)
                return f'P03 {ws.cell(rr, c).coordinate} fill dropped'


def p3(wb):  # SKY JUPITER cell
    ws = wb['SKY JUPITER']
    for r in range(10, 200):
        if isinstance(ws.cell(r, 3).value, float):
            ws.cell(r, 3).value += 1e-6
            return f'SKY JUPITER C{r} +1e-6'


def p4(wb):  # width
    wb['P05'].column_dimensions['D'].width = 12
    return 'P05 width D 11 -> 12'


def p5(wb):  # stray row
    ws = wb['P07']; r, _ = t24_rows(ws, 'T24 HORSE')
    ws.cell(r - 1, 5).value = 'x'
    return f'P07 E{r - 1} stray value'


def p6(wb):  # Jupiter T9 cell
    ws = wb['P09']; _, r = t24_rows(ws, 'T24 HORSE')
    ws.cell(r + 3, 7).value = 167.0
    return f'P09 G{r + 3} Jupiter cell changed'


def p7(wb):  # grey off a stars cell
    ws = wb['P11']; _, r = t24_rows(ws, 'T24 JOCKEY')
    ws.cell(r + 4, 20).fill = PatternFill(fill_type=None)
    return f'P11 T{r + 4} stars-row grey removed'


def p8(wb):  # value in S Jupiter column on a Jupiter row
    ws = wb['P01']; _, r = t24_rows(ws, 'T24 JOCKEY')
    ws.cell(r, 27).value = 1.234567
    return f'P01 AA{r} value in S Jupiter column'


def p9(wb):  # comment on a T1 cell
    ws = wb['P01']
    c = ws.cell(10, 3)
    c.comment = Comment('planted', 'x')
    return 'P01 C10 comment added'


def p10(wb):  # fill on a T24 star row
    ws = wb['P03']; _, r = t24_rows(ws, 'T24 JOCKEY')
    ws.cell(r + 10, 5).fill = PatternFill('solid', fgColor='FFC6EFCE')
    return f'P03 E{r + 10} star-row fill'


def p11(wb):  # unmerge a header
    ws = wb['P05']; t, _ = t24_rows(ws, 'T24 HORSE')
    ws.unmerge_cells(f'C{t + 1}:AA{t + 1}')
    return f'P05 unmerge C{t + 1}:AA{t + 1}'


def p12(wb):  # title count
    ws = wb['P07']; t, _ = t24_rows(ws, 'T24 JOCKEY')
    ws.cell(t, 2).value = ws.cell(t, 2).value.replace(' of 67 rows', ' of 72 rows')
    return f'P07 B{t} title N 67 -> 72'


def p13(wb):  # extra sheet
    wb.create_sheet('P01 RATIOS')
    return 'extra sheet P01 RATIOS'


def p14(wb):  # border on a T5 cell
    ws = wb['P09']
    r = next(r for r in range(1, ws.max_row) if ws.cell(r, 1).value == 'T5' and isinstance(ws.cell(r, 3).value, float))
    ws.cell(r, 3).border = Border(left=Side(style='thick'))
    return f'P09 C{r} T5 border'


def p15(wb):  # POSITIONS freeze
    wb['POSITIONS'].freeze_panes = 'B2'
    return 'POSITIONS freeze B2'


def p16(wb):  # five-row comment text
    ws = wb['P11']; _, r = t24_rows(ws, 'T24 HORSE')
    c = ws.cell(r + 1, 4)
    c.comment = Comment(c.comment.text + '\nrepeat: x', 'x')
    return f'P11 D{r + 1} comment + repeat line'


PLANTS = [p1, p2, p3, p4, p5, p6, p7, p8, p9, p10, p11, p12, p13, p14, p15, p16]


def run(path):
    p = subprocess.run([sys.executable, os.path.join(HERE, 'check_trim_indep.py'), old_def, old_full, path, path],
                       capture_output=True, text=True)
    lines = p.stdout.splitlines()
    mis = [l for l in lines if l.startswith('  [')]
    n = next(l for l in lines if l.startswith('mismatches'))
    return n, mis


# baseline: roundtrip unchanged
wb = openpyxl.load_workbook(new_def)
base = os.path.join(outdir, 'baseline.xlsx'); wb.save(base)
n, mis = run(base)
print('baseline (openpyxl roundtrip):', n)
caught = 0
for f in PLANTS:
    wb = openpyxl.load_workbook(new_def)
    what = f(wb)
    path = os.path.join(outdir, f'{f.__name__}.xlsx'); wb.save(path)
    n, mis = run(path)
    ok = n != 'mismatches: 0'
    caught += ok
    print(f'{f.__name__}: {what} -> {"CAUGHT" if ok else "MISSED"} ({n}); {mis[:2]}')
print(f'caught {caught} of {len(PLANTS)}')
