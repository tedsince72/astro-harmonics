"""Workbook trim (Eddie, 29 Oct 2026): checks one build by the trimmed code against the build before (acf68341), same race and flag.
usage: check_trim.py <before.xlsx> <after.xlsx>
(a) sheets: after = POSITIONS, SKY JUPITER, the runner sheets (in the before order); no other sheet; no removed table anywhere (column A of
    every runner-sheet row is the runner line, the note, 'T1' ... 'T8', 'T24', 'table' or empty);
(b) POSITIONS and SKY JUPITER identical cell for cell (value, fill, font, border, number format, alignment, comment), merged ranges,
    column widths, freeze panes;
(c) runner sheets: rows 1-3, T1 ... T8, T24 HORSE and T24 JOCKEY identical cell for cell to the before blocks (rows shifted), merged
    ranges, every row between blocks empty, the column widths of every column the after sheet uses, freeze panes; T24 HORSE / JOCKEY with
    the five new rows straight after the two header rows;
(d) the five rows: the four Jupiter rows = the before T24 HORSE / JOCKEY (SKY STARS) rows cell for cell from column B to the repeats column
    (column A 'T24'); the Jupiter cell of the last two = the before T9 'S Jupiter – <C> Jupiter' value (6 dp); row 5 'stars' (bold), grey
    D9D9D9 exactly from A to the repeats column, nothing else on it.
Prints counts and 'mismatches: n' last."""
import sys
import openpyxl

KEEP = ('T1', 'T2', 'T3', 'T4', 'T5', 'T6', 'T7', 'T8')
bad = []
n_cells = {'sky': 0, 'runner': 0, 'new': 0}


def err(*a):
    bad.append(' '.join(str(x) for x in a))


def sig(c):
    b = c.border
    return (c.value, c.fill.fill_type, c.fill.start_color.rgb if c.fill.fill_type else None, c.font.b, c.font.color.rgb if c.font.color else None,
            c.number_format, (b.left.style, b.right.style, b.top.style, b.bottom.style), c.comment.text if c.comment else None,
            c.alignment.wrap_text, c.alignment.vertical)


def empty(ws, r, ncol):
    return all(ws.cell(row=r, column=k).value is None and not ws.cell(row=r, column=k).has_style for k in range(1, ncol + 1))


def width(ws, L):
    d = ws.column_dimensions.get(L)
    return (d.width, bool(d.customWidth)) if d is not None and d.customWidth else (None, False)


B = openpyxl.load_workbook(sys.argv[1])
A = openpyxl.load_workbook(sys.argv[2])
AMAX = {s: (A[s].max_row, A[s].max_column) for s in A.sheetnames}    # before any read (reading a cell creates it)
runners = [s for s in B.sheetnames if len(s) == 3 and s[0] == 'P' and s[1:].isdigit()]
if A.sheetnames != ['POSITIONS', 'SKY JUPITER'] + runners:
    err('sheets', A.sheetnames)
print('removed sheets:', [s for s in B.sheetnames if s not in A.sheetnames])

# (b) POSITIONS, SKY JUPITER
for sh in ('POSITIONS', 'SKY JUPITER'):
    b, a = B[sh], A[sh]
    nr, nc = max(a.max_row, b.max_row), max(a.max_column, b.max_column)
    for r in range(1, nr + 1):
        for k in range(1, nc + 1):
            n_cells['sky'] += 1
            if sig(b.cell(row=r, column=k)) != sig(a.cell(row=r, column=k)):
                err(sh, 'cell', r, k)
    if {str(m) for m in a.merged_cells.ranges} != {str(m) for m in b.merged_cells.ranges}:
        err(sh, 'merged ranges')
    for k in range(1, nc + 1):
        L = openpyxl.utils.get_column_letter(k)
        if width(a, L) != width(b, L):
            err(sh, 'width', L, width(b, L), width(a, L))
    if a.freeze_panes != b.freeze_panes:
        err(sh, 'freeze', b.freeze_panes, a.freeze_panes)


def blocks(ws):
    """[(key, title row, last row)] of T1 ... T8 and T24 HORSE / JOCKEY; T24: title, two header rows, then every row with A 'T24'"""
    out, r = [], 1
    while r <= ws.max_row:
        v, t = ws.cell(row=r, column=1).value, str(ws.cell(row=r, column=2).value)
        if v in KEEP:
            e = r
            while ws.cell(row=e + 1, column=1).value in (v, 'table'):
                e += 1
            out.append((v, r, e)); r = e + 1; continue
        if v == 'T24' and t.startswith(('T24 HORSE (natal)', 'T24 JOCKEY (natal)')):
            e = r + 2
            while ws.cell(row=e + 1, column=1).value == 'T24':
                e += 1
            out.append(('T24 ' + t.split()[1], r, e)); r = e + 1; continue
        r += 1
    return out


# the before (SKY STARS) Jupiter rows and the before T9 values, per runner and chart
for P in runners:
    b, a = B[P], A[P]
    amr, amc = AMAX[P]
    nc = max(amc, b.max_column)
    bb, ab = blocks(b), blocks(a)
    if [k for k, _, _ in bb] != [k for k, _, _ in ab] or [k for k, _, _ in ab] != list(KEEP) + ['T24 HORSE', 'T24 JOCKEY']:
        err(P, 'blocks', [k for k, _, _ in bb], [k for k, _, _ in ab])
        continue
    # (a) nothing but kept blocks: column A values
    for r in range(4, amr + 1):
        v = a.cell(row=r, column=1).value
        if v not in KEEP + ('T24', 'table', None):
            err(P, 'unexpected column A', r, v)
    # (c) rows 1-3
    for r in range(1, 4):
        for k in range(1, nc + 1):
            if sig(b.cell(row=r, column=k)) != sig(a.cell(row=r, column=k)):
                err(P, 'top row', r, k)
    used = set(range(1, 4))
    mb = {(m.min_row, m.min_col, m.max_row, m.max_col) for m in b.merged_cells.ranges}
    ma = {(m.min_row, m.min_col, m.max_row, m.max_col) for m in a.merged_cells.ranges}
    want_m = set()
    for (key, b0, b1), (_, a0, a1) in zip(bb, ab):
        t24 = key.startswith('T24')
        rmap = {}                                          # before row -> after row
        for r in range(b0, b1 + 1):
            rmap[r] = a0 + (r - b0) + (5 if t24 and r >= b0 + 3 else 0)
        if a1 != rmap[b1]:
            err(P, key, 'block length', b0, b1, a0, a1)
        for r, ra in rmap.items():
            used.add(ra)
            for k in range(1, nc + 1):
                n_cells['runner'] += 1
                if sig(b.cell(row=r, column=k)) != sig(a.cell(row=ra, column=k)):
                    err(P, key, 'cell', r, '->', ra, k)
        for m in mb:
            if b0 <= m[0] <= b1:
                want_m.add((rmap[m[0]], m[1], rmap[m[2]], m[3]))
        if not t24:
            continue
        # (d) the five rows
        C = 'H' if key == 'T24 HORSE' else 'J'
        crep = next(k for k in range(1, a.max_column + 1) if a.cell(row=a0 + 1, column=k).value == 'repeats')
        ss = next(r for r in range(1, b.max_row + 1) if b.cell(row=r, column=1).value == f'{key} (SKY STARS)'
                  and str(b.cell(row=r, column=2).value).startswith(f'{key} (SKY STARS) (Eddie'))
        labels = ['S Jupiter → S body', f'{C} Jupiter → {C} body', f'S Jupiter → {C} body', f'S body → {C} Jupiter']
        for i, lab in enumerate(labels):
            ra, rb = a0 + 3 + i, ss + 3 + i
            used.add(ra)
            if b.cell(row=rb, column=2).value != lab or a.cell(row=ra, column=2).value != lab or a.cell(row=ra, column=1).value != 'T24':
                err(P, key, 'label', i, b.cell(row=rb, column=2).value, a.cell(row=ra, column=2).value, a.cell(row=ra, column=1).value)
            for k in range(2, crep + 1):
                n_cells['new'] += 1
                if sig(b.cell(row=rb, column=k)) != sig(a.cell(row=ra, column=k)):
                    err(P, key, lab, 'cell', k)
            for k in range(crep + 1, nc + 1):
                if a.cell(row=ra, column=k).value is not None or a.cell(row=ra, column=k).has_style:
                    err(P, key, lab, 'cell beyond the block', k)
        # Jupiter column of the last two rows = the before T9 value
        hdr = {(a.cell(row=a0 + 1, column=k).value if a.cell(row=a0 + 1, column=k).value in ('RA', 'Dec') else None): k for k in range(3, crep)}
        jc = {pl: next(k for k in range(hdr[pl], crep) if a.cell(row=a0 + 2, column=k).value == 'Jupiter') for pl in ('RA', 'Dec')}
        t9 = next(r for r in range(1, b.max_row + 1) if b.cell(row=r, column=1).value == 'T9' and b.cell(row=r, column=2).value == f'S Jupiter – {C} Jupiter')
        for pl, c9 in (('RA', 3), ('Dec', 11)):
            v9 = round(b.cell(row=t9, column=c9).value, 6)
            for i in (2, 3):
                got = a.cell(row=a0 + 3 + i, column=jc[pl]).value
                if got != v9:
                    err(P, key, 'T9', pl, i, got, v9)
            for i in (0, 1):
                if a.cell(row=a0 + 3 + i, column=jc[pl]).value is not None:
                    err(P, key, 'Jupiter cell not blank', pl, i)
        # 'stars' row
        rs = a0 + 7
        used.add(rs)
        if a.cell(row=rs, column=1).value != 'T24' or a.cell(row=rs, column=2).value != 'stars' or not a.cell(row=rs, column=2).font.b:
            err(P, key, 'stars row label')
        for k in range(1, nc + 1):
            c = a.cell(row=rs, column=k)
            grey = c.fill.fill_type == 'solid' and c.fill.start_color.rgb in ('00D9D9D9', 'FFD9D9D9')
            if grey != (k <= crep) or (k > 2 and c.value is not None):
                err(P, key, 'stars row cell', k)
    if ma != want_m:
        err(P, 'merged ranges', sorted(ma ^ want_m)[:5])
    for r in range(1, max(amr, b.max_row) + 1):
        if r not in used and not empty(a, r, nc):
            err(P, 'row outside the blocks not empty', r)
    for k in range(1, amc + 1):
        L = openpyxl.utils.get_column_letter(k)
        if width(a, L) != width(b, L):
            err(P, 'width', L, width(b, L), width(a, L))
    if a.freeze_panes != b.freeze_panes:
        err(P, 'freeze', b.freeze_panes, a.freeze_panes)
    if P == runners[0]:
        gone = sorted({L for L, d in b.column_dimensions.items() if d.customWidth and openpyxl.utils.column_index_from_string(L) > amc})
        print(f'{P}: columns past the last used one ({openpyxl.utils.get_column_letter(amc)}) whose width setting went with the removed tables:',
              len(gone), gone[:3], '...', gone[-2:])

print('cells compared: SKY JUPITER + POSITIONS', n_cells['sky'], '| runner blocks', n_cells['runner'], '| five-row cells', n_cells['new'])
for x in bad[:40]:
    print('  ', x)
print('mismatches:', len(bad))
