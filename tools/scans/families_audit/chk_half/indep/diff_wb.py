"""Cell-by-cell compare of two workbooks (values, fills, borders, number formats, comments, merged ranges, sheet names).
usage: diff_wb.py <before.xlsx> <after.xlsx> [out.json]  - prints diff kinds; writes every differing (sheet, coord, kind, before, after)."""
import sys, json, collections
import openpyxl
a = openpyxl.load_workbook(sys.argv[1]); b = openpyxl.load_workbook(sys.argv[2])
out = []
assert a.sheetnames == b.sheetnames, (a.sheetnames, b.sheetnames)
ncell = 0
def fh(c): return (c.fill.fill_type, c.fill.fgColor.rgb) if c.fill is not None else None
def bd(c): return tuple(getattr(c.border, s).style if getattr(c.border, s) else None for s in ('left', 'right', 'top', 'bottom'))
for s in a.sheetnames:
    wa, wb_ = a[s], b[s]
    if sorted(map(str, wa.merged_cells.ranges)) != sorted(map(str, wb_.merged_cells.ranges)):
        out.append([s, '-', 'merged', '', ''])
    mr, mc = max(wa.max_row, wb_.max_row), max(wa.max_column, wb_.max_column)
    for r in range(1, mr + 1):
        for c in range(1, mc + 1):
            x, y = wa.cell(r, c), wb_.cell(r, c)
            ncell += 1
            co = x.coordinate
            if x.value != y.value: out.append([s, co, 'value', x.value, y.value])
            if fh(x) != fh(y): out.append([s, co, 'fill', str(fh(x)), str(fh(y))])
            if bd(x) != bd(y): out.append([s, co, 'border', str(bd(x)), str(bd(y))])
            if x.number_format != y.number_format: out.append([s, co, 'numfmt', x.number_format, y.number_format])
            cx = x.comment.text if x.comment else None; cy = y.comment.text if y.comment else None
            if cx != cy: out.append([s, co, 'comment', cx, cy])
            if x.font.color != y.font.color or x.font.b != y.font.b: out.append([s, co, 'font', '', ''])
print(sys.argv[2], 'cells compared', ncell, 'differences', collections.Counter(o[2] for o in out))
if len(sys.argv) > 3: json.dump(out, open(sys.argv[3], 'w'), default=str)
