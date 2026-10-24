"""Independent check of the nearest-target wording (Eddie, 24 Oct 2026) over EVERY cell value and comment of a jupiter_tables
workbook: each text in one of the nearest-target forms is matched to its value (comment -> the cell; T-table text -> 2 cells left;
timeline sky-equivalent / FIELD text -> 1 cell left), the expected form recomputed (wfam.nearest_form, own code) and the whole text
compared. Also checks each before -> after difference (diff_wb.py json) is an old-form -> new-form change on the same target and offset.
usage: check_wording.py <after2.xlsx> [diff.json]"""
import sys, re, json, collections
import openpyxl
from wfam import nearest_form, offs

FAM = r'(whole|half|golden|phi power|ninth|sqrt2|silver)'
PAT = {
    'T-table outside': re.compile(rf'^(\S+) \({FAM}\) — outside ±(\S+)$'),
    'T24 outside': re.compile(rf'^nearest: (\S+) · {FAM} — outside ±(\S+) · ([+-]\d+\.\d{{5}})$'),
    'timeline outside': re.compile(rf'^(\S+) \({FAM}\) — outside$'),
    'FIELD outside': re.compile(r'^(\S+) ([+-]\d+\.\d{5}) \(outside ±(\S+)\)$'),
    'T24 not 4-digit': re.compile(rf'^nearest: (\S+) {FAM} — not a 4-digit match \(([+-]\d+\.\d{{5}})\)$'),
    'not 4-digit': re.compile(rf'^(\S+) {FAM} — not a 4-digit match \(([+-]\d+\.\d{{5}})\)$'),
}
LOOSE = re.compile(r'not a 4-digit|— outside ±|\(outside ±|\) — outside$')


def expected(kind, v):
    form, fam, lab, t, tol = nearest_form(v)
    o = offs(v - t)
    if kind.startswith('T24'):
        return f'nearest: {lab} {fam} — not a 4-digit match ({o})' if form == 'notdigit' else \
               f'nearest: {lab} · {fam} — outside ±{tol:.3g} · {o}'
    if form == 'notdigit':
        return f'{lab} {fam} — not a 4-digit match ({o})'
    return {'T-table outside': f'{lab} ({fam}) — outside ±{tol:.3g}', 'timeline outside': f'{lab} ({fam}) — outside',
            'FIELD outside': f'{lab} {o} (outside ±{tol:.3g})'}.get(kind, f'{lab} ({fam}) — outside ±{tol:.3g}')


def classify(s):
    for k, p in PAT.items():
        if p.match(s):
            return k
    return None


OFF = re.compile(r'[+-]\d+\.\d{5}(?!\d)')


def same(a, b):
    """identical up to the 5-dp offsets, which may differ by 1 in the last place (the value cell is rounded to 6 dp)"""
    if OFF.sub('#', a) != OFF.sub('#', b):
        return False
    return all(abs(float(x) - float(y)) <= 1.1e-5 for x, y in zip(OFF.findall(a), OFF.findall(b)))


wb = openpyxl.load_workbook(sys.argv[1])
cnt = collections.Counter(); mism = []; unmatched = []; nchk = 0
RATIO_FIELD = re.compile(r'^FIELD (12|13|14|21|22|23)R?$')     # ratio matrices: interval texts, not distances
for ws in wb.worksheets:
    if RATIO_FIELD.match(ws.title):
        continue
    for row in ws.iter_rows():
        for c in row:
            texts = []
            if isinstance(c.value, str) and LOOSE.search(c.value):
                texts.append(('cell', c.value))
            if c.comment and LOOSE.search(c.comment.text.split('\n')[0]):
                texts.append(('comment', c.comment.text.split('\n')[0]))
            for where, s in texts:
                k = classify(s)
                if k is None:
                    unmatched.append((ws.title, c.coordinate, s)); continue
                if where == 'comment':
                    v = c.value; vk = 'T24 not 4-digit' if 'not a 4-digit' in s else 'T24 outside'
                    if not k.startswith('T24'):
                        mism.append((ws.title, c.coordinate, 'comment shape', s)); continue
                else:
                    if k.startswith('T24'):
                        mism.append((ws.title, c.coordinate, 'cell with comment shape', s)); continue
                    # find the value: T-table 2 left (its own offset at +2), timeline / FIELD 1 left
                    cand = []
                    for dl in (2, 1):
                        if c.column - dl >= 1:
                            x = ws.cell(c.row, c.column - dl).value
                            if isinstance(x, (int, float)) and not isinstance(x, bool):
                                cand.append((dl, x))
                    # T-table: value at -2 and the offset cell at +2
                    tt = ws.cell(c.row, c.column + 2).value
                    if k == 'T-table outside' or (k == 'not 4-digit' and isinstance(tt, str) and re.fullmatch(r'[+-]\d+\.\d{5}', tt or '')
                                                   and any(d == 2 for d, _ in cand)):
                        v = [x for d, x in cand if d == 2]; v = v[0] if v else None; ctx = 'T'
                    else:
                        v = [x for d, x in cand if d == 1]; v = v[0] if v else None; ctx = 'TL/FIELD'
                    if v is None:
                        mism.append((ws.title, c.coordinate, 'no value found', s)); continue
                    if k == 'not 4-digit':
                        # pick the outside-form shape this context would use
                        if ctx == 'T':
                            vk = 'T-table outside'
                        elif ws.title.startswith('FIELD'):
                            vk = 'FIELD outside'
                        else:
                            vk = 'timeline outside'
                    else:
                        vk = k
                    if ctx == 'T':
                        # offset cell must equal the offset from the value
                        form, fam, lab, t, tol = nearest_form(v)
                        if not same(tt or '', offs(v - t)):
                            mism.append((ws.title, c.coordinate, f'offset cell {tt!r} != {offs(v - t)!r} v={v}', s))
                nchk += 1
                if nearest_form(v)[0] == 'HIT':
                    mism.append((ws.title, c.coordinate, f'v={v} is a family hit', s)); continue
                e = expected(vk, v)
                form = 'not 4-digit' if 'not a 4-digit' in s else 'outside'
                cnt[(ws.title if ws.title in ('SKY JUPITER', 'POSITIONS') or ws.title.startswith('FIELD') else re.sub(r'\d+', '#', ws.title), where, form)] += 1
                if not same(s, e):
                    mism.append((ws.title, c.coordinate, f'v={v}', f'{s!r} != {e!r}'))
print(sys.argv[1])
print('  texts checked', nchk, 'mismatches', len(mism), 'unmatched shapes', len(unmatched))
for m in mism[:15]: print('   ', m)
for m in unmatched[:10]: print('    unmatched', m)
tot = collections.Counter()
for (sh, where, form), n in sorted(cnt.items()):
    print(f'  {sh:12s} {where:8s} {form:12s} {n}')
    tot[form] += n
print('  TOTAL', dict(tot))

if len(sys.argv) > 2:
    d = json.load(open(sys.argv[2]))
    bad = collections.Counter(); ok = 0; ex = []
    kinds = collections.Counter(x[2] for x in d)
    for sh, co, kind, a, b in d:
        if kind not in ('value', 'comment'):
            bad['non-text ' + kind] += 1; continue
        a1 = (a or '').split('\n'); b1 = (b or '').split('\n')
        if a1[1:] != b1[1:]:
            bad['other comment lines'] += 1; ex.append((sh, co, a, b)); continue
        ka, kb = classify(a1[0]), classify(b1[0])
        if kb not in ('not 4-digit', 'T24 not 4-digit') or ka is None or 'not 4-digit' in ka:
            bad['not old->new'] += 1; ex.append((sh, co, a, b)); continue
        ma, mb = PAT[ka].match(a1[0]), PAT[kb].match(b1[0])
        if ma.group(1) != mb.group(1):
            bad['target'] += 1; ex.append((sh, co, a, b)); continue
        if ka != 'FIELD outside' and ma.group(2) != mb.group(2):
            bad['family'] += 1; ex.append((sh, co, a, b)); continue
        offa = ma.group(4) if ka == 'T24 outside' else ma.group(2) if ka == 'FIELD outside' else None
        if offa is not None and offa != mb.group(3):
            bad['offset'] += 1; ex.append((sh, co, a, b)); continue
        ok += 1
    print('  diff entries', len(d), dict(kinds), 'old->new same target / family / offset:', ok, 'bad', dict(bad))
    for e in ex[:10]: print('    ', e)
