"""Plant errors in copies of one after workbook. usage: plants_r1.py after.xlsx outdir"""
import sys, os
import openpyxl
from openpyxl.styles import PatternFill

src, out = sys.argv[1], sys.argv[2]
BLUE = PatternFill('solid', start_color='BDD7EE')
NONE = PatternFill(fill_type=None)


def find(wb, pred):
    for ws in wb.worksheets:
        for r in range(1, ws.max_row + 1):
            a, b = ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value
            if a in ('MATCH HORSE', 'MATCH JOCKEY', 'MATCH HORSE | JOCKEY') and pred(ws, r, b):
                return ws, r
    raise SystemExit('no target')


def plant(name, fn):
    wb = openpyxl.load_workbook(src)
    msg = fn(wb)
    p = os.path.join(out, f'plant_{name}.xlsx')
    wb.save(p)
    print(name, msg, p)


def p1(wb):
    ws, r = find(wb, lambda ws, r, b: str(ws.cell(row=r, column=3).value) == '1')
    ws.cell(row=r, column=3).fill = NONE
    return f'cleared blue {ws.title}!C{r}'


def p2a(wb):
    ws, r = find(wb, lambda ws, r, b: b == 'RATIO')
    ws.cell(row=r, column=3).fill = BLUE
    return f'blue on RATIO {ws.title}!C{r} value {ws.cell(row=r, column=3).value!r}'


def p2b(wb):
    ws, r = find(wb, lambda ws, r, b: b == 'BODY')
    ws.cell(row=r, column=3).fill = BLUE
    return f'blue on BODY {ws.title}!C{r}'


def p3a(wb):
    wb['SUMMARY'].freeze_panes = 'C4'
    return 'SUMMARY freeze C4'


def p3b(wb):
    wb['P01'].freeze_panes = 'C4'
    return 'P01 freeze C4'


for n, f in (('1_missing', p1), ('2a_ratio', p2a), ('2b_body', p2b), ('3a_summary', p3a), ('3b_runner', p3b)):
    plant(n, f)
