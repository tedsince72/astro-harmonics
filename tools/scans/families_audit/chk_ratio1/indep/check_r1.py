"""Independent checker (31 Oct 2026): ratio '1' column C light blue in MATCH tables + SUMMARY freeze D4.
usage: check_r1.py before.xlsx after.xlsx   (before may be '-' to run only (a) and (b) on after)
Own code; reads the workbooks only."""
import sys
from collections import Counter, defaultdict
import openpyxl

TABLES = ('MATCH HORSE', 'MATCH JOCKEY', 'MATCH HORSE | JOCKEY')
SECTIONS = ('EXACT', 'RATIO', 'BODY', 'CONJ STAR')


def is_blue(c):
    f = c.fill
    if f is None or f.fill_type != 'solid':
        return False
    rgb = f.fgColor.rgb if f.fgColor is not None else None
    return isinstance(rgb, str) and rgb.upper().endswith('BDD7EE')


def kind(name):
    if name in ('POSITIONS', 'SKY JUPITER', 'SUMMARY'):
        return name
    return 'runner'


def data_rows(ws):
    for r in range(1, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value
        b = ws.cell(row=r, column=2).value
        if a in TABLES and b in SECTIONS:
            yield r, a, b


def check_a(wb, label):
    errs = []
    counts = Counter()          # (table-or-SUMMARY) -> blue count
    secs = defaultdict(Counter)
    n_rows = Counter()
    for ws in wb.worksheets:
        k = kind(ws.title)
        if k not in ('runner', 'SUMMARY'):
            continue
        for r, a, b in data_rows(ws):
            key = 'SUMMARY' if k == 'SUMMARY' else a
            n_rows[key] += 1
            c = ws.cell(row=r, column=3)
            v = c.value
            one = (str(v).strip() == '1') if v is not None else False
            blue = is_blue(c)
            if one:
                counts[key] += 1
                secs[key][b] += 1
            if one != blue:
                errs.append(f'{label} {ws.title}!C{r} table={a} sec={b} value={v!r} blue={blue}')
            if b in ('BODY', 'CONJ STAR') and v not in (None, ''):
                errs.append(f'{label} {ws.title}!C{r} {b} row C not blank: {v!r}')
    return counts, secs, n_rows, errs


def freezes(wb):
    return {ws.title: ws.freeze_panes for ws in wb.worksheets}


def check_b(wb, label):
    errs = []
    for t, f in freezes(wb).items():
        k = kind(t)
        want = None if k in ('POSITIONS', 'SKY JUPITER') else 'D4'
        if f != want:
            errs.append(f'{label} freeze {t}: {f!r} (want {want!r})')
    return errs


def cmt(c):
    return None if c.comment is None else c.comment.text


def compare(wb0, wb1):
    diffs = []
    allowed = Counter()
    if wb0.sheetnames != wb1.sheetnames:
        diffs.append(f'sheet list {wb0.sheetnames} vs {wb1.sheetnames}')
    ncells = 0
    for name in wb0.sheetnames:
        if name not in wb1.sheetnames:
            continue
        s0, s1 = wb0[name], wb1[name]
        if s0.freeze_panes != s1.freeze_panes:
            if name == 'SUMMARY' and s0.freeze_panes is None and s1.freeze_panes == 'D4':
                allowed['SUMMARY freeze'] += 1
            else:
                diffs.append(f'{name} freeze {s0.freeze_panes!r} -> {s1.freeze_panes!r}')
        m0 = sorted(str(m) for m in s0.merged_cells.ranges)
        m1 = sorted(str(m) for m in s1.merged_cells.ranges)
        if m0 != m1:
            diffs.append(f'{name} merges differ ({len(m0)} vs {len(m1)})')
        cols = set(s0.column_dimensions) | set(s1.column_dimensions)
        for L in sorted(cols):
            d0, d1 = s0.column_dimensions[L], s1.column_dimensions[L]
            if (d0.width, d0.customWidth, d0.hidden) != (d1.width, d1.customWidth, d1.hidden):
                diffs.append(f'{name} col {L} width {d0.width} -> {d1.width}')
        if (s0.max_row, s0.max_column) != (s1.max_row, s1.max_column):
            diffs.append(f'{name} dims {s0.max_row}x{s0.max_column} -> {s1.max_row}x{s1.max_column}')
        blue_ok = set()
        if kind(name) in ('runner', 'SUMMARY'):
            for r, a, b in data_rows(s1):
                v = s1.cell(row=r, column=3).value
                if v is not None and str(v).strip() == '1':
                    blue_ok.add(r)
        R = max(s0.max_row, s1.max_row)
        C = max(s0.max_column, s1.max_column)
        for r in range(1, R + 1):
            for col in range(1, C + 1):
                c0, c1 = s0.cell(row=r, column=col), s1.cell(row=r, column=col)
                ncells += 1
                what = []
                if c0.value != c1.value:
                    what.append(f'value {c0.value!r}->{c1.value!r}')
                if repr(c0.font) != repr(c1.font):
                    what.append('font')
                if repr(c0.fill) != repr(c1.fill):
                    if col == 3 and r in blue_ok and is_blue(c1) and not is_blue(c0):
                        allowed['C fill blue'] += 1
                    else:
                        what.append(f'fill {c0.fill.fgColor.rgb}->{c1.fill.fgColor.rgb}')
                if repr(c0.border) != repr(c1.border):
                    what.append('border')
                if c0.number_format != c1.number_format:
                    what.append(f'numfmt {c0.number_format}->{c1.number_format}')
                if repr(c0.alignment) != repr(c1.alignment):
                    what.append('alignment')
                if cmt(c0) != cmt(c1):
                    what.append('comment')
                if what:
                    diffs.append(f'{name}!{c1.coordinate}: ' + ', '.join(what))
    return diffs, allowed, ncells


def main():
    bpath, apath = sys.argv[1], sys.argv[2]
    wa = openpyxl.load_workbook(apath)
    counts, secs, n_rows, errs = check_a(wa, 'after')
    errs += check_b(wa, 'after')
    print('AFTER', apath)
    for key in TABLES + ('SUMMARY',):
        print(f'  (a) {key}: data rows {n_rows[key]}, blue {counts[key]}, sections {dict(secs[key])}')
    print('  (b) freeze after:', dict(Counter((kind(t), f) for t, f in freezes(wa).items())))
    if bpath != '-':
        wb = openpyxl.load_workbook(bpath)
        print('  (b) freeze before:', dict(Counter((kind(t), f) for t, f in freezes(wb).items())))
        c0, _, _, _ = check_a(wb, 'before')
        bb = sum(1 for ws in wb.worksheets for row in ws.iter_rows() for c in row if is_blue(c))
        print(f'  before: ratio-1 C cells {dict(c0)}, BDD7EE cells anywhere {bb}')
        diffs, allowed, ncells = compare(wb, wa)
        print(f'  (c) cells compared {ncells}; allowed {dict(allowed)}; blue expected {sum(counts.values())}')
        if allowed['C fill blue'] != sum(counts.values()):
            errs.append(f"allowed blue diffs {allowed['C fill blue']} != blue count {sum(counts.values())}")
        errs += diffs
    print(f'  MISMATCHES {len(errs)}')
    for e in errs[:15]:
        print('   ', e)
    sys.exit(1 if errs else 0)


if __name__ == '__main__':
    main()
