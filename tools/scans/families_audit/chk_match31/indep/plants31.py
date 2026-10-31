"""Plant errors in copies of the after workbooks and check that check31.py catches each.
usage: python3 plants31.py <before_dir> <after_dir> <work_dir>"""
import copy
import os
import subprocess
import sys

import openpyxl

BD, AD, WD = sys.argv[1:4]
HERE = os.path.dirname(os.path.abspath(__file__))


def find(ws, a, b=None, col=2, start=1):
    for r in range(start, ws.max_row + 1):
        if ws.cell(r, 1).value == a and (b is None or ws.cell(r, col).value == b):
            return r
    raise KeyError((a, b))


def copy_row(ws, src, dst, n=29):
    for c in range(1, n + 1):
        s, d = ws.cell(src, c), ws.cell(dst, c)
        d.value = s.value
        d._style = copy.copy(s._style)


def swap_rows(ws, r1, r2, n=29):
    for c in range(1, n + 1):
        a, b = ws.cell(r1, c), ws.cell(r2, c)
        av, ast = a.value, copy.copy(a._style)
        a.value, a._style = b.value, copy.copy(b._style)
        b.value, b._style = av, ast


def p_missing_conj(wb):        # a CONJ STAR row missing (blanked)
    ws = wb['P01']; r = find(ws, 'MATCH JOCKEY', 'CONJ STAR')
    for c in range(1, 22): ws.cell(r, c).value = None


def p_missing_body(wb):        # a BODY row missing: overwrite it with the next header / row
    ws = wb['P01']; r = find(ws, 'MATCH HORSE', 'BODY')
    ws.cell(r, 2).value = 'RATIO'


def p_dup(wb):                 # duplicated CONJ STAR row
    ws = wb['P01']; r = find(ws, 'MATCH JOCKEY', 'CONJ STAR'); copy_row(ws, r, r + 1, 21)


def p_order(wb):               # two CONJ STAR rows swapped inside the group
    ws = wb['P01']; r = find(ws, 'MATCH JOCKEY', 'CONJ STAR'); swap_rows(ws, r, r + 1, 21)


def p_pairing(wb):             # P05 HJ: the lone jockey CONJ STAR line moved onto the horse line
    ws = wb['P05']; r = find(ws, 'MATCH HORSE | JOCKEY', 'CONJ STAR')
    r2 = find(ws, 'MATCH HORSE | JOCKEY', 'CONJ STAR', start=r + 1)
    src, dst = (r2, r) if ws.cell(r, 22).value is None else (r, r2)
    for c in range(22, 30):
        ws.cell(dst, c).value = ws.cell(src, c).value
        ws.cell(dst, c)._style = copy.copy(ws.cell(src, c)._style)
        ws.cell(src, c).value = None


def p_hdr_count(wb):           # wrong grey header count in MATCH HORSE
    ws = wb['P01']; r = find(ws, 'MATCH HORSE', 'BODY'); h = r - 1
    ws.cell(h, 2).value = ws.cell(h, 2).value.replace('BODY 1', 'BODY 2')


def p_hj_hdr(wb):              # wrong HJ header count
    ws = wb['P01']; r = find(ws, 'MATCH HORSE | JOCKEY', 'BODY'); h = r - 1
    ws.cell(h, 2).value = ws.cell(h, 2).value.replace('BODY 1 | 0', 'BODY 0 | 1')


def p_freeze(wb):
    wb['P03'].freeze_panes = 'C4'


def p_sum_drop(wb):            # last SUMMARY line dropped
    s = wb['SUMMARY']; last = max(r for r in range(1, s.max_row + 1) if s.cell(r, 1).value is not None)
    for c in range(1, 30): s.cell(last, c).value = None


def p_sum_include(wb):         # a non-qualifying HJ line added at the end of the SUMMARY
    s = wb['SUMMARY']; last = max(r for r in range(1, s.max_row + 1) if s.cell(r, 1).value is not None)
    lastrunner = [n for n in wb.sheetnames if n.startswith('P') and n[1:].isdigit()][-1]
    ws = wb[lastrunner]
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == 'MATCH HORSE | JOCKEY' and ws.cell(r, 2).value in ('EXACT', 'RATIO') and \
                (ws.cell(r, 13).value or 0) <= 50:
            for c in range(1, 30):
                s.cell(last + 1, c).value = ws.cell(r, c).value
                s.cell(last + 1, c)._style = copy.copy(ws.cell(r, c)._style)
            return
    raise RuntimeError('no low line')


def p_sum_hdr(wb):             # SUMMARY header count wrong
    s = wb['SUMMARY']
    for r in range(1, s.max_row + 1):
        v = s.cell(r, 2).value
        if isinstance(v, str) and ' — EXACT ' in v and 'RATIO 0 | 0' not in v:
            import re
            s.cell(r, 2).value = re.sub(r'RATIO (\d+)', lambda m: f'RATIO {int(m.group(1)) + 1}', v, count=1)
            return


def p_old_row(wb):             # an existing EXACT row changed
    ws = wb['P01']; r = find(ws, 'MATCH HORSE', 'EXACT'); ws.cell(r, 20).value = (ws.cell(r, 20).value or 0) + 0.01


def p_body_score(wb):          # the BODY row's T24 score changed
    ws = wb['P01']; r = find(ws, 'MATCH HORSE', 'BODY'); ws.cell(r, 20).value = 77.0


def p_title(wb):               # MATCH JOCKEY title CONJ STAR total wrong
    ws = wb['P01']; r = find(ws, 'MATCH JOCKEY')
    ws.cell(r, 2).value = ws.cell(r, 2).value.replace('CONJ STAR 5 rows', 'CONJ STAR 4 rows')


def p_conj_fill(wb):           # fill of a new row's score cell changed
    ws = wb['P01']; r = find(ws, 'MATCH JOCKEY', 'CONJ STAR')
    ws.cell(r, 20)._style = copy.copy(ws.cell(r, 8)._style)


PLANTS = [('newcastle', p_missing_conj), ('ripon', p_missing_body), ('newcastle', p_dup), ('newcastle', p_order),
          ('newcastle', p_pairing), ('ripon', p_hdr_count), ('ripon', p_hj_hdr), ('kelso', p_freeze),
          ('newcastle', p_sum_drop), ('newcastle', p_sum_include), ('ripon', p_sum_hdr), ('lingfield', p_old_row),
          ('ripon', p_body_score), ('newcastle', p_title), ('newcastle', p_conj_fill)]

caught = 0
for race, fn in PLANTS:
    wb = openpyxl.load_workbook(f'{AD}/{race}_jupiter_tables.xlsx')
    fn(wb)
    out = f'{WD}/plant_{fn.__name__}_{race}.xlsx'
    wb.save(out)
    res = subprocess.run([sys.executable, f'{HERE}/check31.py', f'{BD}/{race}_jupiter_tables.xlsx', out, '--quiet'],
                         capture_output=True, text=True)
    lines = res.stdout.strip().splitlines()
    ok = res.returncode == 1
    caught += ok
    print(f"{'CAUGHT' if ok else 'MISSED'} {fn.__name__} ({race}): {lines[0]}; {lines[1].strip() if len(lines) > 1 else ''}")
print(f'{caught} / {len(PLANTS)} caught')
