"""Extra checks: blank-row layout of the MATCH region, and source cells near the 50 threshold.
usage: python3 extra31.py <after.xlsx>"""
import sys
import openpyxl
sys.path.insert(0, __file__.rsplit('/', 1)[0])
import check31 as C

A = openpyxl.load_workbook(sys.argv[1])
sky = C.parse_sky(A)
near = []
for c in sky:
    if c['score'] is not None and 49.9 <= c['score'] <= 50.1:
        near.append(('SKY', c['rel'], c['plane'], c['score']))
cand_sky = [c for c in sky if c['score'] is not None and c['score'] > 50]
print('sky cells > 50:', len(cand_sky), 'of which T24 bodies:',
      sum(1 for c in cand_sky if c['x'] in C.BIDX and c['x'] not in ('Jupiter', 'Moon')),
      'conj stars:', sum(1 for c in cand_sky if c['fam'] == 'conj' and c['x'] not in C.BIDX))
bad = 0
for rn in [n for n in A.sheetnames if len(n) == 3 and n[0] == 'P' and n[1:].isdigit()]:
    ws = A[rn]
    for which, tag in (('HORSE', 'H'), ('JOCKEY', 'J')):
        rows, _ = C.parse_t24(ws, which)
        for lab, d in rows.items():
            for k, cell in d.items():
                if cell['sc'] and 49.9 <= cell['sc']['score'] <= 50.1:
                    near.append((rn, lab, k, cell['sc']['score']))
    # layout: MATCH region rows are table rows, column-names rows or blank; 2 blanks between tables; nothing after
    t = C.find_tables(ws)
    start = t['MATCH HORSE'][0]
    ends = [t[x][-1] for x in C.TABLES]
    starts = [t[x][0] for x in C.TABLES]
    for i in range(2):
        gap = [r for r in range(ends[i] + 1, starts[i + 1])]
        if len(gap) != 2 or any(ws.cell(r, c).value is not None for r in gap for c in range(1, 60)):
            print('LAYOUT', rn, 'gap after', C.TABLES[i], gap); bad += 1
    for r in range(start, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a in C.TABLES:
            continue
        if a == 'table' and ws.cell(r, 2).value == 'section':
            continue
        if any(ws.cell(r, c).value is not None for c in range(1, 60)):
            if r > ends[2] or not any(starts[i] <= r <= ends[i] for i in range(3)):
                if not (ends[0] < r < starts[1] or ends[1] < r < starts[2]):
                    print('LAYOUT stray', rn, r); bad += 1
        elif any(starts[i] < r < ends[i] for i in range(3)):
            print('LAYOUT blank inside table', rn, r); bad += 1
print('near 50:', near)
print('layout problems', bad)
