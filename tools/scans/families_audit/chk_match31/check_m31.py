"""freeze D4; MATCH BODY and CONJ STAR rows; SUMMARY sheet (Eddie, 31 Oct 2026): checks one build by the new code against the build before
(93b54a43), same race.
usage: check_m31.py <before.xlsx> <after.xlsx>
1. POSITIONS, SKY JUPITER identical cell for cell (value + style + comment), freeze unchanged (printed); sheet order = before + 'SUMMARY'.
2. Runner sheets: freeze D4 (before C4); widths and every cell outside the three MATCH tables identical; the MATCH tables after, with the
   BODY / CONJ STAR rows (MATCH HORSE / JOCKEY), the BODY / CONJ STAR lines (MATCH HORSE | JOCKEY) and the grey headers of groups holding only
   those taken out, = the before tables row for row, cell for cell; grey header rows: same sky text, the before counts 'EXACT n, RATIO m'
   (MATCH | : 'EXACT h | j, RATIO h | j') = the first two counts after; title rows (column B) may differ.
3. BODY / CONJ STAR rows rebuilt from the SKY block (SKY JUPITER) and T24 HORSE / JOCKEY cells on the sheets (scores from the cells and comments):
   the set, every cell, their place (after EXACT / RATIO in their sky cell's group, rule order), group order by sky score, header counts.
4. MATCH HORSE | JOCKEY: every BODY / CONJ STAR row of MATCH HORSE / JOCKEY once on its side (cells copied), key and order, header counts.
5. SUMMARY rebuilt from the MATCH HORSE | JOCKEY table on each runner sheet (filter, groups, counts, cells and styles, merges, blank rows).
Prints per runner 'P01 H BODY b CONJ c | J BODY b CONJ c | conj groups H g J g | summary n' and 'mismatches: n' last."""
import sys
import openpyxl

SEC = ('EXACT', 'RATIO', 'BODY', 'CONJ STAR')
NEW = ('BODY', 'CONJ STAR')
T24B = ['Sun', 'Moon', 'Mercury', 'Venus', 'Mars', 'Jupiter', 'Saturn', 'Uranus', 'Neptune', 'Pluto', 'Chiron', 'Ceres', 'Pallas', 'Juno',
        'Vesta', 'Eris', 'Sedna', 'Haumea', 'Makemake', 'Gonggong', 'Quaoar', 'Orcus', 'Transpluto', 'Rahu', 'Ketu']
bad = []


def sig(c):
    return (c.value, repr(c.font), repr(c.fill), repr(c.border), c.number_format, repr(c.alignment), c.comment.text if c.comment else None)


def vals(ws, r, n=29):
    return [ws.cell(row=r, column=k).value for k in range(1, n + 1)]


def last_row(ws):
    m = ws.max_row
    while m > 0 and all(c.value is None for c in ws[m]):
        m -= 1
    return m


B, A = openpyxl.load_workbook(sys.argv[1]), openpyxl.load_workbook(sys.argv[2])
if A.sheetnames != B.sheetnames + ['SUMMARY']:
    bad.append(f'sheets {A.sheetnames}')
for sh in ('POSITIONS', 'SKY JUPITER'):
    b, a = B[sh], A[sh]
    print(f'{sh} freeze: before {b.freeze_panes}, after {a.freeze_panes}')
    if b.freeze_panes != a.freeze_panes:
        bad.append(f'{sh} freeze')
    for r in range(1, max(b.max_row, a.max_row) + 1):
        for k in range(1, max(b.max_column, a.max_column) + 1):
            if sig(b.cell(row=r, column=k)) != sig(a.cell(row=r, column=k)):
                bad.append(f'{sh} {r},{k}')

# SKY block cells: (label, plane) -> (value, target text, score, order)
SK = {}
ws = A['SKY JUPITER']
for r in range(1, ws.max_row + 1):
    if ws.cell(row=r, column=1).value == 'SKY' and str(ws.cell(row=r, column=2).value).startswith('S Jupiter → '):
        for pl, c in (('RA', 3), ('Dec', 12)):
            v, t, s = ws.cell(row=r, column=c).value, ws.cell(row=r, column=c + 2).value, ws.cell(row=r, column=c + 5).value
            if isinstance(v, (int, float)) and isinstance(s, (int, float)):
                SK[(ws.cell(row=r, column=2).value, pl)] = (v, t, s, len(SK))


def t24(ws, who):
    """T24 HORSE / JOCKEY -> {row label: {(body, plane): (value, family, target, score)}} from the cells and their comments"""
    r0 = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value == 'T24'
              and str(ws.cell(row=r, column=2).value).startswith(f'T24 {who}'))
    cols = {}
    pl = None
    for k in range(3, ws.max_column + 1):
        h1, h2 = ws.cell(row=r0 + 1, column=k).value, ws.cell(row=r0 + 2, column=k).value
        if h1 in ('RA', 'Dec'):
            pl = h1
        if h2 in T24B and pl:
            cols[k] = (h2, pl)
        if isinstance(h1, str) and h1.endswith('midpoints'):
            if h1 == 'Dec midpoints':
                break
            pl = None
    out = {}
    r = r0 + 3
    while ws.cell(row=r, column=1).value == 'T24':
        lab = ws.cell(row=r, column=2).value
        d = out.setdefault(lab, {})
        for k, (b, pl) in cols.items():
            c = ws.cell(row=r, column=k)
            if isinstance(c.value, (int, float)) and c.comment and not c.comment.text.startswith('nearest'):
                f = c.comment.text.split('\n')[0].split(' · ')
                fam = 'conj' if f[1].startswith('conjunction') else f[1]
                d[(b, pl)] = (c.value, fam, 'conj' if fam == 'conj' else f[0], float(f[-1]))
        r += 1
    return out, [b for b, pl in cols.values() if pl == 'RA']


def tables(ws):
    """-> {table id: [row numbers]} of the three MATCH tables, title row first"""
    out = {}
    for r in range(1, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value
        if a in ('MATCH HORSE', 'MATCH JOCKEY', 'MATCH HORSE | JOCKEY'):
            out.setdefault(a, []).append(r)
    return out


def head_counts(t):
    return t.split(' — ', 1)[1] if ' — ' in t else None


summ = {}
for P in [s for s in B.sheetnames if s.startswith('P') and s[1:].isdigit()]:
    b, a = B[P], A[P]
    if (b.freeze_panes, a.freeze_panes) != ('C4', 'D4'):
        bad.append(f'{P} freeze {b.freeze_panes} -> {a.freeze_panes}')
    for L, d in b.column_dimensions.items():
        if a.column_dimensions[L].width != d.width:
            bad.append(f'{P} width {L}')
    tb, ta = tables(b), tables(a)
    s0b, s0a = tb['MATCH HORSE'][0], ta['MATCH HORSE'][0]
    if s0b != s0a:
        bad.append(f'{P} MATCH start {s0b} {s0a}')
    for r in range(1, s0b):
        for k in range(1, b.max_column + 1):
            if sig(b.cell(row=r, column=k)) != sig(a.cell(row=r, column=k)):
                bad.append(f'{P} {r},{k} before MATCH')
    if last_row(a) != ta['MATCH HORSE | JOCKEY'][-1]:
        bad.append(f'{P} something after MATCH HORSE | JOCKEY')
    # 2. the tables minus the new rows = before; layout of blank rows between tables
    for tid in ('MATCH HORSE', 'MATCH JOCKEY', 'MATCH HORSE | JOCKEY'):
        rb, ra = tb[tid], ta[tid]
        keep, newrows = [], []
        hj = tid.endswith('JOCKEY') and '|' in tid
        n = 29 if hj else 21
        grp = []
        for r in ra:
            sec = a.cell(row=r, column=2).value
            if sec in NEW:
                newrows.append(r)
                grp[-1][1].append(sec)
            elif sec in SEC:
                keep.append(r)
                grp[-1][1].append(sec)
            elif isinstance(sec, str) and ' — EXACT ' in sec and a.cell(row=r, column=1).fill.fill_type:
                grp.append((r, []))
                keep.append(r)
            else:
                keep.append(r)
        only_new = {r for r, ss in grp if ss and all(s in NEW for s in ss)}
        if any(not ss for r, ss in grp):
            bad.append(f'{P} {tid} empty group')
        keep = [r for r in keep if r not in only_new]
        if len(keep) != len(rb):
            bad.append(f'{P} {tid} rows {len(rb)} -> kept {len(keep)}')
            continue
        for x, y in zip(rb, keep):
            vb, va = vals(b, x, n), vals(a, y, n)
            if x == rb[0]:                                 # title row: column B may differ
                if vb[0] != va[0]:
                    bad.append(f'{P} {tid} title')
                continue
            if isinstance(vb[1], str) and ' — EXACT ' in vb[1]:
                pb, cb = vb[1].split(' — ', 1)
                pa, ca = va[1].split(' — ', 1)
                if pb != pa or not ca.startswith(cb + ', BODY '):
                    bad.append(f'{P} {tid} header {vb[1]!r} -> {va[1]!r}')
                cs = range(1, n + 1)
                if any(repr(b.cell(row=x, column=k).fill) != repr(a.cell(row=y, column=k).fill) for k in cs):
                    bad.append(f'{P} {tid} header fill')
                continue
            for k in range(1, n + 1):
                if sig(b.cell(row=x, column=k)) != sig(a.cell(row=y, column=k)):
                    bad.append(f'{P} {tid} row {x}->{y} col {k}: {b.cell(row=x, column=k).value!r} -> {a.cell(row=y, column=k).value!r}')
        if ra[-1] - ra[0] + 1 != len(ra) + 1:              # + the column-names row (column A 'table')
            bad.append(f'{P} {tid} gap inside')
    if ta['MATCH JOCKEY'][0] != ta['MATCH HORSE'][-1] + 3 or ta['MATCH HORSE | JOCKEY'][0] != ta['MATCH JOCKEY'][-1] + 3:
        bad.append(f'{P} blank rows between tables')
    # 3. BODY / CONJ STAR rebuilt from the sheet cells
    got = {}
    for C, who in (('H', 'HORSE'), ('J', 'JOCKEY')):
        T, bodies = t24(a, who)
        exp = []
        for (lab, pl), (v, tt, sc, i) in SK.items():
            if sc <= 50 or not lab.startswith('S Jupiter → S '):
                continue
            x = lab[len('S Jupiter → S '):]
            sfam, stgt = ('conj', 'conj') if tt == 'conj (conj)' else tt.rsplit(' (', 1)[::-1]
            sfam = sfam.rstrip(')')
            skyc = (lab, pl, v, sfam, stgt, sc)
            if x in bodies and x != 'Jupiter':
                rl = f'{C} Jupiter → {C} body'
                for p2 in ('RA', 'Dec'):
                    c = T[rl].get((x, p2))
                    if c and c[3] > 50:
                        exp.append(('BODY', skyc, (rl, x, p2) + c))
            if sfam == 'conj' and f'{C} {x}' in T:
                for rl in (f'{C} {x}', f'S Jupiter → {C} body'):
                    for bb in bodies:
                        for p2 in ('RA', 'Dec'):
                            c = T[rl].get((bb, p2))
                            if c and c[3] > 50:
                                exp.append(('CONJ STAR', skyc, (rl, bb, p2) + c))
        tid = f'MATCH {who}'
        rows = ta[tid]
        have = []
        cur, gseen = None, []
        for r in rows[1:]:
            v = vals(a, r, 21)
            if v[1] in SEC:
                gseen[-1][2].append(v[1])
                if v[1] in NEW:
                    have.append((v[1], (v[7], v[8], v[9], v[10], v[11], v[12]), tuple(v[13:20]), v[20], r))
                    if (v[7], v[8]) != gseen[-1][1]:
                        bad.append(f'{P} {tid} new row {r} in another sky cell group')
            elif isinstance(v[1], str) and ' — EXACT ' in v[1]:
                nr = vals(a, r + 1, 21)
                gseen.append((r, (nr[7], nr[8]), [], v[1]))
        ex = [(s, (sk[0], sk[1], round(sk[2], 6), sk[3], sk[4], round(sk[5], 2)), (t[0], t[1], t[2], round(t[3], 6), t[4], t[5], round(t[6], 2)),
               round(min(sk[5], t[6]), 2)) for s, sk, t in exp]
        if sorted(map(repr, [h[:4] for h in have])) != sorted(map(repr, ex)):
            bad.append(f'{P} {tid} BODY / CONJ STAR set: sheet {[h[:4] for h in have]} rule {ex}')
        # order inside each group: sections in order, the new rows in rule order (exp is in rule order per sky cell)
        for g, (hr, key, ss, txt) in enumerate(gseen):
            if ss != sorted(ss, key=SEC.index):
                bad.append(f'{P} {tid} section order in {key}')
            mine = [h[:4] for h in have if (h[1][0], h[1][1]) == key]
            rule = [e for e in ex if (e[1][0], e[1][1]) == key]
            if mine != rule:
                bad.append(f'{P} {tid} order / set in {key}: {mine} vs {rule}')
            n4 = [ss.count(s) for s in SEC]
            if not txt.endswith(f' — EXACT {n4[0]}, RATIO {n4[1]}, BODY {n4[2]}, CONJ STAR {n4[3]}'):
                bad.append(f'{P} {tid} header counts {txt!r}')
            if SK[key][1] == 'conj (conj)' and n4[0] + n4[1]:
                bad.append(f'{P} {tid} EXACT / RATIO in a conjunction group')
        sc = [(-SK[k][2], SK[k][3]) for _, k, _, _ in gseen]
        if sc != sorted(sc):
            bad.append(f'{P} {tid} group order')
        for h in have:
            r = h[4]
            if any(a.cell(row=r, column=c).value is not None for c in range(3, 8)):
                bad.append(f'{P} {tid} row {r} ratio cells not blank')
            for c in (13, 20, 21):
                f = a.cell(row=r, column=c).fill
                sv = a.cell(row=r, column=c).value
                want = '00C6EFCE' if sv >= 80 else '00FFEB9C' if sv >= 50 else '00FFF8DC'
                if not f.fill_type or f.start_color.rgb != want:
                    bad.append(f'{P} {tid} row {r} col {c} fill')
        got[C] = (sum(1 for h in have if h[0] == 'BODY'), sum(1 for h in have if h[0] == 'CONJ STAR'),
                  sum(1 for _, k, _, _ in gseen if SK[k][1] == 'conj (conj)'))
    # 4. MATCH HORSE | JOCKEY: the new lines
    src = {C: [vals(a, r, 21) for r in ta[f'MATCH {who}'] if a.cell(row=r, column=2).value in NEW] for C, who in (('H', 'HORSE'), ('J', 'JOCKEY'))}
    key = lambda v: (v[1], v[14], v[15]) if v[1] == 'BODY' else (v[1], 'S Jupiter' if str(v[13]).startswith('S Jupiter →') else 'star', v[14], v[15])
    lines, gsk = [], None
    hjr = ta['MATCH HORSE | JOCKEY']
    hjg = []
    for r in hjr[2:]:
        v = vals(a, r)
        if v[1] in SEC:
            hjg[-1][2].append(r)
            if v[1] in NEW:
                lines.append((r, v, hjg[-1][1]))
        elif isinstance(v[1], str) and ' — EXACT ' in v[1]:
            nv = vals(a, r + 1)
            hjg.append((r, (nv[7], nv[8]), []))
    seen = {'H': [], 'J': []}
    for r, v, g in lines:
        hside, jside = v[13] is not None, v[21] is not None
        if (v[7], v[8]) != g:
            bad.append(f'{P} HJ line {r} in another group')
        ks = []
        if hside:
            ks.append(key(v[:21]))
            seen['H'].append(tuple(v[1:21]))
        if jside:
            vj = v[:13] + v[21:29]
            ks.append(key(vj))
            seen['J'].append(tuple(vj[1:]))
        if len(set(ks)) != 1:
            bad.append(f'{P} HJ line {r} key {ks}')
    for C in 'HJ':
        if sorted(map(repr, seen[C])) != sorted(map(repr, (tuple(x[1:]) for x in src[C]))):
            bad.append(f'{P} HJ {C} new rows differ: {len(seen[C])} vs {len(src[C])}')
    for hr, g, rs in hjg:                                  # order of the new lines inside each group, header counts
        vs = [vals(a, r) for r in rs]
        new = [(r, v) for r, v in zip(rs, vs) if v[1] in NEW]
        ks = [key(v[:21] if v[13] is not None else v[:13] + v[21:29]) for _, v in new]
        bo = lambda k: (SEC.index(k[0]), T24B.index(k[1]), ('RA', 'Dec').index(k[2])) if k[0] == 'BODY' else \
            (SEC.index(k[0]), k[1] != 'star', T24B.index(k[2]), ('RA', 'Dec').index(k[3]))
        if ks != sorted(ks, key=bo) or len(set(ks)) != len(ks):
            bad.append(f'{P} HJ order in {g}: {ks}')
        if [v[1] for v in vs] != sorted([v[1] for v in vs], key=SEC.index):
            bad.append(f'{P} HJ section order in {g}')
        cnt = [(sum(1 for v in vs if v[1] == s and v[13] is not None), sum(1 for v in vs if v[1] == s and v[21] is not None)) for s in SEC]
        txt = a.cell(row=hr, column=2).value
        if not txt.endswith(' — ' + ', '.join(f'{s} {h} | {j}' for s, (h, j) in zip(SEC, cnt))):
            bad.append(f'{P} HJ header {txt!r}')
    # 5. SUMMARY expected for this runner: [(kind, source row or text)]
    big = lambda x: isinstance(x, (int, float)) and x > 50
    exp = [('a1', 1)]
    shown = []
    for hr, g, rs in hjg:
        sel = [r for r in rs if big(a.cell(row=r, column=13).value) and (big(a.cell(row=r, column=20).value) or big(a.cell(row=r, column=28).value))]
        if sel:
            shown.append((hr, sel))
    if shown:
        exp += [('row', hjr[1]), ('row', hjr[1] + 1)]
        for hr, sel in shown:
            cnt = [(sum(1 for r in sel if a.cell(row=r, column=2).value == s and a.cell(row=r, column=14).value is not None),
                    sum(1 for r in sel if a.cell(row=r, column=2).value == s and a.cell(row=r, column=22).value is not None)) for s in SEC]
            exp.append(('head', hr, a.cell(row=hr, column=2).value.split(' — EXACT ')[0] + ' — ' + ', '.join(f'{s} {h} | {j}' for s, (h, j) in zip(SEC, cnt))))
            exp += [('row', r) for r in sel]
    else:
        exp.append(('none',))
    summ[P] = exp
    print(f"{P} H BODY {got['H'][0]} CONJ {got['H'][1]} | J BODY {got['J'][0]} CONJ {got['J'][1]} | conj groups H {got['H'][2]} J {got['J'][2]} | "
          f"summary {sum(1 for e in exp if e[0] == 'row') - (2 if shown else 0)}")

# 5. SUMMARY sheet
S = A['SUMMARY']
r = 1
merges = set()
for P, exp in summ.items():
    a = A[P]
    for e in exp:
        if e[0] == 'a1':
            if sig(S.cell(row=r, column=1)) != sig(a.cell(row=1, column=1)):
                bad.append(f'SUMMARY {r} A1 line of {P}')
            if any(S.cell(row=r, column=k).value is not None for k in range(2, 30)):
                bad.append(f'SUMMARY {r} extra cells')
        elif e[0] == 'none':
            if S.cell(row=r, column=1).value != 'no lines':
                bad.append(f'SUMMARY {r} no lines')
        else:
            for k in range(1, 30):
                s, t = a.cell(row=e[1], column=k), S.cell(row=r, column=k)
                if e[0] == 'head' and k == 2:
                    if t.value != e[2] or (repr(s.font), repr(s.fill)) != (repr(t.font), repr(t.fill)):
                        bad.append(f'SUMMARY {r} head {t.value!r} want {e[2]!r}')
                elif sig(s) != sig(t):
                    bad.append(f'SUMMARY {r},{k} {P} row {e[1]}: {t.value!r} want {s.value!r}')
            if e[1] == A[P].cell(row=e[1], column=1).row and a.cell(row=e[1], column=14).value == 'HORSE':
                merges |= {f'N{r}:U{r}', f'V{r}:AC{r}'}
        r += 1
    for k in range(2):                                     # two blank rows between runners
        if any(S.cell(row=r + k, column=c).value is not None for c in range(1, 30)):
            bad.append(f'SUMMARY {r + k} not blank')
    r += 2
if last_row(S) != r - 3:
    bad.append(f'SUMMARY last row {last_row(S)} want {r - 3}')
if {str(m) for m in S.merged_cells.ranges} != merges:
    bad.append(f'SUMMARY merges {sorted(str(m) for m in S.merged_cells.ranges)} want {sorted(merges)}')
for x in bad[:40]:
    print('  ', x)
print('mismatches:', len(bad))
