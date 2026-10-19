"""check_t24r.py (19 Oct 2026, independent checker): T24 H / J, T24 HORSE RATIOS, T24 JOCKEY RATIOS recomputed from POSITIONS.
Own code; imports only the constant lists from tools/jupiter_tables.py and the interval list / scoring from tools/families.py.
usage: check_t24r.py tables.xlsx [sheet=P01]"""
import sys, os, re
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..', '..'))
import openpyxl
import families
from jupiter_tables import T1_T2_T24_STARS, T24_BODIES

wb = openpyxl.load_workbook(sys.argv[1])
SH = sys.argv[2] if len(sys.argv) > 2 else 'P01'
ws = wb[SH]

# --- positions
pos = {}
for row in wb['POSITIONS'].iter_rows(values_only=True):
    if row[0] in (SH,) and row[2] in ('H', 'J') and isinstance(row[5], (int, float)):
        pos[(row[2], row[3])] = (float(row[5]), float(row[6]))
for c in 'HJ':
    ra, dec = pos[(c, 'Rahu')]
    k = ((ra + 180) % 360, -dec)
    if (c, 'Ketu') in pos:
        assert abs(pos[(c, 'Ketu')][0] - k[0]) < 1e-9 and abs(pos[(c, 'Ketu')][1] - k[1]) < 1e-9, 'Ketu not Rahu mirror'
    pos[(c, 'Ketu')] = k


def dist(a, b, pl):
    if pl == 'RA':
        d = abs(a[0] - b[0]) % 360
        return min(d, 360 - d)
    return abs(a[1] - b[1])


CONJ = {'RA': 1.5, 'Dec': 0.3}
FILL = lambda s: 'C6EFCE' if s >= 80 else 'FFEB9C' if s >= 50 else 'FFF8DC'
AREA_NAME = {'A': 'A', 'P': 'φ and √2', 'B': 'B'}


def sq(tol, dev):
    return (10 * (tol - dev) / tol) ** 2


def deg_cell(v, pl):
    """own: best degree family (conjunction counts) -> (fill, comment)"""
    ev = families.evaluate(v, pl, 'deg')
    if ev:
        b = ev[0]
        fam = f"conjunction ({pl} < {CONJ[pl]})" if b['family'] == 'conj' else b['family']
        tgt = '0' if b['family'] == 'conj' else b['target']
        return FILL(b['score']), f"{tgt} · {fam} · {b['off']:+.5f} · {b['score']:.2f}"
    o = families.outside(v, pl, 'deg')
    return None, f"nearest: {o['target']} · {o['family']} — outside ±{o['tol']:.3g} · {o['off']:+.5f}"


def own_interval(v):
    """own interval scan from families.INTERVALS + interval_tol (not families.evaluate) -> best (score, lab) or None"""
    best = None
    for lab, t, name, grp in families.INTERVALS:
        tol = families.interval_tol(lab)
        dev = abs(v - t) if lab in families.AREA_A else abs(v / t - 1) * 100
        if dev < tol:
            s = sq(tol, dev)
            if best is None or s > best[0]:
                best = (s, lab)
    return best


stats = {'cells': 0, 'dist_cells': 0, 'ratio_cells': 0, 'together': 0, 'assisted': 0, 'automatic': 0, 'mism': []}
hand = []


def ratio_cell(d1, d2, legs, pl, cp, close):
    """d1, d2 the two distances, legs names; cp the two compared points (for q). -> (value, fill, comment, area or None, flags)"""
    tog = [(n, d) for n, d in zip(legs, (d1, d2)) if d < CONJ[pl]]
    if tog:
        txt = '; '.join(f"{n} {d:.6f}" for n, d in tog)
        return 'together', None, f"together ({pl} < {CONJ[pl]}): {txt} — no ratio, no score", None, 'together'
    v = max(d1, d2) / min(d1, d2)
    ev = families.evaluate(v, None, 'ratio')
    mine = own_interval(v)
    assert (mine is None) == (not ev) and (mine is None or (mine[1] == ev[0]['target'] and abs(mine[0] - ev[0]['score']) < 1e-9)), (v, mine, ev[:1])
    if not ev:
        o = families.outside(v, None, 'ratio')
        off = f"{o['off']:+.5f}%" if o['pct'] else f"{o['off']:+.5f}"
        return round(v, 6), None, f"nearest: {o['target']} · {o['name']} — outside · {off}", None, None
    b = ev[0]
    ar = families.area(b['target'])
    off = f"{b['off']:+.5f}%" if b['pct'] else f"{b['off']:+.5f}"
    com = f"{b['target']} · {b['name']} · {b['group']} · {AREA_NAME[ar]} · {off} · {b['score']:.2f}"
    fill, flag = FILL(b['score']), None
    if b['target'] == '1':
        q = dist(cp[0], cp[1], pl) / min(d1, d2)
        if q < 0.0015:
            fill, flag = 'D9D9D9', 'automatic'
            com += f"\nautomatic ({close} close, q {q:.6f}), not counted"
        elif q < 0.02:
            flag = 'assisted'
            com += f"\n{close} assisted (q {q:.6f})"
    return round(v, 6), fill, com, ar, flag


def got(cell):
    f = cell.fill.start_color.rgb[-6:] if cell.fill is not None and cell.fill.fill_type else None
    return cell.value, f, cell.comment.text if cell.comment else None


def cmp(r, c, exp, tag):
    stats['cells'] += 1
    g = got(ws.cell(r, c))
    ev, ef, ec = exp
    vok = (g[0] == ev) if isinstance(ev, str) else (isinstance(g[0], (int, float)) and abs(g[0] - ev) < 5e-7)
    if not (vok and g[1] == ef and g[2] == ec):
        stats['mism'].append((tag, ws.cell(r, c).coordinate, g, exp))


def find(tid):
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == tid and str(ws.cell(r, 2).value or '').startswith(tid + ' ('):
            return r
    raise KeyError(tid)


def merged_at(r):
    return sorted((m.min_col, m.max_col) for m in ws.merged_cells.ranges if m.min_row == r == m.max_row)


def check_title_header(tid, r0, ncol_plane, names2, hits):
    t = ws.cell(r0, 2).value
    m = re.search(r'Hits: A (\d+), φ and √2 (\d+), B (\d+)$', t)
    exp = (hits['A'], hits['P'], hits['B'])
    if not m or tuple(int(x) for x in m.groups()) != exp:
        stats['mism'].append((tid, 'title', t[-40:], exp))
    h1 = [ws.cell(r0 + 1, c).value for c in range(1, 3 + 2 * ncol_plane)]
    e1 = ['table', 'star', 'RA'] + [None] * (ncol_plane - 1) + ['Dec'] + [None] * (ncol_plane - 1)
    if h1 != e1:
        stats['mism'].append((tid, 'header1', h1, e1))
    mg = merged_at(r0 + 1)
    em = [(3, 2 + ncol_plane), (3 + ncol_plane, 2 + 2 * ncol_plane)]
    if mg != em:
        stats['mism'].append((tid, 'merge', mg, em))
    h2 = [ws.cell(r0 + 2, c).value for c in range(3, 3 + 2 * ncol_plane)]
    if h2 != names2 * 2:
        stats['mism'].append((tid, 'header2', h2, names2 * 2))
    if ws.cell(r0 + 3 + len(T1_T2_T24_STARS), 1).value in (tid,):
        stats['mism'].append((tid, 'extra row', r0 + 3 + len(T1_T2_T24_STARS), ''))


# T24 HORSE / JOCKEY Jupiter column cells (for the H / J distance cells)
def t24_jup(block):
    r0 = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == 'T24' and str(ws.cell(r, 2).value).startswith(block + ' '))
    hdr = r0 + 2
    cols = [c for c in range(3, ws.max_column + 1) if ws.cell(hdr, c).value == 'Jupiter']
    out = {}
    r = hdr + 1
    while ws.cell(r, 1).value == 'T24':
        out[ws.cell(r, 2).value] = {'RA': got(ws.cell(r, cols[0])), 'Dec': got(ws.cell(r, cols[1]))}
        r += 1
    return out


T24J = {'H': t24_jup('T24 HORSE'), 'J': t24_jup('T24 JOCKEY')}
counts = {}

# ---- T24 H / J
tid = 'T24 H / J'
r0 = find(tid)
hits = {'A': 0, 'P': 0, 'B': 0}
tc = {'together': 0, 'assisted': 0, 'automatic': 0}
for i, s in enumerate(T1_T2_T24_STARS):
    r = r0 + 3 + i
    if ws.cell(r, 1).value != tid or ws.cell(r, 2).value != s:
        stats['mism'].append((tid, f'row {r}', (ws.cell(r, 1).value, ws.cell(r, 2).value), s))
    for p, pl in enumerate(('RA', 'Dec')):
        c0 = 3 + 3 * p
        dh = dist(pos[('H', s)], pos[('H', 'Jupiter')], pl)
        dj = dist(pos[('J', s)], pos[('J', 'Jupiter')], pl)
        for k, (ch, d) in enumerate((('H', dh), ('J', dj))):
            f, com = deg_cell(d, pl)
            exp = (round(d, 6), f, com)
            cmp(r, c0 + k, exp, f'{tid} {s} {pl} d{ch}')
            stats['dist_cells'] += 1
            t24 = T24J[ch].get(f'{ch} {s}', {}).get(pl)
            g = got(ws.cell(r, c0 + k))
            t24c = (t24[2] or '').split('\n') if t24 else None
            if t24 is None or abs(t24[0] - g[0]) > 5e-7 or t24[1] != g[1] or t24c[0] != g[2]:
                stats['mism'].append((tid, f'{s} {pl} {ch} vs T24 {ch}', g, t24))
            elif len(t24c) > 1:
                stats.setdefault('t24_mid_lines', []).append(f'{ch} {s} {pl}: ' + ' | '.join(t24c[1:]))
        v, f, com, ar, flag = ratio_cell(dh, dj, (f'H {s}, H Jupiter', f'J {s}, J Jupiter'), pl,
                                         (pos[('H', 'Jupiter')], pos[('J', 'Jupiter')]), 'H–J Jupiters')
        cmp(r, c0 + 2, (v, f, com), f'{tid} {s} {pl} ratio')
        stats['ratio_cells'] += 1
        if flag:
            tc[flag] += 1
        if ar and flag != 'automatic':
            hits[ar] += 1
check_title_header(tid, r0, 3, ['d(H star, H Jupiter)', 'd(J star, J Jupiter)', 'H / J'], hits)
counts[tid] = (dict(hits), dict(tc))

# ---- RATIOS
XB = [b for b in T24_BODIES if b not in ('Jupiter', 'Moon')]
for tid, ch in (('T24 HORSE RATIOS', 'H'), ('T24 JOCKEY RATIOS', 'J')):
    r0 = find(tid)
    hits = {'A': 0, 'P': 0, 'B': 0}
    tc = {'together': 0, 'assisted': 0, 'automatic': 0}
    for i, s in enumerate(T1_T2_T24_STARS):
        r = r0 + 3 + i
        if ws.cell(r, 1).value != tid or ws.cell(r, 2).value != f'{ch} {s}':
            stats['mism'].append((tid, f'row {r}', (ws.cell(r, 1).value, ws.cell(r, 2).value), s))
        for p, pl in enumerate(('RA', 'Dec')):
            for j, x in enumerate(XB):
                c = 3 + p * len(XB) + j
                d1 = dist(pos[(ch, s)], pos[(ch, 'Jupiter')], pl)
                d2 = dist(pos[(ch, s)], pos[(ch, x)], pl)
                v, f, com, ar, flag = ratio_cell(d1, d2, (f'{ch} {s}, {ch} Jupiter', f'{ch} {s}, {ch} {x}'), pl,
                                                 (pos[(ch, 'Jupiter')], pos[(ch, x)]), 'Jupiter–X')
                cmp(r, c, (v, f, com), f'{tid} {s} {pl} {x}')
                stats['ratio_cells'] += 1
                if flag:
                    tc[flag] += 1
                if ar and flag != 'automatic':
                    hits[ar] += 1
                if len(hand) < 400 and ar:
                    hand.append((tid, s, pl, x, d1, d2, v, com))
    check_title_header(tid, r0, len(XB), XB, hits)
    counts[tid] = (dict(hits), dict(tc))

# ---- hand re-derivation of a sample (formula written out again, no families scoring)
import math
PHI = (1 + 5 ** 0.5) / 2
VAL = {lab: t for lab, t, *_ in families.INTERVALS}
hand_bad = 0
sample = hand[::7][:40]
for tid, s, pl, x, d1, d2, v, com in sample:
    lab = com.split(' · ')[0]
    t = VAL[lab]
    vv = max(d1, d2) / min(d1, d2)
    if lab in ('1', '3/2', '2', '3', '4', '5', '6', '7', '8', '9'):
        tol, dev = 0.02, abs(vv - t)
    else:
        if lab in ('φ', 'φ²', 'φ³', '1+√2', '√2', '1+1/√2'):
            tol = 0.15
        else:
            k = max(int(z) for z in lab.split('/'))
            tol = 1 / k if k <= 5 else 0.15
        dev = abs(vv / t - 1) * 100
    sc = (10 * (tol - dev) / tol) ** 2
    shown = float(com.split('\n')[0].split(' · ')[-1])
    if abs(round(sc, 2) - shown) > 0.006:
        hand_bad += 1
        print('HAND MISMATCH', tid, s, pl, x, vv, lab, sc, shown)
print(f'hand re-derived scores: {len(sample)}, mismatches {hand_bad}')
for tid in counts:
    print(tid, 'hits', counts[tid][0], 'flags', counts[tid][1])
print('cells checked', stats['cells'], '(distance', stats['dist_cells'], ', ratio', stats['ratio_cells'], ')')
print('T24 cells with extra midpoint lines (not carried into T24 H / J, first line compared):', len(stats.get('t24_mid_lines', [])))
for x in stats.get('t24_mid_lines', []): print('  ', x)
print('mismatches', len(stats['mism']))
for m in stats['mism'][:60]:
    print('  ', m)
