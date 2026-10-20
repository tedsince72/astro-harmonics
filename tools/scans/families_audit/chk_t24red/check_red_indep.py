"""Independent check of the T24 red-conjunction-fill change (own code; no jupiter_tables import).
Usage: python3 check_red_indep.py <before.xlsx> <after.xlsx> [label]"""
import sys
from collections import Counter, defaultdict
import openpyxl

RED = '00FF8080'
NAMES = {'00C6EFCE': 'green', '00FFEB9C': 'yellow', '00FFF8DC': 'light'}


def fillkey(c):
    f = c.fill
    if f is None or f.patternType is None:
        return None
    return (f.patternType, f.fgColor.rgb if f.fgColor is not None else None)


def bordkey(c):
    b = c.border
    return tuple((getattr(b, s).style if getattr(b, s) is not None else None) for s in ('left', 'right', 'top', 'bottom'))


def comtext(c):
    return c.comment.text if c.comment is not None else None


def positions(wb):
    ws = wb['POSITIONS']
    pos = {}
    for r in ws.iter_rows(values_only=True):
        if r[0] is None or r[2] not in ('S', 'H', 'J') or not isinstance(r[5], (int, float)):
            continue
        key = ('SKY', 'S') if r[0] == 'SKY' else (r[0], r[2])
        pos.setdefault(key, {})[r[3]] = (float(r[5]), float(r[6]))
    return pos


def dra(a, b):
    d = abs(a - b) % 360
    return min(d, 360 - d)


def blocks(ws):
    """yield (kind, title_row, hdr1, hdr2, data_rows) for T24 SKY/HORSE/JOCKEY and T24 H / J."""
    rows = list(ws.iter_rows(min_col=1, max_col=2, values_only=True))
    n = len(rows)
    i = 0
    while i < n:
        a, b = rows[i]
        kind = None
        if a == 'T24' and isinstance(b, str):
            for k in ('T24 SKY', 'T24 HORSE', 'T24 JOCKEY'):
                if b.startswith(k):
                    kind = k
        elif a == 'T24 H / J' and isinstance(b, str) and b.startswith('T24 H / J'):
            kind = 'T24 H / J'
        if kind:
            tr = i + 1
            h1, h2 = tr + 1, tr + 2
            assert rows[h1 - 1][0] == 'table', (ws.title, kind, tr)
            data = []
            j = h2 + 1
            while j <= n and rows[j - 1][0] == a and not (isinstance(rows[j - 1][1], str) and rows[j - 1][1].startswith(kind)):
                data.append(j)
                j += 1
            yield kind, tr, h1, h2, data
            i = j - 1
        else:
            i += 1


def main(bp, ap, label):
    wa = openpyxl.load_workbook(ap)
    wbf = openpyxl.load_workbook(bp)
    pos = positions(wa)
    targets = {}  # (sheet,row,col) -> expect_red
    counts = Counter()
    prev = Counter()
    mism = []
    nval = 0
    for sh in wa.sheetnames:
        if sh != 'SKY JUPITER' and not (sh.startswith('P') and sh[1:].isdigit()):
            continue
        ws = wa[sh]
        maxc = ws.max_column
        for kind, tr, h1, h2, data in blocks(ws):
            hdr1 = [ws.cell(h1, c).value for c in range(1, maxc + 1)]
            hdr2 = [ws.cell(h2, c).value for c in range(1, maxc + 1)]
            # plane per column from header row 1 (merged start)
            plane = None
            cols = []
            for c in range(3, maxc + 1):
                v = hdr1[c - 1]
                if v in ('RA', 'Dec'):
                    plane = v
                elif v in ('RA midpoints', 'Dec midpoints'):
                    plane = None
                    continue
                if plane and hdr2[c - 1] is not None:
                    if kind == 'T24 H / J' and hdr2[c - 1] == 'H / J':
                        continue
                    cols.append((c, plane, hdr2[c - 1]))
            if kind == 'T24 SKY':
                chart = ('SKY', 'S'); pref = 'S '
            elif kind == 'T24 HORSE':
                chart = (sh, 'H'); pref = 'H '
            elif kind == 'T24 JOCKEY':
                chart = (sh, 'J'); pref = 'J '
            blk = (label, sh, kind)
            counts[blk] += 0
            for r in data:
                lab = ws.cell(r, 2).value
                for c, pl, name in cols:
                    if kind == 'T24 H / J':
                        ch = 'H' if name.startswith('d(H') else 'J'
                        P = pos[(sh, ch)]
                        star, body = lab, 'Jupiter'
                    else:
                        P = pos[chart]
                        assert lab.startswith(pref), lab
                        star, body = lab[len(pref):], name
                    s, bd = P[star], P[body]
                    d = dra(s[0], bd[0]) if pl == 'RA' else abs(s[1] - bd[1])
                    cell = ws.cell(r, c)
                    nval += 1
                    if not isinstance(cell.value, (int, float)) or round(cell.value, 6) != round(d, 6):
                        mism.append(('value', sh, cell.coordinate, cell.value, round(d, 6)))
                    exp_red = d < (1.5 if pl == 'RA' else 0.3)
                    targets[(sh, r, c)] = exp_red
                    fa = fillkey(cell)
                    fb = fillkey(wbf[sh].cell(r, c))
                    if exp_red:
                        counts[blk] += 1
                        if fa != ('solid', RED):
                            mism.append(('notred', sh, cell.coordinate, fa))
                        prev[NAMES.get(fb[1], fb[1]) if fb else 'unfilled'] += 1
                    elif fa != fb:
                        mism.append(('fillchg', sh, cell.coordinate, fb, fa))
    # every cell of every sheet
    diffs = []
    redelse = []
    ncell = 0
    for sh in wa.sheetnames:
        A, B = wa[sh], wbf[sh]
        mr = max(A.max_row, B.max_row); mc = max(A.max_column, B.max_column)
        for r in range(1, mr + 1):
            for c in range(1, mc + 1):
                ca, cb = A.cell(r, c), B.cell(r, c)
                ncell += 1
                fa, fb = fillkey(ca), fillkey(cb)
                if fb and fb[1] == RED:
                    redelse.append(('before', sh, cb.coordinate))
                if fa and fa[1] == RED and not targets.get((sh, r, c)):
                    redelse.append(('after', sh, ca.coordinate))
                what = []
                if ca.value != cb.value: what.append('value')
                if comtext(ca) != comtext(cb): what.append('comment')
                if bordkey(ca) != bordkey(cb): what.append('border')
                if (sh, r, c) not in targets and fa != fb: what.append('fill')
                if what:
                    diffs.append((sh, ca.coordinate, what, str(A.cell(r, 1).value)))
    print(f'== {label}: {ap}')
    print('body/H-J distance cells checked:', nval, ' all cells compared:', ncell)
    for k, v in sorted(counts.items()):
        print('  red', k[1], k[2], v)
    print('  red total', sum(counts.values()), ' previously:', dict(prev))
    print('mismatches:', len(mism)); [print('  ', m) for m in mism[:20]]
    print('red outside targets:', len(redelse)); [print('  ', m) for m in redelse[:20]]
    print('differing cells:', len(diffs)); [print('  ', d) for d in diffs[:40]]
    return len(mism) + len(redelse)


if __name__ == '__main__':
    sys.exit(1 if main(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else 'race') else 0)
