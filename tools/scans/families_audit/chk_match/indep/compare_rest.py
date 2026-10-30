"""Every non-MATCH cell identical before vs after (values, fills, fonts, borders, alignment, formats, comments), merges, widths, freeze."""
import sys, openpyxl
from openpyxl.utils import get_column_letter


def sig(c):
    f = c.fill; b = c.border; fo = c.font; al = c.alignment
    return (c.value, f.fill_type, f.fgColor.rgb if f.fgColor is not None else None,
            fo.b, fo.i, fo.color.rgb if fo.color is not None else None, fo.sz,
            tuple((getattr(b, s).style, getattr(b, s).color.rgb if getattr(b, s).color is not None else None)
                  for s in ('left', 'right', 'top', 'bottom')),
            al.wrap_text, al.horizontal, al.vertical, c.number_format,
            c.comment.text if c.comment else None)


def main(old, new, tag):
    wo, wn = openpyxl.load_workbook(old), openpyxl.load_workbook(new)
    bad = []
    if wo.sheetnames != wn.sheetnames: bad.append(('sheets', wo.sheetnames, wn.sheetnames))
    ncell = 0; match_rows = 0
    for sn in wo.sheetnames:
        a, b = wo[sn], wn[sn]
        mrows = {r for r in range(1, b.max_row + 1) if b.cell(r, 1).value in ('MATCH HORSE', 'MATCH JOCKEY')}
        match_rows += len(mrows)
        if mrows:
            first = min(mrows)
            # MATCH region = from the first MATCH row to the end; rows before identical, old has nothing there
            if a.max_row >= first - 1 and any(a.cell(r, c).value is not None for r in range(first - 1, a.max_row + 1) for c in range(1, a.max_column + 1)):
                bad.append((sn, 'old has content at/after MATCH start', first))
            # rows inside region that are not MATCH rows must be empty (blank separators)
            for r in range(first, b.max_row + 1):
                if r not in mrows and not (b.cell(r, 1).value == 'table' and (r - 1) in mrows and isinstance(b.cell(r - 1, 2).value, str) and b.cell(r - 1, 2).value.startswith(b.cell(r - 1, 1).value + ' (Eddie')) and any(b.cell(r, c).value is not None for c in range(1, b.max_column + 1)):
                    bad.append((sn, 'non-MATCH content in MATCH region', r))
            last = first - 1
        else:
            last = max(a.max_row, b.max_row)
        mc = max(a.max_column, b.max_column)
        for r in range(1, last + 1):
            for c in range(1, mc + 1):
                ncell += 1
                if sig(a.cell(r, c)) != sig(b.cell(r, c)):
                    bad.append((sn, r, c, sig(a.cell(r, c))[:3], sig(b.cell(r, c))[:3]))
        mo = sorted(str(m) for m in a.merged_cells.ranges)
        mn = sorted(str(m) for m in b.merged_cells.ranges if not (mrows and m.min_row >= min(mrows)))
        if mo != mn: bad.append((sn, 'merges', set(mo) ^ set(mn)))
        cols = set(a.column_dimensions) | set(b.column_dimensions)
        for k in cols:
            if a.column_dimensions[k].width != b.column_dimensions[k].width:
                bad.append((sn, 'width', k, a.column_dimensions[k].width, b.column_dimensions[k].width))
        if a.freeze_panes != b.freeze_panes: bad.append((sn, 'freeze', a.freeze_panes, b.freeze_panes))
    print(tag, 'cells compared', ncell, 'MATCH rows', match_rows, 'differences', len(bad))
    for x in bad[:30]: print('  DIFF', x)
    return bad


if __name__ == '__main__':
    main(*sys.argv[1:4])
