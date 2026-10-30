"""Independent checker for the tightened MATCH RATIO rule (Eddie, 30 Oct 2026).
Usage: check_ratio_rule.py old.xlsx new.xlsx label
 (1) every cell outside MATCH identical old vs new; EXACT rows identical after the column shift (cols 6, 7 blank)
 (2) MATCH rebuilt from POSITIONS (own family code fam.py, own T24 / SKY readers from check_match.py) with the new rules
 (3) every new RATIO row = an old RATIO row (mapped); every old RATIO row missing fails the new rule (by reason)
 (4) counts before -> after."""
import sys, os, math, collections
from fractions import Fraction
import openpyxl
sys.path.insert(0, os.path.dirname(__file__))
from fam import d4s
from check_match import load, sky_cells, t24_cells, fillc
from compare_rest import sig

TABS = ('MATCH HORSE', 'MATCH JOCKEY')
NEWHDR = ['table', 'section', 'ratio p/q', 'ratio', 'shared factor', 'circle factor', 'slice of 360', 'sky: relationship',
          'plane', 'value', 'family', 'target', 'score', 'T24: row', 'body', 'plane', 'value', 'family', 'target', 'score',
          'lower score']
NC = 21


def build(sky, t24):
    ex, ra = [], []
    for s in sky:
        for t in t24:
            if s['fam'] == t['fam']:
                if s['tid'] == t['tid']:
                    ex.append(dict(s=s, t=t, g=None, fr=Fraction(1), cf=None))
                elif s['fam'] == 'whole':
                    a, b = s['tid'], t['tid']
                    cf = math.gcd(math.gcd(a, b), 360)
                    if cf >= 5:
                        ra.append(dict(s=s, t=t, g=math.gcd(a, b), fr=Fraction(max(a, b), min(a, b)), cf=cf))
                elif s['fam'] == 'golden':
                    a, b = s['tid'], t['tid']
                    g = math.gcd(a, b)
                    fr = Fraction(max(a, b), min(a, b))
                    if g >= 2 and fr.numerator <= 6 and fr.denominator <= 6:
                        ra.append(dict(s=s, t=t, g=g, fr=fr, cf=None))
            elif d4s(s['v']) == d4s(t['v']):
                ex.append(dict(s=s, t=t, g=None, fr=Fraction(1), cf=None))
    lou = lambda m: min(m['s']['sc'], m['t']['sc'])
    ex.sort(key=lambda m: -lou(m))
    ra.sort(key=lambda m: (m['fr'], -lou(m)))
    return ex, ra


def exp_row(tab, sec, m):
    s, t, g, fr, cf = m['s'], m['t'], m['g'], m['fr'], m['cf']
    pq = '1' if sec == 'EXACT' else '%d/%d' % (fr.numerator, fr.denominator)
    gf = None if g is None else (str(g) if s['fam'] == 'whole' else '%d×φ' % g)
    return [tab, sec, pq, round(float(fr), 6), gf, cf, None if cf is None else '1/%d' % (360 // cf),
            s['rel'], s['plane'], round(s['v'], 6), s['fam'], s['lab'], round(s['sc'], 2),
            t['rel'], t['body'], t['plane'], round(t['v'], 6), t['fam'], t['lab'], round(t['sc'], 2),
            min(round(s['sc'], 2), round(t['sc'], 2))]


def eq(a, b):
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool):
        return abs(a - b) < 1e-9
    return a == b


def rgb(c):
    f = c.fill
    return f.fgColor.rgb[-6:] if f is not None and f.fill_type == 'solid' else None


def tables(ws):
    """{tab: (title_row, header_row, [data rows])} rows as lists of cells."""
    out = {}
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row))
    for tab in TABS:
        idx = [i for i, r in enumerate(rows) if r[0].value == tab]
        if idx:
            out[tab] = (rows[idx[0]], rows[idx[0] + 1], [rows[i] for i in idx[1:]])
    return out, rows


def oldmap(cells):
    """old 19-col row -> new 21-col layout (insert blank circle factor / slice)."""
    return cells[:5] + [None, None] + cells[5:19]


def rule_fail(r):
    """old RATIO row (new layout values): reason it fails the new rule, or None if it passes."""
    fam = r[10]
    if fam == 'whole':
        a, b = int(r[11]), int(r[18])
        cf = math.gcd(math.gcd(a, b), 360)
        return None if cf >= 5 else 'whole cf %d' % cf
    if fam == 'golden':
        a, b = int(r[11].split('×')[0]), int(r[18].split('×')[0])
        fr = Fraction(max(a, b), min(a, b))
        return None if (fr.numerator <= 6 and fr.denominator <= 6) else 'golden p or q > 6'
    return 'other family ' + str(fam)


def main(old, new, label, quiet=False):
    P = (lambda *a: None) if quiet else print
    msgs = []
    wo = openpyxl.load_workbook(old)
    wn, pos = load(new)
    if wo.sheetnames != wn.sheetnames:
        msgs.append(('sheets', wo.sheetnames, wn.sheetnames))
    sky, bad = sky_cells(wn, pos)
    msgs += bad
    ncell = 0
    tot = collections.Counter()
    for sn in wn.sheetnames:
        a, b = wo[sn], wn[sn]
        ta, rows_a = tables(a)
        tb, rows_b = tables(b)
        # ---- (1) outside MATCH
        fa = min((r[0].row for t in ta.values() for r in [t[0]]), default=None)
        fb = min((r[0].row for t in tb.values() for r in [t[0]]), default=None)
        if fa != fb:
            msgs.append((sn, 'MATCH start', fa, fb))
        last = (fb - 1) if fb else max(a.max_row, b.max_row)
        mc = max(a.max_column, b.max_column)
        for r in range(1, last + 1):
            for c in range(1, mc + 1):
                ncell += 1
                if sig(a.cell(r, c)) != sig(b.cell(r, c)):
                    msgs.append((sn, 'outside cell', r, c, sig(a.cell(r, c))[:3], sig(b.cell(r, c))[:3]))
        if fb:
            # nothing but MATCH title/header/data/blank rows after the start
            for r in range(fb, b.max_row + 1):
                v1 = b.cell(r, 1).value
                if v1 not in TABS and v1 != 'table' and any(b.cell(r, c).value is not None for c in range(1, b.max_column + 1)):
                    msgs.append((sn, 'non-MATCH content in MATCH region', r))
        mo = sorted(str(m) for m in a.merged_cells.ranges if not (fa and m.min_row >= fa))
        mn = sorted(str(m) for m in b.merged_cells.ranges if not (fb and m.min_row >= fb))
        if mo != mn:
            msgs.append((sn, 'merges outside', set(mo) ^ set(mn)))
        mm = [str(m) for m in b.merged_cells.ranges if fb and m.min_row >= fb]
        if mm:
            P('  note', sn, 'merges in MATCH region', mm[:5])
        for k in set(a.column_dimensions) | set(b.column_dimensions):
            if a.column_dimensions[k].width != b.column_dimensions[k].width:
                msgs.append((sn, 'width', k, a.column_dimensions[k].width, b.column_dimensions[k].width))
        if a.freeze_panes != b.freeze_panes:
            msgs.append((sn, 'freeze', a.freeze_panes, b.freeze_panes))
        if not (sn.startswith('P') and sn[1:].isdigit()):
            if ta or tb:
                msgs.append((sn, 'MATCH on non-runner sheet'))
            continue
        for ch, tab in (('H', 'MATCH HORSE'), ('J', 'MATCH JOCKEY')):
            if tab not in ta or tab not in tb:
                msgs.append((sn, tab, 'missing table')); continue
            told, hold, dold = ta[tab]
            tnew, hnew, dnew = tb[tab]
            ov = [oldmap([c.value for c in r[:19]]) for r in dold]
            of = [oldmap([rgb(c) for c in r[:19]]) for r in dold]
            nv = [[c.value for c in r[:NC]] for r in dnew]
            nf = [[rgb(c) for c in r[:NC]] for r in dnew]
            # extra columns beyond 21 empty
            for r in dnew:
                if any(c.value is not None for c in r[NC:]):
                    msgs.append((sn, tab, 'content beyond col 21', r[0].row))
            # (1b) EXACT rows identical after the shift (values + fills, incl. blank cols)
            oe = [(v, f) for v, f in zip(ov, of) if v[1] == 'EXACT']
            ne = [(v, f) for v, f in zip(nv, nf) if v[1] == 'EXACT']
            if len(oe) != len(ne):
                msgs.append((sn, tab, 'EXACT count', len(oe), len(ne)))
            for j, ((v1, f1), (v2, f2)) in enumerate(zip(oe, ne)):
                if not all(eq(x, y) for x, y in zip(v1, v2)) or f1 != f2:
                    msgs.append((sn, tab, 'EXACT row differs', j, v1, v2))
            # (2) rebuild
            t24, b2 = t24_cells(b, pos, sn, ch)
            msgs += b2
            ex, ra = build(sky, t24)
            exp = [exp_row(tab, 'EXACT', m) for m in ex] + [exp_row(tab, 'RATIO', m) for m in ra]
            nw = sum(1 for m in ra if m['s']['fam'] == 'whole'); ng = len(ra) - nw
            want = 'EXACT %d rows, RATIO %d rows (whole %d, golden %d) (SKY block %d scoring cells, %s %d)' % (
                len(ex), len(ra), nw, ng, len(sky), 'T24 HORSE' if ch == 'H' else 'T24 JOCKEY', len(t24))
            title = tnew[1].value or ''
            if not title.endswith(want):
                msgs.append((sn, tab, 'title', title[-110:], want))
            for phrase in ('gcd(A, B, 360) >= 5', 'p, q <= 6', 'slice of 360'):
                if phrase not in title:
                    msgs.append((sn, tab, 'title phrase missing', phrase))
            if [c.value for c in hnew[:NC]] != NEWHDR or any(c.value is not None for c in hnew[NC:]):
                msgs.append((sn, tab, 'header', [c.value for c in hnew[:NC + 1]]))
            if len(nv) != len(exp):
                msgs.append((sn, tab, 'row count', len(nv), len(exp)))
            key = lambda r: (r[1], r[7], r[8], r[13], r[14], r[15])
            gk = collections.Counter(key(r) for r in nv); ek = collections.Counter(key(r) for r in exp)
            for k in gk - ek: msgs.append((sn, tab, 'extra row', k))
            for k in ek - gk: msgs.append((sn, tab, 'missing row', k))
            nd = 0
            for j, (g, f, e) in enumerate(zip(nv, nf, exp)):
                for c in range(NC):
                    if not eq(g[c], e[c]):
                        nd += 1
                        if nd <= 10: msgs.append((sn, tab, 'cell', j, c + 1, g[c], e[c]))
                    wf = fillc(e[c]) if (c + 1) in (13, 20, 21) else None
                    if f[c] != wf:
                        nd += 1
                        if nd <= 10: msgs.append((sn, tab, 'fill', j, c + 1, f[c], wf))
            # (3) after RATIO subset of before RATIO; dropped by reason
            orat = [(v, f) for v, f in zip(ov, of) if v[1] == 'RATIO']
            nrat = [(v, f) for v, f in zip(nv, nf) if v[1] == 'RATIO']
            strip = lambda v: tuple(round(x, 9) if isinstance(x, float) else x for i, x in enumerate(v) if i not in (5, 6))
            ocnt = collections.Counter((strip(v), tuple(f)) for v, f in orat)
            used = collections.Counter()
            for v, f in nrat:
                k = (strip(v), tuple(f))
                if used[k] >= ocnt[k]:
                    msgs.append((sn, tab, 'after RATIO row not in before', v[:5], v[7], v[13], v[14]))
                used[k] += 1
            # order: after RATIO = before RATIO order filtered
            kept_order = []
            rem = collections.Counter(used)
            for v, f in orat:
                k = (strip(v), tuple(f))
                if rem[k] > 0:
                    kept_order.append(k); rem[k] -= 1
            if kept_order != [(strip(v), tuple(f)) for v, f in nrat]:
                msgs.append((sn, tab, 'after RATIO order != before order filtered'))
            dropped = ocnt - used
            reasons = collections.Counter()
            for k, n in dropped.items():
                v = list(k[0][:5]) + [None, None] + list(k[0][5:])
                why = rule_fail(v)
                if why is None:
                    msgs.append((sn, tab, 'dropped row passes new rule', v[:5], v[7], v[13], v[14]))
                    why = 'PASSES (error)'
                reasons[why.split(' cf')[0] + (' cf<5' if 'cf' in why else '')] += n
                reasons['detail ' + why] += n
            ow = sum(1 for v, f in orat if v[10] == 'whole'); og = len(orat) - ow
            P('%s %s %s: EXACT %d -> %d, RATIO whole %d -> %d, golden %d -> %d; dropped %s' % (
                label, sn, tab, len(oe), len(ne), ow, nw, og, ng,
                dict(sorted((k, v) for k, v in reasons.items() if not k.startswith('detail')))))
            P('    dropped detail', dict(sorted((k[7:], v) for k, v in reasons.items() if k.startswith('detail'))))
            tot['EXACT before'] += len(oe); tot['EXACT after'] += len(ne)
            tot['whole before'] += ow; tot['whole after'] += nw; tot['golden before'] += og; tot['golden after'] += ng
            for k, v in reasons.items():
                if not k.startswith('detail'): tot['dropped ' + k] += v
    P(label, 'outside cells compared', ncell, dict(tot))
    for m in msgs[:40]:
        P('  MISMATCH', m)
    P(label, 'mismatches:', len(msgs))
    return msgs


if __name__ == '__main__':
    main(*sys.argv[1:4])
