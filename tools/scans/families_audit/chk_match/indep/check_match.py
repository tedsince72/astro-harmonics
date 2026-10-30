"""Independent checker: MATCH HORSE / MATCH JOCKEY. Usage: check_match.py new.xlsx [label]"""
import sys, math, collections
from fractions import Fraction
import openpyxl
sys.path.insert(0, __import__('os').path.dirname(__file__))
from fam import best, d4s

GREEN, YELLOW, LIGHT = 'C6EFCE', 'FFEB9C', 'FFF8DC'


def fillc(s):
    return GREEN if s >= 80 else YELLOW if s >= 50 else LIGHT


def ra_d(a, b):
    d = abs(a - b) % 360
    return 360 - d if d > 180 else d


def load(path):
    wb = openpyxl.load_workbook(path)
    pos = {}
    for r in wb['POSITIONS'].iter_rows(values_only=True):
        if r[2] in ('S', 'H', 'J') and r[5] is not None:
            pos[(r[0], r[2], r[3])] = (float(r[5]), float(r[6]))
    return wb, pos


def sky_cells(wb, pos):
    ws = wb['SKY JUPITER']
    sj = pos[('SKY', 'S', 'Jupiter')]
    cells, n, bad = [], 0, []
    started = False
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        a, b = row[0].value, row[1].value
        if a == 'SKY' and b == 'S Jupiter → sky points':
            started = True; continue
        if not started:
            continue
        if a is None:
            break
        if a != 'SKY' or not str(b).startswith('S Jupiter → S '):
            continue
        x = b[len('S Jupiter → S '):]
        if x == 'equator':
            vals = {'Dec': abs(sj[1])}
        else:
            p = pos[('SKY', 'S', x)]
            vals = {'RA': ra_d(sj[0], p[0]), 'Dec': abs(sj[1] - p[1])}
        for pl, v in vals.items():
            col = 2 if pl == 'RA' else 11
            sv = row[col].value
            if sv is None or abs(sv - round(v, 6)) > 1.5e-6:
                bad.append(('SKY value', b, pl, sv, v))
            bt = best(v, pl)
            if bt and bt[0] != 'conjunction':
                cells.append(dict(rel=b, plane=pl, v=v, fam=bt[0], tid=bt[1], lab=bt[2], sc=bt[3]))
    return cells, bad


def t24_cells(ws, pos, runner, ch):
    """ch 'H' or 'J'; returns scoring non-conj cells of T24 HORSE/JOCKEY (no S Jupiter col, no S Jupiter → S body row)."""
    name = 'T24 HORSE' if ch == 'H' else 'T24 JOCKEY'
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row))
    st = None
    for i, r in enumerate(rows):
        if r[0].value == 'T24' and isinstance(r[1].value, str) and r[1].value.startswith(name + ' ('):
            st = i; break
    hdr = rows[st + 2]
    cols = []  # (colidx, body, plane)
    plane = 'RA'
    for c in hdr[2:]:
        v = c.value
        if v is None:
            continue
        if v.startswith('star = middle'):
            if plane == 'RA':
                plane = 'Dec'; continue
            break
        if v.startswith('target family'):
            break
        cols.append((c.column - 1, v, plane))
    sj = pos[('SKY', 'S', 'Jupiter')]
    nj = pos[(runner, ch, 'Jupiter')]
    out, bad = [], []
    i = st + 3
    while i < len(rows) and rows[i][0].value == 'T24':
        lab = rows[i][1].value
        r = rows[i]; i += 1
        if lab in ('stars',) or lab == 'S Jupiter → S body':
            continue
        for ci, body, pl in cols:
            if body == 'S Jupiter':
                continue
            if lab == '%s Jupiter → %s body' % (ch, ch):
                if body == 'Jupiter': continue
                a, b = nj, pos[(runner, ch, body)]
            elif lab == 'S Jupiter → %s body' % ch:
                a, b = sj, pos[(runner, ch, body)]
            elif lab == 'S body → %s Jupiter' % ch:
                a, b = (sj if body == 'Jupiter' else pos[('SKY', 'S', body)]), nj
            else:
                assert lab.startswith(ch + ' ')
                a, b = pos[(runner, ch, lab[2:])], pos[(runner, ch, body)]
            v = ra_d(a[0], b[0]) if pl == 'RA' else abs(a[1] - b[1])
            sv = r[ci].value
            if sv is None or abs(sv - round(v, 6)) > 1.5e-6:
                bad.append(('T24 value', runner, lab, body, pl, sv, v))
            bt = best(v, pl)
            if bt and bt[0] != 'conjunction':
                out.append(dict(rel=lab, body=body, plane=pl, v=v, fam=bt[0], tid=bt[1], lab=bt[2], sc=bt[3],
                                inconj=(v < 1.5 if pl == 'RA' else v < 0.3)))
    return out, bad


def build(sky, t24):
    ex, ra = [], []
    for s in sky:
        for t in t24:
            if s['fam'] == t['fam']:
                if s['tid'] == t['tid']:
                    ex.append((s, t, None, Fraction(1), False))
                elif s['fam'] in ('whole', 'golden'):
                    g = math.gcd(s['tid'], t['tid'])
                    if g >= 2:
                        fr = Fraction(max(s['tid'], t['tid']), min(s['tid'], t['tid']))
                        ra.append((s, t, g, fr, False))
            else:
                if d4s(s['v']) == d4s(t['v']):
                    ex.append((s, t, None, Fraction(1), True))
    lo = lambda m: min(round(m[0]['sc'], 2), round(m[1]['sc'], 2))
    lou = lambda m: min(m[0]['sc'], m[1]['sc'])
    ex.sort(key=lambda m: -lou(m))
    ra.sort(key=lambda m: (m[3], -lou(m)))
    return ex, ra


def exp_row(tab, sec, m):
    s, t, g, fr, cross = m
    pq = '1' if sec == 'EXACT' else '%d/%d' % (fr.numerator, fr.denominator)
    gf = None if g is None else (str(g) if s['fam'] == 'whole' else '%d×φ' % g)
    return [tab, sec, pq, round(float(fr), 6), gf, s['rel'], s['plane'], round(s['v'], 6), s['fam'], s['lab'],
            round(s['sc'], 2), t['rel'], t['body'], t['plane'], round(t['v'], 6), t['fam'], t['lab'], round(t['sc'], 2),
            min(round(s['sc'], 2), round(t['sc'], 2))]


FILLCOLS = (10, 17, 18)


def cmp_table(ws, tab, ex, ra, nsky, nt24, msgs, tag):
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row))
    st = [i for i, r in enumerate(rows) if r[0].value == tab]
    if not st:
        msgs.append((tag, tab, 'missing table')); return
    i0 = st[0]
    title = rows[i0][1].value or ''
    want = 'EXACT %d rows, RATIO %d rows (SKY block %d scoring cells, %s %d)' % (
        len(ex), len(ra), nsky, 'T24 HORSE' if tab.endswith('HORSE') else 'T24 JOCKEY', nt24)
    if not title.endswith(want):
        msgs.append((tag, tab, 'title', title[-90:], want))
    hdr = [c.value for c in rows[i0 + 1][:19]]
    if hdr != ['table', 'section', 'ratio p/q', 'ratio', 'shared factor', 'sky: relationship', 'plane', 'value', 'family',
               'target', 'score', 'T24: row', 'body', 'plane', 'value', 'family', 'target', 'score', 'lower score']:
        msgs.append((tag, tab, 'header', hdr))
    got = [r for r in rows[i0 + 2:] if r[0].value == tab]
    exp = [exp_row(tab, 'EXACT', m) for m in ex] + [exp_row(tab, 'RATIO', m) for m in ra]
    if len(got) != len(exp):
        msgs.append((tag, tab, 'row count', len(got), len(exp)))
    # set comparison
    key = lambda r: tuple(r[:2] + r[5:7] + r[11:14])
    gk = collections.Counter(key([c.value for c in r[:19]]) for r in got)
    ek = collections.Counter(key(r) for r in exp)
    for k in (gk - ek): msgs.append((tag, tab, 'extra row', k))
    for k in (ek - gk): msgs.append((tag, tab, 'missing row', k))
    nd = 0
    for j, (g, e) in enumerate(zip(got, exp)):
        gv = [c.value for c in g[:19]]
        for c in range(19):
            a, b = gv[c], e[c]
            okv = (abs(a - b) < 1e-9) if isinstance(a, (int, float)) and isinstance(b, (int, float)) else a == b
            if not okv:
                nd += 1
                if nd <= 15: msgs.append((tag, tab, 'cell', j, c + 1, a, b))
        for c in range(19):
            f = g[c].fill
            rgb = f.fgColor.rgb[-6:] if f is not None and f.fill_type == 'solid' else None
            want_f = fillc(e[c]) if (c + 1) in (11, 18, 19) else None
            if rgb != want_f:
                nd += 1
                if nd <= 15: msgs.append((tag, tab, 'fill', j, c + 1, rgb, want_f))
    if nd > 15: msgs.append((tag, tab, 'more differences', nd))


def main(path, label):
    wb, pos = load(path)
    sky, bad = sky_cells(wb, pos)
    msgs = list(bad)
    stats = collections.Counter()
    for sn in wb.sheetnames:
        if not (sn.startswith('P') and sn[1:].isdigit()):
            continue
        ws = wb[sn]
        for ch, tab in (('H', 'MATCH HORSE'), ('J', 'MATCH JOCKEY')):
            t24, b2 = t24_cells(ws, pos, sn, ch)
            msgs += b2
            ex, ra = build(sky, t24)
            cross = sum(1 for m in ex if m[4])
            inconj = sum(1 for t in t24 if t['inconj'])
            inconj_used = sum(1 for t in t24 if t['inconj'] and any(m[1] is t for m in ex + ra))
            print('%s %s %s: SKY %d x T24 %d -> EXACT %d (cross-family %d), RATIO %d; T24 scoring cells in conj range %d (in a match %d) %s' % (
                label, sn, tab, len(sky), len(t24), len(ex), cross, len(ra), inconj, inconj_used,
                sorted(collections.Counter((t['fam'], t['lab']) for t in t24 if t['inconj']).items())))
            stats['exact'] += len(ex); stats['ratio'] += len(ra); stats['cross'] += cross
            stats['inconj'] += inconj; stats['maxrows'] = max(stats['maxrows'], len(ex) + len(ra))
            cmp_table(ws, tab, ex, ra, len(sky), len(t24), msgs, sn)
    print(label, 'SKY scoring', len(sky), dict(stats))
    for m in msgs[:60]:
        print('  MISMATCH', m)
    print(label, 'mismatches:', len(msgs))
    return msgs


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else sys.argv[1])
