"""Plant errors in copies of the after workbook (Kelso) and confirm check_ratio_rule catches each.
Usage: plants_ratio_rule.py old.xlsx new.xlsx tmpdir"""
import sys, os, copy, openpyxl
from openpyxl.styles import PatternFill
sys.path.insert(0, os.path.dirname(__file__))
import check_ratio_rule

old, new, tmp = sys.argv[1:4]
os.makedirs(tmp, exist_ok=True)


def rows(ws, tab, sec):
    return [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == tab and ws.cell(r, 2).value == sec]


def first_ratio(ws, tab, fam):
    return [r for r in rows(ws, tab, 'RATIO') if ws.cell(r, 11).value == fam][0]


def p1(ws):
    """wrong circle factor"""
    r = first_ratio(ws, 'MATCH HORSE', 'whole'); ws.cell(r, 6).value = ws.cell(r, 6).value * 2
def p2(ws):
    """wrong slice of 360"""
    r = first_ratio(ws, 'MATCH HORSE', 'whole'); ws.cell(r, 7).value = '1/72'
def p3(ws):
    """golden 44/43 row put back (fake, 88×φ / 86×φ)"""
    r = first_ratio(ws, 'MATCH HORSE', 'golden'); ws.insert_rows(r)
    for c in range(1, 22):
        ws.cell(r, c).value = ws.cell(r + 1, c).value; ws.cell(r, c).fill = copy.copy(ws.cell(r + 1, c).fill)
    ws.cell(r, 3).value = '44/43'; ws.cell(r, 4).value = 1.023256; ws.cell(r, 5).value = '2×φ'
    ws.cell(r, 12).value = '88×φ'; ws.cell(r, 19).value = '86×φ'
def p4(ws):
    """whole cf 2 row put back (44 / 66)"""
    r = first_ratio(ws, 'MATCH JOCKEY', 'whole'); ws.insert_rows(r)
    for c in range(1, 22):
        ws.cell(r, c).value = ws.cell(r + 1, c).value; ws.cell(r, c).fill = copy.copy(ws.cell(r + 1, c).fill)
    ws.cell(r, 3).value = '3/2'; ws.cell(r, 4).value = 1.5; ws.cell(r, 5).value = '22'; ws.cell(r, 6).value = 2
    ws.cell(r, 7).value = '1/180'; ws.cell(r, 12).value = '44'; ws.cell(r, 19).value = '66'
def p5(ws):
    """drop a RATIO row"""
    ws.delete_rows(rows(ws, 'MATCH JOCKEY', 'RATIO')[3])
def p6(ws):
    """wrong fill on a score cell"""
    r = rows(ws, 'MATCH HORSE', 'RATIO')[2]; ws.cell(r, 13).fill = PatternFill('solid', fgColor='C6EFCE' if ws.cell(r, 13).value < 80 else 'FFF8DC')
def p7(ws):
    """circle columns filled on a golden row"""
    r = first_ratio(ws, 'MATCH HORSE', 'golden'); ws.cell(r, 6).value = 2; ws.cell(r, 7).value = '1/180'
def p8(ws):
    """EXACT row: value in an inserted blank column"""
    r = rows(ws, 'MATCH JOCKEY', 'EXACT')[0]; ws.cell(r, 6).value = 1
def p9(ws):
    """title whole/golden split wrong"""
    r = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == 'MATCH HORSE'][0]
    ws.cell(r, 2).value = ws.cell(r, 2).value.replace('(whole 60, golden 8)', '(whole 59, golden 9)')
def p10(ws):
    """swap two RATIO rows with different ratios (order)"""
    a, b = rows(ws, 'MATCH HORSE', 'RATIO')[0], rows(ws, 'MATCH HORSE', 'RATIO')[-1]
    for c in range(1, 22):
        x, y = ws.cell(a, c), ws.cell(b, c)
        xv, xf = x.value, copy.copy(x.fill); x.value, x.fill = y.value, copy.copy(y.fill); y.value, y.fill = xv, xf
def p11(ws):
    """shared factor on a whole row replaced by the circle factor where they differ (or +1)"""
    for r in rows(ws, 'MATCH HORSE', 'RATIO'):
        if ws.cell(r, 11).value == 'whole' and str(ws.cell(r, 6).value) != ws.cell(r, 5).value:
            ws.cell(r, 5).value = str(ws.cell(r, 6).value); return
    r = first_ratio(ws, 'MATCH HORSE', 'whole'); ws.cell(r, 5).value = str(int(ws.cell(r, 5).value) + 1)
def p12(ws):
    """changed non-MATCH cell"""
    ws.cell(100, 3).value = (ws.cell(100, 3).value or 0) + 1 if isinstance(ws.cell(100, 3).value, (int, float)) else 'x'

for p in (p1, p2, p3, p4, p5, p6, p7, p8, p9, p10, p11, p12):
    wb = openpyxl.load_workbook(new)
    p(wb['P01'])
    path = os.path.join(tmp, p.__name__ + '.xlsx'); wb.save(path)
    m = check_ratio_rule.main(old, path, p.__name__, quiet=True)
    print(p.__name__, p.__doc__, '->', 'caught' if m else 'MISSED', len(m), m[:1])
