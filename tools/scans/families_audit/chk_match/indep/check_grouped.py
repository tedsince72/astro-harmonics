"""Independent checker: MATCH HORSE / MATCH JOCKEY grouped by sky cell (Eddie, 30 Oct 2026). Order only.
Usage: check_grouped.py before.xlsx after.xlsx label
 (1) every cell outside the MATCH region identical (rows before the MATCH HORSE title, every column; columns > 21 in the
     region empty); merges, widths, freeze; MATCH column-name rows identical; titles: only the ORDER sentence and the totals
 (2) data rows (section EXACT / RATIO) per table = the before rows, same multiset cell for cell (value, fill, font, border, ...)
 (3) expected grouped order rebuilt from POSITIONS (own fam.py scores, unrounded), compared row by row incl. header rows;
     also the order on 2-dp scores; ties reported
 (4) header rows: text, counts, grey D9D9D9 exactly A..U, other cells empty; layout: no blank rows inside, two blank rows
     between HORSE and JOCKEY, nothing after JOCKEY; title totals incl. k sky cells."""
import sys, os, re, collections
import openpyxl
sys.path.insert(0, os.path.dirname(__file__))
from check_match import load, sky_cells, t24_cells
from check_ratio_rule import build, exp_row, eq
from compare_rest import sig

TABS = ('MATCH HORSE', 'MATCH JOCKEY')
NC = 21
GREY = 'D9D9D9'
OLD_ORD = ("ORDER (30 Oct 2026, revised): both sections by the sky cell's score, highest first, each sky cell's matches "
           "together (ties: SKY block order); inside a sky cell, RATIO by the bottom number q of p/q, smallest first "
           "(whole-number ratios first: 2, 3, 4 ..., then 3/2, 5/2 ..., then 4/3, 5/3 ...), then by p; then the T24 score, "
           "highest first, then T24 order.")
NEW_ORD = ("ORDER (30 Oct 2026, grouped by sky cell): one list, grouped by sky cell (SKY block relationship + plane), each group "
           "under a grey header row '<sky relationship> <plane> <value> <target> <family> (<score>) — EXACT n, RATIO m'; "
           "groups by the sky cell's score, highest first (ties: SKY block order, a row's RA cell before its Dec cell); inside "
           "a group the EXACT rows first (T24 score highest first, then T24 order), then the RATIO rows by the bottom number q "
           "of p/q, smallest first (whole-number ratios first: 2, 3, 4 ..., then 3/2, 5/2 ..., then 4/3, 5/3 ...), then p, "
           "then the T24 score highest first, then T24 order; the section column says EXACT or RATIO on every row.")


def rgb(c):
    f = c.fill
    return f.fgColor.rgb[-6:] if f is not None and f.fill_type == 'solid' else None


def rowsig(ws, r):
    return tuple(sig(ws.cell(r, c)) for c in range(1, NC + 1))


def blank(ws, r, mc):
    return all(ws.cell(r, c).value is None for c in range(1, mc + 1))


def expected(sky, t24, rnd):
    """[(sky cell, [exact matches], [ratio matches])] in the expected order."""
    si = {id(s): i for i, s in enumerate(sky)}
    ti = {id(t): i for i, t in enumerate(t24)}
    sc = (lambda x: round(x, 2)) if rnd else (lambda x: x)
    ex, ra = build(sky, t24)
    g = collections.OrderedDict()
    for m in ex:
        g.setdefault(id(m['s']), [m['s'], [], []])[1].append(m)
    for m in ra:
        g.setdefault(id(m['s']), [m['s'], [], []])[2].append(m)
    out = []
    for s, e, r in sorted(g.values(), key=lambda x: (-sc(x[0]['sc']), si[id(x[0])])):
        e = sorted(e, key=lambda m: (-sc(m['t']['sc']), ti[id(m['t'])]))
        r = sorted(r, key=lambda m: (m['fr'].denominator, m['fr'].numerator, -sc(m['t']['sc']), ti[id(m['t'])]))
        out.append((s, e, r))
    return out, ex, ra


def head_text(s, ne, nr):
    return '%s %s %.6f %s %s (%.2f) — EXACT %d, RATIO %d' % (s['rel'], s['plane'], s['v'], s['lab'], s['fam'], s['sc'], ne, nr)


def flat(tab, groups):
    rows = []
    for s, e, r in groups:
        rows.append(('H', head_text(s, len(e), len(r))))
        rows += [('D', exp_row(tab, 'EXACT', m)) for m in e]
        rows += [('D', exp_row(tab, 'RATIO', m)) for m in r]
    return rows


def main(old, new, label, quiet=False):
    P = (lambda *a: None) if quiet else print
    msgs = []
    wo = openpyxl.load_workbook(old)
    wn, pos = load(new)
    if wo.sheetnames != wn.sheetnames:
        msgs.append(('sheets', wo.sheetnames, wn.sheetnames))
    sky, bad = sky_cells(wn, pos)
    msgs += bad
    ut = [(a['rel'], a['plane'], b['rel'], b['plane']) for i, a in enumerate(sky) for b in sky[i + 1:] if a['sc'] == b['sc']]
    rt = [(a['rel'], a['plane'], b['rel'], b['plane'], round(a['sc'], 2)) for i, a in enumerate(sky) for b in sky[i + 1:]
          if round(a['sc'], 2) == round(b['sc'], 2)]
    P(label, 'SKY scoring cells', len(sky), '| sky score ties unrounded', ut, '| 2dp', rt)
    st = collections.Counter()
    for sn in wn.sheetnames:
        a, b = wo[sn], wn[sn]
        ta = {t: [r for r in range(1, a.max_row + 1) if a.cell(r, 1).value == t] for t in TABS}
        tb = {t: [r for r in range(1, b.max_row + 1) if b.cell(r, 1).value == t] for t in TABS}
        has = bool(tb['MATCH HORSE'])
        if bool(ta['MATCH HORSE']) != has:
            msgs.append((sn, 'MATCH presence differs')); continue
        mc = max(a.max_column, b.max_column)
        start = tb['MATCH HORSE'][0] if has else None
        if has and ta['MATCH HORSE'][0] != start:
            msgs.append((sn, 'MATCH start moved', ta['MATCH HORSE'][0], start))
        last = start - 1 if has else max(a.max_row, b.max_row)
        # (1) outside
        for r in range(1, last + 1):
            for c in range(1, mc + 1):
                st['outside cells'] += 1
                if sig(a.cell(r, c)) != sig(b.cell(r, c)):
                    msgs.append((sn, 'outside cell', r, c))
        if has:
            for r in range(start, b.max_row + 1):
                for c in range(NC + 1, mc + 1):
                    if b.cell(r, c).value is not None or rgb(b.cell(r, c)) is not None:
                        msgs.append((sn, 'content beyond U in MATCH region', r, c))
        fm = lambda ws, s: sorted(str(m) for m in ws.merged_cells.ranges if not (s and m.min_row >= s))
        if fm(a, start) != fm(b, start):
            msgs.append((sn, 'merges outside'))
        if has and [str(m) for m in b.merged_cells.ranges if m.min_row >= start]:
            msgs.append((sn, 'merges in MATCH region'))
        for k in set(a.column_dimensions) | set(b.column_dimensions):
            if a.column_dimensions[k].width != b.column_dimensions[k].width:
                msgs.append((sn, 'width', k))
        if a.freeze_panes != b.freeze_panes:
            msgs.append((sn, 'freeze'))
        if not has:
            continue
        # layout of the region: title, colnames, rows..., 2 blank, title, colnames, rows..., end
        h0 = tb['MATCH HORSE'][0]
        j0 = tb['MATCH JOCKEY'][0]
        hend = max(tb['MATCH HORSE'])
        if not (blank(b, hend + 1, mc) and blank(b, hend + 2, mc) and hend + 3 == j0):
            msgs.append((sn, 'not exactly two blank rows between HORSE and JOCKEY', hend, j0))
        jend = max(tb['MATCH JOCKEY'])
        for r in range(jend + 1, b.max_row + 1):
            if not blank(b, r, mc):
                msgs.append((sn, 'content after MATCH JOCKEY', r))
        for ch, tab in (('H', 'MATCH HORSE'), ('J', 'MATCH JOCKEY')):
            t0 = tb[tab][0]
            o0 = ta[tab][0]
            # column-name row
            if rowsig(a, o0 + 1) != rowsig(b, t0 + 1) or b.cell(t0 + 1, 1).value != 'table':
                msgs.append((sn, tab, 'column-name row differs'))
            # rows of the table: contiguous from t0+2 to last tab row
            body = list(range(t0 + 2, max(tb[tab]) + 1))
            for r in body:
                if b.cell(r, 1).value != tab:
                    msgs.append((sn, tab, 'row inside table not of the table (blank row?)', r))
            data_b = [r for r in body if b.cell(r, 2).value in ('EXACT', 'RATIO')]
            head_b = [r for r in body if b.cell(r, 2).value not in ('EXACT', 'RATIO')]
            data_a = [r for r in ta[tab][1:] if a.cell(r, 2).value in ('EXACT', 'RATIO')]
            if len(data_a) != len(ta[tab]) - 1:
                msgs.append((sn, tab, 'before has non-data rows'))
            # (2) multiset
            ca = collections.Counter(rowsig(a, r) for r in data_a)
            cb = collections.Counter(rowsig(b, r) for r in data_b)
            if ca != cb:
                msgs.append((sn, tab, 'data multiset differs', sum((ca - cb).values()), sum((cb - ca).values())))
            # header row format
            for r in head_b:
                for c in range(1, NC + 1):
                    if rgb(b.cell(r, c)) != GREY:
                        msgs.append((sn, tab, 'header grey missing', r, c))
                    if c > 2 and b.cell(r, c).value is not None:
                        msgs.append((sn, tab, 'header cell not empty', r, c))
                if rgb(b.cell(r, NC + 1)) is not None:
                    msgs.append((sn, tab, 'header grey beyond U', r))
            for r in data_b:
                if any(rgb(b.cell(r, c)) == GREY for c in range(1, NC + 1)):
                    msgs.append((sn, tab, 'grey on data row', r))
            # (3) expected order
            t24, b2 = t24_cells(b, pos, sn, ch)
            msgs += b2
            gu, ex, ra = expected(sky, t24, False)
            gr, _, _ = expected(sky, t24, True)
            eu, er = flat(tab, gu), flat(tab, gr)
            got = [('H' if b.cell(r, 2).value not in ('EXACT', 'RATIO') else 'D',
                    b.cell(r, 2).value if b.cell(r, 2).value not in ('EXACT', 'RATIO') else
                    [b.cell(r, c).value for c in range(1, NC + 1)]) for r in body]
            if len(got) != len(eu):
                msgs.append((sn, tab, 'row count incl headers', len(got), len(eu)))

            def same(x, y):
                if x[0] != y[0]:
                    return False
                return x[1] == y[1] if x[0] == 'H' else all(eq(p, q) for p, q in zip(x[1], y[1]))
            nd = 0
            for j, (g, e) in enumerate(zip(got, eu)):
                if not same(g, e):
                    nd += 1
                    if nd <= 4:
                        msgs.append((sn, tab, 'row', j, str(g[1])[:160], '| expected', str(e[1])[:160]))
            if nd:
                msgs.append((sn, tab, 'rows differing from expected', nd))
            d2 = sum(1 for x, y in zip(eu, er) if not same(x, y))
            g2 = sum(1 for x, y in zip(got, er) if not same(x, y))
            st['tables'] += 1; st['groups'] += len(gu); st['rows'] += len(data_b); st['headers'] += len(head_b)
            st['2dp-order differs (expected)'] += d2
            # T24 ties inside a group
            for s, e, r in gu:
                for lst, kf in ((e, lambda m: 0), (r, lambda m: m['fr'])):
                    for x, y in zip(lst, lst[1:]):
                        if kf(x) == kf(y):
                            if x['t']['sc'] == y['t']['sc']:
                                st['T24 exact score ties (T24 order decides)'] += 1
                                if not quiet:
                                    print('   tie', sn, tab, s['rel'], s['plane'], x['t']['rel'], x['t']['body'], x['t']['plane'],
                                          '|', y['t']['rel'], y['t']['body'], y['t']['plane'], round(x['t']['sc'], 4))
                            elif round(x['t']['sc'], 2) == round(y['t']['sc'], 2):
                                st['T24 2dp-only ties'] += 1
            # title
            to, tn = a.cell(o0, 2).value, b.cell(t0, 2).value
            want = to.replace(OLD_ORD, NEW_ORD) if to.count(OLD_ORD) == 1 else None
            nw = sum(1 for m in ra if m['s']['fam'] == 'whole')
            ng = sum(1 for m in ra if m['s']['fam'] == 'golden')
            tail_old = 'EXACT %d rows, RATIO %d rows (whole %d, golden %d) (SKY block ' % (len(ex), len(ra), nw, ng)
            tail_new = 'EXACT %d rows, RATIO %d rows (whole %d, golden %d), %d sky cells (SKY block ' % (len(ex), len(ra), nw, ng, len(gu))
            if want is None or want.count(tail_old) != 1:
                msgs.append((sn, tab, 'old title not as expected'))
            elif want.replace(tail_old, tail_new) != tn:
                msgs.append((sn, tab, 'title differs', tn[-200:]))
            if sig(a.cell(o0, 2))[1:] != sig(b.cell(t0, 2))[1:]:
                msgs.append((sn, tab, 'title format differs'))
            for c in list(range(1, mc + 1)):
                if c != 2 and sig(a.cell(o0, c)) != sig(b.cell(t0, c)):
                    msgs.append((sn, tab, 'title row cell differs', c))
            P('%s %s %s: groups %d, data rows %d (EXACT %d, RATIO %d), vs expected %d off; 2dp expected order differs on %d '
              'rows, sheet vs 2dp order %d' % (label, sn, tab, len(gu), len(data_b), len(ex), len(ra), nd, d2, g2))
    P(label, dict(st))
    for m in msgs[:40]:
        P('  MISMATCH', m)
    P(label, 'mismatches:', len(msgs))
    return msgs


if __name__ == '__main__':
    main(*sys.argv[1:4])
