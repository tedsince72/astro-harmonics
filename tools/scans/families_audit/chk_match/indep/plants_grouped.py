"""Planted errors for check_grouped.py on Ripon P01 MATCH HORSE (after build). Usage: plants_grouped.py before.xlsx after.xlsx workdir"""
import sys, os, copy
import openpyxl
from openpyxl.styles import PatternFill
sys.path.insert(0, os.path.dirname(__file__))
from check_grouped import main

NC = 21


def swap_rows(ws, r1, r2):
    for c in range(1, NC + 2):
        a, b = ws.cell(r1, c), ws.cell(r2, c)
        va, vb = a.value, b.value
        sa, sb = copy.copy(a._style), copy.copy(b._style)
        a.value, b.value = vb, va
        a._style, b._style = sb, sa


def find(ws, text):
    return [r for r in range(1, ws.max_row + 1) if isinstance(ws.cell(r, 2).value, str) and ws.cell(r, 2).value.startswith(text)][0]


def move_block(ws, src, n, dst):
    """move rows src..src+n-1 to just after row dst (dst > src) by repeated adjacent swaps."""
    for k in range(n):
        r = src + n - 1 - k
        while r < dst - k:
            swap_rows(ws, r, r + 1); r += 1


def p_swap_groups(ws):
    w = find(ws, 'S Jupiter → S Wezen Dec'); m = find(ws, 'S Jupiter → S Murzim RA')
    move_block(ws, w, 2, m + 43)   # Wezen group (header + 1 row) after the Murzim group (header + 43)

def p_swap_exact_ratio(ws):
    m = find(ws, 'S Jupiter → S Murzim RA'); swap_rows(ws, m + 1, m + 2)

def p_swap_ratio_q(ws):
    m = find(ws, 'S Jupiter → S Murzim RA'); swap_rows(ws, m + 11, m + 12)   # 18/1 and 3/2

def p_swap_exact_t24(ws):
    f = find(ws, 'S Jupiter → S Facies Dec'); swap_rows(ws, f + 1, f + 2)

def p_header_count(ws):
    m = find(ws, 'S Jupiter → S Murzim RA'); c = ws.cell(m, 2); c.value = c.value.replace('RATIO 42', 'RATIO 41')

def p_header_grey(ws):
    m = find(ws, 'S Jupiter → S Murzim RA'); ws.cell(m, 21).fill = PatternFill(fill_type=None)

def p_grey_beyond(ws):
    m = find(ws, 'S Jupiter → S Murzim RA'); ws.cell(m, 22).fill = PatternFill('solid', fgColor='D9D9D9')

def p_blank_between_groups(ws):
    m = find(ws, 'S Jupiter → S Murzim RA'); ws.insert_rows(m)

def p_cell_value(ws):
    m = find(ws, 'S Jupiter → S Murzim RA'); ws.cell(m + 3, 17).value = ws.cell(m + 3, 17).value + 0.000001

def p_cell_fill(ws):
    m = find(ws, 'S Jupiter → S Murzim RA'); ws.cell(m + 3, 20).fill = PatternFill('solid', fgColor='C6EFCE')

def p_section(ws):
    m = find(ws, 'S Jupiter → S Murzim RA'); ws.cell(m + 1, 2).value = 'RATIO'

def p_title_skycells(ws):
    t = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == 'MATCH HORSE'][0]
    ws.cell(t, 2).value = ws.cell(t, 2).value.replace('13 sky cells', '12 sky cells')

def p_one_blank_before_jockey(ws):
    j = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == 'MATCH JOCKEY'][0]; ws.delete_rows(j - 1)

def p_header_text_value(ws):
    m = find(ws, 'S Jupiter → S Murzim RA'); c = ws.cell(m, 2); c.value = c.value.replace('90.000720', '90.000721')

def p_drop_header(ws):
    m = find(ws, 'S Jupiter → S Wezen Dec'); ws.delete_rows(m)


PLANTS = [p_swap_groups, p_swap_exact_ratio, p_swap_ratio_q, p_swap_exact_t24, p_header_count, p_header_grey, p_grey_beyond,
          p_blank_between_groups, p_cell_value, p_cell_fill, p_section, p_title_skycells, p_one_blank_before_jockey,
          p_header_text_value, p_drop_header]


if __name__ == '__main__':
    before, after, wd = sys.argv[1:4]
    os.makedirs(wd, exist_ok=True)
    wb = openpyxl.load_workbook(after)
    ctrl = os.path.join(wd, 'control.xlsx'); wb.save(ctrl)
    n0 = len(main(before, ctrl, 'control', quiet=True))
    print('control (resaved, unmodified) mismatches:', n0)
    caught = 0
    for p in PLANTS:
        wb = openpyxl.load_workbook(after)
        p(wb['P01'])
        f = os.path.join(wd, p.__name__ + '.xlsx'); wb.save(f)
        m = main(before, f, p.__name__, quiet=True)
        caught += bool(m)
        print('%-26s caught=%s (%d mismatches) first: %s' % (p.__name__, bool(m), len(m), str(m[0])[:150] if m else ''))
    print('plants caught %d of %d' % (caught, len(PLANTS)))
