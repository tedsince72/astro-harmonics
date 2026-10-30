"""Plant order errors in copies of the after workbook (Kelso) and confirm check_order catches each.
Usage: plants_order.py old.xlsx new.xlsx tmpdir"""
import sys, os, copy
from fractions import Fraction
import openpyxl
sys.path.insert(0, os.path.dirname(__file__))
import check_order

old, new, tmp = sys.argv[1:4]
os.makedirs(tmp, exist_ok=True)
NC = 21


def drows(ws, tab):
    return [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == tab][1:]


def snap(ws, r):
    return [(ws.cell(r, c).value, copy.copy(ws.cell(r, c)._style)) for c in range(1, NC + 1)]


def permute(ws, rows, new_order):
    """rows: positions; new_order: list of source positions to place there."""
    data = {r: snap(ws, r) for r in rows}
    for dst, src in zip(rows, new_order):
        for c, (v, st) in enumerate(data[src], 1):
            ws.cell(dst, c).value = v; ws.cell(dst, c)._style = st


def swap(ws, a, b):
    permute(ws, [a, b], [b, a])


def info(ws, r):
    v = [ws.cell(r, c).value for c in range(1, NC + 1)]
    return dict(sec=v[1], fr=Fraction(v[2]), sky=(v[7], v[8]), t24sc=v[19], row=r)


def groups(ws, tab, sec):
    g = {}
    for r in drows(ws, tab):
        i = info(ws, r)
        if i['sec'] == sec:
            g.setdefault(i['sky'], []).append(i)
    return g


def p1(ws):
    """swap two RATIO rows inside one sky group (same p/q, different T24 score)"""
    for lst in groups(ws, 'MATCH HORSE', 'RATIO').values():
        for x, y in zip(lst, lst[1:]):
            if x['fr'] == y['fr'] and x['t24sc'] != y['t24sc']:
                return swap(ws, x['row'], y['row'])
def p2(ws):
    """move a sky group (first RATIO sky group to the end of RATIO)"""
    rows = [r for r in drows(ws, 'MATCH HORSE') if ws.cell(r, 2).value == 'RATIO']
    first = groups(ws, 'MATCH HORSE', 'RATIO')
    k0 = (ws.cell(rows[0], 8).value, ws.cell(rows[0], 9).value)
    g0 = [i['row'] for i in first[k0]]
    permute(ws, rows, [r for r in rows if r not in g0] + g0)
def p3(ws):
    """wrong q order: a q=2 row placed before a q=1 row of the same sky group"""
    for lst in groups(ws, 'MATCH JOCKEY', 'RATIO').values():
        a = [i for i in lst if i['fr'].denominator == 1]; b = [i for i in lst if i['fr'].denominator == 2]
        if a and b:
            return swap(ws, a[-1]['row'], b[0]['row'])
def p4(ws):
    """wrong T24 tie order (equal T24 scores; T24 order decides): the J Jupiter T9 duplicate pair in EXACT"""
    for lst in groups(ws, 'MATCH JOCKEY', 'EXACT').values():
        for x, y in zip(lst, lst[1:]):
            if x['t24sc'] == y['t24sc']:
                return swap(ws, x['row'], y['row'])
    raise SystemExit('no tie pair')
def p5(ws):
    """move an EXACT row (first EXACT row to the end of EXACT)"""
    rows = [r for r in drows(ws, 'MATCH HORSE') if ws.cell(r, 2).value == 'EXACT']
    permute(ws, rows, rows[1:] + rows[:1])
def p6(ws):
    """wrong p order inside one q (e.g. 5/2 before 3/2)"""
    for lst in groups(ws, 'MATCH HORSE', 'RATIO').values():
        for x in lst:
            for y in lst:
                if x['row'] < y['row'] and x['fr'].denominator == y['fr'].denominator > 1 and x['fr'] != y['fr']:
                    return swap(ws, x['row'], y['row'])
    raise SystemExit('no p pair')
def p7(ws):
    """swap two adjacent EXACT sky groups"""
    rows = [r for r in drows(ws, 'MATCH JOCKEY') if ws.cell(r, 2).value == 'EXACT']
    keys = []
    for r in rows:
        k = (ws.cell(r, 8).value, ws.cell(r, 9).value)
        if k not in keys: keys.append(k)
    g = lambda k: [r for r in rows if (ws.cell(r, 8).value, ws.cell(r, 9).value) == k]
    a, b = g(keys[0]), g(keys[1])
    permute(ws, rows, b + a + [r for r in rows if r not in a + b])

for p in (p1, p2, p3, p4, p5, p6, p7):
    wb = openpyxl.load_workbook(new)
    p(wb['P01'])
    path = os.path.join(tmp, p.__name__ + '.xlsx'); wb.save(path)
    m = check_order.main(old, path, p.__name__, quiet=True)
    print(p.__name__, p.__doc__, '->', 'caught' if m else 'MISSED', len(m), m[:2])
