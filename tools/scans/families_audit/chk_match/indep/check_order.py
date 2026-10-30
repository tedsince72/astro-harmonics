"""Independent checker: MATCH HORSE / MATCH JOCKEY re-ordered (Eddie, 30 Oct 2026). Rules, rows, cells unchanged.
Usage: check_order.py old.xlsx new.xlsx label
 (1) every cell outside the MATCH rows identical (titles: only the order wording may differ), merges, widths, freeze
 (2) same multiset of MATCH data rows per table, each row cell for cell (value, fill, font, border, number format, comment)
 (3) expected order rebuilt from POSITIONS (own fam.py scores, unrounded) and compared row by row; also with 2-dp scores
 (4) counts and moved rows."""
import sys, os, collections
import openpyxl
sys.path.insert(0, os.path.dirname(__file__))
from check_match import load, sky_cells, t24_cells
from check_ratio_rule import build, exp_row, eq
from compare_rest import sig

TABS = ('MATCH HORSE', 'MATCH JOCKEY')
NC = 21
OLD_EX = 'sorted by the lower score, highest first. RATIO'
NEW_EX = 'order below. RATIO'
OLD_RA = 'ratio = larger / smaller in lowest terms; sorted by ratio, lowest first, ties by the lower score. Ninths'
NEW_RA = ('ratio = larger / smaller in lowest terms. ORDER (30 Oct 2026, revised): both sections by the sky cell\'s score, '
          'highest first, each sky cell\'s matches together (ties: SKY block order); inside a sky cell, RATIO by the bottom '
          'number q of p/q, smallest first (whole-number ratios first: 2, 3, 4 ..., then 3/2, 5/2 ..., then 4/3, 5/3 ...), '
          'then by p; then the T24 score, highest first, then T24 order. Ninths')


def order(ex, ra, sky, t24, rnd):
    si = {id(s): i for i, s in enumerate(sky)}
    ti = {id(t): i for i, t in enumerate(t24)}
    sc = (lambda x: round(x, 2)) if rnd else (lambda x: x)
    k_ex = lambda m: (-sc(m['s']['sc']), si[id(m['s'])], -sc(m['t']['sc']), ti[id(m['t'])])
    k_ra = lambda m: (-sc(m['s']['sc']), si[id(m['s'])], m['fr'].denominator, m['fr'].numerator,
                      -sc(m['t']['sc']), ti[id(m['t'])])
    return sorted(ex, key=k_ex), sorted(ra, key=k_ra)


def match_rows(ws):
    """{tab: (title_row_idx, header_row_idx, [data row idx])} (1-based rows)."""
    out = {}
    for tab in TABS:
        idx = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == tab]
        if idx:
            out[tab] = (idx[0], idx[0] + 1, idx[1:])
    return out


def rowsig(ws, r):
    return tuple(sig(ws.cell(r, c)) for c in range(1, NC + 1))


def main(old, new, label, quiet=False):
    P = (lambda *a: None) if quiet else print
    msgs = []
    wo = openpyxl.load_workbook(old)
    wn, pos = load(new)
    if wo.sheetnames != wn.sheetnames:
        msgs.append(('sheets', wo.sheetnames, wn.sheetnames))
    sky, bad = sky_cells(wn, pos)
    msgs += bad
    ncell = 0
    stats = collections.Counter()
    tie_notes = []
    # sky-cell score ties (unrounded and 2 dp)
    scs = [s['sc'] for s in sky]
    u_ties = [(a['rel'], a['plane'], b['rel'], b['plane']) for i, a in enumerate(sky) for b in sky[i + 1:] if a['sc'] == b['sc']]
    r_ties = [(a['rel'], a['plane'], round(a['sc'], 2), b['rel'], b['plane']) for i, a in enumerate(sky) for b in sky[i + 1:]
              if round(a['sc'], 2) == round(b['sc'], 2)]
    P(label, 'SKY scoring cells', len(sky), 'sky-score ties unrounded', u_ties, '2dp', r_ties)
    same_row_ties = [t for t in r_ties if t[0] == t[3]]
    for sn in wn.sheetnames:
        a, b = wo[sn], wn[sn]
        ma, mb = match_rows(a), match_rows(b)
        if set(ma) != set(mb):
            msgs.append((sn, 'MATCH tables', set(ma), set(mb)))
        mrows_a = {r for t in ma.values() for r in t[2]}
        mrows_b = {r for t in mb.values() for r in t[2]}
        title_rows = {t[0] for t in mb.values()}
        if {t[0] for t in ma.values()} != title_rows:
            msgs.append((sn, 'title rows moved'))
        if mrows_a != mrows_b:
            msgs.append((sn, 'MATCH data row positions differ'))
        # (1) every cell outside MATCH data rows; title rows checked separately
        mr = max(a.max_row, b.max_row); mc = max(a.max_column, b.max_column)
        for r in range(1, mr + 1):
            if r in mrows_b:
                for c in range(NC + 1, mc + 1):
                    ncell += 1
                    if sig(a.cell(r, c)) != sig(b.cell(r, c)):
                        msgs.append((sn, 'beyond col 21 in MATCH row', r, c))
                continue
            for c in range(1, mc + 1):
                ncell += 1
                sa, sb = sig(a.cell(r, c)), sig(b.cell(r, c))
                if sa == sb:
                    continue
                if r in title_rows and c == 2:
                    t_old = sa[0]
                    want = t_old.replace(OLD_EX, NEW_EX).replace(OLD_RA, NEW_RA)
                    if t_old.count(OLD_EX) == 1 and t_old.count(OLD_RA) == 1 and want == sb[0] and sa[1:] == sb[1:]:
                        stats['titles changed (order wording only)'] += 1
                        continue
                    msgs.append((sn, 'title differs beyond order wording', r))
                    continue
                msgs.append((sn, 'outside cell', r, c, sa[:3], sb[:3]))
        if sorted(map(str, a.merged_cells.ranges)) != sorted(map(str, b.merged_cells.ranges)):
            msgs.append((sn, 'merges'))
        for k in set(a.column_dimensions) | set(b.column_dimensions):
            if a.column_dimensions[k].width != b.column_dimensions[k].width:
                msgs.append((sn, 'width', k))
        if a.freeze_panes != b.freeze_panes:
            msgs.append((sn, 'freeze'))
        if not (sn.startswith('P') and sn[1:].isdigit()):
            continue
        for ch, tab in (('H', 'MATCH HORSE'), ('J', 'MATCH JOCKEY')):
            if tab not in mb:
                msgs.append((sn, tab, 'missing')); continue
            da, db = ma[tab][2], mb[tab][2]
            # (2) multiset
            ca = collections.Counter(rowsig(a, r) for r in da)
            cb = collections.Counter(rowsig(b, r) for r in db)
            if ca != cb:
                msgs.append((sn, tab, 'row multiset differs', sum((ca - cb).values()), sum((cb - ca).values())))
            moved = sum(1 for ra_, rb_ in zip(da, db) if rowsig(a, ra_) != rowsig(b, rb_))
            nex = sum(1 for r in db if b.cell(r, 2).value == 'EXACT')
            mex = sum(1 for ra_, rb_ in zip(da, db) if b.cell(rb_, 2).value == 'EXACT' and rowsig(a, ra_) != rowsig(b, rb_))
            # sections contiguous: EXACT then RATIO
            secs = [b.cell(r, 2).value for r in db]
            if secs != sorted(secs, key=lambda s: 0 if s == 'EXACT' else 1):
                msgs.append((sn, tab, 'sections not EXACT then RATIO'))
            # (3) expected order
            t24, b2 = t24_cells(b, pos, sn, ch)
            msgs += b2
            ex, ra = build(sky, t24)
            exu, rau = order(ex, ra, sky, t24, False)
            exr, rar = order(ex, ra, sky, t24, True)
            expu = [exp_row(tab, 'EXACT', m) for m in exu] + [exp_row(tab, 'RATIO', m) for m in rau]
            expr = [exp_row(tab, 'EXACT', m) for m in exr] + [exp_row(tab, 'RATIO', m) for m in rar]
            got = [[b.cell(r, c).value for c in range(1, NC + 1)] for r in db]
            if len(got) != len(expu):
                msgs.append((sn, tab, 'row count', len(got), len(expu)))
            nd = 0
            for j, (g, e) in enumerate(zip(got, expu)):
                if not all(eq(x, y) for x, y in zip(g, e)):
                    nd += 1
                    if nd <= 5:
                        msgs.append((sn, tab, 'order/cell row', j, g[1], g[7], g[8], g[13], g[14], g[15], '| expected', e[7], e[8], e[13], e[14], e[15]))
            if nd:
                msgs.append((sn, tab, 'rows differing from expected order', nd))
            diff_ur = sum(1 for x, y in zip(expu, expr) if not all(eq(p, q) for p, q in zip(x, y)))
            if diff_ur:
                tie_notes.append((sn, tab, diff_ur))
            # groups contiguous per sky cell within each section
            for sec in ('EXACT', 'RATIO'):
                keys = [(g[7], g[8]) for g in got if g[1] == sec]
                seen, prev = set(), None
                for k in keys:
                    if k != prev:
                        if k in seen:
                            msgs.append((sn, tab, sec, 'sky group split', k))
                        seen.add(k); prev = k
            # T24 score ties inside one group decided by T24 order (unrounded / 2dp)
            for lst in (exu, rau):
                for x, y in zip(lst, lst[1:]):
                    if x['s'] is y['s'] and x['fr'] == y['fr']:
                        if x['t']['sc'] == y['t']['sc']:
                            stats['T24 exact-score ties (T24 order decides)'] += 1
                        elif round(x['t']['sc'], 2) == round(y['t']['sc'], 2):
                            stats['T24 2dp-only ties'] += 1
            nsky_groups = len({(id(m['s']), sec) for sec, lst in (('E', ex), ('R', ra)) for m in lst})
            P('%s %s %s: rows %d (EXACT %d, RATIO %d), sky groups %d, moved %d (EXACT %d, RATIO %d), vs expected %d off; '
              'unrounded vs 2dp expected order differ on %d rows' % (
                  label, sn, tab, len(db), nex, len(db) - nex, nsky_groups, moved, mex, moved - mex, nd, diff_ur))
            stats['rows'] += len(db); stats['moved'] += moved
    P(label, 'outside cells compared', ncell, dict(stats))
    P(label, 'same-row (RA vs Dec) 2dp sky ties:', same_row_ties)
    P(label, 'tables where 2-dp order != unrounded order:', tie_notes)
    for m in msgs[:40]:
        P('  MISMATCH', m)
    P(label, 'mismatches:', len(msgs))
    return msgs


if __name__ == '__main__':
    main(*sys.argv[1:4])
