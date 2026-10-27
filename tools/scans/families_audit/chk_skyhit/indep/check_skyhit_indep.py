"""Independent checker (27 Oct 2026): 'sky hit' moved to the end of T24 HORSE / JOCKEY (SKY STARS) and the SKY STARS SUMMARY,
widths 11 / 40 / 60 / 36, freeze panes unchanged. Own code (openpyxl only).
usage: python3 check_skyhit_indep.py BEFORE.xlsx AFTER.xlsx [BEFORE_pairs.xlsx AFTER_pairs.xlsx]"""
import sys
import openpyxl
from openpyxl.utils import get_column_letter as L

SS = ('T24 HORSE (SKY STARS)', 'T24 JOCKEY (SKY STARS)')
mism = []
counts = {}


def bump(k, n=1):
    counts[k] = counts.get(k, 0) + n


def bad(msg):
    mism.append(msg)


def sig(c):
    """full signature of a cell; None for an empty cell with default style and no comment"""
    if c is None:
        return None
    com = c.comment
    s = (c.value, repr(c.font), repr(c.fill), repr(c.border), c.number_format, repr(c.alignment), repr(c.protection),
         (com.text, com.author, com.width, com.height) if com else None)
    if c.value is None and not c.has_style and com is None:
        return None
    return s


def cellsig(ws, r, c):
    x = ws._cells.get((r, c))
    return sig(x)


def diffsig(a, b):
    names = ['value', 'font', 'fill', 'border', 'number_format', 'alignment', 'protection', 'comment']
    if a is None or b is None:
        return 'one side empty: %r | %r' % (a and a[0], b and b[0])
    return ','.join(n for n, x, y in zip(names, a, b) if x != y)


def tables(ws):
    """[(name, first row, last row)] of the (SKY STARS) tables: title row (col A = name), two header rows, then rows with col A = name"""
    out = []
    r = 1
    mr = ws.max_row
    while r <= mr:
        a = ws.cell(r, 1).value
        if a in SS:
            r0 = r
            r += 2
            while r + 1 <= mr and ws.cell(r + 1, 1).value == a:
                r += 1
            out.append((a, r0, r))
        r += 1
    return out


def t24block(ws, who):
    """natal T24 block: title row (col A 'T24', col B starts 'T24 <WHO> (natal)'), header rows, star rows -> (title, {label: row})"""
    for r in range(1, ws.max_row + 1):
        b = ws.cell(r, 2).value
        if ws.cell(r, 1).value == 'T24' and isinstance(b, str) and b.startswith(f'T24 {who} (natal)'):
            rows = {}
            k = r + 3
            while ws.cell(k, 1).value == 'T24':
                rows[ws.cell(k, 2).value] = k
                k += 1
            return r, rows
    return None, {}


def roles(ws, hr):
    """column roles from the two header rows starting at hr: row 1 group labels (merged), row 2 names"""
    top = {}
    for m in ws.merged_cells.ranges:
        if m.min_row == hr == m.max_row:
            for c in range(m.min_col, m.max_col + 1):
                top[c] = ws.cell(hr, m.min_col).value
    for c in range(1, ws.max_column + 1):
        v = ws.cell(hr, c).value
        if v is not None and c not in top:
            top[c] = v
    R = {}
    for c, g in top.items():
        if c <= 2:
            continue
        if g in ('RA', 'Dec', 'SCORES RA', 'SCORES Dec'):
            R[c] = 'value' if g in ('RA', 'Dec') else 'scores'
        elif g in ('RA midpoints', 'Dec midpoints'):
            R[c] = 'mid'
        elif g == 'repeats':
            R[c] = 'rep'
        elif g == 'sky hit':
            R[c] = 'hit'
        else:
            R[c] = '?' + str(g)
    return R


WANT = {'value': 11, 'scores': 11, 'mid': 40, 'rep': 60, 'hit': 36}


def width(ws, c):
    d = ws.column_dimensions.get(L(c))
    return d.width if d is not None and d.customWidth else None


def merges(ws, r0, r1):
    return {(m.min_row, m.min_col, m.max_row, m.max_col) for m in ws.merged_cells.ranges if r0 <= m.min_row and m.max_row <= r1}


def check(bp, ap):
    B = openpyxl.load_workbook(bp)
    A = openpyxl.load_workbook(ap)
    if B.sheetnames != A.sheetnames:
        bad(f'sheet order {B.sheetnames} vs {A.sheetnames}')
    width_changes = []
    for name in A.sheetnames:
        wb_, wa = B[name], A[name]
        if wb_.freeze_panes != wa.freeze_panes:
            bad(f'{name} freeze {wb_.freeze_panes} -> {wa.freeze_panes}')
        bump('freeze checked')
        # widths changed
        cols = set(k for k in wb_.column_dimensions) | set(k for k in wa.column_dimensions)
        for k in sorted(cols, key=lambda s: (len(s), s)):
            ci = openpyxl.utils.column_index_from_string(k)
            wb0, wa0 = width(wb_, ci), width(wa, ci)
            if wb0 != wa0:
                width_changes.append((name, k, wb0, wa0))
        is_sum = name == 'SKY STARS SUMMARY'
        tbs = tables(wa) if (name.startswith('P') and name[1:].isdigit()) or is_sum else []
        tbb = tables(wb_) if tbs else []
        if [(t[0], t[1], t[2]) for t in tbs] != [(t[0], t[1], t[2]) for t in tbb]:
            bad(f'{name} table row ranges differ {tbb} vs {tbs}')
        inside = set()
        if is_sum and tbs:
            inside.update((1, 2))
            nm, r0, _ = tbs[0]
            hc = max(c for (r, c) in wa._cells if r == r0 + 2 and wa._cells[(r, c)].value is not None)
            for r in (1, 2):
                for bc in range(1, hc + 3):
                    ac = bc if (bc <= 2 or bc > hc) else (hc if bc == 3 else bc - 1)
                    bump('summary rows 1-2 cells vs before')
                    if cellsig(wb_, r, bc) != cellsig(wa, r, ac):
                        bad(f'{name} {L(ac)}{r} (before {L(bc)}{r}): {diffsig(cellsig(wb_, r, bc), cellsig(wa, r, ac))}')
            mb = {(a, b if b < 3 else b - 1, c, d if d < 3 else d - 1) for a, b, c, d in merges(wb_, 1, 2)}
            if mb != merges(wa, 1, 2):
                bad(f'{name} rows 1-2 merges {sorted(mb)} vs {sorted(merges(wa, 1, 2))}')
        for (nm, r0, r1), (_, b0, b1) in zip(tbs, tbb):
            inside.update(range(r0, r1 + 1))
            # hitc: the last column of header row 2 (before: SCORES Dec end)
            hitc_b = max(c for (r, c) in wb_._cells if r == b0 + 2 and wb_._cells[(r, c)].value is not None)
            hitc = max(c for (r, c) in wa._cells if r == r0 + 2 and wa._cells[(r, c)].value is not None)
            if hitc != hitc_b:
                bad(f'{name} {nm} r{r0}: last column {hitc_b} -> {hitc}')
            if wa.cell(r0 + 1, hitc).value != 'sky hit' or wb_.cell(b0 + 1, 3).value != 'sky hit':
                bad(f'{name} {nm} r{r0}: sky hit header not where expected')
            def amap(bc):                                  # before column -> after column, rows of the table
                if bc <= 2 or bc > hitc:
                    return bc
                return hitc if bc == 3 else bc - 1
            maxc = max(hitc + 3, wb_.max_column, wa.max_column)
            for r in range(r0, r1 + 1):
                seen = set()
                for bc in range(1, maxc + 1):
                    ac = amap(bc)
                    seen.add(ac)
                    sb, sa = cellsig(wb_, r, bc), cellsig(wa, r, ac)
                    if r == r0 and bc >= 3:               # title row: only A, B
                        sb, sa = cellsig(wb_, r, bc), cellsig(wa, r, bc)
                        ac = bc
                    bump('table cells compared')
                    if sb != sa:
                        bad(f'{name} {nm} {L(ac)}{r} (before {L(bc)}{r}): {diffsig(sb, sa)}')
                for ac in range(1, maxc + 1):
                    if ac not in seen and cellsig(wa, r, ac) is not None:
                        bad(f'{name} {nm} {L(ac)}{r}: extra cell')
            # merges
            mb = {(a, amap(b) if b != 3 else None, c, amap(d) if d != 3 else None) for a, b, c, d in merges(wb_, b0, b1)}
            ma = merges(wa, r0, r1)
            if mb != ma:
                bad(f'{name} {nm} r{r0}: merges before(mapped) {sorted(mb)} after {sorted(ma)}')
            bump('table merges', len(ma))
            # sky hit blank on Jupiter rows and stars row; stars row grey A..hitc
            for r in range(r0 + 3, r1 + 1):
                b = wa.cell(r, 2).value
                star = isinstance(b, str) and (b.startswith('H ') or b.startswith('J ')) and '→' not in b
                if not star and wa.cell(r, hitc).value is not None:
                    bad(f'{name} {nm} {L(hitc)}{r}: sky hit not blank on row {b!r}')
                if b == 'stars':
                    for c in range(1, hitc + 1):
                        if wa.cell(r, c).fill.fgColor.rgb not in ('00D9D9D9', 'FFD9D9D9'):
                            bad(f'{name} {nm} {L(c)}{r}: stars row not grey ({wa.cell(r, c).fill.fgColor.rgb})')
                    if cellsig(wa, r, hitc + 1) is not None:
                        bad(f'{name} {nm} {L(hitc + 1)}{r}: grey past the last column')
                    bump('stars rows')
                elif star:
                    bump('star rows')
                    if wa.cell(r, hitc).value is None:
                        bad(f'{name} {nm} {L(hitc)}{r}: star row without sky hit')
            # widths by role (header rows of the after table)
            R = roles(wa, r0 + 1)
            if R.get(hitc) != 'hit':
                bad(f'{name} {nm}: column {L(hitc)} role {R.get(hitc)}')
            for c, ro in R.items():
                w = width(wa, c)
                bump('width checks')
                if ro not in WANT or w != WANT[ro]:
                    bad(f'{name} {nm}: column {L(c)} role {ro} width {w} (want {WANT.get(ro)})')
            if not is_sum:
                # star rows and header rows = the natal T24 block, column for column from C to repeats
                who = nm.split()[1]
                t0, srows = t24block(wa, who)
                last = max(c for c, ro in R.items() if ro == 'rep')
                RT = roles(wa, t0 + 1)
                if {c: ro for c, ro in R.items() if c <= last} != RT:
                    bad(f'{name} {nm}: column roles differ from T24 {who}')
                for c, ro in RT.items():
                    bump('width checks T24')
                    if width(wa, c) != WANT.get(ro):
                        bad(f'{name} T24 {who}: column {L(c)} role {ro} width {width(wa, c)}')
                for k in range(2):
                    for c in range(2, last + 1):
                        bump('header cells vs T24')
                        if cellsig(wa, r0 + 1 + k, c) != cellsig(wa, t0 + 1 + k, c):
                            bad(f'{name} {nm} {L(c)}{r0 + 1 + k} header != T24 {L(c)}{t0 + 1 + k}: '
                                f'{diffsig(cellsig(wa, r0 + 1 + k, c), cellsig(wa, t0 + 1 + k, c))}')
                tm = {(a - t0, b, c - t0, d) for a, b, c, d in merges(wa, t0 + 1, t0 + 2) if d <= last}
                sm = {(a - r0, b, c - r0, d) for a, b, c, d in merges(wa, r0 + 1, r0 + 2) if d <= last}
                if tm != sm:
                    bad(f'{name} {nm}: header merges {sorted(sm)} vs T24 {sorted(tm)}')
                for r in range(r0 + 3, r1 + 1):
                    b = wa.cell(r, 2).value
                    if b in srows:
                        for c in range(2, last + 1):
                            bump('star cells vs T24')
                            if cellsig(wa, r, c) != cellsig(wa, srows[b], c):
                                bad(f'{name} {nm} {L(c)}{r} != T24 {L(c)}{srows[b]}: {diffsig(cellsig(wa, r, c), cellsig(wa, srows[b], c))}')
                if name == 'P01' or True:
                    for col in ('AA', 'AZ', 'BA'):
                        ci = openpyxl.utils.column_index_from_string(col)
                        if width(wa, ci) != 11:
                            bad(f'{name}: column {col} width {width(wa, ci)}')
        # every other row unchanged
        rows = set(r for r, c in wb_._cells) | set(r for r, c in wa._cells)
        for r in sorted(rows - inside):
            cs = set(c for rr, c in wb_._cells if rr == r) if False else None
        cellsB, cellsA = {}, {}
        for (r, c), x in wb_._cells.items():
            if r not in inside:
                cellsB[(r, c)] = x
        for (r, c), x in wa._cells.items():
            if r not in inside:
                cellsA[(r, c)] = x
        for k in set(cellsB) | set(cellsA):
            sb, sa = sig(cellsB.get(k)), sig(cellsA.get(k))
            bump('other cells compared')
            if sb != sa:
                bad(f'{name} {L(k[1])}{k[0]}: {diffsig(sb, sa)}')
        mo_b = {m for m in map(lambda m: (m.min_row, m.min_col, m.max_row, m.max_col), wb_.merged_cells.ranges) if m[0] not in inside}
        mo_a = {m for m in map(lambda m: (m.min_row, m.min_col, m.max_row, m.max_col), wa.merged_cells.ranges) if m[0] not in inside}
        if mo_b != mo_a:
            bad(f'{name}: merges outside tables differ: {sorted(mo_b ^ mo_a)[:10]}')
        if is_sum:
            check_summary(A, wa)
    return width_changes


def check_summary(A, out):
    """summary = header rows of the first table (rows 1-2), then per runner its A1 line, H table, 1 blank, J table, 2 blanks"""
    runners = [s for s in A.sheetnames if s.startswith('P') and s[1:].isdigit()]
    row = 3
    first = None
    for P in runners:
        ws = A[P]
        tb = {t[0]: t for t in tables(ws)}
        bump('summary runners')
        if cellsig(out, row, 1) != cellsig(ws, 1, 1):
            bad(f'SUMMARY A{row} != {P}!A1')
        row += 1
        for C in SS:
            _, r0, r1 = tb[C]
            if first is None:
                first = (P, r0)
            for r in range(r0, r1 + 1):
                for c in range(1, max(ws.max_column, out.max_column) + 1):
                    bump('summary cells')
                    if cellsig(out, row + r - r0, c) != cellsig(ws, r, c):
                        bad(f'SUMMARY {L(c)}{row + r - r0} != {P}!{L(c)}{r}: {diffsig(cellsig(out, row + r - r0, c), cellsig(ws, r, c))}')
            ms = {(a - r0 + row, b, c - r0 + row, d) for a, b, c, d in merges(ws, r0, r1)}
            mo = merges(out, row, row + r1 - r0)
            if ms != mo:
                bad(f'SUMMARY merges rows {row}-{row + r1 - r0} vs {P}: {sorted(ms ^ mo)}')
            row += r1 - r0 + 1
            gap = 1 if C == SS[0] else 2
            for g in range(gap):
                for c in range(1, out.max_column + 1):
                    if cellsig(out, row + g, c) is not None:
                        bad(f'SUMMARY {L(c)}{row + g}: blank row not blank')
            row += gap
    if row - 1 < out.max_row and any(cellsig(out, r, c) is not None for r in range(row, out.max_row + 1) for c in range(1, out.max_column + 1)):
        bad('SUMMARY: rows after the last runner')
    P, r0 = first
    ws = A[P]
    for k in range(2):
        for c in range(1, out.max_column + 1):
            bump('summary header cells')
            if cellsig(out, 1 + k, c) != cellsig(ws, r0 + 1 + k, c):
                bad(f'SUMMARY {L(c)}{1 + k} != {P}!{L(c)}{r0 + 1 + k}')
    ms = {(a - r0 - 1 + 1, b, c - r0 - 1 + 1, d) for a, b, c, d in merges(ws, r0 + 1, r0 + 2)}
    if ms != merges(out, 1, 2):
        bad('SUMMARY rows 1-2 merges differ')
    if out.freeze_panes != 'C3':
        bad(f'SUMMARY freeze {out.freeze_panes}')


def pairs_same(bp, ap):
    B, A = openpyxl.load_workbook(bp), openpyxl.load_workbook(ap)
    if B.sheetnames != A.sheetnames:
        bad('pairs sheet order')
    n = 0
    for s in B.sheetnames:
        x, y = B[s], A[s]
        for k in set(x._cells) | set(y._cells):
            n += 1
            if sig(x._cells.get(k)) != sig(y._cells.get(k)):
                bad(f'pairs {s} {L(k[1])}{k[0]}')
        if {str(m) for m in x.merged_cells.ranges} != {str(m) for m in y.merged_cells.ranges}:
            bad(f'pairs {s} merges')
    bump('pairs cells compared', n)


if __name__ == '__main__':
    wc = check(sys.argv[1], sys.argv[2])
    if len(sys.argv) > 4:
        pairs_same(sys.argv[3], sys.argv[4])
    for k, v in sorted(counts.items()):
        print(f'{k}: {v}')
    print('width changes vs before:', len(wc))
    for x in wc:
        print('  ', x)
    print('mismatches:', len(mism))
    for m in mism[:200]:
        print('  ', m)
