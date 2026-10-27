"""plant errors in a copy of the AFTER workbook: python3 plant.py AFTER.xlsx OUT.xlsx"""
import sys, openpyxl
from openpyxl.styles import PatternFill
from openpyxl.comments import Comment
wb = openpyxl.load_workbook(sys.argv[1])
def find(ws, name):
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == name:
            return r
ws = wb['P03']; r = find(ws, 'T24 HORSE (SKY STARS)')
ws.cell(r + 9, 10).value = 1.234567                                  # 1 value in a star row (J10 area)
print('P03 value', ws.cell(r + 9, 10).coordinate)
ws.cell(r + 9, 60).fill = PatternFill('solid', fgColor='FF00FF00')  # 2 fill on a SCORES cell
print('P03 fill', ws.cell(r + 9, 60).coordinate)
ws.cell(r + 3, 106).value = 'oops'                                   # 3 sky hit on a Jupiter row
print('P03 jup hit', ws.cell(r + 3, 106).coordinate)
ws = wb['SKY STARS SUMMARY']
for row in ws.iter_rows(min_row=3):
    hit = [c for c in row if c.comment]
    if hit:
        print('SUMMARY comment removed', hit[0].coordinate); hit[0].comment = None; break   # 4
wb['P05'].column_dimensions['AZ'].width = 40                         # 5 width
print('P05 AZ 40')
ws = wb['P07']; r = find(ws, 'T24 JOCKEY (SKY STARS)')
m = [m for m in ws.merged_cells.ranges if m.min_row == r + 1][0]
ws.unmerge_cells(str(m)); print('P07 unmerged', m)                    # 6 merge
wb['SKY JUPITER'].cell(10, 3).value = 'x'; print('SKY JUPITER C10')   # 7 other cell
wb['P09'].freeze_panes = 'B4'; print('P09 freeze B4')                 # 8 freeze
ws = wb['P11']; r = find(ws, 'T24 HORSE (SKY STARS)')
for k in range(r, r + 10):
    if ws.cell(k, 2).value == 'stars':
        ws.cell(k, 106).fill = PatternFill(); print('P11 grey removed', ws.cell(k, 106).coordinate); break   # 9
ws = wb['P01']; r = find(ws, 'T24 JOCKEY (SKY STARS)')
c = ws.cell(r + 12, 20); c.border = openpyxl.styles.Border(); print('P01 border cleared', c.coordinate)   # 10 outline
wb.save(sys.argv[2])
