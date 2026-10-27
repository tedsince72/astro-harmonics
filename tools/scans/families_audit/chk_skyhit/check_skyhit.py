"""'sky hit' moved to the end of T24 HORSE / JOCKEY (SKY STARS) and exact widths (Eddie, 27 Oct 2026): checks one build against the build
before (the display commit).
usage: check_skyhit.py <before.xlsx> <after.xlsx>
(a) every sheet other than the runner sheets and the summary identical; runner sheets and summary: every row outside the (SKY STARS)
    tables identical; inside them (summary: every row) cell before (r, c) = cell after (r, c - 1) for c >= 4, and before (r, 3) ('sky hit') =
    after (r, last), last = the table's last column (after the SCORES Dec columns) - values, fills (the 'stars' grey incl.), fonts, borders,
    number formats, comments, wrap; merged ranges mapped the same way; freeze panes as before;
(b) the copied T24 part in the same columns as T24 HORSE / JOCKEY: header rows and every star row = the T24 row, column for column (C to
    the repeats column);
(c) widths, runner sheets and summary: every T24 HORSE / JOCKEY and (SKY STARS) value, S Jupiter and SCORES column 11, RA / Dec midpoints
    40, repeats 60, sky hit 36, exactly; runner-sheet AA, AZ, BA 11.
Prints 'mismatches: n' last."""
import sys, re
import openpyxl
from openpyxl.utils import get_column_letter as L

bad = []
err = lambda *a: bad.append(' '.join(str(x) for x in a))


def sig(c):
    b = c.border
    return (c.value, c.fill.fill_type, c.fill.start_color.rgb if c.fill.fill_type else None, c.font.b, c.number_format,
            (b.left.style, b.right.style, b.top.style, b.bottom.style), c.comment.text if c.comment else None, c.alignment.wrap_text)


B = openpyxl.load_workbook(sys.argv[1])
A = openpyxl.load_workbook(sys.argv[2])
if A.sheetnames != B.sheetnames:
    err('sheet order')
SUM = 'SKY STARS SUMMARY'
runners = [s for s in A.sheetnames if re.fullmatch(r'P\d\d', s)]


def ss_rows(ws):
    """rows of the (SKY STARS) tables (title, headers, rows) -> set"""
    out = set()
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(row=r, column=1).value).endswith('(SKY STARS)') and str(ws.cell(row=r, column=2).value).startswith(
                f'{ws.cell(row=r, column=1).value} ('):
            out |= {r + 1, r + 2}
            k = r + 3
            while ws.cell(row=k, column=1).value == ws.cell(row=r, column=1).value:
                out.add(k); k += 1
            out.add(r)
    return out


nmoved = nt24 = nw = 0
for sh in A.sheetnames:
    a, b = A[sh], B[sh]
    nc = max(a.max_column, b.max_column) + 1
    if a.max_row != b.max_row:
        err(sh, 'rows', a.max_row, b.max_row)
    if a.freeze_panes != b.freeze_panes:
        err(sh, 'freeze', a.freeze_panes, b.freeze_panes)
    if sh not in runners and sh != SUM:
        if any(sig(a.cell(row=r, column=k)) != sig(b.cell(row=r, column=k)) for r in range(1, b.max_row + 1) for k in range(1, nc)):
            err(sh, 'sheet differs')
        if {str(m) for m in a.merged_cells.ranges} != {str(m) for m in b.merged_cells.ranges}:
            err(sh, 'merges differ')
        continue
    rows = set(range(1, b.max_row + 1)) if sh == SUM else ss_rows(b)
    # the table's last column after the move: the 'sky hit' header in row 1 of the headers (after)
    hits = {}
    for r in range(1, a.max_row + 1):
        for k in range(3, a.max_column + 1):
            if a.cell(row=r, column=k).value == 'sky hit':
                hits[r] = k
    lastc = max(hits.values()) if hits else None
    if len(set(hits.values())) > 1:
        err(sh, 'sky hit in more than one column', set(hits.values()))
    if b.cell(row=min(hits) if hits else 1, column=3).value != 'sky hit':
        err(sh, 'before: sky hit not in C')
    for r in range(1, b.max_row + 1):
        if r not in rows:
            if any(sig(a.cell(row=r, column=k)) != sig(b.cell(row=r, column=k)) for k in range(1, nc)):
                err(sh, 'row differs', r)
            continue
        if any(b.cell(row=r, column=k).value is not None for k in range(lastc + 1, nc + 1)):
            err(sh, 'before: cell beyond the table', r)
        for k in range(1, lastc + 1):                      # the table spans the same columns before and after (A to lastc)
            ka = k if k < 3 else lastc if k == 3 else k - 1
            if sig(a.cell(row=r, column=ka)) != sig(b.cell(row=r, column=k)):
                err(sh, 'moved cell differs', r, L(k), '->', L(ka)); break
            nmoved += k == 3 and b.cell(row=r, column=3).value is not None
        if any(sig(a.cell(row=r, column=k)).__getitem__(0) is not None or a.cell(row=r, column=k).fill.fill_type
               for k in range(lastc + 1, nc + 1)):
            err(sh, 'cell beyond the table', r)
    mp = lambda m: (m.min_row, m.min_col, m.max_row, m.max_col)
    sh_ = lambda c: c if c < 3 else lastc if c == 3 else c - 1
    mb = {(x[0], sh_(x[1]), x[2], sh_(x[3])) if x[0] in rows else x for x in map(mp, b.merged_cells.ranges)}
    if {mp(m) for m in a.merged_cells.ranges} != mb:
        err(sh, 'merged ranges', sorted({mp(m) for m in a.merged_cells.ranges} ^ mb)[:4])
    # (b) T24 part in the same columns as T24 HORSE / JOCKEY (runner sheets)
    W = lambda k: a.column_dimensions[L(k)].width if L(k) in a.column_dimensions and a.column_dimensions[L(k)].customWidth else None
    if sh in runners:
        for who, C in (('HORSE', 'H'), ('JOCKEY', 'J')):
            t = next(r for r in range(1, a.max_row + 1) if a.cell(row=r, column=1).value == 'T24'
                     and str(a.cell(row=r, column=2).value).startswith(f'T24 {who} (natal)'))
            rep = next(k for k in range(3, a.max_column + 1) if a.cell(row=t + 1, column=k).value == 'repeats')
            s = next(r for r in range(1, a.max_row + 1) if a.cell(row=r, column=1).value == f'T24 {who} (SKY STARS)')
            for k in range(2):
                if any(sig(a.cell(row=t + 1 + k, column=c)) != sig(a.cell(row=s + 1 + k, column=c)) for c in range(2, rep + 1)):
                    err(sh, who, 'header row not as T24', k)
            star = {a.cell(row=r, column=2).value: r for r in range(t + 3, t + 80) if a.cell(row=r, column=1).value == 'T24'}
            r = s + 3
            while a.cell(row=r, column=1).value == f'T24 {who} (SKY STARS)':
                lab = a.cell(row=r, column=2).value
                if lab in star:
                    nt24 += 1
                    if any(sig(a.cell(row=r, column=c)) != sig(a.cell(row=star[lab], column=c)) for c in range(2, rep + 1)):
                        err(sh, who, lab, 'not as its T24 row')
                r += 1
            # (c) widths: T24 header columns, then SCORES and sky hit
            for c in range(3, lastc + 1):
                h1, h2 = a.cell(row=s + 1, column=c).value, a.cell(row=s + 2, column=c).value
                want = 60 if h1 == 'repeats' else 40 if str(h1).endswith(' midpoints') else 36 if h1 == 'sky hit' else 11
                nw += 1
                if W(c) != want:
                    err(sh, who, 'width', L(c), W(c), want, h1, h2)
        for col in ('AA', 'AZ', 'BA'):                     # by name (Eddie, 27 Oct 2026): 11
            nw += 1
            if W(openpyxl.utils.column_index_from_string(col)) != 11:
                err(sh, col, W(openpyxl.utils.column_index_from_string(col)))
    else:
        for c in range(3, lastc + 1):
            h1 = a.cell(row=1, column=c).value
            want = 60 if h1 == 'repeats' else 40 if str(h1).endswith(' midpoints') else 36 if h1 == 'sky hit' else 11
            nw += 1
            if W(c) != want:
                err(sh, 'width', L(c), W(c), want, h1)
print('sky hit cells moved:', nmoved, '| star rows = T24 rows:', nt24, '| widths checked:', nw)
for x in bad[:40]:
    print('  ', x)
print('mismatches:', len(bad))
