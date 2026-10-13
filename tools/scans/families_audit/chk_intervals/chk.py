import sys, math, re
sys.path.insert(0, '/Users/eddieplace/astro-harmonics/rebuild_kit')
sys.path.insert(0, '/private/tmp/claude-501/-Users-eddieplace-astro-harmonics/0a4b9192-0277-48b6-a3a8-817a486af8bb/scratchpad/chk_int')
import openpyxl
from active_vibrations import (score_whole, score_phi_interval, score_phi_powers,
                               score_ninth, _NINTH_RANGES, _PHI_POWER_TARGETS)
import eph2
REPO = '/Users/eddieplace/astro-harmonics/'
INP = REPO + 'reference/classic/race_batch_20220102_newcastle_aw_1330.xlsx'
OUT = REPO + 'reference/classic/race_batch_20220102_newcastle_aw_1330_jupiter_tables.xlsx'
SHEET='P09'; HTAB='P09'; JTAB='P10'
PHI=(1+5**0.5)/2; R2=2**0.5

wi = openpyxl.load_workbook(INP, data_only=True)
def load(tab):
    return {r[1]:(r[4],r[8]) for r in wi[tab].iter_rows(min_row=2, values_only=True) if r[3]=='transit_self'}
POS={'S':load('TRANS'),'H':load(HTAB),'J':load(JTAB)}
# +5 sky: TRANS + independent swisseph delta
s0=eph2.sky(0); s5=eph2.sky(5)
S5={}
for n,(ra,dec) in POS['S'].items():
    if n in s0 and s0[n] is not None:
        dra=((s5[n][0]-s0[n][0]+180)%360)-180; S5[n]=((ra+dra)%360, dec+s5[n][1]-s0[n][1])
    else:
        S5[n]=(ra,dec)  # Transpluto, equator, plain royal names: no delta (Transpluto ~static)
POS['S5']=S5
for k in ('S','H','J','S5'):
    ra,dec=POS[k]['Rahu']; POS[k]['KetuMirror']=((ra+180)%360,-dec)

def pt(chart,name,t5=False):
    if chart=='S' and t5: chart='S5'
    if name=='equator': return (None,0.0)
    if name=='Ketu': return (None,POS[chart]['KetuMirror'][1])
    if chart in 'HJ' and name=='Moon': raise KeyError('natal Moon')
    d=POS[chart]
    return d[name+'_B'] if name+'_B' in d else d[name]
def ra_sep(a,b):
    x=abs(a-b)%360; return 360-x if x>180 else x
def gc(p,q):
    r1,d1=map(math.radians,p); r2,d2=map(math.radians,q)
    h=math.sin((d2-d1)/2)**2+math.cos(d1)*math.cos(d2)*math.sin((r2-r1)/2)**2
    return math.degrees(2*math.asin(min(1,math.sqrt(h))))
def planes(p,q):
    out={'Dec':abs(p[1]-q[1])}
    if p[0] is not None and q[0] is not None:
        ra=ra_sep(p[0],q[0]); out['RA']=ra; out['Flat']=math.hypot(ra,out['Dec']); out['GC']=gc(p,q)
    return out

# ---------- degree families ----------
def phi_label(i):
    if i<11: p=(-1,1,2,3,4,5,6,7,8,9,10)[i]; f=''
    elif i<19: p=(-1,1,2,3,4,5,6,7)[i-11]; f='10×'
    else: p=(-1,1)[i-19]; f='100×'
    return '%sφ^%d'%(f,p)
SQ=[('1.414',R2,'sqrt2'),('14.14',10*R2,'sqrt2'),('141.4',100*R2,'sqrt2'),('0.707',R2/2,'sqrt2'),('7.07',5*R2,'sqrt2'),('70.7',50*R2,'sqrt2'),
    ('2.414',1+R2,'silver'),('24.14',10*(1+R2),'silver'),('0.414',R2-1,'silver'),('4.14',10*(R2-1),'silver'),('41.4',100*(R2-1),'silver')]
SQX=[('1.414',1.414,'sqrt2'),('14.14',14.14,'sqrt2'),('141.4',141.4,'sqrt2'),('0.707',0.707,'sqrt2'),('7.07',7.07,'sqrt2'),('70.7',70.7,'sqrt2'),
    ('2.414',2.414,'silver'),('24.14',24.14,'silver'),('0.414',0.414,'silver'),('4.14',4.14,'silver'),('41.4',41.4,'silver')]
SQUSE = SQX if '--exactlabels' in sys.argv else SQ
def fams_deg(v,plane):
    res=[]
    s=score_whole(v)
    if s>0: n=round(v); res.append(('whole',str(n),v-n,s,n))
    s=score_phi_interval(v)
    if s>0: k=round(v/PHI); res.append(('golden','%d×φ'%k,v-k*PHI,s,k*PHI))
    best=None
    for i,t in enumerate(_PHI_POWER_TARGETS):
        tol=max(t*0.005,0.025); dev=abs(v-t)
        if dev<=tol:
            sc=(10*(tol-dev)/tol)**2
            if best is None or sc>best[3]: best=('phi power',phi_label(i),v-t,sc,t)
    if best and best[3]>0:
        assert abs(best[3]-score_phi_powers(v))<1e-9; res.append(best)
    s=score_ninth(v)
    if s>0:
        for lo,hi,tol,t in _NINTH_RANGES:
            if lo<=v<hi: res.append(('ninth','%g'%t,v-t,s,t)); break
    for fam in ('sqrt2','silver'):
        b=None
        for lab,t,f in SQUSE:
            if f!=fam: continue
            tol=max(0.005*t,0.025); dev=abs(v-t)
            if dev<=tol:
                sc=(10*(tol-dev)/tol)**2
                if b is None or sc>b[3]: b=(fam,lab,v-t,sc,t)
        if b and b[3]>0: res.append(b)
    if plane in ('RA','Dec','Equator'):
        tol=1.5 if plane=='RA' else 0.3
        if v<tol:
            sc=(10*(tol-v)/tol)**2
            if sc>0: res.append(('conj','conj',v,sc,0.0))
    res.sort(key=lambda x:-x[3]); return res

# ---------- intervals ----------
IV=[('Perfect',[('1',1),('2',2),('3/2',1.5),('4/3',4/3),('3',3),('4',4)]),
    ('Imperfect',[('5/4',1.25),('6/5',1.2),('5/3',5/3),('8/5',1.6),('5/2',2.5),('8/3',8/3),('9/4',2.25)]),
    ('Dissonant',[('√2',R2),('1+1/√2',1+1/R2),('9/8',9/8),('16/9',16/9),('9/5',1.8),('15/8',15/8),('16/15',16/15)]),
    ('Septimal',[('7/4',1.75),('7/3',7/3),('7/2',3.5),('7/5',1.4),('7/6',7/6),('8/7',8/7)]),
    ('Harmonic',[('5',5),('6',6),('7',7),('8',8),('9',9)]),
    ('High',[('11/5',2.2),('11/6',11/6),('11/3',11/3),('11/8',11/8),('11/4',2.75),('13/5',2.6),('13/8',13/8),('13/4',3.25)]),
    ('Phi',[('φ',PHI),('φ²',PHI**2),('φ³',PHI**3),('1+√2',1+R2),('φ√5',PHI*5**0.5),('2−1/φ',2-1/PHI),('φ³+1',PHI**3+1),('2/φ',2/PHI)])]
GROUPS=[g for g,_ in IV]
NAMES={}  # filled from the sheet, checked for consistency
def fams_iv(v):
    res=[]
    for g,lst in IV:
        for lab,t in lst:
            dev=abs(v/t-1)*100
            if dev<0.15:
                sc=(10*(0.15-dev)/0.15)**2
                if sc>0: res.append(('interval',lab,(v/t-1)*100,sc,t,g))
    res.sort(key=lambda x:-x[3]); return res

def colour(s):
    if s>=80: return 'C6EFCE'
    if s>=50: return 'FFEB9C'
    if s>0: return 'FFF8DC'
    return None
def offs(f): return ('%+.5f%%'%f[2]) if f[0]=='interval' else '%+.5f'%f[2]
def tlabel(f, inv):
    if f[0]=='interval':
        return '%s · %s · %s'%(f[1],NAMES.get(f[1],'?'),f[5])+(' (inverted)' if inv else '')
    return '%s (%s)'%(f[1],f[0])+(' (inverted)' if inv else '')
def fams(v,kind,plane): return fams_deg(v,plane) if kind=='deg' else fams_iv(v)
def expected(v,kind,plane,inv):
    f=fams(v,kind,plane); e={'fams':f}
    if f:
        b=f[0]; e['best']=b
        e['target']=tlabel(b,inv); e['group']=b[5] if b[0]=='interval' else None
        e['offset']=offs(b); e['score']=round(b[3],2); e['colour']=colour(b[3])
        e['also']=['%s (%s) %s [%.1f]'%(x[1],x[0],offs(x),x[3]) for x in f[1:]]
        e['ties']=[x for x in f[1:] if abs(x[3]-b[3])<1e-9]
    else:
        e.update(target=None,group=None,colour=None,also=[],best=None,ties=[])
    c=[]
    for tag,sv in (('x10',10*v),('/10',v/10)):
        ff=[x for x in fams(sv,kind,plane) if x[0]!='conj']
        if ff: c.append((tag,ff[0]))
    if c:
        c.sort(key=lambda z:-z[1][3]); tag,b=c[0]
        if b[0]=='interval': s='%s: %s · %s · %s %s [%.1f]'%(tag,b[1],NAMES.get(b[1],'?'),b[5],offs(b),b[3])
        else: s='%s: %s (%s) %s [%.1f]'%(tag,b[1],b[0],offs(b),b[3])
        if tag=='/10' and kind=='deg': s+=' (widens 10x in degrees)'
        e['scaled']=s; e['scaled_colour']=colour(b[3]); e['scaled_tie']=len(c)==2 and abs(c[0][1][3]-c[1][1][3])<1e-9
    else:
        e.update(scaled=None,scaled_colour=None,scaled_tie=False)
    return e

def P(s,t5=False):
    s=s.strip(); return pt(s[0],s[2:].strip(),t5)
def has_S(lab): return bool(re.search(r'(^|[\s(|/–])S ',lab))
def mid_offsets(a,b,c):
    out={'Dec':c[1]-(a[1]+b[1])/2}
    if a[0] is not None and b[0] is not None and c[0] is not None:
        d=((b[0]-a[0]+180)%360)-180; m=(a[0]+d/2)%360; out['RA']=((c[0]-m+180)%360)-180
    return out
def ratio_planes(num,den):
    out={}
    for pl in num:
        if pl in den and den[pl]>0 and num[pl]>0:
            n,d=num[pl],den[pl]; out[pl]=(max(n,d)/min(n,d), n<d)
    return out
def compute(table,label,t5=False):
    t=int(table[1:])
    lab=re.sub(r'\s+\(RA span.*|\s+\(Dec span.*','',label).strip()
    Q=lambda s:P(s,t5)
    if t<=8:
        a,b=lab.split('→'); return 'deg',{k:(v,False) for k,v in planes(Q(a),Q(b)).items()}
    if t==9:
        if lab.endswith('Equator'):
            p=Q(lab[:-len('Equator')]); return 'deg',{'Dec':(abs(p[1]),False)}
        a,b=lab.split(' – '); return 'deg',{k:(v,False) for k,v in planes(Q(a),Q(b)).items()}
    if t in (10,11):
        ab,c=lab.split('|'); a,b=ab.split(' / '); return 'mid',mid_offsets(Q(a),Q(b),Q(c))
    if t in (12,13):
        n,d=lab.split(' / '); n1,n2=n.split('–'); d1,d2=d.split('–')
        return 'ratio',ratio_planes(planes(Q(n1),Q(n2)),planes(Q(d1),Q(d2)))
    if t==14:
        inv=lab.endswith('(inverted)'); l2=lab.replace('(inverted)','').strip(); a,b=l2.split(' / ')
        va=abs(Q(a.replace('|Dec|',''))[1]); vb=abs(Q(b.replace('|Dec|',''))[1])
        return 'ratio1',{'Equator':(max(va,vb)/min(va,vb),va<vb,inv)}
    if t==15:
        inv=lab.endswith('(inverted)'); l2=lab.replace('(inverted)','').strip()
        m=re.match(r'\((.*)\) / \((.*)\)$',l2); n,d=m.group(1),m.group(2)
        n1,n2=n.split('–'); d1,d2=d.split('–')
        if d2.strip()=='equator':
            v=planes(Q(n1),Q(n2))['Dec']; den=abs(Q(d1)[1])
            return 'ratio1',{'Dec':(max(v,den)/min(v,den),v<den,inv)}
        return 'ratio',ratio_planes(planes(Q(n1),Q(n2)),planes(Q(d1),Q(d2)))
    raise ValueError(table)

def fill(c):
    if c.fill and c.fill.fill_type:
        rgb=c.fill.fgColor.rgb
        if rgb and rgb!='00000000': return rgb[-6:]
    return None

wo=openpyxl.load_workbook(OUT)
# names from sheet (tool's names) - collect & check consistency
for n in wo.sheetnames:
    if re.match(r'^P\d\d$',n):
        for row in wo[n].iter_rows(values_only=True):
            for v in row:
                if isinstance(v,str) and ' · ' in v:
                    m=re.match(r'^(?:x10: |/10: )?(.+?) · (.+?) · (\w+)',v)
                    if m:
                        if NAMES.get(m.group(1),m.group(2))!=m.group(2): print('NAME INCONSISTENT',m.groups())
                        NAMES[m.group(1)]=m.group(2)
                        g=[gg for gg,l in IV if any(x[0]==m.group(1) for x in l)]
                        if g!=[m.group(3)]: print('GROUP WRONG',v,g)
ws=wo[SHEET]
cells=0; mism=[]; notes=[]; rows_done={}; ratio_hits=[]; EXP={}
PLANE_COL={}
group_hits={g:0 for g in GROUPS}
for r in range(1,ws.max_row+1):
    a=ws.cell(r,1).value
    if a=='table':
        PLANE_COL={}
        for i in range(4):
            h=ws.cell(r,3+9*i).value
            if h: PLANE_COL[h]=3+9*i
        hdr=[ws.cell(r,c).value for c in range(3,12)]
        if hdr!=[PLANE_COL and list(PLANE_COL)[0],'+5 min','nearest target','group','offset','score','off->+5','also (score)','scaled x10 / /10']:
            mism.append((r,'header',hdr))
        continue
    if not (isinstance(a,str) and re.match(r'^T\d+$',a)): continue
    label=ws.cell(r,2).value
    if all(ws.cell(r,c).value is None for c in range(3,40)): continue
    table=a
    try:
        mode,vals=compute(table,label); mode5,vals5=compute(table,label,True)
    except Exception as e:
        mism.append((r,table,label,'PARSE',repr(e))); continue
    rows_done[table]=rows_done.get(table,0)+1
    sflag=has_S(label)
    for plane,c0 in PLANE_COL.items():
        C=lambda k: ws.cell(r,c0+k)
        sv,v5s,tgt,grp,off,sc,arr,also,scaled=[C(k).value for k in range(9)]
        col=fill(C(0)); scol=fill(C(8))
        def M(what,sheet,mine): mism.append((r,table,label[:75],plane,what,sheet,mine))
        cells+=9
        if plane not in vals:
            if any(x is not None for x in (sv,v5s,tgt,grp,off,sc,arr,also,scaled)): M('cell present but plane n/a',sv,None)
            continue
        # +5
        if mode=='mid': v=vals[plane]; v5=vals5[plane]
        elif mode=='ratio1': v,inv,labinv=vals[plane]; v5=vals5[plane][0]
        else: v,inv=vals[plane]; v5=vals5[plane][0]
        if sflag:
            if v5s is None or abs(v5s-v5)>2e-5: M('+5 value',v5s,round(v5,6))
        else:
            if v5s is not None: M('+5 present on natal-only row',v5s,None)
        if sv is None or abs(sv-v)>2e-6: M('value',sv,round(v,6))
        if mode=='mid':
            if tgt!='midpoint (midpoint)': M('target',tgt,'midpoint (midpoint)')
            if grp is not None: M('group',grp,None)
            if off!='%+.5f'%v: M('offset',off,'%+.5f'%v)
            tt=round(abs(v)/0.05,2)
            if sc is None or abs(sc-tt)>0.011: M('score',sc,tt)
            av=abs(v); ec='C6EFCE' if av<=0.05 else 'FFEB9C' if av<=0.15 else 'FFF8DC' if av<=0.45 else None
            if col!=ec: M('colour',col,ec)
            if also or scaled: M('also/scaled present',(also,scaled),None)
            if v5s is not None:
                ea='-> toward' if abs(v5s)<abs(sv) else '<- away' if abs(v5s)>abs(sv) else '?='
                if arr!=ea: M('arrow',arr,ea)
            continue
        kind='deg' if mode=='deg' else 'ratio'
        if mode=='ratio1':
            if labinv!=inv: M('row (inverted) flag',labinv,inv)
            inv=False
        e=expected(v,kind,plane,inv)
        if e['ties']: notes.append((r,table,label,plane,'tie',[x[:2] for x in e['ties']]))
        if e['scaled_tie']: notes.append((r,table,label,plane,'scaled tie'))
        if tgt!=e['target']: M('target',tgt,e['target'])
        if grp!=e['group']: M('group',grp,e['group'])
        if e['target']:
            if off!=e['offset']: M('offset',off,e['offset'])
            if sc is None or abs(sc-e['score'])>0.011: M('score',sc,e['score'])
            b=e['best']
            if v5s is not None:
                if b[0]=='interval': d0,d1=abs(sv/b[4]-1),abs(v5s/b[4]-1)
                else: d0,d1=abs(sv-b[4]),abs(v5s-b[4])
                ea='-> toward' if d1<d0 else '<- away' if d1>d0 else '?='
                if arr!=ea: M('arrow',arr,ea)
            elif arr is not None: M('arrow without +5',arr,None)
            if kind=='ratio':
                group_hits[b[5]]+=1
                ratio_hits.append((b[5],table,label,plane,round(v,6),b[1],e['offset'],round(b[3],2),b[3],inv))
        else:
            if off is not None or sc is not None or arr is not None: M('offset/score/arrow present',(off,sc,arr),None)
        if col!=e['colour']: M('colour',col,e['colour'])
        sh_also=[x.strip() for x in re.split(r';\s*|,\s(?=[^\]]*\[)',also)] if also else []
        sh_also=[x for x in sh_also if x]
        if sorted(sh_also)!=sorted(e['also']): M('also',also,e['also'])
        if scaled!=e['scaled']: M('scaled',scaled,e['scaled'])
        if scol!=e['scaled_colour']: M('scaled colour',scol,e['scaled_colour'])
        EXP[(table,re.sub(r'\s*\(inverted\)$','',label),plane)]=(v,e)
print('P09 rows per table',rows_done)
print('P09 cells compared',cells)
print('P09 mismatches',len(mism))
for m in mism[:80]: print('  ',m)
print('notes',len(notes))
for n in notes[:30]: print('  ',n)
print('ratio cells with interval score>0 by group',group_hits, 'total',sum(group_hits.values()))
print('interval names used by tool:',NAMES)

# ---------- RATIOS sheet ----------
wr=wo[SHEET+' RATIOS']; rcells=0; rm=[]
sheet_rows=[]; cur=None
for r in range(1,wr.max_row+1):
    row=[wr.cell(r,c).value for c in range(1,12)]
    if row[0] in GROUPS and row[1] is None: cur=row[0]; sheet_rows.append(('HDR',cur)); continue
    if row[0]=='group' or row[0] is None or r==1: continue
    sheet_rows.append(('ROW',row,fill(wr.cell(r,5)),[c for c in range(1,12) if c!=5 and fill(wr.cell(r,c)) not in (None,'D9E1F2')]))
hdrs=[x[1] for x in sheet_rows if x[0]=='HDR']
print('RATIOS group sections in sheet:',hdrs)
exp_sorted=[]
for g in GROUPS:
    exp_sorted+=sorted([h for h in ratio_hits if h[0]==g],key=lambda h:-h[8])
got=[x for x in sheet_rows if x[0]=='ROW']
print('RATIOS rows: sheet',len(got),'expected',len(exp_sorted))
for i in range(max(len(got),len(exp_sorted))):
    if i>=len(got) or i>=len(exp_sorted): rm.append(('length',i)); continue
    row,colr,extra=got[i][1],got[i][2],got[i][3]; h=exp_sorted[i]
    rcells+=11
    lab=re.sub(r'\s*\(inverted\)$','',h[2])
    expv=[h[0],h[1],lab,h[3],h[4],None,h[5],NAMES.get(h[5]),h[6],h[7],None]
    for j in (0,1,2,3,6,7,8):
        if row[j]!=expv[j]: rm.append((i,'col%d'%(j+1),row[j],expv[j]))
    if abs(row[4]-h[4])>2e-6: rm.append((i,'value',row[4],h[4]))
    if abs(row[9]-h[7])>0.011: rm.append((i,'score',row[9],h[7]))
    # +5 & arrow vs P09 sheet cells
    if colr!=colour(h[8]): rm.append((i,'colour',colr,colour(h[8])))
    if extra: rm.append((i,'extra fills',extra))
    if h[9] : notes.append(('RATIOS row drops (inverted) marker',h[2],h[3]))
# cross-check +5 and arrow in RATIOS vs P09 sheet values
p09={}
for r in range(1,ws.max_row+1):
    a=ws.cell(r,1).value
    if a=='table':
        PC={ws.cell(r,3+9*i).value:3+9*i for i in range(4) if ws.cell(r,3+9*i).value}; continue
    if isinstance(a,str) and re.match(r'^T1[2-5]$',a) and ws.cell(r,2).value:
        for pl,c0 in PC.items():
            p09[(a,re.sub(r'\s*\(inverted\)$','',ws.cell(r,2).value),pl)]=(ws.cell(r,c0+1).value,ws.cell(r,c0+6).value)
for i,g in enumerate(got):
    row=g[1]; k=(row[1],row[2],row[3])
    if k not in p09: rm.append((i,'not in P09',k)); continue
    if (row[5],row[10])!=p09[k]: rm.append((i,'+5/arrow differ from P09',(row[5],row[10]),p09[k]))
print('RATIOS cells compared',rcells,'mismatches',len(rm))
for m in rm: print('  ',m)

# ---------- FIELD sheets ----------
fcells=0; fm=[]
for t in (12,13,14,15):
    wf=wo['FIELD %d'%t]; plane=None; vcol=None; nrow=0
    for r in range(1,wf.max_row+1):
        a=wf.cell(r,1).value
        if a in ('RA','Dec','Flat','GC','Equator') and isinstance(wf.cell(r,2).value,str) and wf.cell(r,2).value.startswith('P01'):
            plane=a; vcol=[c for c in range(2,wf.max_column+1) if isinstance(wf.cell(r,c).value,str) and wf.cell(r,c).value.startswith('P09')][0]; continue
        if plane is None or not a: continue
        sv=wf.cell(r,vcol).value; st=wf.cell(r,vcol+1).value; col=fill(wf.cell(r,vcol))
        fcells+=3; nrow+=1
        try: mode,vals=compute('T%d'%t,a)
        except Exception as ex: fm.append((t,r,a,'PARSE',repr(ex))); continue
        if plane not in vals:
            if sv is not None or st is not None: fm.append((t,r,a,plane,'present n/a',sv))
            continue
        x=vals[plane]; v=x[0]
        if sv is None or abs(sv-v)>2e-6: fm.append((t,r,a,plane,'value',sv,round(v,6))); continue
        e=expected(v,'ratio',plane,False)
        b=e['best']
        es=None if b is None else '%s %s [%.1f] [%s]'%(b[1],offs(b),b[3],b[5])
        if st!=es: fm.append((t,r,a,plane,'string',st,es))
        if col!=e['colour']: fm.append((t,r,a,plane,'colour',col,e['colour']))
        k=('T%d'%t,a,plane)
        if k not in EXP: fm.append((t,r,a,plane,'row not found on P09'))
    print('FIELD %d rows for Venturous'%t,nrow)
print('FIELD cells compared',fcells,'mismatches',len(fm))
for m in fm[:50]: print('  ',m)
print('RATIOS rows whose P09 cell/row is inverted:',[(h[1],h[2],h[3]) for h in exp_sorted if h[9] or h[2].endswith('(inverted)')])
