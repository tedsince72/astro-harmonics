"""check46_engine.py (independent checker, 20 Oct 2026): the 11 new stars in the engine, the 46-star RA order, the rebuilt batch vs
the old one, and the POSITIONS rows of the 11 = the batch values.
usage: check46_engine.py <old batch.xlsx> <rebuilt batch.xlsx> <tables.xlsx>"""
import os, sys
import openpyxl
REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..', '..', '..'))
sys.path[:0] = [os.path.join(REPO, 'rebuild_kit'), os.path.join(REPO, 'tools')]
import swisseph as swe
import celestial_bodies_swisseph as E
from jupiter_tables import T1_T2_T24_STARS, STARS                      # constant lists only

swe.set_ephe_path(os.path.expanduser('~/astronomy-project/ephe')); swe.set_topo(0, 51.5, 0)
CAT = {'Alpheratz': (2.097, 29.090), 'Algenib': (3.309, 15.184), 'Difda': (10.897, -17.987), 'Mirach': (17.433, 35.621),
       'Menkar': (45.570, 4.090), 'Phact': (84.912, -34.074), 'Adhara': (104.656, -28.972), 'Wezen': (107.098, -26.393),
       'Menkent': (211.671, -36.370), 'Unukalhai': (236.067, 6.426), 'Rasalhague': (263.734, 12.560)}
NEW11 = list(CAT)
bad = []
B = dict(E.BEHENIAN_STARS)
print('engine names:', {x: B.get(x) for x in NEW11})
for x, (ra, dec) in CAT.items():
    r, nm, _ = swe.fixstar2_ut(B[x], 2451545.0, E._FLAG)
    ok = abs(r[0] - ra) < 0.02 and abs(r[1] - dec) < 0.02
    print(f'  {x:11s} {nm:26s} {r[0]:11.6f} {r[1]:+10.6f}  dRA {r[0] - ra:+.4f} dDec {r[1] - dec:+.4f} {"ok" if ok else "OFF"}')
    if not ok:
        bad.append(('catalogue', x, r[:2]))
nm = swe.fixstar2_ut(B['Menkar'], 2451545.0, E._FLAG)[1]
plain = swe.fixstar2_ut('Menkar', 2451545.0, E._FLAG)
print('Menkar as used:', nm, '| plain "Menkar":', plain[1], plain[0][:2])
if 'alCet' not in nm:
    bad.append(('Menkar not alpha Ceti', nm))
# 46 by engine RA
ra46 = {}
for x in T1_T2_T24_STARS:
    r = swe.fixstar2_ut(B[x], 2451545.0, E._FLAG)[0]
    ra46[x] = r
order = sorted(T1_T2_T24_STARS, key=lambda x: ra46[x][0])
print('46 in engine RA order = T1_T2_T24_STARS:', order == T1_T2_T24_STARS, '| max |Dec|', max(abs(v[1]) for v in ra46.values()))
if order != T1_T2_T24_STARS or len(set(order)) != 46 or set(order) != (set(STARS) - {'Polaris'}) | (set(T1_T2_T24_STARS) - set(STARS)):
    bad.append(('order', order))
print('  ' + ', '.join(f'{x} {ra46[x][0]:.3f}' for x in order))

# batch: old + Galactic Centre_B + 13 + 11, everything else identical
old, new = openpyxl.load_workbook(sys.argv[1], read_only=True), openpyxl.load_workbook(sys.argv[2], read_only=True)
added_exp = {'Galactic Centre_B'} | {f'{x}_B' for x in T1_T2_T24_STARS if x not in STARS}
batch = {}
for sh in old.sheetnames:
    ro = [tuple(r) for r in old[sh].iter_rows(values_only=True)]
    rn = [tuple(r) for r in new[sh].iter_rows(values_only=True)]
    keyo = {r[1]: r for r in ro if r and r[1]}
    rest = [r for r in rn if not (r and r[1] in added_exp)]
    added = sorted({r[1] for r in rn if r and r[1] in added_exp})
    same = rest == ro
    print(f'batch {sh}: old rows {len(ro)}, new {len(rn)}, added {len(added)} (expected {len(added_exp)}), rest identical {same}')
    if not same or set(added) != added_exp:
        bad.append(('batch', sh, same, sorted(added_exp - set(added))))
    batch[sh] = {r[1]: r for r in rn if r and r[1]}
# POSITIONS rows of the 11 = batch (TRANS -> S, P01 -> H, P02 -> J); col E RA, I Dec
wb = openpyxl.load_workbook(sys.argv[3], read_only=True)
SHMAP = {'S': 'TRANS', 'H': 'P01', 'J': 'P02'}
n = 0
for r in wb['POSITIONS'].iter_rows(values_only=True):
    if r[2] in SHMAP and r[3] in NEW11:
        b = batch[SHMAP[r[2]]][f'{r[3]}_B']
        n += 1
        if abs(float(r[5]) - float(b[4])) > 1e-9 or abs(float(r[6]) - float(b[8])) > 1e-9:
            bad.append(('POSITIONS', r[2], r[3], r[5:7], (b[4], b[8])))
print('POSITIONS rows of the 11 checked against the batch:', n)
print('MISMATCHES:', len(bad))
for x in bad:
    print('  ', x)
