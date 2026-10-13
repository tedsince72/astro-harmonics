#!/usr/bin/env python3
"""t12_perfect_harmonic_scan.py OUTDIR (read-only, 13 Oct 2026, Eddie): races with known results; per race, runners with a T12 interval hit in group Perfect or
Harmonic at score >= 90 (jupiter_tables RATIOS sheets; automatic unisons excluded). Winners vs beaten; chance = sum of 1/field over holders.
Winners: older races from the race csv finish_pos (results already opened; Redcar from blind_test.md as field_only_rates.py);
12 Oct 2026 races from tools/scans/jr_tally.py RACES (winner cloth, non-runner cloths)."""
import csv, glob, os, re, sys, subprocess
import openpyxl
REPO = os.path.expanduser('~/astro-harmonics'); os.chdir(REPO)
OUT = sys.argv[1]                      # a scratch folder for the rebuilt tables (not the repo)
sys.path.insert(0, 'tools/scans')
from jr_tally import RACES as R12
W12 = {r: (w, nr) for r, w, nr, _ in R12}
REDCAR = {'Present Moment': '1'}

def inputs():
    for f in sorted(glob.glob('reference/classic/race_batch_2*_[0-9][0-9][0-9][0-9].xlsx')):
        m = re.match(r'race_batch_(\d{8})_(.+)_(\d{4})\.xlsx', os.path.basename(f))
        yield f'{m.group(2)}_{m.group(3)}_{m.group(1)}', f, f[:-5] + '.csv'
    yield 'goodwood_1853_20210827', 'reference/classic/goodwood_1853_20210827', 'reference/classic/race_batch_20210827_goodwood_1853.csv'

def finish(race, csvp):
    if race in W12:
        return None
    if race.startswith('redcar_1515'):
        return {'Present Moment': '1'}
    return {re.sub(r'^\d+ ', '', r['name']): r['finish_pos'] for r in csv.DictReader(open(csvp)) if r['subject_type'] == 'horse'}

rows, tot = [], {}
for race, inp, csvp in inputs():
    out = f'{OUT}/{race}.xlsx'
    subprocess.run([sys.executable, 'tools/jupiter_tables.py', inp, '--out', out], capture_output=True, check=True)
    wb = openpyxl.load_workbook(out, read_only=True)
    run = {}
    for s in wb.sheetnames:
        if not s.endswith(' RATIOS'):
            continue
        it = wb[s].iter_rows(values_only=True)
        title = next(it)[0]
        m = re.match(r'(P\d\d) (\d+) (.+?) / ', title)
        hits = [r for r in it if r and r[1] == 'T12' and r[0] in ('Perfect', 'Harmonic') and isinstance(r[9], (int, float)) and r[9] >= 90
                and not str(r[11] or '').startswith('automatic')]
        run[m.group(1)] = (m.group(2), m.group(3), hits)
    if len(run) < 2:
        continue
    fin = finish(race, csvp)
    if fin is None:
        wc, nrc = W12[race]
        win = [P for P, v in run.items() if v[0] == wc]; nr = [P for P, v in run.items() if v[0] in nrc]
    else:
        win = [P for P, v in run.items() if fin.get(v[1]) == '1']; nr = []
    n = len(run); n2 = n - len(nr)
    hold = [P for P, v in run.items() if v[2]]
    grp = 'blind 12 Oct 2026' if race in W12 else 'older (read; results known)'
    t = tot.setdefault(grp, [0, 0, 0.0, 0, 0.0, 0])
    t[0] += 1; t[5] += 1 if win else 0
    for P in hold:
        w = P in win
        t[1] += w; t[2] += 1 / n; t[3] += (not w)
        if P not in nr:
            t[4] += 1 / n2
    lines = [f"{race}  field {n}{f' (NR {len(nr)}: ' + ','.join(nr) + ')' if nr else ''}  winner {', '.join(f'{P} {run[P][1]}' for P in win) or '?'}"]
    for P in hold:
        lines.append(f"   {'W' if P in win else 'b'} {P} {run[P][1]}{' NR' if P in nr else ''}: " +
                     '; '.join(f"{h[2]} {h[3]} {h[6]} {h[7]} {h[4]} [{h[9]}]" for h in sorted(run[P][2], key=lambda h: -h[9])))
    if not hold:
        lines.append('   (no runner)')
    rows.append('\n'.join(lines))
print('\n'.join(rows))
print('\nTotals (chance = sum of 1/field over holders, field as read; w/o NR = 1/(field - non-runners), non-runner holders left out)')
for g, t in tot.items():
    print(f"  {g}: races {t[0]} (winner identified {t[5]}); holders {t[1] + t[3]}: winners {t[1]} (chance {t[2]:.2f}; w/o NR {t[4]:.2f}), beaten {t[3]}")
