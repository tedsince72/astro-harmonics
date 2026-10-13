"""motion.py - JR v2 section F: sky rates from the chart engine (off and off + 5 min), checked against the workbook's TRANS.

Engine: rebuild_kit/celestial_bodies_swisseph.py (v2.2, same file as ~/astronomy-project), imported, called with the same
flags as the race workbook (J2000 equatorial, apparent, topocentric at the racecourse, Europe/London clock time).
Ephemeris folder: ~/astronomy-project/ephe. Off time and course: the batch workbook's META sheet; else the race folder name
(<course>_<hhmm>_<yyyymmdd>) + RACECOURSE_COORDS from ~/astronomy-project/racingpost_excel_charts_swiss.py;
--off / --site override both.

Reproduction check (design F): every sky point the reader uses, engine at the off vs the workbook value, RA and Dec, to
0.0001 deg. A point that fails gets no rate: every hit with that leg is "a/s n/a" and keeps its level from the off alone.
Rate = engine(off + 5) - engine(off), applied to the workbook value at the off (sky value at +5 = workbook + that change)."""
import os, re, sys, ast, glob, csv
from datetime import datetime, timedelta

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
ENGINE_DIR = os.path.join(REPO, 'rebuild_kit')
EPHE = os.path.expanduser('~/astronomy-project/ephe')
BUILDER = os.path.expanduser('~/astronomy-project/racingpost_excel_charts_swiss.py')
RACECARDS = os.path.expanduser('~/astronomy-project/racecards')
REPRO_TOL = 0.0001
STEP_MIN = 5.0
DEFAULT_DUR_S = 300.0


def race_parts(race):
    """'kempton_aw_1440_20210904' -> ('kempton_aw', '1440', '20210904')"""
    m = re.match(r'^(.*)_(\d{4})_(\d{8})$', race)
    if not m:
        raise SystemExit(f'race name not <course>_<hhmm>_<yyyymmdd>: {race}')
    return m.group(1), m.group(2), m.group(3)


def read_meta(batch):
    import openpyxl
    if not batch or not os.path.exists(batch):
        return {}
    wb = openpyxl.load_workbook(batch, read_only=True, data_only=True)
    if 'META' not in wb.sheetnames:
        return {}
    out = {}
    for r in wb['META'].iter_rows(values_only=True):
        if r and r[0] and r[1] is not None and r[0] != 'tab':
            out[str(r[0])] = str(r[1])
        if r and r[0] == 'tab':
            break
    return out


def course_coords(course):
    src = open(BUILDER).read()
    m = re.search(r'RACECOURSE_COORDS\s*=\s*\{', src)
    i = m.end() - 1; depth = 0
    for j in range(i, len(src)):
        depth += src[j] == '{'; depth -= src[j] == '}'
        if depth == 0:
            break
    table = ast.literal_eval(src[i:j + 1])
    if course not in table:
        raise KeyError(f'course {course!r} not in RACECOURSE_COORDS')
    return table[course]


def off_and_site(race, trans_label, batch, off_arg=None, site_arg=None):
    meta = read_meta(batch)
    course, hhmm, ymd = race_parts(race)
    src = []
    if off_arg:
        off = off_arg; src.append('off: --off')
    elif 'race_local_time' in meta:
        off = meta['race_local_time']; src.append('off: META race_local_time')
    else:
        off = f'{ymd[:4]}-{ymd[4:6]}-{ymd[6:]} {hhmm[:2]}:{hhmm[2:]}'; src.append(f'off: race folder name (TRANS label {trans_label!r})')
    if site_arg:
        lat, lon, el = (float(x) for x in site_arg.split(',')); src.append('site: --site')
    elif 'racecourse_lat_lon_elev' in meta:
        lat, lon, el = (float(x) for x in meta['racecourse_lat_lon_elev'].split(',')); src.append('site: META')
    else:
        name = course.replace('_', ' ')
        lat, lon, el = course_coords(name); src.append(f'site: RACECOURSE_COORDS[{name!r}]')
    return off, (lat, lon, el), '; '.join(src)


def race_duration(race):
    """expected_dur_s from the day's race_info.csv (key <yyyymmdd>_<course>_<hhmm>); default 300 s."""
    course, hhmm, ymd = race_parts(race)
    key = f'{ymd}_{course}_{hhmm}'
    for f in sorted(glob.glob(f'{RACECARDS}/{ymd}_*/race_info.csv')):
        with open(f) as fh:
            for r in csv.DictReader(fh):
                if r.get('race') == key and r.get('expected_dur_s'):
                    return float(r['expected_dur_s']), f'{f} ({key})'
    return DEFAULT_DUR_S, 'default 5 min (race not in any race_info.csv)'


def engine_rates(sky, off, site):
    """-> rates {body: (dra_per_min, ddec_per_min) or None}, report lines. sky = {body: (ra, dec, kind)} from positions.csv."""
    lines = []
    try:
        sys.path.insert(0, ENGINE_DIR)
        import swisseph as swe
        import celestial_bodies_swisseph as E
        swe.set_ephe_path(EPHE)
        t0 = E._parse_time(off, 'Europe/London')
        jd0 = E._to_jd_ut(t0)
        jd5 = E._to_jd_ut(t0 + timedelta(minutes=STEP_MIN))
        b0 = E.bodies_at(jd0, site[0], site[1], site[2], False)
        b5 = E.bodies_at(jd5, site[0], site[1], site[2], False)
    except Exception as e:                      # engine cannot run -> every a/s n/a
        lines.append(f'ENGINE NOT RUN: {type(e).__name__}: {e} -> every hit a/s n/a')
        return {b: None for b in sky}, lines
    lines.append(f'engine v{E.ENGINE_VERSION} {os.path.join(ENGINE_DIR, "celestial_bodies_swisseph.py")}, ephe {EPHE}')
    lines.append(f'off {off} Europe/London = {t0.astimezone(__import__("pytz").utc):%Y-%m-%d %H:%M} UTC; site {site}; step {STEP_MIN:g} min')
    rates, bad = {}, []
    for b, (ra, de, kind) in sky.items():
        if b == 'equator':
            rates[b] = (0.0, 0.0); continue
        eb = 'Ketu' if b == 'Ketu' else (b + '_B' if kind == 'star' else b)
        if eb not in b0 or b0[eb]['ra'] is None:
            rates[b] = None; bad.append(f'{b}: not in engine output'); continue
        dra = abs((b0[eb]['ra'] - ra + 180.0) % 360.0 - 180.0)
        dde = abs(b0[eb]['dec'] - de)
        ok = dra <= REPRO_TOL and dde <= REPRO_TOL
        r_ra = ((b5[eb]['ra'] - b0[eb]['ra'] + 180.0) % 360.0 - 180.0) / STEP_MIN
        r_de = (b5[eb]['dec'] - b0[eb]['dec']) / STEP_MIN
        rates[b] = (r_ra, r_de) if ok else None
        lines.append(f'  {b:12s} off: workbook RA {ra:.6f} Dec {de:+.6f} | engine dRA {dra:.2e} dDec {dde:.2e} '
                     f'{"OK" if ok else "FAILS (> 0.0001) -> a/s n/a"} | rate RA {r_ra:+.6f}/min Dec {r_de:+.6f}/min')
        if not ok:
            bad.append(b)
    n_ok = sum(1 for v in rates.values() if v is not None)
    lines.insert(2, f'reproduction: {n_ok} of {len(rates)} sky points within {REPRO_TOL} deg; failing: {", ".join(bad) if bad else "none"}')
    return rates, lines


def moved(pos, rate, minutes=STEP_MIN):
    if rate is None:
        return None
    return ((pos[0] + rate[0] * minutes) % 360.0, pos[1] + rate[1] * minutes, pos[2])
