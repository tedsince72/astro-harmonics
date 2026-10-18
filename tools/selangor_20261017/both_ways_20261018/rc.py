import openpyxl,sys,re
sys.path.insert(0,'tools')
import families as fm
F=None; wb=None; cache={}; P={}; N={}; RUN=[]
def load(hhmm):
    global F,wb,RUN
    F=f'reference/selangor_20261017/race_batch_20261017_selangor_{hhmm}_jupiter_tables.xlsx'
    wb=openpyxl.load_workbook(F,read_only=True); cache.clear(); P.clear(); N.clear()
    for r in rows('POSITIONS')[4:]:
        if r[0] and r[5] is not None:
            P[(r[0],r[2],r[3])]=(r[5],r[6],r[7],r[8])
            if r[2]=='H': N[r[0]]=r[1].split(' / ')[0]
    RUN[:]=sorted(N)
def rows(s):
    if s not in cache: cache[s]=list(wb[s].iter_rows(values_only=True))
    return cache[s]
def trim(r):
    r=list(r)
    while r and r[-1] is None: r.pop()
    return r
def pos(run,ch,b): return P[('SKY' if ch=='S' else run,ch,b)]
def d(a,b,pl):
    if pl=='RA': x=abs(a[0]-b[0])%360; return min(x,360-x)
    if pl=='Dec': return abs(a[1]-b[1])
    return abs(abs(a[1])-abs(b[1]))
def straddle(*ps): s=[p[1]>=0 for p in ps]; return not(all(s) or not any(s))
def fam(v,pl,kind='deg'):
    b=fm.best(v,pl,kind)
    if b: return f"{b['target']} ({b['family']},{b['score']:.2f})"
    o=fm.outside(v,pl,kind); return f"[out {o['target']} {o['off']:+.5f}{'%' if o['pct'] else ''}]" if o else '-'
def ratio3(J,x,y,pl):
    if pl=='Equator' and not straddle(J,x,y): return 'Equator n/a'
    a=d(J,x,pl); b=d(J,y,pl); c=d(x,y,pl); q=max(a,b)/min(a,b)
    return f"{a:.6f} {fam(a,pl)} | {b:.6f} {fam(b,pl)} | dXY {c:.6f} {fam(c,pl)} | ratio {q:.6f} {fam(q,pl,'ratio')} | clos {min(a,b)/max(a,b)*100:.3f}%"

def load_file(path):
    global F,wb,RUN
    F=path
    wb=openpyxl.load_workbook(F,read_only=True); cache.clear(); P.clear(); N.clear()
    for r in rows('POSITIONS')[4:]:
        if r[0] and r[5] is not None:
            P[(r[0],r[2],r[3])]=(r[5],r[6],r[7],r[8])
            if r[2]=='H': N[r[0]]=r[1]
    RUN[:]=sorted(N)
