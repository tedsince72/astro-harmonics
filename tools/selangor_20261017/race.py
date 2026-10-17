import openpyxl,re,glob,sys,collections
U='/root/.claude/uploads/a6920965-38c7-5676-b509-936bc1207077/'
P={'13:45':'afee33f2','14:15':'929b5b49','14:45':'8a883cbf','15:15':'792f2545','15:45':'49a23bd4','16:15':'4f275ecf','16:45':'61dc3d7a','17:15':'692f0455','17:45':'a6b9a230','13:15':'e96aa387'}
hhmm=sys.argv[1]; lo,hi=float(sys.argv[2]),float(sys.argv[3])
wb=openpyxl.load_workbook(glob.glob(U+P[hhmm]+'-*jupiter_tables.xlsx')[0],read_only=True)
isJ=lambda l: bool(re.search(r'(^|[ /|→:])J ',l))
lre=re.compile(r'→ (.+?) \((\d+\.\d+)\) \[(T\d+)')
rows=list(wb['THREADS'].iter_rows(values_only=True)); hdr=[r for r in rows if r and r[0]=='body'][0]
runners={i:str(h).split(' / ')[0] for i,h in enumerate(hdr) if h and str(h).startswith('P')}
links=collections.defaultdict(list) # (runner, body) -> (table, score, text)
for r in rows:
    if not r or not r[0] or r is hdr: continue
    for i,n in runners.items():
        if i<len(r) and r[i]:
            for l in str(r[i]).split('\n'):
                m=lre.search(l)
                if m and not isJ(l) and m[1]!='conj': links[(n,r[0])].append((m[3],float(m[2]),l))
# T17 horse lines
f17=list(wb['FIELD 17'].iter_rows(values_only=True)); h17=[r for r in f17 if r and r[0]=='Jupiter: X = Y'][0]
c17={i:str(h).split(' / ')[0] for i,h in enumerate(h17) if h and str(h).startswith('P')}
for r in f17:
    if not r or not r[0] or not isinstance(r[0],str) or ':' not in r[0] or isJ(r[0]) or r is h17: continue
    m=re.match(r'(\S+) Jupiter: H (\S+) = H (\S+)',r[0])
    if not m: continue
    for i,n in c17.items():
        if i<len(r) and r[i]:
            s=re.search(r'\[(\d+\.\d+)\]',str(r[i]))
            if s and 'not counted' not in str(r[i]):
                for b in (m[2],m[3]):
                    bb='Nodes (Rahu / Ketu)' if b in('Rahu','Ketu') else b
                    links[(n,bb)].append(('T17',float(s[1]),f'{r[0]} {r[1]} {str(r[i])[:80]}'))
DUMP=len(sys.argv)>4
print("#### THREADS (same body, >=70, >=3 tables incl T17)")
for (n,b),L in sorted(links.items()):
    L70=[x for x in L if x[1]>=70]; tabs={x[0] for x in L70}
    if len(tabs)>=3:
        print(' ',n,'|',b,'|',sorted(tabs))
        for x in L70: print('      ',x[2])
print('#### BODY FIELD COUNT: runners with a horse link >=90 on the body (any table incl T17)')
bc=collections.defaultdict(set)
for (n,b),L in links.items():
    if any(x[1]>=90 for x in L): bc[b].add(n)
for b in ['Chiron','Orcus','Saturn','Neptune','Quaoar','Makemake']:
    print(' ',b,len(bc[b]),'/',len(runners),':',sorted(bc[b]))
print('#### TIMED events (area A / phi, score>=85) exact between',lo,hi,'min from scheduled off')
for s in [x for x in wb.sheetnames if re.fullmatch(r'P\d\d',x)]:
    seen=set()
    for r in wb[s].iter_rows(values_only=True):
        if r and r[0] in('TIMELINE H','LONG TIMELINE H') and r[3] and not str(r[3]).startswith('—') and 'H + J' not in str(r[12] or ''):
            try: m=float(r[2])
            except:
                try: m=float(r[-1])
                except: continue
            if lo<=m<=hi and r[10] is not None and float(r[10])>=85:
                k=(r[3],r[4],r[5],r[7])
                if k in seen: continue
                seen.add(k)
                print(' ',s,f'{r[1]} ({m:+.1f})',r[3],r[4],r[5],'→',str(r[7])[:34],'val',r[9],'score',r[10],r[11])
if DUMP:
    for (n,b),L in sorted(links.items()):
        if any(k in n for k in sys.argv[4].split(',')) and b in ('Chiron','Orcus','Saturn','Neptune','Quaoar','Makemake'):
            for x in L:
                if x[1]>=float(sys.argv[5] if len(sys.argv)>5 else 70): print('  BODY',n,'|',b,'|',x[2])
