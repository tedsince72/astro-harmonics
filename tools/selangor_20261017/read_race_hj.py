import openpyxl, re, sys, collections
p=sys.argv[1]
wb=openpyxl.load_workbook(p,read_only=True)
def isJ(s): return bool(re.search(r'(^|[^A-Za-z])J (Jupiter|[A-Z])',s or ''))
sc=re.compile(r'\[(\d+(?:\.\d+)?)\]|\((\d+\.\d+)\)|(?:near|far|RA|Dec|Equator) (\d+\.\d+)$')
out=collections.defaultdict(list)
names={}
for sh,MIN in [('FIELD 5',90),('FIELD 6',90),('FIELD 7',90),('FIELD 8',90),('FIELD 10',50),('FIELD 11',50),('FIELD 12',85),('FIELD 13',85),('FIELD 14',0.01),('FIELD 16',85),('FIELD 17',85),('FIELD 20',50)]:
    if sh not in wb.sheetnames: continue
    rows=list(wb[sh].iter_rows(values_only=True)); cols={}; plane=''
    for rr in rows:
        if rr is None: continue
        hdr={i:str(x) for i,x in enumerate(rr) if isinstance(x,str) and re.match(r'P\d\d ',x)}
        if len(hdr)>=3:
            cols=hdr; 
            for i,x in hdr.items(): names[x[:3]]=x.split(' /')[0]
            if rr[0] in ('RA','Dec','Equator'): plane=rr[0]
            continue
        if isinstance(rr[0],str) and rr[0].split(' ')[0] in ('RA','Dec','Equator') and ('AREA' in rr[0] or 'NO HIT' in rr[0]):
            plane=rr[0].split(' ')[0]; nohit='NO HIT' in rr[0]; continue
        if not cols: continue
        label=' | '.join(str(x) for x in rr[:2] if x is not None and not (isinstance(x,str) and re.match(r'P\d\d ',x)))[:70]
        pass
        for c,name in cols.items():
            if c>=len(rr) or rr[c] is None: continue
            t=str(rr[c])+((' '+str(rr[c+1])) if c+1<len(rr) and (c+1) not in cols and rr[c+1] is not None else '')
            if 'outside' in t: continue
            m=[g for tup in sc.findall(t) for g in tup if g]
            if not m: continue
            s=max(float(x) for x in m)
            if s>=MIN: out[name[:3]].append((s,sh,plane,label.split(' | ')[0],t.replace('\n',' ')[:100]))
# THREADS horse-only mutual
ws=wb['THREADS']; rows=list(ws.iter_rows(values_only=True))
hdr=[r for r in rows if r and r[0]=='body' and r[1] and str(r[1]).startswith('P01')][0]
mut=collections.defaultdict(list)
for r in rows:
    if not r or not r[0] or r is hdr: continue
    for i in range(1,len(hdr)):
        h=hdr[i]
        if not h or not str(h).startswith('P') or r[i] is None: continue
        L=str(r[i]).split('\n')
        a=[l for l in L if re.match(r'S Jup→[HJ] ',l) and re.search(r'\((\d+\.\d+)\)',l) and float(re.search(r'\((\d+\.\d+)\)',l).group(1))>=50]
        b=[l for l in L if re.match(r'S \S+( \S+)?→[HJ] Jup',l) and float(re.search(r'\((\d+\.\d+)\)',l).group(1))>=50]
        if a and b: mut[str(h)[:3]].append(r[0])
# timeline in race
tl=collections.defaultdict(list)
for s in [n for n in wb.sheetnames if re.fullmatch(r'P\d\d',n)]:
    for r in wb[s].iter_rows(values_only=True):
        if r and r[0] in ('TIMELINE H','TIMELINE J') and r[3] and not str(r[3]).startswith('—'):
            try: m=float(r[2])
            except: continue
            if -3<=m<=8 and str(r[3]) in ('T5','T6','T7','T8','T11','T12','T13','T14','T17','T20','T16','T10') :
                tl[s].append(f"{r[1]} ({m:+.1f}) {r[3]} {r[4]} {r[5]} → {str(r[7])[:25]} val {r[9]} score {r[10]} {r[11]}")
for k in sorted(names):
    print('##',names[k],'| horse mutual:',mut.get(k,[]))
    for x in sorted(out.get(k,[]),reverse=True)[:16]: print(f'   {x[0]:6.2f} {x[1]:9s} {x[2]:7s} {x[3]} :: {x[4]}')
    for t in tl.get(k,[]): print('   TL',t)
