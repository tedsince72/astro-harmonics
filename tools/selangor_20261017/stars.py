import openpyxl,re,glob,sys,collections
U='/root/.claude/uploads/a6920965-38c7-5676-b509-936bc1207077/'
R={'12:15':'7566745e','12:45':'2673848e','13:15':'e96aa387'}
for hhmm in sys.argv[1:]:
    wb=openpyxl.load_workbook(glob.glob(U+R[hhmm]+'-*jupiter_tables.xlsx')[0],read_only=True)
    print('########',hhmm)
    # B: field star matrices, H hits
    for sh in ['FIELD 21','FIELD 22','FIELD 21R','FIELD 22R']:
        rows=list(wb[sh].iter_rows(values_only=True)); hdr=[r for r in rows if r and r[0] in('star','star pair')][0]
        area=''
        for r in rows:
            if not r or r is hdr or not r[0]: continue
            if str(r[0]).startswith('AREA') or 'NO HIT' in str(r[0]): area=str(r[0])[:12]; continue
            hold=[]
            for i,h in enumerate(hdr):
                if i==0 or not h or not str(h).startswith('P') or i>=len(r) or not r[i]: continue
                for l in str(r[i]).split('\n'):
                    m=re.match(r'H (\S+) (.+) \((\d+\.\d+)\)',l)
                    if m: hold.append((str(h).split(' / ')[0][4:],l,float(m[3])))
            n=len({h[0] for h in hold})
            for h in hold:
                if h[2]>=70: print(f'  {sh:9s} {area:12s} {r[0]:24s} {h[0]:20s} {h[1]}   H-holders {n}')
    # A: same-as-sky tags and LONG TIMELINE star rows
    for s in [n for n in wb.sheetnames if re.fullmatch(r'P\d\d',n)]:
        rows=list(wb[s].iter_rows(values_only=True)); col=None
        name=[r for r in rows if r and r[0]=='TIMELINE H'][:1]
        for r in rows:
            if not r: continue
            if 'same as sky' in r: col=r.index('same as sky'); continue
            if r[0] and isinstance(r[0],str) and (r[0].startswith('T') and len(r[0])>12 or r[0].startswith('TIMELINE')): col=None if not r[0].startswith('T19 AREA') and not r[0].startswith('T21') and not r[0].startswith('T19 NO') else col
            if col is not None and col<len(r) and r[col] and str(r[col]).strip() and str(r[col])!='same as sky':
                lab=' | '.join(str(x) for x in r[:4] if x is not None)
                if re.search(r'(^|[ /|→:])J ',lab) and not str(r[col]).startswith('H'): continue
                print(f'  TAG {s} :: {lab[:90]} :: {str(r[col])[:120]}')
            if r[0]=='LONG TIMELINE H' and r[3] in ('T21','T22','T21R','T22R') and r[10] and float(r[10])>=90:
                print(f'  STAR-TL {s} :: {r[3]} {r[4]} {r[5]} → {str(r[7])[:28]} val {r[9]} score {r[10]} {r[11]} {r[1]}')
