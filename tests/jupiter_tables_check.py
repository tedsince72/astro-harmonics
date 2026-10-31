#!/usr/bin/env python3
"""jupiter_tables_check.py - tests for tools/jupiter_tables.py (13 Oct 2026). Builds the three test workbooks and checks cells.
13 Oct 2026, Eddie's correction (tools/families.py): scores replace tolerance bands (green >= 80, yellow 50-80, light < 50);
ninths = the 48 _NINTH_RANGES only. Test 1's ninth check (Dec 15.222681 -> 137/9, not one of the 48) replaced by T5 S Jupiter -> H Mercury
RA 14.4444; expected values changed: T6 RA 18 green -> yellow (whole score 62.9), Pholas phi yellow -> green (golden 1xphi),
Venturous S Uranus Dec 7.008205 yellow -> light (whole score 34.78).
Same day, Eddie's interval list for ratios (0.15 %, squared): Pholas phi now an interval, phi · golden ratio · Phi, +0.05151 %, score 43.11,
light (green under the first families.py build, yellow under JR v3 bands); RATIOS BY GROUP sheet checked.
Test 6 (Eddie 13 Oct): Jupiter return line (information); per-plane automatic unisons (d(S Jup, natal Jup) / d(natal Jup, X) < 0.0015) greyed, not counted;\nreturn-assisted tag for 0.0015 <= d1/d2 < 0.02; empty RATIOS group headers hidden.\nTest 5 (Eddie 13 Oct): unison counts, with its field count.\nTest 4 (T16 RA–Dec crossing, Eddie 13 Oct): Sandown 14:25 2021 Nube Negra / Harry Skelton, Jupiter–Orcus J RA = H Dec.
Test 7 (Eddie 13 Oct): tapered ratio tolerances (unison 1 %, 1 % / k, floor 0.15 %; Phi group φ, φ², φ³, 1+√2 only) and the nearest
target + offset shown outside every tolerance (plain text, no score, no colour). Expected values changed by it: test 1 Dec 15.222681
now shows '15.5556 (ninth) — outside ±0.106' -0.33287 (still no score, no colour); test 5 P05 now has 5 unison rows (the Rahu ratio is a
unison in RA, Flat and GC); test 6 Cobden J Pluto GC 1.001353 score 0.96 light -> 74.77 yellow, Cobden 41 -> 61 unisons (20 automatic,
all RA; 41 genuine: Flat 20, GC 20, RA 1, all return-assisted, two at 2/6).
Revised same day (Eddie): k = the LARGER number of the ratio (3/2 -> 0.333 %, 4/3 -> 0.25 %, 5/x -> 0.2 %, 6/5 and up -> 0.15 %) and the
nearest target outside the tolerance by tightness (offset / own tolerance) for ratios too. Test 7 changed: Flat 1.0445 nearest 16/15
-2.07813 % -> unison +4.44999 % (outside ±1 %); Haumea 1.506122 3/2 score 3.37 -> outside ±0.333 %; the 3/2 hit is now Makemake 1.499624
on 0.333 % (85.52). Tests 1-6 unchanged.
Test 8 (Eddie 13 Oct): THREADS sheet, Nottingham 15:30 2022 Uranus row; expected values set from the rule (hits >= 50 for tags, counts
and cell lines; mutual same chart / across charts), not tuned to the winner. Same day: scaled repeats of a main hit (same family, same
digits) dropped and score-0 hits left out of '+ n weaker': Dynamic Force Uranus '+ 11 weaker' -> '+ 9 weaker', 5 scaled lines -> 1.
Test 9 (Eddie 13 Oct): NUMBERS sheet per runner + FIELD NUMBERS, Wincanton 14:20 2022 P07 7.777 (ninth) row.
Same day (Eddie): a repeat (same family, same digits) keeps the higher form at the shown precision (tie = main), in THREADS and NUMBERS;
H Jupiter – J Jupiter in its own H–J column, once. Changed: test 8 '+ 9 weaker' -> '+ 8 weaker' (J Jup→H Uranus 14.14 shown as /10 1.414,
99.1); test 9 J Haumea 77.798638 (83.67) -> its /10 form 7.779864 (83.68); P07 silver 0.414 sky + horse + jockey -> two of three.
14 Oct (Eddie): keep-higher compares unrounded scores, display 2 dp in THREADS and NUMBERS. Changed: Dynamic Force S Jup→H Uranus 11.10793
now shown as /10 1.110793 (94.7963 > 94.7947), '+ 8 weaker' -> '+ 9 weaker'; Wincanton P07 H Rigel 7.777589 now shown as x10 77.775893
(98.5224 > 98.5209); scores in THREADS lines at 2 dp.
14 Oct, revised (Eddie): the main form is kept unless the scaled form scores > 0.5 higher (unrounded); all scores at 2 dp everywhere
(scaled, also, FIELD matrices). Changed back: Dynamic Force 11.10793 main (94.80), Wincanton P07 H Rigel 7.777589 main (98.52) and
J Haumea 77.798638 main (83.67); FIELD 16 tag '[92.3]' -> '[92.30]'.
Test 10 (Eddie 14 Oct): midpoints - Equator plane (|Dec|, only when the three are not on one side), RA far (span 360 - near), tolerance
green = max(0.05, 0.5 % of span), yellow 3x, light 9x. No earlier test checked a midpoint value; THREADS test 8 now also lists midpoint
lines (they carry span and tolerance); its tags and counts are unchanged.
Revised 14 Oct (Eddie): green = 0.1 % of span, floor 0.05, cap 0.15 (yellow 3x, light 9x); FIELD 11 added. Test 10 changed: tolerance
cases; Dynamic Force T10 Equator tol ±0.1108 -> ±0.0500 (still green).
14 Oct (Eddie): Jupiter-only midpoint lines in T11 only; NUMBERS HORSE = T1, JOCKEY = T2 (no T3 / T4, no H–J column). Changed: test 9's
H–J check replaced; Wincanton FIELD NUMBERS 6 -> 5 numbers (sqrt2 0.707 / 70.7 drops out; 166.667 3/6 -> 1/6; silver 2.414 2/6 -> 1/6).
Test 11 (Eddie 14 Oct): EQUIDISTANCE on SKY JUPITER (Beverley 14:05 2024, sky-only workbook copied from ~/astronomy-project/sky_0608_wb),
T17 + FIELD 17 (Wincanton). Same day: "X, Y together" (d(X, Y) < 1.5 RA / 0.3 other planes) tagged, at the bottom, not counted in
FIELD 17; a Ketu line with the same numbers as a Rahu line dropped. Changed: test 11 best-first check now above the together lines.
Test 12 (Eddie 14 Oct): EQUIDISTANCE family hits for d(Jup, X), d(Jup, Y), d(X, Y) and timing (ratio +5, arrow, applying / separating,
exact clock time, linear) - Lingfield AW 14:35 2021 batch (copied from ~/astronomy-project/racecards/20210402_lingfield_aw; xlsx only).
Same day: each family hit split into a value cell and a family cell (coloured by its score); test 12 column positions changed.
14 Oct (Eddie): THREADS every family hit >= 50 (not only the best), cell colour = best line - test 8 '+ 9 weaker' -> '+ 11 weaker';
T19 / PAIRS H–J across = same body only; T18 never a 'together' plane as the closest plane.
Test 13 (Eddie 14 Oct): T18 - equidistance pair side by side for horse and jockey (block 1 S Jupiter, block 2 natal Jupiters);
'both' tag and count. Test 14: T19 sky Jupiter on natal pairs, P## PAIRS, SKY STARS. (No Messi workbook in the repo: that check skipped.)
29 Oct 2026 (Eddie, workbook trim): the workbook keeps POSITIONS, SKY JUPITER and T1-T8, T24 HORSE / JOCKEY on the runner sheets (plus five
Jupiter rows on top of T24 HORSE / JOCKEY); no --full-only sheets, no SKY STARS SUMMARY, no _pairs workbook. Every check of a removed table,
sheet or the pairs workbook is retired in place ('# retired 29 Oct 2026 (Eddie, workbook trim): ...'); T24 HORSE / JOCKEY readers skip the
Jupiter rows (column B with '→' or 'stars')."""
import os, re, subprocess, sys
import openpyxl
from openpyxl.cell.rich_text import CellRichText

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(REPO, "tools"))
os.chdir(REPO)
GREEN, YELLOW, LIGHT = '00C6EFCE', '00FFEB9C', '00FFF8DC'
COL = {'RA': 2, 'Dec': 10}                                # 0-based column of each plane's value (RA, Dec; 8 columns each since the scaled column went, 16 Oct 2026)
RES = []


def check(name, ok, detail):
    RES.append(ok); print(f"  {'PASS' if ok else 'FAIL'}  {name}: {detail}")


def old_code(commit):
    """tools/jupiter_tables.py of an earlier commit, for the comparisons with a build before. It runs on today's families.py (tools/); since
    24 Oct 2026 (half family) NUMBERS lists 'half' too, so the old code gets the same NUM_FAMILIES line - each comparison still checks its
    own change only"""
    src = subprocess.run(['git', 'show', f'{commit}:tools/jupiter_tables.py'], capture_output=True, text=True).stdout
    return src.replace("NUM_FAMILIES = ('whole', 'golden',", "NUM_FAMILIES = ('whole', 'half', 'golden',")


import time
T0, NBUILT = time.time() - 1, [0]


def build(inp):
    NBUILT[0] += 1
    r = subprocess.run([sys.executable, 'tools/jupiter_tables.py', inp, '--full'], capture_output=True, text=True)   # every sheet (19 Oct 2026)
    print(r.stdout.strip() or r.stderr)
    return r.stdout.split(' -> ')[1].split(' (')[0]


def row(ws, tid, label):
    for r in ws.iter_rows():
        if r[0].value == tid and r[1].value == label:
            return r


def sheet_of(wb, horse):
    return next(wb[s] for s in wb.sheetnames if s.startswith('P') and horse in str(wb[s]['A1'].value))


def cell(r, pl, col0=None):
    c0 = COL[pl] if col0 is None else col0
    c = r[c0]
    return c.value, r[c0 + 2].value, r[c0 + 4].value, (c.fill.start_color.rgb if c.fill.fill_type else None), r[c0 + 5].value


line = lambda out, k: next((l for l in out.splitlines() if l.startswith(k)), '')

# ---------------------------------------------------------------------------------------------------------------------------------
# Cached old builds (Eddie, 31 Oct 2026). Every workbook built by an old commit's code goes to tests/.cache/<commit>/ and is built only when
# missing. Key: commit, race (input name), options, a hash of the input's content (the cell values of every sheet, without the META row
# 'created_utc' that a card-pipeline rebuild stamps; or every file of a folder), and for old_code() builds (old jupiter_tables.py on today's tools/) a hash of
# today's tools/families.py, tools/jup2/*.py and rebuild_kit/*.py. run_build() copies the cached workbook (and a _pairs sibling) to the path
# the test asks for and returns the build's stdout, so the checks below are unchanged.
import hashlib, types, zipfile, glob, shutil as _shutil
CACHE = os.path.join(REPO, 'tests', '.cache')
ARCH = {}                                                  # extracted tools/ folder -> commit


def _h(b):
    return hashlib.sha1(b).hexdigest()[:12]


_KEYS = {}


def input_key(inp):
    p = os.path.abspath(inp)
    st = os.stat(p)
    k = (p, st.st_mtime, st.st_size)
    if k not in _KEYS:
        h = hashlib.sha1()
        if os.path.isdir(p):
            for f in sorted(os.listdir(p)):
                if f.endswith('.xlsx') and not f.startswith('~$'):
                    h.update(f.encode()); h.update(input_key(os.path.join(p, f)).encode())
        else:                                              # cell values of every sheet; the META row 'created_utc' (a card-pipeline rebuild's time) left out
            _w = openpyxl.load_workbook(p, read_only=True)
            for _ws in _w.worksheets:
                h.update(_ws.title.encode())
                for _r in _ws.iter_rows(values_only=True):
                    if not (_r and _r[0] == 'created_utc'):
                        h.update(repr(_r).encode())
            _w.close()
        _KEYS[k] = h.hexdigest()[:12]
    return _KEYS[k]


def env_key():
    files = ['tools/families.py'] + sorted(glob.glob('tools/jup2/*.py')) + sorted(glob.glob('rebuild_kit/*.py'))
    return _h(b''.join(open(os.path.join(REPO, f), 'rb').read() for f in files))


def extract(commit, d):
    """git archive <commit> tools rebuild_kit -> d (as before; cheap), remembered so run_build() knows the commit of d/tools/jupiter_tables.py"""
    os.makedirs(d, exist_ok=True)
    if not os.path.exists(os.path.join(d, 'tools', 'jupiter_tables.py')):
        subprocess.run(f'git archive {commit} tools rebuild_kit | tar -x -C {d}', shell=True, check=True)
    ARCH[os.path.abspath(d)] = commit


NBUILT_OLD = [0, 0]                                        # [built, from the cache]


def run_build(argv, **kw):
    """an old build: argv = [script, input, options..., '--out', out]; script = tools/_jt_before_<commit>.py (old_code) or <extracted>/tools/
    jupiter_tables.py (git archive). -> object with .stdout / .returncode, as subprocess.run"""
    script, inp, rest = argv[0], argv[1], list(argv[2:])
    i = rest.index('--out')
    out, opts = rest[i + 1], rest[:i] + rest[i + 2:]
    m = re.search(r'_jt_before_(\w+)\.py$', script)
    if m:
        commit, extra = m.group(1), '_env' + env_key()
    else:
        commit, extra = ARCH[os.path.abspath(os.path.dirname(os.path.dirname(script)))], ''
    stem = os.path.basename(inp.rstrip('/')).replace('.xlsx', '')
    key = f"{stem}_{input_key(inp)}_{'-'.join(o.strip('-') for o in opts) or 'def'}{extra}"
    d = os.path.join(CACHE, commit)
    path = os.path.join(d, f'{key}_jupiter_tables.xlsx')
    if not os.path.exists(path):
        os.makedirs(d, exist_ok=True)
        r = subprocess.run([sys.executable, script, inp] + opts + ['--out', path], capture_output=True, text=True)
        if r.returncode != 0 or not os.path.exists(path):
            print(f'  old build failed: {commit} {stem} {opts}: {r.stderr[-400:]}')
            return types.SimpleNamespace(stdout=r.stdout, stderr=r.stderr, returncode=r.returncode or 1)
        open(path + '.log', 'w').write(r.stdout)
        NBUILT_OLD[0] += 1
    else:
        NBUILT_OLD[1] += 1
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    _shutil.copy(path, out)
    pairs = path.replace('_jupiter_tables.xlsx', '_pairs.xlsx')
    if os.path.exists(pairs):
        _shutil.copy(pairs, out.replace('_jupiter_tables', '_pairs'))
    log = open(path + '.log').read() if os.path.exists(path + '.log') else ''
    return types.SimpleNamespace(stdout=log.replace(path, out), stderr='', returncode=0)



# ---------------------------------------------------------------------------------------------------------------------------------
# QUICK / FULL (Eddie, 31 Oct 2026). The rest of this file is cut into segments by '#@ quick: <why>' / '#@ full: <why>' lines (the code before the
# first marker is quick). Default run: the quick segments - every check on the current build (by-hand cells from POSITIONS, family rules,
# layout, order, counts) and the comparison with the commit before the current change. --full: every segment, as before (every old-commit
# comparison). All segments run in this module's namespace, in file order, with their own line numbers.
FULL = '--full' in sys.argv
if not globals().get('_SEGMENTED'):
    _SEGMENTED = True
    _src = open(os.path.abspath(__file__), encoding='utf-8').read().split('\n')
    _me = next(i for i, l in enumerate(_src) if l.startswith('#@@ end of the segment runner'))
    _segs, _mode, _why, _start = [], 'quick', 'tests 1 - 30 on the current builds', _me + 1
    for _i in range(_me + 1, len(_src) + 1):
        if _i == len(_src) or _src[_i].startswith('#@ '):
            _segs.append((_mode, _why, _start, _i))
            if _i < len(_src):
                _mode, _why = _src[_i][3:].split(':', 1)
                _mode, _why, _start = _mode.strip(), _why.strip(), _i + 1
    _skip = [s for s in _segs if s[0] == 'full' and not FULL]
    print(f"{'FULL' if FULL else 'QUICK'} run: {len(_segs) - len(_skip)} of {len(_segs)} segments"
          + ('' if FULL else f" ({len(_skip)} full-only segments skipped; --full runs them)"))
    for _mode, _why, _a, _b in _segs:
        if _mode == 'full' and not FULL:
            continue
        exec(compile('\n' * _a + '\n'.join(_src[_a:_b]), os.path.abspath(__file__), 'exec'), globals())
    sys.exit(0 if all(RES) else 1)
#@@ end of the segment runner
print('Test 1: Nottingham 15:30 2022 (batch xlsx) — Dynamic Force')
wb = openpyxl.load_workbook(build('reference/classic/race_batch_20220601_nottingham_1530.xlsx'))
ws = sheet_of(wb, 'Dynamic Force')
v, t, o, f, sc = cell(row(ws, 'T5', 'S Jupiter → H Mercury'), 'RA')
# 24 Oct 2026 (Eddie, 4-digit rule): was ninth 14.4444 (-0.00459, 95.14) green; 14.43... vs 14.44... -> outside the digit match
check('T5 S Jupiter → H Mercury RA 14.439854: ninth 14.4444 outside the 4-digit match (1443 vs 1444; was 95.14 green, 24 Oct 2026) -> '
      'shown outside, nearest 14.4444 (ninth), -0.00459, no score, no colour',
      round(v, 6) == 14.439854 and t == '14.4444 ninth — not a 4-digit match (-0.00459)' and o == '-0.00459' and sc is None and f is None, f'{v} {t} {o} {sc} {f}')
r = row(ws, 'T6', 'S Jupiter → J Vesta')
v, t, o, f, sc = cell(r, 'RA')
check('T6 S Jupiter → J Vesta RA 18 (+0.00414, whole score 62.9), yellow', t == '18 (whole)' and o == '+0.00414' and sc == 62.9 and f == YELLOW,
      f'{v} {t} {o} {sc} {f}')
v, t, o, f, sc = cell(r, 'Dec')
check('T6 S Jupiter → J Vesta Dec 15.222681: no family (137/9 is not one of the 48), no colour, no score; nearest shown: '
      '15.5556 (ninth) — outside ±0.106, -0.33287', round(v, 6) == 15.222681 and t == '15.5556 (ninth) — outside ±0.106' and o == '-0.33287'
      and f is None and sc is None,
      f'{v} {t} {o} {sc} {f}')

print('\nTest 2: Pholas, Lingfield 2 Apr 2021 (split folder, older workbook) — T7')
wb = openpyxl.load_workbook(build('reference/classic/lingfield_aw_1435_20210402'))
ws = sheet_of(wb, 'Pholas')
# retired 29 Oct 2026 (Eddie, workbook trim): T14 |Dec| S Jupiter / |Dec| J Jupiter 1.61887 -> φ (43.11, light, Group Phi), and the P## RATIOS
# (RATIOS BY GROUP) sheet listing it under Phi, marked (inverted) - T14 and the RATIOS sheets are gone
r = row(ws, 'T7', 'S Rahu → H Jupiter')
check('arrow n/a where the engine does not reproduce (S Rahu)', r[COL['RA'] + 1].value == 'n/a' and r[COL['RA'] + 6].value == 'n/a',
      f"+5 {r[COL['RA'] + 1].value}, arrow {r[COL['RA'] + 6].value}")

print('\nTest 3: Newcastle AW 13:30 2022 (batch xlsx) — Venturous, T7')
wb = openpyxl.load_workbook(build('reference/classic/race_batch_20220102_newcastle_aw_1330.xlsx'))
ws = sheet_of(wb, 'Venturous')
v, t, o, f, sc = cell(row(ws, 'T7', 'S Venus → H Jupiter'), 'Dec')
check('S Venus → H Jupiter Dec 40.00191 (whole score 81.81), green', round(v, 5) == 40.00191 and t == '40 (whole)' and sc == 81.81 and f == GREEN,
      f'{v} {t} {o} {sc} {f}')
v, t, o, f, sc = cell(row(ws, 'T7', 'S Uranus → H Jupiter'), 'Dec')
check('S Uranus → H Jupiter Dec 7.008205 (whole score 34.78), light yellow', round(v, 4) == 7.0082 and t == '7 (whole)' and sc == 34.78 and f == LIGHT,
      f'{v} {t} {o} {sc} {f}')

print('\nTest 4: Sandown 14:25 2021 (batch xlsx) — build only (test 7 reads its T5)')
wb = openpyxl.load_workbook(build('reference/classic/race_batch_20211204_sandown_1425.xlsx'))
# retired 29 Oct 2026 (Eddie, workbook trim): T16 Nube Negra Jupiter–Orcus J RA 30.55841 = H Dec 30.55995, RA − Dec −0.001541, 92.3 green,
# and FIELD 16 showing it ('crossing -0.00154 [92.30]') - T16 and the FIELD sheets are gone

# retired 29 Oct 2026 (Eddie, workbook trim): Test 5, Newcastle AW 12:20 2021 - T12 P05 / P11 S Jupiter–S Rahu / H Jupiter–H Rahu RA unison
# with "unison: field 2/6", P05 RATIOS Rahu RA unison row 1.005972 (49.20) field 2/6, FIELD 12 area A "field 2/6" for P05 and P11 - T12, the
# RATIOS and FIELD sheets are gone (its build removed)

# retired 29 Oct 2026 (Eddie, workbook trim): Test 6, Wincanton 14:20 2022 P09 Cobden Jupiter return - T13 return line kept as information,
# T13 RA unison J Makemake automatic (grey, not counted), P09 RATIOS 25 unisons (20 automatic, 5 genuine return-assisted 1/6), P03 RATIOS no
# empty group header - T13 and the RATIOS sheets are gone (its build removed)

print('\nTest 7: Sandown 14:25 2021 (batch xlsx) — nearest target outside the tolerance, families tolerances')
wb = openpyxl.load_workbook('reference/classic/race_batch_20211204_sandown_1425_jupiter_tables.xlsx')   # built in test 4
ws = wb['P05']
# retired 29 Oct 2026 (Eddie, workbook trim): P05 T12 RA Mars 1.002631 unison (75.42, yellow), P07 T12 RA Vesta 1.499361 3/2 (93.71, green),
# P07 T13 RA J Mars 1.125881 9/8 (22.83, light) - T12 / T13 are gone (the tolerances stay checked at families level below)
v, t, o, f, sc = cell(row(ws, 'T5', 'S Jupiter → H Mars'), 'Dec')
check('degrees outside every family: T5 S Jupiter → H Mars Dec 8.398276 -> 5×φ (golden) — outside ±0.0324, +0.30811, plain',
      t == '5×φ (golden) — outside ±0.0324' and o == '+0.30811' and sc is None and f is None, f'{v} {t} {o} {sc} {f}')
import families
TK = ('1', '2', '3/2', '3', '4/3', '4', '5/4', '5', '5/2', '5/3', '6/5', '8/5', '8/3', '9/4', '7/4', '13/4', '11/5', '7/2', '√2', '1+1/√2', 'φ')
check('Phi group = φ, φ², φ³, 1+√2 only; area A (1, 3/2, 2, 3 ... 9) ±0.02 absolute (15 Oct); area B by the larger number (1 % / k, 0.15 % from k = 6)',
      [i[0] for i in families.INTERVALS if i[3] == 'Phi'] == ['φ', 'φ²', 'φ³', '1+√2']
      and families.AREA_A == ('1', '3/2', '2', '3', '4', '5', '6', '7', '8', '9')
      and [round(families.TOL[k], 3) for k in TK] == [0.02, 0.02, 0.02, 0.02, 0.25, 0.02, 0.2, 0.02, 0.2, 0.2] + [0.15] * 11
      and all(families.TOL[k] == 0.15 for k in families.TOL
              if k[0].isdigit() and '√' not in k and k not in families.AREA_A and max(int(x) for x in k.split('/')) >= 6), str(families.TOL))

# retired 29 Oct 2026 (Eddie, workbook trim): Test 8, Nottingham 15:30 2022 THREADS - Dynamic Force Uranus cell (T8 S Uranus→J Jup Dec
# 18.995206 → 19 (57.80) listed, T5 ninth 11.10793 gone), Transpluto tags 'directions 1', cells list hits >= 50 on the value as it is
# (J Jup→H Uranus Dec 14.141031 → 14.14 (96.90) [T4]) - the THREADS sheet is gone

# retired 29 Oct 2026 (Eddie, workbook trim): Test 9, Wincanton 14:20 2022 NUMBERS - P01 NUMBERS 'ninth: 11.1111' HORSE only, P11
# 'whole: 10' two of three, NUMBERS columns (HORSE = T1, JOCKEY = T2, no H–J column), FIELD NUMBERS without sqrt2 0.707 / 70.7 and with
# no number held sky + horse + jockey, every NUMBERS hit >= 40 - the NUMBERS and FIELD NUMBERS sheets are gone (its build removed)


print('\nTest 10: midpoints (Eddie 14 Oct 2026, revised scoring) — closeness half + distance half (each x 0.5), RA closer of near / far')
# rewritten 14 Oct 2026: the span-scaled tolerance (0.1 % of span, floor 0.05, cap 0.15, yellow 3x, light 9x) is removed
import jupiter_tables as JT
g = JT.mid_geom(10, 20, -15, 'Equator')
check('geometry: Equator Dec +10, +20, −15 -> |Dec| midpoint 15, offset 0, distances 5 / 5; all on one side -> not listed',
      g == (0.0, '', 5.0, 5.0) and JT.mid_geom(10, 20, 15, 'Equator') is None, str(g))
check('geometry: RA 350 / 30 -> c = 10 at the near midpoint (offset 0, near); c = 190 at the far one (offset 0, far); never both',
      JT.mid_geom(350, 30, 10, 'RA')[:2] == (0.0, 'near') and JT.mid_geom(350, 30, 190, 'RA')[:2] == (0.0, 'far')
      and JT.mid_geom(350, 30, 100, 'RA')[1] == 'near' and JT.mid_geom(350, 30, 101, 'RA')[1] == 'far', '')
m = JT.mid_score(10.0, 10.05, 0.1, 'Dec')
check('a hit needs both halves: Dec 10 / 10.05 closeness 99.5025 % (45.15) and offset 0.1 on 0.15 (5.56) -> total 50.70; offset 0.16 -> '
      'no score; 10 / 11 closeness 90.9091 % still inside (0.41); 10 / 11.2 = 89.2857 % under 90 % -> no score; RA 100 / 100.5, offset 0.25 '
      'on 0.5 -> 45.15 + 12.50 = 57.65, offset 0.51 -> no score',
      m['hit'] and round(m['score'], 2) == 50.70 and round(m['rscore'], 2) == 45.15 and round(m['dscore'], 2) == 5.56
      and JT.mid_score(10.0, 10.05, 0.16, 'Dec')['score'] is None and round(JT.mid_score(10.0, 11.0, 0.1, 'Dec')['rscore'], 2) == 0.41
      and JT.mid_score(10.0, 11.2, 0.1, 'Dec')['score'] is None and round(JT.mid_score(100.0, 100.5, 0.25, 'RA')['score'], 2) == 57.65
      and not JT.mid_score(100.0, 100.5, 0.51, 'RA')['hit'], str(m))
# retired 29 Oct 2026 (Eddie, workbook trim): Nottingham Dynamic Force T10 S Jupiter / J Uranus | H Uranus Equator (offset +0.027875,
# 99.4981 %, 45.11 + 33.14 = 78.25), the T10 header planes and FIELD 10 / 11 planes, and the three Jupiter-only midpoint lines in T11 only
# - T10, T11 and the FIELD sheets are gone (mid_geom / mid_score stay checked above)

print('\nTest 11 (Eddie 14 Oct): EQUIDISTANCE — Beverley 14:05 2024 (sky-only workbook), SKY JUPITER')
wb = openpyxl.load_workbook(build('reference/classic/race_batch_20240608_beverley_1405.xlsx'))
eq = [r for r in wb['SKY JUPITER'].iter_rows(values_only=True) if r[0] == 'SKY EQ' and r[2]]
# retired 14 Oct 2026 (GC removed): Beverley S Jupiter: S Ceres = S Pallas GC 133.11392 / 133.19155, 0.999417, far side. Beverley's SKY EQ
# has no non-together line left; the value case is moved to Lingfield 14:35 2021 SKY EQ (Equator, found by the rule), checked in test 12.
tg = [r for r in eq if r[10].endswith('X, Y together')]     # columns (14 Oct, revised): closeness %, midpoint offset, ratio half, distance half, total
nt = [r for r in eq if not r[10].endswith('X, Y together')]
# 14 Oct 2026 (midpoint scoring): Beverley's together lines (S Sun = S Pluto Equator, d(X, Y) 0.003798) are no longer hits - Jupiter is
# not near the midpoint of two bodies in conjunction in Dec / Equator - so that case is retired; moved to Newcastle 13:30 2022 P07 (RA far)
check('"X, Y together" lines tagged, d(X, Y) under the conjunction tolerance, and listed at the bottom (Beverley SKY EQ; none left there)',
      (eq[-len(tg):] == tg if tg else True)
      and all(r[11] < (1.5 if r[2] == 'RA' else 0.3) for r in tg) and all(r[11] >= (1.5 if r[2] == 'RA' else 0.3) for r in nt), str(tg))
# retired 29 Oct 2026 (Eddie, workbook trim): '"X, Y together" (moved)': Newcastle 13:30 P07 T17 S Jupiter: H Orcus = H Transpluto RA
# 179.975528 / 179.599503, far −0.212485, d(X, Y) 0.376025, tagged and at the bottom - T17 is gone
check('hits only (closeness >= 90 %, total scored), best first above the together lines; no Rahu = Ketu Equator line',
      all(r[5] >= 90 and r[9] is not None for r in eq) and [r[9] for r in nt] == sorted([r[9] for r in nt], reverse=True)
      and not any('Rahu' in r[1] and 'Ketu' in r[1] and r[2] == 'Equator' for r in eq), str(len(eq)))
# retired 29 Oct 2026 (Eddie, workbook trim): Wincanton P07 T17 present (H, J and S Jupiter lines) with FIELD 17 holder counts; P05
# Ketu line repeating a Rahu line dropped (S Jupiter: H Juno = H Ketu Dec gone, H Juno = H Rahu Equator kept); FIELD 17 holder counts
# leave out together cells (Newcastle 13:30 H Orcus = H Transpluto RA 0/5) - T17 and the FIELD sheets are gone

print('\nTest 12 (Eddie 14 Oct): EQUIDISTANCE family hits and timing — Lingfield AW 14:35 2021 (batch), P09 Pholas / Hollie Doyle')
wb = openpyxl.load_workbook(build('reference/classic/race_batch_20210402_lingfield_aw_1435.xlsx'))
sq = [r for r in wb['SKY JUPITER'].iter_rows(values_only=True) if r[0] == 'SKY EQ' and r[1] == 'S Jupiter: S Mars = S Neptune' and r[2] == 'Equator']
check('SKY EQ (moved from Beverley GC, test 11): S Jupiter: S Mars = S Neptune Equator 9.894 / 9.906583, closeness 99.873 %, between, d(X, Y) 19.800583',
      len(sq) == 1 and sq[0][3] == 9.894 and sq[0][4] == 9.906583 and sq[0][5] == 99.873 and sq[0][10] == 'between' and sq[0][11] == 19.800583, str(sq))
# retired 29 Oct 2026 (Eddie, workbook trim): P09 T17 S Jupiter: H Uranus = H Pluto Equator 7.004077 / 6.988434, between, d 13.992511; its
# family cells (7 whole 63.38 · applying 16:10, 7 whole 17.78 · 19:06, 14 whole 39.13 · fixed) and their fills (yellow, light, light); its
# ratio +5 0.997827, toward, unison applying 17:38; natal-only T17 lines "fixed" - T17 is gone

# retired 29 Oct 2026 (Eddie, workbook trim): Test 13, T18 horse | jockey (Lingfield P09 Uranus–Pluto horse Equator 99.7767 % | jockey Dec
# 99.4741 %; block 1 sort; no together side, block 2 present; "both" tag and title count; P11 Vesta–Ketu kept) - T18 is gone



print('\nTest 14 (Eddie 14 Oct): SKY STARS, SKY PAIRS — Lingfield AW 14:35 2021')
# retired 29 Oct 2026 (Eddie, workbook trim): P09 T19 S Jupiter on H Uranus – H Pluto Equator (7 / 7 / 14 whole, unison 1.002238 (78.87),
# between, (also T17), midpoint +0.007822, 99.7767 %, 47.79 + 44.92 = 92.71) and P09 T19 H Pluto – H Sedna Equator (unison 1.011174 (19.47),
# midpoint 98.8949 %, 39.56 + 27.36 = 66.92) - T19 is gone
best = lambda r: max(-1 if r[13] is None else r[13], -1 if len(r) < 22 or r[21] is None else r[21])


def blocks(ws, tid):
    """rows of a pair table by area block (15 Oct 2026): {'A': [...], 'B': [...], 'rest': [...]}"""
    out, cur = {}, None
    for r in ws.iter_rows(values_only=True):
        if r and isinstance(r[0], str) and r[0].startswith(tid + ' ') and ': ' in r[0]:
            cur = 'A' if 'AREA A' in r[0] else 'P' if 'AREA φ' in r[0] else 'B' if 'AREA B' in r[0] else 'rest'
        elif r and r[0] == tid:
            out.setdefault(cur, []).append(r)
    return out


sorted_blocks = lambda bl: all([best(r) for r in v] == sorted([best(r) for r in v], reverse=True) for v in bl.values())
# retired 29 Oct 2026 (Eddie, workbook trim): P09 T19 filter (interval / midpoint / all three families, no together, best first per area
# block, stars royal only); P## PAIRS in the separate _pairs.xlsx workbook; one timing method T17 = T19; H–J across = the same body only
# (T19 and PAIRS); P09 PAIRS unfiltered - T17, T19 and the _pairs workbook are gone
sk = [r for r in wb['SKY JUPITER'].iter_rows(values_only=True) if r[0] == 'SKY STARS' and r[3]]
check('SKY JUPITER: S Jupiter on two sky stars (background, star–star only)', sk and all(r[1] == 'S–S' for r in sk), f'{len(sk)} rows')

# Eddie 14 Oct: SKY PAIRS (S Jupiter on two sky bodies, Moon included, no stars, filtered as T19) and the T17 pair rule
STARN = ('Algol', 'Pleiades', 'Capella', 'Sirius', 'Procyon', 'Alkaid', 'Algorab', 'Spica', 'Arcturus', 'Alphecca', 'Vega', 'Deneb Algedi',
         'Rigel', 'Betelgeuse', 'Polaris', 'Castor', 'Altair', 'Bellatrix', 'Aldebaran', 'Regulus', 'Antares', 'Fomalhaut')
sp = [r for r in wb['SKY JUPITER'].iter_rows(values_only=True) if r[0] == 'SKY PAIRS' and r[3]]
spc = [best(r) for r in sp]
ROYALN = ('Aldebaran', 'Regulus', 'Antares', 'Fomalhaut')
# 15 Oct 2026 (Eddie): body – royal rows added (one sky body, one royal star), in the same blocks and sort
check('SKY PAIRS: S Jupiter on two sky bodies (S–S, no stars) or a sky body and a royal star ("body – royal", 15 Oct), no together, interval hit, '
      'midpoint hit or all three families, best first by the higher of ratio score and midpoint total',
      sp and all((r[1] == 'S–S' and not any(q[2:] in STARN for q in r[2].split(' – '))) or
                 (r[1] == 'body – royal' and r[2].split(' – ')[0][2:] not in STARN + ('Jupiter',) and r[2].split(' – ')[1][2:] in ROYALN) for r in sp)
      and {r[1] for r in sp} == {'S–S', 'body – royal'}
      and not any(r[16] and 'together' in r[16] for r in sp) and sorted_blocks(blocks(wb['SKY JUPITER'], 'SKY PAIRS'))
      and all(r[13] is not None or r[21] is not None or all(c and c.startswith('→') for c in (r[5], r[7], r[9])) for r in sp), f'{len(sp)} rows')
# retired 29 Oct 2026 (Eddie, workbook trim): T17 S Jupiter pairs H–H, J–J or the same body only (Lingfield, Nottingham) - T17 is gone

own = []
# 29 Oct 2026 (Eddie, workbook trim): T17, T18, T19 and the _pairs workbook (PAIRS) are gone - SKY EQ and SKY PAIRS on SKY JUPITER checked
w = openpyxl.load_workbook('reference/classic/race_batch_20210402_lingfield_aw_1435_jupiter_tables.xlsx', read_only=True)
for n in w.sheetnames:
    for r in w[n].iter_rows(values_only=True):
        if r and r[0] in ('SKY EQ', 'SKY PAIRS') and isinstance(r[1], str) and isinstance(r[2], str):
            t = r[2] if r[0] == 'SKY PAIRS' else r[1]
            if 'S Rahu' in t and 'S Ketu' in t:
                own.append((n, r[0], t))
check("the sky's own Rahu–Ketu pair never paired, any plane (SKY EQ, SKY PAIRS; Lingfield)", not own, str(own[:3]))

print('\nTest 15 (Eddie 14 Oct 2026): Haydock 14:20 2021, P03 Golden Flame / Joe Fanning - POSITIONS distances')
wb = openpyxl.load_workbook(build('reference/classic/race_batch_20210904_haydock_1420.xlsx'))
# retired 29 Oct 2026 (Eddie, workbook trim): T11 H Jupiter / J Jupiter | S Uranus RA far (offset 0.032036, 99.963 %, 49.63 + 43.80 = 93.43)
# - T11 is gone; the POSITIONS distances behind it stay checked below
pos = {(x[0], x[2], x[3]): (x[5], x[6]) for x in wb['POSITIONS'].iter_rows(min_row=5, values_only=True) if x and x[3]}
u, hj, jj = pos[('SKY', 'S', 'Uranus')], pos[('P03', 'H', 'Jupiter')], pos[('P03', 'J', 'Jupiter')]
arc = lambda a, b: abs((a - b + 180) % 360 - 180)
check('its distances from POSITIONS: d(S Uranus, H Jupiter) RA 173.135571, d(S Uranus, J Jupiter) RA 173.071498 (the hand values '
      '173.135586 / 173.071514 were 0.000015 high; the offset, ratio and scores are unchanged)',
      round(arc(u[0], hj[0]), 6) == 173.135571 and round(arc(u[0], jj[0]), 6) == 173.071498, f'{arc(u[0], hj[0]):.6f} {arc(u[0], jj[0]):.6f}')
# retired 29 Oct 2026 (Eddie, workbook trim): THREADS Golden Flame Uranus cell lists the T11 midpoint (93.43); Nottingham P03 T11 H Jupiter /
# J Jupiter | S Sun Equator short span (−0.002998, 97.5353 %, 28.39 + 48.02 = 76.41); every T10 / T11 midpoint hit total = ratio half +
# distance half (Haydock, all runners) - THREADS, T10 and T11 are gone

# Eddie 14 Oct 2026: planes RA, Dec, Equator only - no Flat or GC column, row or text in any sheet of either workbook
import glob, re
pat, fg = re.compile(r'\b(Flat|GC)\b'), []
# 29 Oct 2026 (Eddie, workbook trim): no _pairs workbook any more - the tables workbooks built in this run only (stale pairs files left out)
for f in sorted(glob.glob('reference/classic/race_batch_*_jupiter_tables.xlsx')):
    if os.path.getmtime(f) < T0:
        continue                                           # built in this run only
    w = openpyxl.load_workbook(f, read_only=True)
    fg += [(os.path.basename(f), n) for n in w.sheetnames for r in w[n].iter_rows(values_only=True) for v in (r or ())
           if isinstance(v, str) and pat.search(v)]
check('no Flat or GC column, row or label in any sheet of the tables workbooks built in this run',
      not fg and NBUILT[0] >= 6, f'{NBUILT[0]} tables workbooks; {fg[:3]}')

# Eddie 15 Oct 2026: no conjunction family hit in the Equator plane (since 29 Oct 2026 (workbook trim): SKY EQ, SKY PAIRS on SKY JUPITER;
# T17, T18, T19 and the _pairs workbook are gone)
eqc = []
for f in sorted(glob.glob('reference/classic/race_batch_*_jupiter_tables.xlsx')):
    if os.path.getmtime(f) < T0:
        continue
    w = openpyxl.load_workbook(f, read_only=True)
    eqc += [(os.path.basename(f), n, r[2]) for n in w.sheetnames for r in w[n].iter_rows(values_only=True)
            if r and 'Equator' in r and any(isinstance(v, str) and '(conj,' in v for v in r)]
wbl = openpyxl.load_workbook('reference/classic/race_batch_20210402_lingfield_aw_1435_jupiter_tables.xlsx', read_only=True)
# retired 29 Oct 2026 (Eddie, workbook trim): the second half, Lingfield P01 T19 H Sun – H Ceres Equator gone from T19 - T19 is gone
check('no conjunction in the Equator plane (Eddie 15 Oct): no Equator row with a conj family cell in any sheet built in this run',
      not eqc, f'{eqc[:3]}')

# retired 29 Oct 2026 (Eddie, workbook trim): T15 removed (no T15 row, no FIELD 15 sheet, no T15 line in RATIOS / THREADS) - every one of
# those sheets and tables is gone

# retired 29 Oct 2026 (Eddie, workbook trim): Test 16, T20 (Lingfield P03 Sun horse Dec 0.02218, 98.3287 %, 34.68 + 36.31 = 70.99, applying
# 17:21; T20 layout; bodies only; FIELD 20 'H Dec 70.99' 1/6; THREADS T20 line; THREADS midpoints at >= 50 only) - T20, FIELD and THREADS
# are gone

# retired 29 Oct 2026 (Eddie, workbook trim): Test 17, T21 / T22 fixed stars (P03 T21 H Algorab – H Castor Dec unison 98.29 + midpoint
# 98.45; P11 T22 Aldebaran RA unison 51.53, far midpoint 53.41; T22 lists every star; FIELD 21 / 22 cells and k/n; T21 filter and no empty
# FIELD 21 row) - T21, T22 and the FIELD sheets are gone
# retired 29 Oct 2026 (Eddie, workbook trim): THREADS and NUMBERS carry no T21 / T22 - those sheets are gone

# Eddie 15 Oct 2026: SKY cells in T1 / T2 (S Jupiter -> S X, same plane), sky side in T21 (S Jupiter: S A – S B), tag 'same as sky'
print('\nTest 18 (Eddie 15 Oct): SKY cells in T1 / T2 — Lingfield AW 14:35 2021')
t1 = [r for r in wbl['P03'].iter_rows(values_only=True) if r and r[0] == 'T1' and r[1] == 'H Jupiter → H Sun']
check('T1 P03 H Jupiter → H Sun, SKY Dec cell S Jupiter → S Sun (by hand from POSITIONS: Dec S Jupiter −14.389918, S Sun 5.028708; rates '
      '+0.0000426 / +0.0002664 per min): 19.418627, +5 min 19.419746, 12×φ golden +0.00222 (86.76), away, separating, exact 14:25 (9.9 min ago); '
      'natal Dec 2.609925 (24 Oct 2026: φ^2 2.618 now outside the 4-digit match, 2609 vs 2618 - nearest φ^2 shown outside, was a φ^2 hit); '
      'no same as sky tag',
      len(t1) == 1 and t1[0][10] == 2.609925 and t1[0][12] == 'φ^2 phi power — not a 4-digit match (-0.00811)' and t1[0][15] is None and t1[0][27] == 19.418627 and t1[0][28] == 19.419746
      and t1[0][29] == '12×φ (golden)' and t1[0][31] == '+0.00222' and t1[0][32] == 86.76 and t1[0][33] == '<- away'
      and t1[0][35] == 'separating, exact 14:25 (9.9 min ago)' and (len(t1[0]) < 37 or t1[0][36] is None), str(t1[0][26:37] if t1 else t1))
# retired 29 Oct 2026 (Eddie, workbook trim): T21 P11 Alphecca – Betelgeuse RA sky side (92.120174 / 122.994551 φ^10 / 144.885275, 4/3
# 1.335153 (20.61), applying 23:43, far side); its "same as sky" tag 'H: ratio 4/3 · fourth'; T21 listing 244 lines; T21 "same as sky"
# never on d(A, B) - T21 is gone

# Eddie 15 Oct 2026: every ratio table split into area A "whole + 3/2" (1, 3/2, 2, 3 ... 9 on ±0.02 absolute, or a midpoint hit) and
# area B "other" (every other interval); "no hit" at the bottom where every row is listed
print('\nTest 19 (Eddie 15 Oct): areas A / φ and √2 / B — SKY PAIRS / SKY STARS, Lingfield AW 14:35 2021 and Ascot 14:30 2012')
AA, AP = set(families.AREA_A), set(families.AREA_P)    # 15 Oct 2026: third area "φ and √2" between A and B
iv = lambda t: str(t).split(' ')[0]                    # '4 · double octave · Perfect' -> '4'
iarea = lambda lab: 'A' if lab in AA else 'P' if lab in AP else 'B'
tarea = lambda t: 'A' if 'AREA A' in t else 'P' if 'AREA φ' in t else 'B' if 'AREA B' in t else 'rest'


def side_area(tgt, score, mid, rmin=0.0):
    """every area a cell / side has a hit in: the midpoint (A) and the ratio"""
    out = ['A'] if mid is not None else []
    if score is not None and score >= rmin:
        out.append(iarea(iv(tgt)))
    return out


def row_ar(areas):
    """-> (highest area A, φ and √2 (P), B, or 'rest'; hits in more than one area)"""
    areas = [a for x in areas for a in x]
    return next((a for a in ('A', 'P', 'B') if a in areas), 'rest'), len(set(areas)) > 1


def walk(ws, tids, row_area, title=lambda r: None):
    """-> (n data rows, [mismatches]): block from the title rows, area of each data row from its own cells"""
    cur, n, bad, seen, nm, tm = None, 0, [], set(), 0, 0
    for r in ws.iter_rows(values_only=True):
        if not r:
            continue
        t = title(r)
        if t is None:
            for v in r[:2]:
                if isinstance(v, str) and ': ' in v and ('AREA ' in v or 'NO HIT' in v or 'NO RATIO' in v) \
                        and any(v.startswith(x + ' ') for x in tids):    # this table's block titles only
                    t = v
        if t:
            if nm != tm:                               # the previous block title's "n carry hits in more than one area"
                bad.append((ws.title, 'multi count', cur, tm, nm))
            cur = tarea(t)
            m = re.search(r'\((\d+) carry hits in more than one area\)', t)
            nm, tm = 0, int(m.group(1)) if m else 0
            seen.add(cur)
            continue
        if r[0] in tids and isinstance(r[1], str):
            a = row_area(r)
            if a == 'skip':                            # a table's main title row
                continue
            n += 1
            a, mu = a
            nm += mu
            if a != cur:
                bad.append((ws.title, r[0], r[1], r[2] if len(r) > 2 else None, cur, a))
    if nm != tm:
        bad.append((ws.title, 'multi count', cur, tm, nm))
    return n, bad, seen


PL = ('RA', 'Dec', 'Equator')
pairs = lambda r: 'skip' if len(r) < 4 or r[3] not in PL else row_ar([side_area(r[11], r[13], r[21] if len(r) > 21 else None)])
build('reference/classic/race_batch_20120619_ascot_1430.xlsx')     # 19 Oct 2026: Ascot rebuilt here (the checks below had read an older build)
# 29 Oct 2026 (Eddie, workbook trim): T12-T14, T19, T21, T22, P## RATIOS, the _pairs workbook (P## PAIRS) and FIELD 12-14 / 21 / 22 are
# gone - the area blocks are checked on SKY PAIRS and SKY STARS (SKY JUPITER) only
res, tot = [], {}
for f in ('reference/classic/race_batch_20210402_lingfield_aw_1435_jupiter_tables.xlsx', 'reference/classic/race_batch_20120619_ascot_1430_jupiter_tables.xlsx'):
    w = openpyxl.load_workbook(f, read_only=True)
    k, b, _ = walk(w['SKY JUPITER'], ('SKY PAIRS', 'SKY STARS'), pairs)
    tot['SKY'] = tot.get('SKY', 0) + k
    res += b
check('every area A row has an A interval (1, 3/2, 2, 3 ... 9) or a midpoint hit, every φ and √2 row a φ / √2 interval and no A hit, '
      'every area B row a B interval (and no A or φ / √2 hit), every no-hit row neither; block titles count the rows with hits in more than one area; '
      '(15 Oct: third area φ and √2) (SKY PAIRS, SKY STARS; 29 Oct 2026: the runner-sheet ratio tables, P## PAIRS, P## RATIOS and FIELD sheets are gone)',
      not res and all(tot.values()), f'{tot}; mismatches {res[:5]}')

# Eddie 15 Oct 2026: SKY PAIRS body – royal rows
print('\nTest 20 (Eddie 15 Oct): SKY PAIRS body – royal — Lingfield AW 14:35 2021')
wbl = openpyxl.load_workbook('reference/classic/race_batch_20210402_lingfield_aw_1435_jupiter_tables.xlsx', read_only=True)
hf = [r for r in wbl['SKY JUPITER'].iter_rows(values_only=True) if r and r[0] == 'SKY PAIRS' and r[2] == 'S Haumea – S Fomalhaut' and r[3] == 'Dec']
check('SKY PAIRS body – royal S Haumea – S Fomalhaut Dec (by hand from POSITIONS: Dec S Jupiter −14.389918, S Haumea 16.069041, '
      'S Fomalhaut −29.623914): 30.458960 / 15.233996 / 45.692955, ratio 1.999407 on 2 (offset −0.00059, (10 (0.02 − 0.000593) / 0.02)² = 94.16), '
      'between, in the area A block',
      len(hf) == 1 and hf[0][1] == 'body – royal' and abs(hf[0][4] - 30.45896) < 2e-6 and hf[0][6] == 15.233996 and hf[0][8] == 45.692955
      and hf[0][10] == 1.999407 and hf[0][11] == '2 · octave · Perfect' and hf[0][12] == '-0.00059' and hf[0][13] == 94.16 and hf[0][15] == 'between'
      and any(x[2] == 'S Haumea – S Fomalhaut' and x[3] == 'Dec' for x in blocks(wbl['SKY JUPITER'], 'SKY PAIRS').get('A', [])),
      str(hf[0][:16] if hf else hf))

# retired 29 Oct 2026 (Eddie, workbook trim): Test 21, T19 sky side and "same as sky" (P07 J Chiron – J Juno RA sky side 41.375538 'no family'
# / 62.173090 / 103.548628 64×φ (68.66), 3/2 1.502653 (75.23), area A; its tag 'ratio 3/2 · fifth'; never on d(S Jupiter, royal star), P05
# H Gonggong – H Regulus Equator 5/4, 8 together sides plain; every H–J sky side '—', listing 1334 rows) - T19 is gone
RY = ('Aldebaran', 'Regulus', 'Antares', 'Fomalhaut')
# Eddie 16 Oct 2026: timing on the SKY JUPITER "SKY" block (the scaled timing retired the same day: values are never scaled)
print('\nTest 22 (Eddie 16 Oct): SKY block timing — Lingfield AW 14:35 2021')
wbl = openpyxl.load_workbook('reference/classic/race_batch_20210402_lingfield_aw_1435_jupiter_tables.xlsx', read_only=True)
sm = [r for r in wbl['SKY JUPITER'].iter_rows(values_only=True) if r and r[0] == 'SKY' and r[1] == 'S Jupiter → S Mercury']
# 24 Oct 2026 (Eddie, 4-digit rule): the Dec hit φ^5 11.090170 (applying, exact 15:32) is gone - 11.063496 is 1106 vs 1109, outside the
# digit match; both planes now outside, no timing cell (the timing-where-a-hit-is check below keeps the rule)
check('SKY S Jupiter → S Mercury (by hand from POSITIONS: Dec S Jupiter −14.389918, S Mercury −3.326422): Dec 11.063496 → 11.065810 at +5, '
      'φ^5 11.090170 outside the 4-digit match (1106 vs 1109; was applying, exact 15:32 before 24 Oct 2026) -> shown outside, no score, no timing; '
      'RA 32.268875 outside (20×φ 32.360680 outside ±0.0324): no timing. Timing inside each plane block, after its cells: RA cols 3-11, Dec cols 12-20',
      len(sm) == 1 and sm[0][11] == 11.063496 and sm[0][13] == 'φ^5 phi power — not a 4-digit match (-0.02667)' and sm[0][16] is None and sm[0][19] is None
      and sm[0][2] == 32.268875 and sm[0][4].startswith('20×φ (golden) — outside') and sm[0][7] is None and sm[0][10] is None,
      str(sm[0][2:21] if sm else sm))
skr = [list(r) + [None] * 30 for r in wbl['SKY JUPITER'].iter_rows(values_only=True) if r and r[0] == 'SKY' and r[1] != 'S Jupiter → sky points']
check('SKY block: a timing exactly where the best hit (score) is, per plane (RA, Dec); no scaled column',
      skr and all((r[c + 5] is not None) == (r[c + 8] is not None) for r in skr for c in (2, 11))
      and not any(isinstance(x, str) and ('x10:' in x or '/10:' in x) for r in skr for x in r), f'{len(skr)} rows')

# Eddie 16 Oct 2026: no together side shows a score (since 29 Oct 2026 (workbook trim): SKY ROYAL rows only; the T19 sky sides and T21 /
# T21R sides are gone)
tsc = []
for f in ('reference/classic/race_batch_20210402_lingfield_aw_1435_jupiter_tables.xlsx', 'reference/classic/race_batch_20120619_ascot_1430_jupiter_tables.xlsx'):
    w = openpyxl.load_workbook(f, read_only=True)
    for r in w['SKY JUPITER'].iter_rows(values_only=True):
        if not r or r[0] != 'SKY ROYAL':
            continue
        nc, sc_, m0 = 16, 13, 19
        if len(r) > m0 + 2 and r[nc] and 'together' in str(r[nc]) and any(x is not None for x in (r[sc_], r[m0], r[m0 + 1], r[m0 + 2])):
            tsc.append(('SKY JUPITER', r[0], r[1], r[2], r[3] if len(r) > 3 else None))
check('no together side shows a score (Eddie 16 Oct): SKY ROYAL rows; Lingfield and Ascot (T19 / T21 / T21R sides gone 29 Oct 2026)', not tsc, str(tsc[:4]))
# retired 29 Oct 2026 (Eddie, workbook trim): T19 P07 J Orcus – J Gonggong Dec sky side together (1.007991, no score) - T19 is gone

# Eddie 16 Oct 2026: no value scaling anywhere (no scaled column, scaled timing, scaled THREADS / NUMBERS / TIMELINE hits, 'widens' note)
scl = []
for f in ('reference/classic/race_batch_20210402_lingfield_aw_1435_jupiter_tables.xlsx', 'reference/classic/race_batch_20120619_ascot_1430_jupiter_tables.xlsx'):
    w = openpyxl.load_workbook(f, read_only=True)
    for n in w.sheetnames:
        for r in w[n].iter_rows(values_only=True):
            for x in r or ():
                if isinstance(x, str) and ('x10:' in x or '/10:' in x or 'scaled x10 / /10' == x or '[scaled]' in x or ', scaled]' in x
                                           or 'widens 10x' in x or x.startswith('distance (scaled') or x.startswith('ratio (scaled') or x.startswith('timing (scaled')):
                    scl.append((n, x[:60]))
check('no value scaling anywhere in the Lingfield and Ascot workbooks (Eddie 16 Oct): no scaled column or timing, no scaled THREADS / NUMBERS / '
      'TIMELINE hit, no "widens 10x" note', not scl, str(scl[:4]))

# Eddie 16 Oct 2026: royal-star-only tables T21R, T22R, SKY ROYAL, FIELD 21R / 22R
print('\nTest 23 (Eddie 16 Oct): royal-star-only tables — SKY ROYAL, Lingfield AW 14:35 2021 and Ascot 14:30 2012')
RY = ('Aldebaran', 'Regulus', 'Antares', 'Fomalhaut')
# 29 Oct 2026 (Eddie, workbook trim): T21R, T22R (runner sheets) and FIELD 21R / 22R are gone - SKY ROYAL (SKY JUPITER) checked
bad = []
for f in ('reference/classic/race_batch_20210402_lingfield_aw_1435_jupiter_tables.xlsx', 'reference/classic/race_batch_20120619_ascot_1430_jupiter_tables.xlsx'):
    w = openpyxl.load_workbook(f, read_only=True)
    sk = [r for r in w['SKY JUPITER'].iter_rows(values_only=True) if r and len(r) > 3 and r[3] in PL]
    spb = {(r[2], r[3]): r[1:22] for r in sk if r[0] == 'SKY PAIRS' and r[1] == 'body – royal'}
    sr = {(r[2], r[3]): r[1:22] for r in sk if r[0] == 'SKY ROYAL'}
    bad += [('SKY ROYAL', k) for k in spb if sr.get(k) != spb[k]]
    if not all(r[1] in ('royal – royal', 'body – royal') for r in sk if r[0] == 'SKY ROYAL'):
        bad.append('SKY ROYAL blocks')
    if not sr:
        bad.append((f[-40:], 'no SKY ROYAL rows'))
wa = openpyxl.load_workbook('reference/classic/race_batch_20120619_ascot_1430_jupiter_tables.xlsx', read_only=True)
tg = [r for r in wa['SKY JUPITER'].iter_rows(values_only=True) if r and r[0] == 'SKY ROYAL' and len(r) > 16 and r[16] and 'together' in r[16]]
check('SKY ROYAL "X, Y together" rows plain with no score (Eddie 16 Oct): Ascot S Transpluto – S Regulus RA 0.929582 (unison, was 24.22) '
      'and Dec, values kept, no ratio score, no midpoint scores',
      len(tg) == 2 and any(r[3] == 'RA' and r[8] == 0.929582 and r[10] is not None for r in tg)
      and all(r[13] is None and all(x is None for x in r[19:22]) for r in tg), str([r[2:4] + r[13:14] for r in tg]))
# retired 29 Oct 2026 (Eddie, workbook trim), the T21R / T22R / FIELD 21R / 22R half: T21R / T22R lines = the royal lines of T21 / T22 (same
# values), FIELD 21R / 22R present - those tables and sheets are gone
check('SKY ROYAL holds the SKY PAIRS body – royal rows unchanged, blocks "royal – royal" / "body – royal" only (Lingfield, Ascot)',
      not bad, f"{bad[:5]}")
# Eddie 16 Oct 2026: TIMELINE on the SKY JUPITER sheet
print('\nTest 24 (Eddie 16 Oct): TIMELINE — Lingfield AW 14:35 2021 and Ascot 14:30 2012')
import datetime


def tl_check(f, off):
    """-> (events, [problems]): window, order, and each event's exact time = its source cell (a timing cell on the sheet with the same text;
    SKY cells outside at the off have none and are recomputed from POSITIONS, linear off -> +5 min)"""
    w = openpyxl.load_workbook(f, read_only=True)
    allr = [r for r in w['SKY JUPITER'].iter_rows(values_only=True) if r]
    pos = {r[3]: r for r in w['POSITIONS'].iter_rows(values_only=True) if r and r[0] == 'SKY'}
    tl = [r for r in allr if r[0] == 'TIMELINE' and isinstance(r[2], (int, float))]
    ev = [r for r in tl if r[4]]
    bad, last = [], -1e9
    for r in tl:
        if r[2] < last:
            bad.append(('order', r[1:5]))
        last = r[2]
    for r in ev:
        if not -30 <= r[2] <= 60:
            bad.append(('window', r[1:5]))
        txt = (f'applying, exact {r[1]} (in {r[2]:.1f} min)' if r[2] >= 0 else f'separating, exact {r[1]} ({-r[2]:.1f} min ago)')
        names = [x for x in re.split(r'S Jupiter: |S Jupiter → | – ', r[4]) if x]
        src = 'SKY EQ' if r[3] == 'SKY EQ' else r[3]
        cand = [x for x in allr if x[0] == src and all(any(isinstance(c, str) and nm in c for c in x[1:3]) for nm in names)]
        if any(isinstance(c, str) and txt in c for x in cand for c in x):
            continue
        if r[3] == 'SKY' and r[12] and 'outside at the off' in r[12]:
            b = r[4].split('→ S ')[1]
            p0, q0 = pos['Jupiter'], pos[b]
            ix = 5 if r[5] == 'RA' else 6
            dd = lambda m: (abs((p0[5] + m * p0[7]) - (q0[5] + m * q0[7])) % 360) if ix == 5 else abs((p0[6] + m * p0[8]) - (q0[6] + m * q0[8]))
            d = lambda m: min(dd(m), 360 - dd(m)) if ix == 5 else dd(m)
            t = (r[8] - d(0)) / ((d(5) - d(0)) / 5)
            clock = (off + datetime.timedelta(minutes=t)).strftime('%H:%M')
            if abs(t - r[2]) < 0.06 and clock == r[1]:
                continue
        bad.append(('source', r[1:9]))
    return ev, bad


bad, n = [], 0
for f, off in (('reference/classic/race_batch_20210402_lingfield_aw_1435_jupiter_tables.xlsx', datetime.datetime(2021, 4, 2, 14, 35)),
               ('reference/classic/race_batch_20120619_ascot_1430_jupiter_tables.xlsx', datetime.datetime(2012, 6, 19, 14, 30))):
    ev, b = tl_check(f, off)
    n += len(ev); bad += b
check('TIMELINE: every event inside the off − 30 / + 60 min, rows (and the OFF / +5 MIN markers) in time order, each exact time = its source '
      'cell on the sheet (SKY outside at the off: recomputed from POSITIONS)', not bad and n > 0, f'{n} events; {bad[:4]}')
wbl = openpyxl.load_workbook('reference/classic/race_batch_20210402_lingfield_aw_1435_jupiter_tables.xlsx', read_only=True)
tm = [r for r in wbl['SKY JUPITER'].iter_rows(values_only=True) if r and r[0] == 'TIMELINE' and r[4] == 'S Jupiter → S Moon' and r[5] == 'Dec']
check('TIMELINE S Jupiter → S Moon Dec (by hand from POSITIONS: Dec S Jupiter −14.389918 (+0.0000425872 /min), S Moon −24.052038 '
      '(−0.000983196 /min)): d 9.662120 growing 0.00102578 /min, nearest 6×φ 9.708204 outside at the off, reached in 44.9 min -> 15:19, '
      'applying, no score',
      len(tm) == 1 and tm[0][1] == '15:19' and tm[0][2] == 44.9 and tm[0][7] == '6×φ (golden)' and tm[0][9] == 9.66212 and tm[0][10] is None
      and tm[0][11] == 'applying' and 'outside at the off' in tm[0][12], str(tm[0] if tm else tm))
mk = [r for r in wbl['SKY JUPITER'].iter_rows(values_only=True) if r and r[0] == 'TIMELINE' and isinstance(r[3], str) and r[3].startswith('———')]
check('TIMELINE markers: OFF 14:35 (0 min) and OFF + 5 MIN 14:40', [(r[1], r[2]) for r in mk] == [('14:35', 0), ('14:40', 5)], str([r[1:4] for r in mk]))
# Eddie 16 Oct 2026: TIMELINE H / TIMELINE J on every runner sheet
# retired 29 Oct 2026 (Eddie, workbook trim): Test 25, TIMELINE H / J (every runner event recomputed from POSITIONS: exact time, clock,
# window, applying / separating, order, "H + J"; P05 TIMELINE H S Jupiter → H Neptune RA φ^6 separating 14:14 (94.19); P05 TIMELINE J
# S Jupiter → J Venus Dec 33 applying 14:47 (94.54)) - the runner timelines are gone
# Eddie 16 Oct 2026: LONG TIMELINE (score >= 80, any time) and the sky equivalent cells on TIMELINE H / J and LONG TIMELINE H / J
print('\nTest 26 (Eddie 16 Oct): LONG TIMELINE and sky equivalents — Lingfield AW 14:35 2021 (every runner) and Ascot 14:30 2012')
import families as FAM


def long_check(f, off):
    """LONG rows: score >= 80, SEPARATING (t < 0, most recent first) then APPLYING (soonest first); every sky equivalent cell (TIMELINE H / J
    and LONG H / J) recomputed from POSITIONS: value, target (best family / interval or nearest outside), score, timing minutes.
    29 Oct 2026 (Eddie, workbook trim): SKY JUPITER only - the runner TIMELINE H / J and LONG TIMELINE H / J (and their sky equivalents) are gone"""
    w = openpyxl.load_workbook(f, read_only=True)
    posr = list(w['POSITIONS'].iter_rows(values_only=True))
    sky = {r[3]: r for r in posr if r and r[0] == 'SKY'}
    bad, nl, ns = [], 0, 0

    def pos(pt, m):
        r = sky[pt.split(' ', 1)[1]]
        return ((r[5] + m * r[7]) % 360, r[6] + m * r[8])

    def d(a, b, pl, m):
        p, q = pos(a, m), pos(b, m)
        if pl == 'RA':
            x = abs(p[0] - q[0]) % 360
            return min(x, 360 - x)
        return abs(p[1] - q[1]) if pl == 'Dec' else abs(abs(p[1]) - abs(q[1]))

    mins = lambda txt: (float(re.search(r'\(in ([\d.]+) min\)', txt).group(1)) if '(in ' in txt else
                        -float(re.search(r'\(([\d.]+) min ago\)', txt).group(1)))
    for n in [x for x in w.sheetnames if x == 'SKY JUPITER']:
        sec, last = None, None
        for r in w[n].iter_rows(values_only=True):
            if not r or not isinstance(r[0], str):
                continue
            if r[0].startswith('LONG TIMELINE') and isinstance(r[3], str) and r[3].startswith('———'):
                sec, last = r[3].split()[1].rstrip(':'), None
                continue
            ti = 13 if n == 'SKY JUPITER' else 19            # 'minutes from the off' (last column: after the sky cells on runner sheets)
            if r[0].startswith('LONG TIMELINE') and r[4] and len(r) > ti and isinstance(r[ti], (int, float)):
                nl += 1
                t = r[ti]
                key = (-r[10], -t if sec == 'SEPARATING' else t) if r[10] is not None else None   # 16 Oct: score first, ties by time
                # an event within 0.05 min of the off shows -0.0 / 0.0: its section follows its applying / separating (24 Oct 2026: P07 LONG J
                # S Jupiter → J Spica RA 124.500003 half, 2 s before the off)
                near0 = abs(t) < 0.05 and (sec == 'SEPARATING') == (r[11] == 'separating')
                if r[10] is None or r[10] < 80 or ((sec == 'SEPARATING') != (t < 0) or r[11] != ('applying' if t >= 0 else 'separating')) and not near0 \
                        or (last is not None and key < last):
                    bad.append((n, 'long', r[1:6], sec))
                last = key
            # (the sky-equivalent cells of TIMELINE H / J and LONG TIMELINE H / J: gone 29 Oct 2026, workbook trim)
    return nl, ns, bad


bad, nl, ns = [], 0, 0
for f, off in (('reference/classic/race_batch_20210402_lingfield_aw_1435_jupiter_tables.xlsx', None),
               ('reference/classic/race_batch_20120619_ascot_1430_jupiter_tables.xlsx', None)):
    a_, b_, c_ = long_check(f, off)
    nl += a_; ns += b_; bad += c_
check('LONG TIMELINE (sky; the runner H / J blocks and their sky equivalent cells gone 29 Oct 2026): every row scores >= 80, SEPARATING then '
      'APPLYING, each by score at the off (highest first; ties: separating most recent first, applying soonest first, Eddie 16 Oct)',
      not bad and nl > 0, f'{nl} long rows; {bad[:3]}')
# 16 Oct 2026 (Eddie): real exact times from the ephemeris beyond 60 min (engine set up here; used below and in test 33)
import sys as _s
_s.path.insert(0, 'rebuild_kit'); _s.path.insert(0, 'tools/jup2')
import swisseph as swe, celestial_bodies_swisseph as EN, motion as MO
swe.set_ephe_path(MO.EPHE)
JD0 = {'L': EN._to_jd_ut(EN._parse_time('2021-04-02 14:35', 'Europe/London'))}
SITE = MO.off_and_site('lingfield_aw_1435_20210402', 'TRANS', 'reference/classic/race_batch_20210402_lingfield_aw_1435.xlsx', None, None)[1]
b0 = EN.bodies_at(JD0['L'], *SITE, False)
# retired 29 Oct 2026 (Eddie, workbook trim): LONG TIMELINE H P05 S Betelgeuse → H Jupiter RA 110.940467 (86.51 on 10×φ^5) not listed (< 90),
# with its ±1 year Betelgeuse ephemeris scan - LONG TIMELINE H is gone
pr = [r for r in wbl['SKY JUPITER'].iter_rows(values_only=True) if r and r[0] == 'LONG TIMELINE' and r[4] == 'S Jupiter: S Pallas – S Rahu' and r[5] == 'Dec']
skyp = {r[3]: r for r in wbl['POSITIONS'].iter_rows(values_only=True) if r and r[0] == 'SKY'}


def dec_at(b, m):
    eb = b + '_B' if skyp[b][4] == 'star' else b
    return skyp[b][6] + EN.bodies_at(JD0['L'] + m / 1440.0, *SITE, False)[eb]['dec'] - b0[eb]['dec']


qv = lambda m: abs(dec_at('Jupiter', m) - dec_at('Rahu', m)) / abs(dec_at('Jupiter', m) - dec_at('Pallas', m))   # larger / smaller at the off
check('LONG TIMELINE (sky) S Jupiter: S Pallas – S Rahu Dec on 2 (93.55): real exact 11:32 (−182.8 min) vs linear 11:31 (−183.4 min); direct '
      'ephemeris lookup: d(S Jup, S Rahu) / d(S Jup, S Pallas) (larger / smaller) crosses 2 between one minute before and one minute after the found time',
      len(pr) == 1 and pr[0][1] == '11:32' and pr[0][13] == -182.8 and pr[0][14] == '11:31 (-183.4 min)'
      and (qv(pr[0][13] - 1) - 2) * (qv(pr[0][13] + 1) - 2) < 0, f"{pr[0][:15] if pr else pr}; {qv(-183.8):.7f} {qv(-181.8):.7f}")
# retired 29 Oct 2026 (Eddie, workbook trim): sky equivalent "—" on P05 TIMELINE H / LONG TIMELINE H T12 / T13 / T22 ratio rows - the runner
# timelines are gone
# retired 29 Oct 2026 (Eddie, workbook trim): Test 27, T23 (rows = every body, star and equator in T1 and T2, S / H and S / J cells = T12;
# P03 Haumea RA H / J 3.000249 on 3 (97.53); FIELD 23 = the T23 H / J cells; H / J unisons automatic / assisted by the rule (Lingfield 7, Ascot
# 14 / 11 assisted); the automatic branch on a constructed case (JT.t23_cells); T23 distance cells = the T1 / T2 natal and SKY cells) - T23,
# FIELD 23 and t23_cells are gone. The Ascot input and site below stay: test 33 uses _site.
sys.path.insert(0, 'tools')
import jupiter_tables as JT
_r, _b, _l, _sky, _run = JT.load_input('reference/classic/race_batch_20120619_ascot_1430.xlsx')
_off, _site, _ = MO.off_and_site(_r, _l, _b, None, None)
# (T23 distance cells check retired 29 Oct 2026: see above)
# Test 28 (Eddie, 19 Oct 2026): Galactic Centre as the 23rd star (engine + card pipeline) and T24 stars against bodies
import shutil, tempfile, math
print('test 28: Galactic Centre (card pipeline rebuild of Ascot 14:30 2012) and T24')
AP = os.path.expanduser('~/astronomy-project')
tmp = tempfile.mkdtemp(prefix='jt_gc_')
shutil.copy(f'{AP}/racecards/20120619_ascot/race_batch_20120619_ascot_1430.csv', tmp)
r = subprocess.run([sys.executable, 'run_card_swiss.py', '--folder', tmp, '--skip-scrape', '--skip-vibrations', '--skip-pipeline',
                    '--skip-detail', '--skip-predict', '--tables'], cwd=AP, capture_output=True, text=True)
gcf = f'{tmp}/race_batch_20120619_ascot_1430_jupiter_tables.xlsx'
check('T28 card pipeline rebuild of Ascot 14:30 2012 writes the tables workbook', r.returncode == 0 and os.path.exists(gcf), r.stdout[-300:] + r.stderr[-300:])
gb = openpyxl.load_workbook(f'{tmp}/race_batch_20120619_ascot_1430.xlsx', read_only=True)
ob = openpyxl.load_workbook('reference/classic/race_batch_20120619_ascot_1430.xlsx', read_only=True)
bad = []
for sh in ('TRANS', 'P01', 'P02'):
    g = {x[1]: x for x in gb[sh].iter_rows(min_row=2, values_only=True) if x and x[1]}
    o = {x[1]: x for x in ob[sh].iter_rows(min_row=2, values_only=True) if x and x[1]}
    if set(g) - set(o) != {x + '_B' for x in ['Galactic Centre'] + JT.NEW_STARS} or any(g.get(k) != o[k] for k in o):   # + 13 stars (19 Oct 2026)
        bad.append(sh)
gsky = next(x for x in gb['TRANS'].iter_rows(min_row=2, values_only=True) if x and x[1] == 'Galactic Centre_B')
check('T28 batch: Galactic Centre_B (and, 19 Oct 2026, the 13 new stars) on TRANS, horse and jockey sheets, every other row identical to the old batch; sky RA / Dec = Sgr A* '
      '(apparent J2000; ~266.417 / -29.008 to 0.01)', not bad and abs(gsky[4] - 266.417) < 0.01 and abs(gsky[8] + 29.008) < 0.01,
      f'{bad}; sky {gsky[4]:.6f} / {gsky[8]:.6f}')
gw = openpyxl.load_workbook(gcf)


def t24_blocks(ws, mids=None):
    """{block title start: (header bodies [(body, plane)], {star label: [cells]})}. Layout (Eddie, 19 Oct 2026, revised): header row 1
    'RA' merged over the RA body columns, 'RA midpoints', 'Dec' merged over the Dec body columns, 'Dec midpoints'; row 2 body names.
    mids: filled with {(block, star label, plane): midpoint cell}"""
    out, cur, hdr, rows = {}, None, None, None
    it = list(ws.iter_rows())
    for i, x in enumerate(it):
        if x[0].value == 'T24' and str(x[1].value).startswith('T24 '):
            cur = x[1].value.split(' (')[0]; rows = {}
            h1, h2 = it[i + 1], it[i + 2]
            hdr, cols, mcol, pl = [], [], {}, None
            for c in range(2, len(h1)):
                if h1[c].value in ('RA', 'Dec'):
                    pl = h1[c].value
                if h1[c].value == 'repeats':                 # T24 repeats column (Eddie, 22 Oct 2026; test 41)
                    break
                if str(h1[c].value or '').endswith(' midpoints'):
                    mcol[h1[c].value.split()[0]] = c
                elif pl and h2[c].value and h2[c].value != 'S Jupiter':   # S Jupiter column (27 Oct 2026; test 46): not a body
                    hdr.append((h2[c].value, pl)); cols.append(c)
            out[cur] = (hdr, rows)
        elif cur and x[0].value == 'T24' and ('→' in str(x[1].value) or x[1].value == 'stars'):
            continue                                         # the Jupiter rows and the 'stars' separator of T24 HORSE / JOCKEY (29 Oct 2026)
        elif cur and x[0].value == 'T24':
            rows[x[1].value] = [x[c] for c in cols]
            if mids is not None:
                for p, c in mcol.items():
                    mids[(cur, x[1].value, p)] = x[c]
        elif cur and x[0].value != 'table' and x[0].value is None and rows:
            cur = None
    return out


gc_in = {}
for sh in ('SKY JUPITER', 'P01'):
    for x in gw[sh].iter_rows():
        if x[1].value and 'Galactic Centre' in str(x[1].value) and x[0].value:
            gc_in.setdefault(x[0].value, 0); gc_in[x[0].value] += 1
# 29 Oct 2026 (Eddie, workbook trim): T19, T21 and T22 are gone - the Galactic Centre checked in SKY, T1, T2, T24
check('T28 Galactic Centre in T1, T2, T24 (and SKY) of the rebuilt Ascot (T21 / T22 / T19 gone 29 Oct 2026)',
      all(gc_in.get(t) for t in ('SKY', 'T1', 'T2', 'T24')), str(gc_in))
# (b) five cells by hand from POSITIONS
pos = {}
for x in gw['POSITIONS'].iter_rows(min_row=5, values_only=True):
    if x and x[3]:
        pos[(x[0] if x[0] != 'SKY' else 'SKY', x[2], x[3])] = (x[5], x[6])
PHI = (1 + 5 ** 0.5) / 2
sq = lambda tol, dev: (10 * (tol - dev) / tol) ** 2
cases = [  # (sheet, block, chart, star, body, plane, expected value, score formula, fill, comment)
    ('SKY JUPITER', 'T24 SKY', ('SKY', 'S'), 'Galactic Centre', 'Ceres', 'RA', 150.475762, lambda v: sq(0.03236, abs(v - 93 * PHI)), GREEN,
     '93×φ · golden · {o} · {s}', lambda v: v - 93 * PHI),
    ('P01', 'T24 HORSE', ('P01', 'H'), 'Galactic Centre', 'Mars', 'RA', 177.745864, lambda v: 100 * (1.277778 - abs(v - 1600 / 9)) / 1.277778,
     GREEN, '177.778 · ninth · {o} · {s}', lambda v: v - 1600 / 9),
    ('P01', 'T24 JOCKEY', ('P01', 'J'), 'Galactic Centre', 'Sedna', 'Dec', 32.362221, lambda v: sq(0.03236, abs(v - 20 * PHI)), GREEN,
     '20×φ · golden · {o} · {s}', lambda v: v - 20 * PHI),
    ('P01', 'T24 HORSE', ('P01', 'H'), 'Algorab', 'Makemake', 'RA', 0.064638, lambda v: sq(1.5, v), '00FF8080',   # conjunction red (20 Oct 2026)
     '0 · conjunction (RA < 1.5) · {o} · {s}', lambda v: v),
    ('P01', 'T24 HORSE', ('P01', 'H'), 'Galactic Centre', 'Sun', 'RA', 57.972709, None, None, 'nearest: 58 · whole — outside ±0.02 · {o}',
     lambda v: v - 58),
]
for sh, blk, (who, C), star, body, pl, exp, sc, fill, txt, off in cases:
    a, b = pos[(who, C, star)], pos[(who, C, body)]
    d = abs(a[0] - b[0]) % 360.0
    v = (min(d, 360.0 - d)) if pl == 'RA' else abs(a[1] - b[1])
    hdr, rows = t24_blocks(gw[sh])[blk]
    c = rows[f'{C[-1]} {star}'][hdr.index((body, pl))]
    want = txt.format(o=f'{off(v):+.5f}', s=f'{sc(v):.2f}' if sc else '')
    got_fill = c.fill.start_color.rgb if c.fill and c.fill.fill_type else None
    check(f'T28 {blk} {C[-1]} {star} – {C[-1]} {body} {pl} by hand from POSITIONS', abs(c.value - round(v, 6)) < 1e-9 and abs(v - exp) < 1e-6
          and got_fill == fill and c.comment and c.comment.text.split('\n')[0] == want, f'{c.value} / {v:.6f}, {got_fill}, {c.comment and c.comment.text!r} vs {want!r}')
# (c) grid shapes: rebuilt Ascot 23 stars; the old Lingfield and Ascot batches 22 (Galactic Centre skipped quietly)
# 19 Oct 2026 (Eddie): T24 on T1_T2_T24_STARS (no Polaris, + 13 stars): rebuilt Ascot 35 rows, the old batches 21
bad, n = [], 0
for f, nstar in ((gcf, 67), ('reference/classic/race_batch_20210402_lingfield_aw_1435_jupiter_tables.xlsx', 21),
                 ('reference/classic/race_batch_20120619_ascot_1430_jupiter_tables.xlsx', 21)):
    w = gw if f == gcf else openpyxl.load_workbook(f)
    for sh in [s for s in w.sheetnames if s == 'SKY JUPITER' or re.fullmatch(r'P\d\d', s)]:
        bl = t24_blocks(w[sh])
        want = ['T24 SKY'] if sh == 'SKY JUPITER' else ['T24 HORSE', 'T24 JOCKEY']
        if sorted(bl) != sorted(want):
            bad.append((f, sh, list(bl))); continue
        for k, (hdr, rows) in bl.items():
            n += 1
            nb = 25 if k == 'T24 SKY' else 24
            moon = any(b == 'Moon' for b, _ in hdr)
            cells = [c for cs in rows.values() for c in cs]
            if (len(rows) != nstar or len(hdr) != 2 * nb or moon != (k == 'T24 SKY')
                    or any(c.value is None or c.comment is None for c in cells) or ('Galactic Centre' in str(list(rows))) != (nstar == 67)):
                bad.append((f[-40:], sh, k, len(rows), len(hdr), moon))
check('T28 every T24 grid = stars x (2 x bodies): 67 x 50 sky (Moon incl.), 67 x 48 natal on the rebuilt Ascot (35 until 20 Oct 2026, 46 until 21 Oct); 21 rows on the old batches; '
      'every cell a value and a comment', not bad and n == 3 + 13 + 3, f'{n} grids (rebuilt Ascot 3, Lingfield 13, old Ascot 3); {bad[:3]}')
# Test 29 (Eddie, 19 Oct 2026): T24 columns RA then Dec, and T24 midpoints (the star the middle point of two bodies of the block)
print('test 29: T24 RA-then-Dec columns and midpoints')
bad, n = [], 0
for f in (gcf, 'reference/classic/race_batch_20210402_lingfield_aw_1435_jupiter_tables.xlsx'):
    w = gw if f == gcf else openpyxl.load_workbook(f)
    for sh in [s for s in w.sheetnames if s == 'SKY JUPITER' or re.fullmatch(r'P\d\d', s)]:
        for r0 in w[sh].iter_rows():
            if r0[0].value == 'T24' and str(r0[1].value).startswith('T24 '):
                n += 1
                h1 = [c.value for c in w[sh][r0[0].row + 1]]; h2 = [c.value for c in w[sh][r0[0].row + 2]]
                nb = 25 if r0[1].value.startswith('T24 SKY') else 24     # (the HORSE / JOCKEY note mentions the SKY block, 27 Oct 2026)
                body = [b for b in JT.T24_BODIES if nb == 25 or b != 'Moon'] + ([] if nb == 25 else ['S Jupiter'])   # S Jupiter column (27 Oct 2026)
                nb = len(body)
                want1 = ['RA'] + [None] * (nb - 1) + ['RA midpoints', 'Dec'] + [None] * (nb - 1) + ['Dec midpoints']
                if h1[2:2 + 2 * nb + 2] != want1 or h2[2:2 + nb] != body or h2[3 + nb:3 + 2 * nb] != body:
                    bad.append((f[-40:], sh, r0[1].value[:12]))
check('T29 every T24 block: header row 1 RA (merged) .. RA midpoints, Dec (merged) .. Dec midpoints; row 2 the bodies in order (HORSE / JOCKEY: then '
      'S Jupiter, 27 Oct 2026), RA half then Dec half',
      not bad and n == 3 + 13, f'{n} blocks; {bad[:3]}')
gp = {}
for x in gw['POSITIONS'].iter_rows(min_row=5, values_only=True):
    if x and x[3]:
        gp[(x[0], x[2], x[3])] = (x[5], x[6])


def by_hand(who, C, star, a, b, pl):
    """midpoint from POSITIONS, own code: RA near / far midpoint (closer), arc distances; Dec signed"""
    s, pa, pb = (gp[(who, C, z)] for z in (star, a, b))
    if pl == 'RA':
        arc = lambda u, v: abs((u - v + 180) % 360 - 180)
        mid = (pa[0] + ((pb[0] - pa[0] + 180) % 360 - 180) / 2) % 360
        near, far = (s[0] - mid + 180) % 360 - 180, (s[0] - mid) % 360 - 180
        side, off = ('near', near) if abs(near) <= abs(far) else ('far', far)
        d1, d2, dab, tol = arc(s[0], pa[0]), arc(s[0], pb[0]), arc(pa[0], pb[0]), 0.5
    else:
        side, off = '', s[1] - (pa[1] + pb[1]) / 2
        d1, d2, dab, tol = abs(s[1] - pa[1]), abs(s[1] - pb[1]), abs(pa[1] - pb[1]), 0.15
    close = 100 * min(d1, d2) / max(d1, d2)
    rh, dh = 0.5 * (10 * (10 - (100 - close)) / 10) ** 2, 0.5 * (10 * (tol - abs(off)) / tol) ** 2
    pos = 'between' if abs(dab - (d1 + d2)) <= 0.01 * (d1 + d2) else 'far side'
    return d1, d2, off, side, close, rh, dh, rh + dh, pos


gm = {}
for sh in ('SKY JUPITER', 'P01'):
    t24_blocks(gw[sh], gm)
cases = [  # (block, who, chart, star, A, B, plane, d1, d2, offset, side, closeness, ratio half, distance half, total, between / far side)
    ('T24 SKY', 'SKY', 'S', 'Galactic Centre', 'Mercury', 'Jupiter', 'RA', 153.457407, 153.219387, 0.119010, 'far', 99.8449, 48.46, 29.03, 77.49, 'far side'),
    ('T24 HORSE', 'P01', 'H', 'Arcturus', 'Jupiter', 'Orcus', 'RA', 69.091938, 69.107610, 0.007836, 'near', 99.9773, 49.77, 48.45, 98.22, 'between'),
    ('T24 JOCKEY', 'P01', 'J', 'Regulus', 'Chiron', 'Orcus', 'Dec', 5.534608, 5.506910, -0.013849, '', 99.4996, 45.12, 41.19, 86.31, 'between'),
]
for blk, who, C, star, a, b, pl, e1, e2, eo, es, ec, er, ed, et, ep in cases:
    d1, d2, off, side, close, rh, dh, tot, pos = by_hand(who, C, star, a, b, pl)
    txt = f"{a} = {b} {tot:.2f} ({side + ', ' if side else ''}{pos})"
    mc = gm[(blk, f'{C} {star}', pl)]
    hdr, rows = t24_blocks(gw['SKY JUPITER' if who == 'SKY' else 'P01'])[blk]
    ca, cb = (rows[f'{C} {star}'][hdr.index((z, pl))] for z in (a, b))
    ok = (abs(d1 - e1) < 1e-6 and abs(d2 - e2) < 1e-6 and abs(off - eo) < 1e-6 and side == es and abs(close - ec) < 1e-4
          and f'{rh:.2f}' == f'{er:.2f}' and f'{dh:.2f}' == f'{ed:.2f}' and f'{tot:.2f}' == f'{et:.2f}' and pos == ep
          and txt in str(mc.value) and ca.border.left.style == 'thick' == cb.border.right.style
          and f'midpoint with {b}: total {tot:.2f}' in ca.comment.text and f'midpoint with {a}: total {tot:.2f}' in cb.comment.text)
    check(f'T29 {blk} {C} {star} {pl} {a} = {b} by hand from POSITIONS', ok,
          f'{d1:.6f} / {d2:.6f}, offset {off:+.6f} {side}, {close:.4f} %, {rh:.2f} + {dh:.2f} = {tot:.2f}, {pos}; cell {mc.value!r}')
# not listed: "A, B together" (Ascot P01 H Capella RA Juno – Quaoar, d 1.3896 < 1.5, would score 64.50) and the own Rahu – Ketu pair
# (Lingfield P03 J Alkaid RA, d 180, would score 46.33); no 'Rahu = Ketu' anywhere
d1, d2, off, side, close, rh, dh, tot, pos = by_hand('P01', 'H', 'Capella', 'Juno', 'Quaoar', 'RA')
dab = abs((gp[('P01', 'H', 'Juno')][0] - gp[('P01', 'H', 'Quaoar')][0] + 180) % 360 - 180)
cap = gm[('T24 HORSE', 'H Capella', 'RA')].value
lw = openpyxl.load_workbook('reference/classic/race_batch_20210402_lingfield_aw_1435_jupiter_tables.xlsx')
lm = {}
t24_blocks(lw['P03'], lm)
alk = lm[('T24 JOCKEY', 'J Alkaid', 'RA')].value
rk = sum('Rahu = Ketu' in str(c.value) for w in (gw, lw) for sh in w.sheetnames for r in w[sh].iter_rows() if r[0].value == 'T24' for c in r)
check('T29 not listed: together pair Ascot P01 H Capella RA Juno – Quaoar (by hand d(A, B) < 1.5, a midpoint hit otherwise) and own Rahu – Ketu '
      '(Lingfield P03 J Alkaid RA); no "Rahu = Ketu" in any T24 cell',
      dab < 1.5 and tot >= 50 and 'Juno = Quaoar' not in str(cap) and 'Rahu = Ketu' not in str(alk) and rk == 0,
      f'd(Juno, Quaoar) {dab:.4f}, total {tot:.2f}; Capella cell {cap!r}; Alkaid cell {alk!r}; Rahu = Ketu cells {rk}')
# Test 30 (Eddie, 19 Oct 2026): 13 more stars in the engine; T1, T2 and T24 on T1_T2_T24_STARS (35 stars, |Dec| <= 50: no Polaris),
# every other table on STARS as before; the default workbook = POSITIONS, SKY JUPITER and the runner sheets (--full: every sheet as before)
print('test 30: T1 / T2 / T24 star list and the trimmed default workbook')
NEWRE = re.compile(r'\b(' + '|'.join(x if x != 'Deneb' else r'Deneb(?! Algedi)' for x in JT.NEW_STARS) + r')\b')
runner_tabs = lambda w: [s for s in w.sheetnames if re.fullmatch(r'P\d\d', s)]
lf = 'reference/classic/race_batch_20210402_lingfield_aw_1435_jupiter_tables.xlsx'      # built with --full above
ld = f'{tmp}/lingfield_default_jupiter_tables.xlsx'
r = subprocess.run([sys.executable, 'tools/jupiter_tables.py', 'reference/classic/race_batch_20210402_lingfield_aw_1435.xlsx', '--out', ld],
                   capture_output=True, text=True)
ldw, lfw = openpyxl.load_workbook(ld, read_only=True), openpyxl.load_workbook(lf, read_only=True)
vals = lambda ws: [tuple(x) for x in ws.iter_rows(values_only=True)]
same = [sh for sh in ldw.sheetnames if sh in lfw.sheetnames and vals(ldw[sh]) == vals(lfw[sh])]
# 29 Oct 2026 (Eddie, workbook trim): no SKY STARS SUMMARY, and --full = the default (no --full-only sheets)
# 31 Oct 2026 (Eddie): + the SUMMARY sheet after the runner sheets
check('T30 default workbook = POSITIONS, SKY JUPITER, the P## sheets and SUMMARY only (rebuilt Ascot via run_card_swiss.py --tables, Lingfield); '
      'each equal to the --full workbook\'s sheet, and --full adds no sheet (29 Oct 2026; SUMMARY 31 Oct 2026)',
      gw.sheetnames == ['POSITIONS', 'SKY JUPITER', 'P01', 'SUMMARY']
      and ldw.sheetnames == ['POSITIONS', 'SKY JUPITER'] + runner_tabs(lfw) + ['SUMMARY']
      and same == ldw.sheetnames and lfw.sheetnames == ldw.sheetnames,
      f'Ascot {gw.sheetnames}; Lingfield {ldw.sheetnames} ({len(same)} equal to --full, which has {len(lfw.sheetnames)} sheets)')
# T1, T2 and T24 on the rebuilt Ascot: the 35 stars in order, the 13 new ones in, no Polaris
gs = {}
for sh in ('SKY JUPITER', 'P01'):
    for x in gw[sh].iter_rows(values_only=True):
        if x[0] in ('T1', 'T2') and x[1] and ' → ' in str(x[1]) and x[1].split(' → ')[1].split(' ', 1)[1] in JT.T1_T2_T24_STARS + ['Polaris']:
            gs.setdefault(x[0], []).append(x[1].split(' → ')[1].split(' ', 1)[1])
for k, (hdr, rows) in {**t24_blocks(gw['SKY JUPITER']), **t24_blocks(gw['P01'])}.items():
    gs[k] = [lab.split(' ', 1)[1] for lab in rows]
check('T30 rebuilt Ascot T1, T2, T24 SKY / HORSE / JOCKEY: the 67 stars of T1_T2_T24_STARS in order (13 + 11 + 21 new ones in, no Polaris)',
      len(JT.T1_T2_T24_STARS) == 67 and all(gs.get(k) == JT.T1_T2_T24_STARS for k in ('T1', 'T2', 'T24 SKY', 'T24 HORSE', 'T24 JOCKEY')),
      str({k: len(v) for k, v in gs.items()}))
# no other table carries the 13 (rebuilt Ascot, --full); 29 Oct 2026 (workbook trim): no pairs workbook, no T23 (that half retired)
gf = f'{tmp}/full_jupiter_tables.xlsx'
r = subprocess.run([sys.executable, 'tools/jupiter_tables.py', f'{tmp}/race_batch_20120619_ascot_1430.xlsx', '--full', '--out', gf],
                   capture_output=True, text=True)
gfw = openpyxl.load_workbook(gf, read_only=True)
# 20 Oct 2026 (Eddie, 11 more stars): the comparisons of tests 31 - 33 with their builds before run on a 35-star copy of the rebuilt batch
# (the 11 '_B' rows taken out), so they keep checking their own change; test 34 checks 35 -> 46
NEW11 = ['Alpheratz', 'Algenib', 'Difda', 'Mirach', 'Menkar', 'Phact', 'Adhara', 'Wezen', 'Menkent', 'Unukalhai', 'Rasalhague']
# 21 Oct 2026 (Eddie, 20 more stars and the Orion Nebula Cluster): likewise a 46-star copy (the 21 '_B' rows taken out) for tests 34 and 35;
# test 36 checks 46 -> 67
NEW21 = ['Ankaa', 'Al Rescha', 'Mirfak', 'El Nath', 'Meissa', 'Alnilam', 'Murzim', 'Alhena', 'Acubens', 'Alphard', 'Alkes', 'Zosma',
         'Vindemiatrix', 'Diadem', 'Ras Algethi', 'Aculeus', 'Acumen', 'Facies', 'Rukbat', 'Sualocin', 'Orion Nebula Cluster']


def batch_copy(out, drop):
    _w = openpyxl.load_workbook(f'{tmp}/race_batch_20120619_ascot_1430.xlsx')
    for _ws in _w.worksheets:
        for _i in [x[0].row for x in _w[_ws.title].iter_rows() if len(x) > 1 and x[1].value in {n + '_B' for n in drop}][::-1]:
            _ws.delete_rows(_i)
    _w.save(out)


os.makedirs(f'{tmp}/s35', exist_ok=True)
os.makedirs(f'{tmp}/s46', exist_ok=True)
b35, b46 = f'{tmp}/s35/race_batch_20120619_ascot_1430.xlsx', f'{tmp}/s46/race_batch_20120619_ascot_1430.xlsx'
batch_copy(b35, NEW11 + NEW21)
batch_copy(b46, NEW21)
gf35, gf46 = f'{tmp}/s35/full_jupiter_tables.xlsx', f'{tmp}/s46/full_jupiter_tables.xlsx'
for _b, _o in ((b35, gf35), (b46, gf46)):
    subprocess.run([sys.executable, 'tools/jupiter_tables.py', _b, '--full', '--out', _o], capture_output=True, text=True)
gw46 = openpyxl.load_workbook(gf46)
#@ full: old_code builds of 93e73bc4, ad0defaf, 18aff98a, 632447ff, 00b6db9e, 5e10f7bf: inputs of the old-against-old comparisons of tests 31-41
# 20 Oct 2026 (Eddie, T24 conjunction red): the comparisons of tests 31 - 34 with their builds before run on builds by the code of commit
# 93e73bc4 (before the red fill), so they keep checking their own change; test 35 checks 93e73bc4 -> now (red fill only)
bf = 'tools/_jt_before_93e73bc4.py'
open(bf, 'w').write(old_code('93e73bc4'))
os.makedirs(f'{tmp}/h93', exist_ok=True)
lf_h, gf_h, gf35_h = (f'{tmp}/h93/{n}_jupiter_tables.xlsx' for n in ('lingfield', 'ascot', 'ascot35'))
for inp, out in (('reference/classic/race_batch_20210402_lingfield_aw_1435.xlsx', lf_h), (b46, gf_h),          # 46-star copy (21 Oct 2026)
                 (b35, gf35_h)):
    run_build([bf, inp, '--full', '--out', out], capture_output=True, text=True)
os.remove(bf)
# 21 Oct 2026 (Eddie, 20 more stars): builds by the code of commit ad0defaf (46 stars, red fill) - test 35 compares 93e73bc4 with these,
# test 36 compares these with now (lf_r, gf_r on the full 66-star batch)
bf = 'tools/_jt_before_ad0defaf.py'
open(bf, 'w').write(old_code('ad0defaf'))
os.makedirs(f'{tmp}/had', exist_ok=True)
lf_r, gf_r, gf46_r, gf35_r = (f'{tmp}/had/{n}_jupiter_tables.xlsx' for n in ('lingfield', 'ascot', 'ascot46', 'ascot35'))
for inp, out in (('reference/classic/race_batch_20210402_lingfield_aw_1435.xlsx', lf_r), (f'{tmp}/race_batch_20120619_ascot_1430.xlsx', gf_r),
                 (b46, gf46_r), (b35, gf35_r)):
    run_build([bf, inp, '--full', '--out', out], capture_output=True, text=True)
os.remove(bf)
# 22 Oct 2026 (Eddie, every star in every timeline, LONG >= 90): builds by the code of commit 18aff98a (67 stars, before that change) - tests 34
# and 36 count and compare on these, test 37 compares them with now
bf = 'tools/_jt_before_18aff98a.py'
open(bf, 'w').write(old_code('18aff98a'))
os.makedirs(f'{tmp}/h18', exist_ok=True)
lf_s, gf_s, gf46_s = (f'{tmp}/h18/{n}_jupiter_tables.xlsx' for n in ('lingfield', 'ascot', 'ascot46'))
for inp, out in (('reference/classic/race_batch_20210402_lingfield_aw_1435.xlsx', lf_s), (f'{tmp}/race_batch_20120619_ascot_1430.xlsx', gf_s),
                 (b46, gf46_s)):
    run_build([bf, inp, '--full', '--out', out], capture_output=True, text=True)
os.remove(bf)
gw46_s, gw_s = openpyxl.load_workbook(gf46_s), openpyxl.load_workbook(gf_s)
# 22 Oct 2026 (Eddie, LONG STAR TIMELINE): builds by the code of commit 632447ff (star timelines, before LONG STAR TIMELINE) - test 37 compares
# 18aff98a with these, test 38 compares these with now
bf = 'tools/_jt_before_632447ff.py'
open(bf, 'w').write(old_code('632447ff'))
os.makedirs(f'{tmp}/h63', exist_ok=True)
lf_t, gf_t = (f'{tmp}/h63/{n}_jupiter_tables.xlsx' for n in ('lingfield', 'ascot'))
for inp, out in (('reference/classic/race_batch_20210402_lingfield_aw_1435.xlsx', lf_t), (f'{tmp}/race_batch_20120619_ascot_1430.xlsx', gf_t)):
    run_build([bf, inp, '--full', '--out', out], capture_output=True, text=True)
os.remove(bf)
# 22 Oct 2026 (Eddie, sky star timelines trimmed, T24 STAR SELECT): builds by the code of commit 00b6db9e (LONG STAR TIMELINE) - test 38
# compares 632447ff with these, test 39 compares these with now
bf = 'tools/_jt_before_00b6db9e.py'
open(bf, 'w').write(old_code('00b6db9e'))
os.makedirs(f'{tmp}/h00', exist_ok=True)
lf_u, gf_u = (f'{tmp}/h00/{n}_jupiter_tables.xlsx' for n in ('lingfield', 'ascot'))
lf_ud, gf_ud = (f'{tmp}/h00/{n}_def_jupiter_tables.xlsx' for n in ('lingfield', 'ascot'))
for inp, out, fl in (('reference/classic/race_batch_20210402_lingfield_aw_1435.xlsx', lf_u, ['--full']), (f'{tmp}/race_batch_20120619_ascot_1430.xlsx', gf_u, ['--full']),
                     ('reference/classic/race_batch_20210402_lingfield_aw_1435.xlsx', lf_ud, []), (f'{tmp}/race_batch_20120619_ascot_1430.xlsx', gf_ud, [])):
    run_build([bf, inp] + fl + ['--out', out], capture_output=True, text=True)
os.remove(bf)
# 29 Oct 2026 (Eddie, workbook trim): the builds by 0ad7d2e5's code (with T24 STAR SELECT) fed only test 40's comparison, retired - removed
# 22 Oct 2026 (Eddie, T24 repeats flag): builds by the code of commit 5e10f7bf (T24 STAR SELECT removed, before the repeats), default and
# --full - test 41 compares these with now
bf = 'tools/_jt_before_5e10f7bf.py'
open(bf, 'w').write(old_code('5e10f7bf'))
os.makedirs(f'{tmp}/h5e', exist_ok=True)
H5E = {(n, f): f'{tmp}/h5e/{n}_{f}_jupiter_tables.xlsx' for n in ('lingfield', 'ascot') for f in ('full', 'def')}
for (n, f), out in H5E.items():
    inp = 'reference/classic/race_batch_20210402_lingfield_aw_1435.xlsx' if n == 'lingfield' else f'{tmp}/race_batch_20120619_ascot_1430.xlsx'
    run_build([bf, inp] + (['--full'] if f == 'full' else []) + ['--out', out], capture_output=True, text=True)
os.remove(bf)
#@ quick: T30 new-star leak check on the current rebuilt Ascot build
leak = []
for w in (gfw,):                                              # 29 Oct 2026 (workbook trim): no pairs workbook
    for sh in w.sheetnames:
        if sh == 'POSITIONS':
            continue
        for x in w[sh].iter_rows(values_only=True):
            if x and (x[0] in ('T1', 'T2', 'T24', 'MATCH HORSE', 'MATCH JOCKEY', 'MATCH HORSE | JOCKEY')   # MATCH tables name SKY block / T24 cells (Eddie, 30 Oct 2026; tests 51, 52)
                      or sh == 'SKY JUPITER' and x[0] in ('SKY', 'SKY STARS', 'TIMELINE', 'LONG TIMELINE', 'STAR TIMELINE', 'LONG STAR TIMELINE')):   # SKY block (test 31); SKY STARS, sky timelines (test 33)
                continue
            hit = [v for v in x if isinstance(v, str) and NEWRE.search(v)]
            if hit:
                leak.append((sh, x[0], hit[0][:60]))
# retired 29 Oct 2026 (Eddie, workbook trim), the T23 half: T23 on STARS as before (Polaris kept, every row = the old Ascot build's + Galactic
# Centre) - T23 is gone; and the pairs workbook part of the leak check
check('T30 rebuilt Ascot --full: none of the 13 new stars in any table but T1, T2, T24 and (19 Oct 2026) the SKY block (every sheet)',
      not leak, f'{len(gfw.sheetnames)} sheets; leaks {leak[:3]}')
# two T1, two T2 and three T24 cells of new stars by hand from POSITIONS (rebuilt Ascot)
rates = {x[3]: (x[7], x[8]) for x in gw['POSITIONS'].iter_rows(min_row=5, values_only=True) if x and x[0] == 'SKY' and x[3]}
ra_d = lambda u, v: min(abs(u - v) % 360, 360 - abs(u - v) % 360)
dist = lambda p, q, pl: ra_d(p[0], q[0]) if pl == 'RA' else abs(p[1] - q[1])
cases = [  # (table, chart, star, plane, value, target text, offset text, score by hand, fill)
    ('T1', 'H', 'Pollux', 'RA', 166.673131, '166.667 (ninth)', lambda v: v - 500 / 3, lambda v: 100 * (1.166667 - abs(v - 500 / 3)) / 1.166667, GREEN),
    ('T1', 'H', 'Markab', 'Dec', 37.994684, '38 (whole)', lambda v: v - 38, lambda v: sq(0.02, abs(v - 38)), YELLOW),
    # 24 Oct 2026 (Eddie, 4-digit rule): was 29×φ (golden) -0.00360 (78.98) yellow; 46.919384 vs 29×φ 46.923 (4691 vs 4692) -> outside, no score
    ('T2', 'J', 'Hamal', 'Dec', 46.919384, '29×φ golden — not a 4-digit match (-0.00360)', lambda v: v - 29 * PHI, None, None),
    ('T2', 'J', 'Kaus Australis', 'RA', 0.188691, 'conj (conj)', lambda v: v, lambda v: sq(1.5, v), YELLOW),
]
for tid, C, star, pl, exp, tt, off, sc, fill in cases:
    x = row(gw['P01'], tid, f'{C} Jupiter → {C} {star}')
    c0 = 2 if pl == 'RA' else 10
    v = dist(gp[('P01', C, 'Jupiter')], gp[('P01', C, star)], pl)
    got_fill = x[c0].fill.start_color.rgb if x[c0].fill and x[c0].fill.fill_type else None
    check(f'T30 {tid} {C} Jupiter → {C} {star} {pl} by hand from POSITIONS', abs(v - exp) < 1e-6 and x[c0].value == round(v, 6)
          and x[c0 + 2].value == tt and x[c0 + 4].value == f'{off(v):+.5f}' and got_fill == fill
          and (x[c0 + 5].value is None if sc is None else f'{x[c0 + 5].value:.2f}' == f'{sc(v):.2f}'),
          f'{v:.6f}; {x[c0 + 2].value}, {x[c0 + 4].value}, {x[c0 + 5].value} (by hand {f"{sc(v):.2f}" if sc else None}), {got_fill}')
# its SKY cell: S Jupiter → S Pollux RA, at the off and +5 min from the POSITIONS rates; nearest 35×φ outside
x = row(gw['P01'], 'T1', 'H Jupiter → H Pollux')
sj, sp = gp[('SKY', 'S', 'Jupiter')], gp[('SKY', 'S', 'Pollux')]
v0 = ra_d(sj[0], sp[0]); v5 = ra_d(sj[0] + 5 * rates['Jupiter'][0], sp[0] + 5 * rates['Pollux'][0])
check('T30 T1 H Pollux row, SKY cell S Jupiter → S Pollux RA by hand from POSITIONS (+5 min from the rates): 56.678406 -> 56.677607, nearest 35×φ outside',
      abs(v0 - 56.678406) < 1e-6 and x[18].value == round(v0, 6) and x[19].value == round(v5, 6) and x[20].value == '35×φ (golden) — outside ±0.0324'
      and x[22].value == f'{v0 - 35 * PHI:+.5f}', f'{v0:.6f} -> {v5:.6f}; {x[18].value}, {x[19].value}, {x[20].value}, {x[22].value}')
cases = [  # (sheet, block, who, chart, star, body, plane, value, comment, offset, score)
    ('SKY JUPITER', 'T24 SKY', 'SKY', 'S', 'Deneb', 'Makemake', 'Dec', 17.799198, '11×φ · golden', lambda v: v - 11 * PHI, lambda v: sq(0.03236, abs(v - 11 * PHI))),
    ('P01', 'T24 HORSE', 'P01', 'H', 'Pollux', 'Eris', 'Dec', 32.999538, '33 · whole', lambda v: v - 33, lambda v: sq(0.02, abs(v - 33))),
    ('P01', 'T24 JOCKEY', 'P01', 'J', 'Praesepe', 'Neptune', 'RA', 138.999718, '139 · whole', lambda v: v - 139, lambda v: sq(0.02, abs(v - 139))),
]
for sh, blk, who, C, star, body, pl, exp, tt, off, sc in cases:
    v = dist(gp[(who, C, star)], gp[(who, C, body)], pl)
    hdr, rows = t24_blocks(gw[sh])[blk]
    c = rows[f'{C} {star}'][hdr.index((body, pl))]
    want = f'{tt} · {off(v):+.5f} · {sc(v):.2f}'
    got_fill = c.fill.start_color.rgb if c.fill and c.fill.fill_type else None
    check(f'T30 {blk} {C} {star} – {C} {body} {pl} by hand from POSITIONS', abs(v - exp) < 1e-6 and c.value == round(v, 6)
          and c.comment.text.split('\n')[0] == want and got_fill == GREEN, f'{v:.6f}; {c.comment.text.splitlines()[0]!r} vs {want!r}; {got_fill}')
# Test 31 (Eddie, 19 Oct 2026): T1, T2, T24 and the SKY block star rows on T1_T2_T24_STARS in RA order (engine J2000 RA at J2000.0);
# the SKY block on the 35 stars (ctx.x), everything else as before (TIMELINE, LONG TIMELINE on STARS). "Before" = commit 1d37f4a9.
print('test 31: RA order of the 35 stars; SKY block on the 35 stars')
sys.path.insert(0, 'rebuild_kit')
import swisseph as swe, datetime
import celestial_bodies_swisseph as ENG
swe.set_ephe_path(os.path.expanduser('~/astronomy-project/ephe')); swe.set_topo(0, 51.5, 0)
era = {x: swe.fixstar2_ut(dict(ENG.BEHENIAN_STARS)[x], 2451545.0, ENG._FLAG)[0][0] for x in JT.T1_T2_T24_STARS}
check('T31 T1_T2_T24_STARS = the 67 stars (35 until 20 Oct 2026, 46 until 21 Oct) sorted by engine J2000 RA at J2000.0 (Alpheratz 2.10 ... Markab 346.19)',
      JT.T1_T2_T24_STARS == sorted(era, key=era.get) and len(era) == 67,
      ', '.join(f'{x} {era[x]:.2f}' for x in JT.T1_T2_T24_STARS[:3]) + ' ... ' + f'Markab {era["Markab"]:.2f}')
sky_star_rows = lambda w: [x[1].split(' → S ')[1] for x in w['SKY JUPITER'].iter_rows(values_only=True)
                           if x[0] == 'SKY' and x[1] and ' → S ' in str(x[1]) and x[1].split(' → S ')[1] in set(JT.STARS) | set(JT.T1_T2_T24_STARS)]
lfw = openpyxl.load_workbook(lf)
ok, det = True, []
for nm, w, nst in (('Ascot', gw, 67), ('Lingfield', lfw, 21)):
    want = [x for x in JT.T1_T2_T24_STARS if nst == 67 or x in JT.STARS and x != 'Galactic Centre']
    got = {'SKY': sky_star_rows(w)}
    for sh in ['SKY JUPITER'] + runner_tabs(w):
        for x in w[sh].iter_rows(values_only=True):
            if x[0] in ('T1', 'T2') and x[1] and ' → ' in str(x[1]) and x[1].split(' → ')[1][2:] in JT.T1_T2_T24_STARS + ['Polaris']:
                got.setdefault(f'{sh} {x[0]}', []).append(x[1].split(' → ')[1][2:])
        for k, (hdr, rows) in t24_blocks(w[sh]).items():
            got[f'{sh} {k}'] = [lab.split(' ', 1)[1] for lab in rows]
    ok &= len(want) == nst and all(v == want for v in got.values()) and len(got) == 1 + 1 + 4 * len(runner_tabs(w))
    det.append(f'{nm} {len(got)} lists of {nst}')
check('T31 T1, T2, T24 SKY / HORSE / JOCKEY and SKY block star rows in RA order: rebuilt Ascot 67 (35 until 20 Oct 2026, 46 until 21 Oct), Lingfield 21 (its old batch: 22 stars '
      'with Polaris, no Galactic Centre)', ok, '; '.join(det))
#@ full: T31 comparison 1d37f4a9 -> 93e73bc4 builds (old against old)
# the build before (commit 1d37f4a9) and the comparison, both races (--full)
bf = 'tools/_jt_before_1d37f4a9.py'
open(bf, 'w').write(old_code('1d37f4a9'))
cmp = {}
for nm, inp, after in (('Lingfield', 'reference/classic/race_batch_20210402_lingfield_aw_1435.xlsx', lf_h),
                       ('Ascot', b35, gf35_h)):               # 35-star copy (20 Oct 2026), code of 93e73bc4
    bo = f'{tmp}/{nm}_before_jupiter_tables.xlsx'
    run_build([bf, inp, '--full', '--out', bo], capture_output=True, text=True)
    r = subprocess.run([sys.executable, 'tools/scans/families_audit/chk_raorder/compare_raorder.py', bo, after], capture_output=True, text=True)
    cmp[nm] = (r.stdout, bo)
os.remove(bf)
line = lambda out, k: next((l for l in out.splitlines() if l.startswith(k)), '')
check('T31 every T1 / T2 / T24 star row identical to the build before, matched by star (value, fill, comment, outline; only the order '
      'differs); the rest of the workbook identical apart from the SKY block star rows',
      all(line(o, 'differences beyond those:') == 'differences beyond those: 0' for o, _ in cmp.values())
      and line(cmp['Lingfield'][0], 'T1 / T2 / T24 star rows matched') == 'T1 / T2 / T24 star rows matched by label: 525'
      and line(cmp['Ascot'][0], 'T1 / T2 / T24 star rows matched') == 'T1 / T2 / T24 star rows matched by label: 175',
      ' | '.join(f"{n}: {line(o, 'T1 / T2 / T24')}; {line(o, 'differences')}" for n, (o, _) in cmp.items()))
sd = {n: line(o, 'SKY block star rows:') for n, (o, _) in cmp.items()}
check('T31 SKY block star rows vs before: Lingfield Polaris gone; Ascot Polaris gone + the 13 new stars; no other row changed',
      sd['Lingfield'] == "SKY block star rows: {'SKY JUPITER': {'gone': ['S Jupiter → S Polaris']}}"
      and "'gone': ['S Jupiter → S Polaris']" in sd['Ascot'] and 'changed' not in sd['Ascot']
      and all(f'S Jupiter → S {x}' in sd['Ascot'] for x in JT.NEW_STARS if x not in NEW11 + NEW21), str(sd)[:300])   # the 13 (35-star copy)
# 20 Oct 2026 (Eddie): the sky timelines now take their star events from the 35 stars (test 33): before minus the Polaris events = after
# minus the new-star events (section count rows left out)
tl = lambda f, drop: [x for x in openpyxl.load_workbook(f, read_only=True)['SKY JUPITER'].iter_rows(values_only=True)
                      if x and x[0] in ('TIMELINE', 'LONG TIMELINE') and not str(x[3] or '').startswith('———') and not drop(str(x))]
tls = {n: (tl(b, lambda t: 'Polaris' in t), tl(a, lambda t: bool(NEWRE.search(t))))
       for n, (b, a) in (('Lingfield', (cmp['Lingfield'][1], lf_h)), ('Ascot', (cmp['Ascot'][1], gf35_h)))}
check('T31 sky TIMELINE and LONG TIMELINE rows identical to the build before apart from the Polaris events (gone) and the new-star events '
      '(20 Oct 2026; both races)', all(b == a and b for b, a in tls.values()),
      ', '.join(f'{n} {len(a)} rows' for n, (b, a) in tls.items()))
#@ quick: T31 SKY block cells and by-hand cases on the current build
# SKY block new-star cells = the T1 SKY cells (rebuilt Ascot)
sb = {x[1].value.split(' → S ')[1]: x for x in gw['SKY JUPITER'].iter_rows() if x[0].value == 'SKY' and x[1].value and ' → S ' in str(x[1].value)}
t1 = {x[1].value.split(' → H ')[1]: x for x in gw['P01'].iter_rows() if x[0].value == 'T1' and x[1].value and ' → H ' in str(x[1].value)}
sig = lambda cs: [c.value for c in cs] + [cs[0].fill.start_color.rgb if cs[0].fill and cs[0].fill.fill_type else None]
bad = [(x, pl) for x in JT.NEW_STARS for i, pl in enumerate(('RA', 'Dec'))
       if sig(sb[x][2 + 9 * i:11 + 9 * i]) != sig(t1[x][18 + 9 * i:27 + 9 * i])]
check('T31 every SKY block new-star cell on the rebuilt Ascot = the T1 SKY cell for the same star and plane (value, +5, target, group, '
      'offset, score, arrow, also, timing, fill)', not bad and len(JT.NEW_STARS) == 45, f'90 cell groups (13 + 11 + 21 stars); {bad[:3]}')
# two SKY block new-star cells by hand from POSITIONS (rates for +5 and timing; linear, off 2012-06-19 14:30)
off = datetime.datetime(2012, 6, 19, 14, 30)
cases = [  # (star, plane, value, target text, target, tol, score, timing text)
    # 24 Oct 2026 (Eddie, 4-digit rule): both were hits (Praesepe √2 70.7 (0.35), separating 18 Jun 03:50; Kaus Australis ninth 144.444 (9.85),
    # applying 23 Jun 07:15); 70.378057 (7037 vs 7071) and 143.593046 (1435 vs 1444) are outside the digit match -> shown outside, no score, no timing
    ('Praesepe', 'RA', 70.378057, '70.7 sqrt2 — not a 4-digit match (-0.33262)', 50 * 2 ** 0.5, 0.005 * 50 * 2 ** 0.5, 'out', 'separating'),
    ('Kaus Australis', 'RA', 143.593046, '144.444 ninth — not a 4-digit match (-0.85140)', 1300 / 9, 0.944444, 'out', 'applying'),
]
for star, pl, exp, tt, tv, tol, kind, way in cases:
    sj, sx = gp[('SKY', 'S', 'Jupiter')], gp[('SKY', 'S', star)]
    v0 = ra_d(sj[0], sx[0]); v5 = ra_d(sj[0] + 5 * rates['Jupiter'][0], sx[0] + 5 * rates[star][0])
    dev = abs(v0 - tv)
    sc = (10 * (tol - dev) / tol) ** 2 if kind == 'sq' else 100 * (tol - dev) / tol
    t = (tv - v0) / ((v5 - v0) / 5)
    clk = (off + datetime.timedelta(minutes=t)).strftime('%d %b %Y %H:%M')
    txt = f"{way}, exact {clk} ({abs(t):.1f} min {'ago' if t < 0 else ''})".replace(' )', ')') if t < 0 else f"{way}, exact {clk} (in {t:.1f} min)"
    x = sb[star]
    if kind == 'out':
        sc, txt = None, None
    check(f'T31 SKY block S Jupiter → S {star} {pl} by hand from POSITIONS', abs(v0 - exp) < 1e-6 and x[2].value == round(v0, 6)
          and x[3].value == round(v5, 6) and x[4].value == tt and x[6].value == f'{v0 - tv:+.5f}'
          and (x[7].value is None if sc is None else f'{x[7].value:.2f}' == f'{sc:.2f}') and x[10].value == txt, f'{v0:.6f} -> {v5:.6f}; {x[4].value}, {x[6].value}, {x[7].value} (by hand {sc}); {x[10].value!r} vs {txt!r}')
# retired 29 Oct 2026 (Eddie, workbook trim): Test 32, T24 ratio tables T24 H / J, T24 HORSE RATIOS, T24 JOCKEY RATIOS (layout and every
# distance / ratio cell from the T24 distances; T24 H / J = the T23 H / J star cells; 7 cells by hand from POSITIONS (H / J Deneb RA 5/4, Altair
# RA outside, Vega Dec unison assisted; HORSE H Algol / Sun RA 3/2, H Denebola / Uranus RA φ; JOCKEY J Bellatrix / Makemake RA 2, J Alkaid /
# Gonggong Dec unison assisted); H / J Nunki RA together; automatic / assisted counts (Ascot 43, Lingfield 156); every other row identical to
# the 10b7d76a build (chk_t24r/compare_t24r.py)) - those tables are gone (the 10b7d76a builds removed)
fillrgb = lambda c: c.fill.start_color.rgb if c.fill and c.fill.fill_type else None
first = lambda c: c.comment.text.split('\n')[0] if c.comment else None
# Test 33 (Eddie, 20 Oct 2026): SKY STARS and the sky TIMELINE / LONG TIMELINE on the 35 stars (ctx.x); everything else unchanged
print('test 33: SKY STARS and the sky timelines on the 35 stars')
import itertools
SS_ORDER = [x for x in JT.STARS if x in JT.T1_T2_T24_STARS] + [x for x in JT.T1_T2_T24_STARS if x not in JT.STARS]


def pos_ctx(w, off):
    """a Ctx from the POSITIONS sky rows (positions and rates at the off), for the listing rule"""
    sky, rt = {}, {}
    for x in w['POSITIONS'].iter_rows(min_row=5, values_only=True):
        if x and x[0] == 'SKY' and x[3]:
            sky[x[3]] = (x[5], x[6], x[4]); rt[x[3]] = None if x[7] in (None, 'n/a') else (x[7], x[8])
    c = JT.Ctx(sky, rt); c.off = off; c.off_str = off.strftime('%Y-%m-%d %H:%M')
    return c


def ss_rows(w):
    blk, out = None, {}
    for x in w['SKY JUPITER'].iter_rows(values_only=True):
        if x and str(x[0] or '').startswith('SKY STARS') and x[1] is None:
            blk = x[0]
        elif x and x[0] == 'SKY STARS':
            out[(x[2], x[3])] = (blk, x)
    return out


bad, det = [], []
for nm, w, off, nst in (('Ascot', gw, datetime.datetime(2012, 6, 19, 14, 30), 67), ('Lingfield', lfw, datetime.datetime(2021, 4, 2, 14, 35), 21)):
    cx = pos_ctx(w, off)
    stars = [x for x in SS_ORDER if x in cx.S0]
    rows = JT.pair_rows(cx, {'H': {}, 'J': {}}, stars, tid='SKY STARS', sky=True)
    jd = cx.S0['Jupiter'][1]
    want_cand = {(f'S {a} – S {b}', pl) for a, b in itertools.combinations(stars, 2) for pl in ('RA', 'Dec', 'Equator')
                 if pl != 'Equator' or not ((jd >= 0) == (cx.S0[a][1] >= 0) == (cx.S0[b][1] >= 0))}
    got_cand = {(h['pair'], h['plane']) for h in rows}
    rule = {(h['pair'], h['plane']) for h in rows if not h['together'] and (h['ev'] or h['mid'] is not None or h['allfam'])}
    sheet = ss_rows(w)
    ok_ = len(stars) == nst and got_cand == want_cand and set(sheet) == rule and not any('Polaris' in k[0] for k in sheet)
    bad += [] if ok_ else [(nm, len(stars), len(got_cand ^ want_cand), sorted(set(sheet) ^ rule)[:3])]
    det.append(f'{nm} {nst} stars, {len(want_cand)} pair-planes, {len(rule)} listed')
check('T33 SKY STARS complete: every pair of the batch\'s stars (67 / 21, no Polaris; 35 until 20 Oct 2026, 46 until 21 Oct) and plane (Equator: straddle) measured, and the rows '
      'on the sheet = exactly those passing the filter (ratio on an interval, midpoint hit, or all three families; together out)',
      not bad, '; '.join(det) + f' {bad[:2]}')
#@ full: T33 comparison bb23ade2 -> 93e73bc4 builds (old against old)
# against the build before (commit bb23ade2): old-star rows unchanged, Polaris rows / events gone, new-star rows / events added, rest identical
bf = 'tools/_jt_before_bb23ade2.py'
open(bf, 'w').write(old_code('bb23ade2'))
out33 = {}
for nm, inp, after in (('Lingfield', 'reference/classic/race_batch_20210402_lingfield_aw_1435.xlsx', lf_h),
                       ('Ascot', b35, gf35_h)):               # 35-star copy (20 Oct 2026), code of 93e73bc4
    bo = f'{tmp}/{nm}_before33_jupiter_tables.xlsx'
    run_build([bf, inp, '--full', '--out', bo], capture_output=True, text=True)
    out33[nm] = subprocess.run([sys.executable, 'tools/scans/families_audit/chk_sky35/compare_sky35.py', bo, after], capture_output=True, text=True).stdout
os.remove(bf)
ok_ = True
for nm, o in out33.items():
    ls = o.splitlines()
    ok_ &= ('old rows changed: 0 []' in line(o, '  gone') and 'all Polaris: True' in line(o, '  gone') and 'all with a new star: True' in line(o, '  gone')
            and all(l.endswith('True') for l in ls if l.startswith(('   - gone', '   + new ')))
            and not any(l.startswith('   ~ changed') for l in ls)
            and all(', changed 0' in l for l in ls if l.startswith(('TIMELINE:', 'LONG TIMELINE:')))
            and line(o, 'differences outside') == 'differences outside SKY STARS / TIMELINE / LONG TIMELINE: 0 []')
check('T33 against the build before (commit bb23ade2): every old-star SKY STARS row identical (cells, fills, block); the rows dropped all '
      'Polaris, the rows added all with a new star; TIMELINE / LONG TIMELINE: events dropped all Polaris, added all new-star, every other event '
      'unchanged (time, cells, also); every other row of every sheet identical (both races)', ok_,
      ' | '.join(f"{n}: {line(o, 'SKY STARS rows per area')[24:]}; {line(o, 'TIMELINE:')}; {line(o, 'LONG TIMELINE:')}" for n, o in out33.items()))
#@ quick: T33 new-star events re-searched with the engine on the current build
# new-star events: exact time = the source cell's timing; beyond 60 min the real time re-searched with the engine
A_OFF = datetime.datetime(2012, 6, 19, 14, 30)
JA = EN._to_jd_ut(EN._parse_time('2012-06-19 14:30', 'Europe/London'))
SA = _site
ba = EN.bodies_at(JA, *SA, False)
skp = {x[3]: x for x in gw['POSITIONS'].iter_rows(min_row=5, values_only=True) if x and x[0] == 'SKY' and x[3]}
ecache = {}


def eng(m):
    if m not in ecache:
        ecache[m] = EN.bodies_at(JA + m / 1440.0, *SA, False)
    return ecache[m]


def sky_at(p, m):
    eb = p + '_B' if skp[p][4] == 'star' else p
    b = eng(m)
    return ((skp[p][5] + (b[eb]['ra'] - ba[eb]['ra'] + 180) % 360 - 180) % 360, skp[p][6] + b[eb]['dec'] - ba[eb]['dec'])


def pdist(p, q, pl):
    if pl == 'RA':
        return ra_d(p[0], q[0])
    return abs(p[1] - q[1]) if pl == 'Dec' else abs(abs(p[1]) - abs(q[1]))


ss = ss_rows(gw)
tlr = [x for x in gw['SKY JUPITER'].iter_rows(values_only=True) if x and x[0] in ('TIMELINE', 'LONG TIMELINE') and x[4] and NEWRE.search(str(x[4]))]
bad, nreal, nnf = [], 0, 0
for e in tlr:
    rel, pl, kind, tv = e[4], e[5], e[6], e[8]
    if kind in ('ratio', 'midpoint'):                      # 22 Oct 2026: a midpoint event (unison) of a SKY STARS row, e.g. ONC – Phact (LONG >= 90)
        a, b = rel.split(': ')[1].split(' – ')
        src = ss[(f'{a} – {b}', pl)][1][14]                 # ratio timing cell
        # orientation fixed at the off, as the table (21 Oct 2026: larger / smaller can never cross a unison)
        big = pdist(sky_at('Jupiter', 0), sky_at(a[2:], 0), pl) >= pdist(sky_at('Jupiter', 0), sky_at(b[2:], 0), pl)
        f = lambda m: (lambda dx, dy: (dx / dy if big else dy / dx) - tv)(*(pdist(sky_at('Jupiter', m), sky_at(z[2:], m), pl) for z in (a, b)))
    elif ' → ' in rel:                                     # a S Jupiter leg of a SKY STARS row (20 Oct 2026: e.g. S Jupiter → S Menkent)
        a, b = rel.split(' → ')
        r_ = next(r for k, (_, r) in ss.items() if k[1] == pl and re.search(rf'{re.escape(b)}( –|$)', k[0]))
        src = r_[5] if r_[2].startswith(b + ' –') else r_[7]
        f = lambda m: pdist(sky_at('Jupiter', m), sky_at(b[2:], m), pl) - tv
    else:
        a, b = rel.split(' – ')
        r_ = next(r for k, (_, r) in ss.items() if k[1] == pl and a in k[0] and b in k[0])
        src = r_[9] if f'{a} – {b}' == r_[2] else None     # d(X, Y) family cell: '→ T (fam, sc) · timing'
        f = lambda m: pdist(sky_at(a[2:], m), sky_at(b[2:], m), pl) - tv
    if e[1] == 'not within 1 year':
        nnf += 1
        days = [f(d * 1440) for d in range(-366, 367)]
        ok_ = all((x > 0) == (days[0] > 0) for x in days)
    elif e[0] == 'TIMELINE' or abs(e[13]) <= 60:
        mins = e[2] if e[0] == 'TIMELINE' else e[13]
        ok_ = src and f'exact {e[1]}' in src and f'{abs(mins):.1f} min' in src
    else:
        nreal += 1
        ok_ = f(e[13] - 1) * f(e[13] + 1) < 0
    if not ok_:
        bad.append((e[0], rel, pl, e[1], e[2], src))
check('T33 every new-star event on the rebuilt Ascot: within 60 min the exact time = its SKY STARS source cell\'s timing; beyond 60 min the '
      'real time re-searched with the engine (the value crosses the target between one minute before and after; "not within 1 year": no '
      'crossing in a daily scan of ±1 year)', not bad and len(tlr) > 0, f'{len(tlr)} event rows, {nreal} real times, {nnf} not within 1 year; {bad[:3]}')
# by hand from POSITIONS (rates for +5): two new-star SKY STARS rows, the TIMELINE event, the LONG TIMELINE event
for a, b, pl, area, tv, tt, kind in (('Algol', 'Sadalsuud', 'Equator', 'A', 1.5, '3/2 · fifth · Perfect', 'abs'),
                                     ('Vega', 'Nunki', 'Dec', 'φ and √2', 1 + 2 ** 0.5, '1+√2 · silver ratio · Phi', 'pct')):
    J0, J5 = gp[('SKY', 'S', 'Jupiter')], (gp[('SKY', 'S', 'Jupiter')][0] + 5 * rates['Jupiter'][0], gp[('SKY', 'S', 'Jupiter')][1] + 5 * rates['Jupiter'][1])
    X0, Y0 = gp[('SKY', 'S', a)], gp[('SKY', 'S', b)]
    X5, Y5 = ((p[0] + 5 * rates[z][0], p[1] + 5 * rates[z][1]) for p, z in ((X0, a), (Y0, b)))
    dx, dy, dxy = pdist(J0, X0, pl), pdist(J0, Y0, pl), pdist(X0, Y0, pl)
    r0 = max(dx, dy) / min(dx, dy)
    r5 = (pdist(J5, X5, pl) / pdist(J5, Y5, pl)) if dx > dy else (pdist(J5, Y5, pl) / pdist(J5, X5, pl))
    off_ = r0 - tv if kind == 'abs' else (r0 / tv - 1) * 100
    sc = sq(0.02, abs(off_)) if kind == 'abs' else sq(0.15, abs(off_))
    t = (tv - r0) / ((r5 - r0) / 5)
    clk = (A_OFF + datetime.timedelta(minutes=t)).strftime('%H:%M')
    tim = f"{'applying' if t > 0 else 'separating'}, exact {clk} (in {t:.1f} min)" if t > 0 else f"separating, exact {clk} ({-t:.1f} min ago)"
    pos_ = 'between' if abs(dxy - (dx + dy)) <= 0.01 * (dx + dy) else 'far side'
    blk, x = ss[(f'S {a} – S {b}', pl)]
    offt = f'{off_:+.5f}' + ('%' if kind == 'pct' else '')
    check(f'T33 SKY STARS S {a} – S {b} {pl} by hand from POSITIONS', [round(v, 6) for v in (dx, dy, dxy, r0)] == [x[4], x[6], x[8], x[10]]
          and x[11] == tt and x[12] == offt and f'{x[13]:.2f}' == f'{sc:.2f}' and x[14] == tim and x[15] == pos_ and area.upper() in blk.upper(),
          f'{dx:.6f} / {dy:.6f} / {dxy:.6f}, ratio {r0:.6f} -> {r5:.6f}, {tt} {offt} ({sc:.2f}), {tim}, {pos_}; {blk[:30]}; sheet {x[10:16]}')
ev = [x for x in gw['SKY JUPITER'].iter_rows(values_only=True) if x and x[0] == 'TIMELINE' and x[4] == 'S Jupiter: S Vega – S Nunki']
check('T33 TIMELINE new-star event S Jupiter: S Vega – S Nunki Dec 1+√2, exact 14:32 (+2.8 min), applying, 99.13 - inside −30 / +60',
      len(ev) == 1 and ev[0][1:4] == ('14:32', 2.8, 'SKY STARS') and ev[0][5] == 'Dec' and ev[0][10] == 99.13 and ev[0][11] == 'applying', str(ev))
lv = [x for x in gw['SKY JUPITER'].iter_rows(values_only=True) if x and x[0] == 'LONG TIMELINE' and x[4] == 'S Jupiter: S Pleiades – S Nunki']
fpn = lambda m: (lambda dx, dy: max(dx, dy) / min(dx, dy) - 1.5)(*(pdist(sky_at('Jupiter', m), sky_at(z, m), 'Equator') for z in ('Pleiades', 'Nunki')))
check('T33 LONG TIMELINE new-star event S Jupiter: S Pleiades – S Nunki Equator 3/2 (95.80 at the off): real 12:29 (−120.9 min) by the engine, '
      'linear 12:28 (−121.1 min); direct engine lookup: the ratio crosses 3/2 between one minute before and after',
      len(lv) == 1 and lv[0][1] == '12:29' and lv[0][13] == -120.9 and lv[0][14] == '12:28 (-121.1 min)' and lv[0][10] == 95.8
      and fpn(-121.9) * fpn(-119.9) < 0, f'{lv[0][:15] if lv else lv}; {fpn(-121.9):+.7f} {fpn(-119.9):+.7f}')
# Test 34 (Eddie, 20 Oct 2026): 11 more stars (Alpheratz, Algenib, Difda, Mirach, Menkar = alpha Ceti, Phact, Adhara, Wezen, Menkent,
# Unukalhai, Rasalhague); T1_T2_T24_STARS 46 in engine RA order, the same tables as the 13; everything else unchanged. "Before" = 7d88f329.
print('test 34: 11 more stars (46 in T1, T2, T24, T24 ratio tables, SKY block, SKY STARS, sky timelines)')
NEW11 = ['Alpheratz', 'Algenib', 'Difda', 'Mirach', 'Menkar', 'Phact', 'Adhara', 'Wezen', 'Menkent', 'Unukalhai', 'Rasalhague']
CAT11 = {'Alpheratz': (2.097, 29.090), 'Algenib': (3.309, 15.184), 'Difda': (10.897, -17.987), 'Mirach': (17.433, 35.621),
         'Menkar': (45.570, 4.090), 'Phact': (84.912, -34.074), 'Adhara': (104.656, -28.972), 'Wezen': (107.098, -26.393),
         'Menkent': (211.671, -36.370), 'Unukalhai': (236.067, 6.426), 'Rasalhague': (263.734, 12.560)}
EXPECT46 = ['Alpheratz', 'Algenib', 'Difda', 'Mirach', 'Hamal', 'Menkar', 'Algol', 'Pleiades', 'Aldebaran', 'Rigel', 'Capella', 'Bellatrix',
            'Phact', 'Betelgeuse', 'Sirius', 'Adhara', 'Wezen', 'Castor', 'Procyon', 'Pollux', 'Praesepe', 'Regulus', 'Denebola', 'Algorab',
            'Spica', 'Alkaid', 'Menkent', 'Arcturus', 'Zubenelgenubi', 'Zubeneschamali', 'Alphecca', 'Unukalhai', 'Antares', 'Rasalhague',
            'Galactic Centre', 'Kaus Australis', 'Vega', 'Nunki', 'Altair', 'Deneb', 'Sadalsuud', 'Deneb Algedi', 'Sadalmelik', 'Fomalhaut',
            'Scheat', 'Markab']
B46 = dict(ENG.BEHENIAN_STARS)
e11 = {x: swe.fixstar2_ut(B46[x], 2451545.0, ENG._FLAG) for x in NEW11}
off11 = {x: (e11[x][0][0] - CAT11[x][0], e11[x][0][1] - CAT11[x][1]) for x in NEW11}
check('T34 engine: the 11 stars at J2000.0 within 0.02° of the catalogue (largest Adhara dRA +0.0070); Menkar looked up as ",alCet" = alpha '
      'Ceti (plain "Menkar" would give lambda Ceti, Dec +8.9)', all(abs(a) < 0.02 and abs(b) < 0.02 for a, b in off11.values())
      and B46['Menkar'] == ',alCet' and e11['Menkar'][1] == 'Menkar,alCet' and e11['Difda'][1] == 'Difda,beCet' and e11['Algenib'][1] == 'Algenib,gaPeg',
      ', '.join(f'{x} {a:+.4f}/{b:+.4f}' for x, (a, b) in off11.items()) + f'; Menkar -> {e11["Menkar"][1]} {e11["Menkar"][0][1]:+.4f}')
era46 = {x: swe.fixstar2_ut(B46[x], 2451545.0, ENG._FLAG)[0] for x in EXPECT46}
check('T34 T1_T2_T24_STARS = the 46 stars (|Dec| <= 50, no Polaris) sorted by engine J2000 RA at J2000.0 = Eddie\'s list (no pair differs)',
      [x for x in JT.T1_T2_T24_STARS if x in EXPECT46] == EXPECT46 == sorted(EXPECT46, key=lambda x: era46[x][0]) and all(abs(era46[x][1]) <= 50 for x in EXPECT46),
      ', '.join(f'{x} {era46[x][0]:.2f}' for x in EXPECT46[:6]) + ' ...')
# the 46 in order in every table that uses the list (46-star copy of the rebuilt Ascot since 21 Oct 2026; Lingfield 21)
ok, det = True, []
for nm, w, nst in (('Ascot', gw46, 46), ('Lingfield', lfw, 21)):
    want = [x for x in EXPECT46 if nst == 46 or x in JT.STARS and x != 'Galactic Centre']
    got = {'SKY': sky_star_rows(w)}
    for sh in ['SKY JUPITER'] + runner_tabs(w):
        for x in w[sh].iter_rows(values_only=True):
            if x[0] in ('T1', 'T2') and x[1] and ' → ' in str(x[1]) and x[1].split(' → ')[1][2:] in JT.T1_T2_T24_STARS + ['Polaris']:
                got.setdefault(f'{sh} {x[0]}', []).append(x[1].split(' → ')[1][2:])
        for k, (hdr, rows) in t24_blocks(w[sh]).items():
            got[f'{sh} {k}'] = [lab.split(' ', 1)[1] for lab in rows]
        # 29 Oct 2026 (Eddie, workbook trim): T24 H / J and T24 HORSE / JOCKEY RATIOS are gone (7 -> 4 lists per runner)
    ok &= len(want) == nst and all(v == want for v in got.values()) and len(got) == 1 + 1 + 4 * len(runner_tabs(w))
    det.append(f'{nm} {len(got)} lists of {nst}')
check('T34 the 46 (Ascot) / 21 (Lingfield) in RA order in T1, T2, T24 SKY / HORSE / JOCKEY and the SKY block (T24 H / J and the T24 RATIOS '
      'tables gone 29 Oct 2026)', ok, '; '.join(det))
# against the build before (commit 7d88f329): the after workbook minus the rows naming one of the 11 = the before workbook, row for row
# (values, fills, comments), except title rows with counts; the 11 only in the allowed tables. 29 Oct 2026 (Eddie, workbook trim): the
#@ full: T34 comparison 7d88f329 -> 93e73bc4, and timeline counts on the 18aff98a build (old)
# pairs-workbook part retired (it compared with the current build's pairs workbook, no longer written); both builds are by older code
bf = 'tools/_jt_before_7d88f329.py'
open(bf, 'w').write(old_code('7d88f329'))
out34 = {}
for nm, inp, after in (('Lingfield', 'reference/classic/race_batch_20210402_lingfield_aw_1435.xlsx', lf_h),
                       ('Ascot', b46, gf_h)):     # 46-star copy (21 Oct 2026)
    bo = f'{tmp}/{nm}_before34_jupiter_tables.xlsx'
    run_build([bf, inp, '--full', '--out', bo], capture_output=True, text=True)
    o = subprocess.run([sys.executable, 'tools/scans/families_audit/chk_stars46/compare_46.py', bo, after], capture_output=True, text=True).stdout
    out34[nm] = (o, True)
os.remove(bf)
ALLOW = {'POSITIONS': {'SKY', 'P01'}, 'SKY JUPITER': {'SKY', 'SKY STARS', 'T24', 'TIMELINE', 'LONG TIMELINE'},
         'P01': {'T1', 'T2', 'T24', 'T24 H / J', 'T24 HORSE RATIOS', 'T24 JOCKEY RATIOS'}}
rem = {}
for l in out34['Ascot'][0].splitlines():
    m = re.match(r'(.+?): rows before \d+, after \d+, new-star rows (\d+) (\{.*?\});', l)
    if m:
        rem[m.group(1)] = (int(m.group(2)), set(eval(m.group(3))))
ok_ = (all(line(o, 'sheets with differences other than title rows:') == 'sheets with differences other than title rows: 0' and ps for o, ps in out34.values())
       and all(n == 0 or ids <= ALLOW.get(sh, set()) for sh, (n, ids) in rem.items())
       and not any('new-star rows 0 {}' not in l for l in out34['Lingfield'][0].splitlines() if ': rows before' in l))
check('T34 against the build before (commit 7d88f329), both races --full: every row not naming one of the 11 identical (values, fills, '
      'comments) - the earlier 35 stars\' rows, SKY STARS pairs and timeline events included - only count titles differ; the 11 only in '
      'POSITIONS, T1, T2, T24, T24 H / J, T24 RATIOS, the SKY block, SKY STARS, TIMELINE, LONG TIMELINE; '
      'Lingfield (old batch) gains nothing (builds by 7d88f329 / 93e73bc4 code; pairs part retired 29 Oct 2026)', ok_,
      f"Ascot new-star rows { {k: v for k, v in rem.items() if v[0]} }")
# TIMELINE / LONG TIMELINE = before plus the new-star events (count titles)
tlt = lambda w: [str(x[0]) for x in w['SKY JUPITER'].iter_rows(values_only=True) if x and str(x[0] or '').startswith(('TIMELINE (', 'LONG TIMELINE ('))]
tln = lambda w, t: sum(1 for x in w['SKY JUPITER'].iter_rows(values_only=True) if x and x[0] == t and x[4] and re.search(r'\b(' + '|'.join(NEW11) + r')\b', str(x[4])))
# 24 Oct 2026 (Eddie, half + 4-digit rule; these builds run on today's families.py): TIMELINE 50 (12 / 4 / 34) -> 50 (12 / 5 / 33), new-star 12;
# LONG TIMELINE 67 (33 / 34), new-star 20 -> 53 (23 / 30), new-star 17
# 29 Oct 2026 (Eddie, half removed; these builds run on today's families.py): TIMELINE 50 (12 / 5 / 33) -> 47 (10 / 4 / 33), LONG 53 (23 / 30) ->
# 49 (21 / 28); new-star 12 / 17 unchanged
check('T34 rebuilt Ascot: TIMELINE 47 events (12 new-star; 10 before the off, 4 in the first 5 min, 33 after), LONG TIMELINE 49 '
      '(17 new-star; 21 separating, 28 applying; built by the code of 18aff98a, LONG >= 80; on the families of 29 Oct 2026 (no half) - before: 50 (12 / 4 / 34), '
      '67 (20 new-star; 33 / 34))', tln(gw46_s, 'TIMELINE') == 12 and tln(gw46_s, 'LONG TIMELINE') == 17
      and tlt(gw46_s)[0].endswith('47 events, 10 before the off, 4 in the first 5 min, 33 after') and tlt(gw46_s)[1].endswith('49 events, 21 separating, 28 applying'),
      f"{tln(gw46_s, 'TIMELINE')} / {tln(gw46_s, 'LONG TIMELINE')}; {tlt(gw46_s)[0][-70:]}; {tlt(gw46_s)[1][-40:]}")
#@ quick: T34 by-hand cells from POSITIONS on the current builds
# by hand from POSITIONS (rebuilt Ascot): two cells each of T1, T2, T24 (one ratio table cell), the SKY block and SKY STARS; Menkar in each table
NT = {122.222: 0.722222, 13.3333: 0.083333}                 # the _NINTH_RANGES tolerances of these two bands
cases = [  # (table, chart, star, plane, column, value, target text, offset, score by hand, fill)
    # 24 Oct 2026 (Eddie, 4-digit rule): Menkar was ninth 122.222 (52.76), Adhara 10×φ^-1 (74.64); 122.563385 (1225 vs 1222) and 6.184544 (6184 vs
    # 6180) are outside the digit match -> shown outside, no score, no fill
    ('T1', 'H', 'Menkar', 'RA', 2, 122.563385, '122.222 ninth — not a 4-digit match (+0.34116)', lambda v: v - 1100 / 9, None, None),
    ('T1', 'H', 'Adhara', 'Dec', 10, 6.184544, '10×φ^-1 phi power — not a 4-digit match (+0.00420)', lambda v: v - 10 / PHI, None, None),
    ('T2', 'J', 'Rasalhague', 'Dec', 10, 36.018968, '36 (whole)', lambda v: v - 36, lambda v: sq(0.02, abs(v - 36)), LIGHT),
]
for tid, C, star, pl, c0, exp, tt, off, sc, fill in cases:
    x = row(gw46['P01'], tid, f'{C} Jupiter → {C} {star}')
    v = dist(gp[('P01', C, 'Jupiter')], gp[('P01', C, star)], pl)
    got_fill = x[c0].fill.start_color.rgb if x[c0].fill and x[c0].fill.fill_type else None
    check(f'T34 {tid} {C} Jupiter → {C} {star} {pl} by hand from POSITIONS', abs(v - exp) < 1e-6 and x[c0].value == round(v, 6)
          and x[c0 + 2].value == tt and x[c0 + 4].value == f'{off(v):+.5f}' and got_fill == fill
          and (x[c0 + 5].value is None if sc is None else f'{x[c0 + 5].value:.2f}' == f'{sc(v):.2f}'),
          f'{v:.6f}; {x[c0 + 2].value}, {x[c0 + 4].value}, {x[c0 + 5].value} (by hand {f"{sc(v):.2f}" if sc else None}), {got_fill}')
# T2's SKY cell S Jupiter → S Unukalhai Dec (the J Unukalhai row), +5 min from the POSITIONS rates
x = row(gw46['P01'], 'T2', 'J Jupiter → J Unukalhai')
sj, su = gp[('SKY', 'S', 'Jupiter')], gp[('SKY', 'S', 'Unukalhai')]
v0, v5 = abs(sj[1] - su[1]), abs(sj[1] + 5 * rates['Jupiter'][1] - su[1] - 5 * rates['Unukalhai'][1])
check('T34 T2 J Unukalhai row, SKY cell S Jupiter → S Unukalhai Dec by hand from POSITIONS: 13.296270, ninth 13.3333 −0.03706 outside the 4-digit '
      'match (1329 vs 1333; was 55.52 before 24 Oct 2026) -> shown outside, no score',
      abs(v0 - 13.29627) < 1e-6 and x[27].value == round(v0, 6) and x[28].value == round(v5, 6) and x[29].value == '13.3333 ninth — not a 4-digit match (-0.03706)'
      and x[31].value == f'{v0 - 120 / 9:+.5f}' and x[32].value is None,
      f'{v0:.6f} -> {v5:.6f}; {x[27].value}, {x[28].value}, {x[29].value}, {x[31].value}, {x[32].value}')
# T24 HORSE H Phact – H Makemake Dec; T24 HORSE RATIOS H Menkar Dec Jupiter / Mercury
v = dist(gp[('P01', 'H', 'Phact')], gp[('P01', 'H', 'Makemake')], 'Dec')
hdr, rws = t24_blocks(gw['P01'])['T24 HORSE']
c = rws['H Phact'][hdr.index(('Makemake', 'Dec'))]
want = f'39×φ · golden · {v - 39 * PHI:+.5f} · {sq(0.03236, abs(v - 39 * PHI)):.2f}'
check('T34 T24 HORSE H Phact – H Makemake Dec by hand from POSITIONS: 63.104540, 39×φ +0.00121 (92.64)', abs(v - 63.10454) < 1e-6
      and c.value == round(v, 6) and c.comment.text.split('\n')[0] == want and fillrgb(c) == GREEN, f'{v:.6f}; {first(c)!r} vs {want!r}')
# retired 29 Oct 2026 (Eddie, workbook trim): T24 HORSE RATIOS H Menkar Dec d(Menkar, Jupiter) / d(Menkar, Mercury) 1.501372 3/2 (86.75) -
# the T24 RATIOS tables are gone
# SKY block: S Jupiter → S Adhara RA (45 whole, timing = the TIMELINE event 15:19) and S Jupiter → S Menkar RA (√2 14.14)
# 24 Oct 2026 (Eddie, 4-digit rule): S Menkar RA 14.076783 was √2 14.14 (0.57); 1407 vs 1414 -> outside, no score
for star, tt, tv, tol, timing in (('Adhara', '45 (whole)', 45, 0.02, 'applying'), ('Menkar', '14.14 sqrt2 — not a 4-digit match (-0.06535)', 10 * 2 ** 0.5, None, None)):
    x = row(gw['SKY JUPITER'], 'SKY', f'S Jupiter → S {star}')
    sx = gp[('SKY', 'S', star)]
    v0 = ra_d(sj[0], sx[0]); v5 = ra_d(sj[0] + 5 * rates['Jupiter'][0], sx[0] + 5 * rates[star][0])
    t = (tv - v0) / ((v5 - v0) / 5)
    tim = f"applying, exact {(A_OFF + datetime.timedelta(minutes=t)).strftime('%H:%M')} (in {t:.1f} min)" if timing else None
    check(f'T34 SKY block S Jupiter → S {star} RA by hand from POSITIONS: {v0:.6f} -> {tt} {v0 - tv:+.5f} ({sq(tol, abs(v0 - tv)) if tol else None})'
          + (f', {tim}' if tim else ''), x[2].value == round(v0, 6) and x[3].value == round(v5, 6) and x[4].value == tt
          and x[6].value == f'{v0 - tv:+.5f}' and (x[7].value is None and x[10].value is None if tol is None else f'{x[7].value:.2f}' == f'{sq(tol, abs(v0 - tv)):.2f}')
          and (not timing or x[10].value == tim),
          f'{x[2].value}, {x[3].value}, {x[4].value}, {x[6].value}, {x[7].value}, {x[10].value}')
# SKY STARS: S Denebola – S Unukalhai RA 3/2 and S Menkar – S Praesepe RA 5 (area A, absolute)
ss = ss_rows(gw)
for a, b, pl, tv, tt in (('Denebola', 'Unukalhai', 'RA', 1.5, '3/2 · fifth · Perfect'), ('Menkar', 'Praesepe', 'RA', 5, '5 · 5th harmonic · Harmonic')):
    J0, J5 = gp[('SKY', 'S', 'Jupiter')], (gp[('SKY', 'S', 'Jupiter')][0] + 5 * rates['Jupiter'][0], gp[('SKY', 'S', 'Jupiter')][1] + 5 * rates['Jupiter'][1])
    X0, Y0 = gp[('SKY', 'S', a)], gp[('SKY', 'S', b)]
    X5, Y5 = ((p[0] + 5 * rates[z][0], p[1] + 5 * rates[z][1]) for p, z in ((X0, a), (Y0, b)))
    dx, dy, dxy = pdist(J0, X0, pl), pdist(J0, Y0, pl), pdist(X0, Y0, pl)
    r0 = max(dx, dy) / min(dx, dy)
    r5 = (pdist(J5, X5, pl) / pdist(J5, Y5, pl)) if dx > dy else (pdist(J5, Y5, pl) / pdist(J5, X5, pl))
    t = (tv - r0) / ((r5 - r0) / 5)
    clk = (A_OFF + datetime.timedelta(minutes=t)).strftime('%H:%M')
    tim = f"applying, exact {clk} (in {t:.1f} min)" if t > 0 else f"separating, exact {clk} ({-t:.1f} min ago)"
    pos_ = 'between' if abs(dxy - (dx + dy)) <= 0.01 * (dx + dy) else 'far side'
    blk, x = ss[(f'S {a} – S {b}', pl)]
    check(f'T34 SKY STARS S {a} – S {b} {pl} by hand from POSITIONS', [round(v, 6) for v in (dx, dy, dxy, r0)] == [x[4], x[6], x[8], x[10]]
          and x[11] == tt and x[12] == f'{r0 - tv:+.5f}' and f'{x[13]:.2f}' == f'{sq(0.02, abs(r0 - tv)):.2f}' and x[14] == tim and x[15] == pos_
          and 'AREA A' in blk.upper(), f'{dx:.6f} / {dy:.6f} / {dxy:.6f}, ratio {r0:.6f} -> {r5:.6f}, {tt} {r0 - tv:+.5f} '
          f'({sq(0.02, abs(r0 - tv)):.2f}), {tim}, {pos_}; sheet {x[10:16]}')
# ---- 35. T24 conjunction red fill (Eddie, 20 Oct 2026; display only): every T24 SKY / HORSE / JOCKEY body cell and T24 H / J distance
# cell red (FF8080) exactly when RA < 1.5 / Dec < 0.3, otherwise the fill of the build before (commit 93e73bc4); values, comments, borders
# and every other cell identical (only the T24 title note changes); pairs workbooks identical
#@ full: T35 comparison 93e73bc4 -> ad0defaf builds (old against old)
out35 = {}
for nm, bo, after in (('Lingfield', lf_h, lf_r), ('Ascot', gf_h, gf46_r), ('Ascot35', gf35_h, gf35_r)):   # builds by ad0defaf's code (21 Oct 2026)
    o = subprocess.run([sys.executable, 'tools/scans/families_audit/chk_t24red/compare_red.py', bo, after], capture_output=True, text=True).stdout
    psame = True                                             # 29 Oct 2026 (Eddie, workbook trim): the pairs-workbook part retired
    reds = {(m.group(1), m.group(2)): int(m.group(3)) for m in re.finditer(r'^(\S+(?: JUPITER)?)\s+(T24 \S+(?: / J)?)\s+red\s+(\d+) of', o, re.M)}
    out35[nm] = (o, psame, reds)
W35 = {'Ascot': {('SKY JUPITER', 'T24 SKY'): 20, ('P01', 'T24 HORSE'): 16, ('P01', 'T24 JOCKEY'): 13, ('P01', 'T24 H / J'): 2},
       'Lingfield': {('SKY JUPITER', 'T24 SKY'): 7, **{(p, t): n for p, ns in (('P01', (13, 3, 1)), ('P03', (9, 11, 2)), ('P05', (7, 1, 0)),
                     ('P07', (9, 2, 0)), ('P09', (10, 5, 1)), ('P11', (7, 3, 1))) for t, n in zip(('T24 HORSE', 'T24 JOCKEY', 'T24 H / J'), ns)}}}
check('T35 T24 conjunction red, both races --full, commit 93e73bc4 vs commit ad0defaf (the red fill): every T24 SKY / HORSE / JOCKEY body cell and T24 '
      'H / J distance cell red exactly when RA < 1.5 / Dec < 0.3, otherwise the fill it had; values, comments, outlines and every other cell '
      'identical; only the T24 title notes changed (Ascot 3, Lingfield 13); red cells per block as reported (builds by 93e73bc4 / ad0defaf code; '
      'pairs part retired 29 Oct 2026); the 35-star '
      'copy of tests 31 - 33 the same',
      all('mismatches: 0' in o and ps for o, ps, _ in out35.values()) and 'title cells changed: 3' in out35['Ascot'][0]
      and 'title cells changed: 13' in out35['Lingfield'][0] and 'title cells changed: 3' in out35['Ascot35'][0]
      and all(out35[k][2] == W35[k] for k in W35),
      f"{ {k: (line(v[0], 'mismatches:'), v[1], v[2]) for k, v in out35.items()} }")
# retired 29 Oct 2026 (Eddie, workbook trim): T24 H / J Nunki RA d(H Nunki, H Jupiter) 0.805294 red, H / J ratio "together" no fill - T24 H / J
# is gone
# Test 36 (Eddie, 21 Oct 2026): 20 more stars (Ankaa ... Sualocin; Aculeus = M6, Acumen = M7, Facies = M22) and the Orion Nebula Cluster
# (M42), built as one change; T1_T2_T24_STARS 67 in engine RA order, the same tables as the 11; everything else unchanged. "Before" = ad0defaf.
#@ quick: T36 star order and lists on the current builds
print('test 36: 20 more stars and the Orion Nebula Cluster (67 in T1, T2, T24, T24 ratio tables, SKY block, SKY STARS, sky timelines)')
CAT21 = {'Ankaa': (6.571, -42.306), 'Al Rescha': (30.512, 2.764), 'Mirfak': (51.081, 49.861), 'El Nath': (81.573, 28.607),
         'Meissa': (83.784, 9.934), 'Alnilam': (84.053, -1.202), 'Murzim': (95.675, -17.956), 'Alhena': (99.428, 16.399),
         'Acubens': (134.622, 11.858), 'Alphard': (141.897, -8.659), 'Alkes': (164.944, -18.299), 'Zosma': (168.527, 20.524),
         'Vindemiatrix': (195.544, 10.959), 'Diadem': (197.497, 17.529), 'Ras Algethi': (258.662, 14.390), 'Aculeus': (265.083, -32.253),
         'Acumen': (268.463, -34.793), 'Facies': (279.100, -23.905), 'Rukbat': (290.972, -40.616), 'Sualocin': (309.910, 15.912),
         'Orion Nebula Cluster': (83.822, -5.391)}
EXPECT67 = ['Alpheratz', 'Algenib', 'Ankaa', 'Difda', 'Mirach', 'Al Rescha', 'Hamal', 'Menkar', 'Algol', 'Mirfak', 'Pleiades', 'Aldebaran',
            'Rigel', 'Capella', 'Bellatrix', 'El Nath', 'Meissa', 'Orion Nebula Cluster', 'Alnilam', 'Phact', 'Betelgeuse', 'Murzim', 'Alhena', 'Sirius', 'Adhara',
            'Wezen', 'Castor', 'Procyon', 'Pollux', 'Praesepe', 'Acubens', 'Alphard', 'Regulus', 'Alkes', 'Zosma', 'Denebola', 'Algorab',
            'Vindemiatrix', 'Diadem', 'Spica', 'Alkaid', 'Menkent', 'Arcturus', 'Zubenelgenubi', 'Zubeneschamali', 'Alphecca', 'Unukalhai',
            'Antares', 'Ras Algethi', 'Rasalhague', 'Aculeus', 'Galactic Centre', 'Acumen', 'Kaus Australis', 'Facies', 'Vega', 'Nunki',
            'Rukbat', 'Altair', 'Sualocin', 'Deneb', 'Sadalsuud', 'Deneb Algedi', 'Sadalmelik', 'Fomalhaut', 'Scheat', 'Markab']
B67 = dict(ENG.BEHENIAN_STARS)
e21 = {x: swe.fixstar2_ut(B67[x], 2451545.0, ENG._FLAG) for x in NEW21}
off21 = {x: (e21[x][0][0] - CAT21[x][0], e21[x][0][1] - CAT21[x][1]) for x in NEW21}
check('T36 engine: the 21 stars at J2000.0 within 0.02° of the catalogue (largest Rukbat dRA −0.0078); sefstars names as given (Murzim = '
      'Mirzam,beCMa; clusters Aculeus M6, Acumen M7, Facies M22; Orion Nebula Cluster = Orion Nebula,M42); Mirfak Dec +49.864 <= 50',
      all(abs(u) < 0.02 and abs(v) < 0.02 for u, v in off21.values()) and e21['Murzim'][1] == 'Mirzam,beCMa'
      and [e21[x][1] for x in ('Aculeus', 'Acumen', 'Facies', 'Orion Nebula Cluster')] == ['Aculeus,M6', 'Acumen,M7', 'Facies,M22', 'Orion Nebula,M42'] and abs(e21['Mirfak'][0][1]) <= 50,
      ', '.join(f'{x} {u:+.4f}/{v:+.4f}' for x, (u, v) in off21.items()) + f"; Mirfak Dec {e21['Mirfak'][0][1]:+.6f}")
era67 = {x: swe.fixstar2_ut(B67[x], 2451545.0, ENG._FLAG)[0] for x in EXPECT67}
check('T36 T1_T2_T24_STARS = the 67 stars (|Dec| <= 50, no Polaris) sorted by engine J2000 RA at J2000.0 = Eddie\'s list (no pair differs; '
      'closest Meissa 83.7901 / Orion Nebula Cluster 83.8276)', JT.T1_T2_T24_STARS == EXPECT67 == sorted(EXPECT67, key=lambda x: era67[x][0])
      and all(abs(era67[x][1]) <= 50 for x in EXPECT67) and len(EXPECT67) == 67,
      f"Facies {era67['Facies'][0]:.4f} / Vega {era67['Vega'][0]:.4f}; Meissa {era67['Meissa'][0]:.4f} / ONC {era67['Orion Nebula Cluster'][0]:.4f} / Alnilam {era67['Alnilam'][0]:.4f}; "
      f"Bellatrix {era67['Bellatrix'][0]:.4f} / El Nath {era67['El Nath'][0]:.4f}")
ok, det = True, []
for nm, w, nst in (('Ascot', gw, 67), ('Lingfield', lfw, 21)):
    want = [x for x in EXPECT67 if nst == 67 or x in JT.STARS and x != 'Galactic Centre']
    got = {'SKY': sky_star_rows(w)}
    for sh in ['SKY JUPITER'] + runner_tabs(w):
        for x in w[sh].iter_rows(values_only=True):
            if x[0] in ('T1', 'T2') and x[1] and ' → ' in str(x[1]) and x[1].split(' → ')[1][2:] in JT.T1_T2_T24_STARS + ['Polaris']:
                got.setdefault(f'{sh} {x[0]}', []).append(x[1].split(' → ')[1][2:])
        for k, (hdr, rows) in t24_blocks(w[sh]).items():
            got[f'{sh} {k}'] = [lab.split(' ', 1)[1] for lab in rows]
        # 29 Oct 2026 (Eddie, workbook trim): T24 H / J and T24 HORSE / JOCKEY RATIOS are gone (7 -> 4 lists per runner)
    ok &= len(want) == nst and all(v == want for v in got.values()) and len(got) == 1 + 1 + 4 * len(runner_tabs(w))
    det.append(f'{nm} {len(got)} lists of {nst}')
check('T36 the 67 (Ascot) / 21 (Lingfield) in RA order in T1, T2, T24 SKY / HORSE / JOCKEY and the SKY block (T24 H / J and the T24 RATIOS '
      'tables gone 29 Oct 2026)', ok, '; '.join(det))
# against the build before (commit ad0defaf), both races --full: the after workbook minus the rows naming one of the 20 = the before
# workbook, row for row (values, fills, comments; every row of the 46 earlier stars, SKY STARS pairs and timeline events), only count titles
# differ; the 21 only in the allowed tables; pairs identical
#@ full: T36 comparison ad0defaf -> 18aff98a, and timeline counts on the 18aff98a build (old)
out36 = {}
for nm, bo, after in (('Lingfield', lf_r, lf_s), ('Ascot', gf_r, gf_s)):   # 22 Oct 2026: the after = built by the code of 18aff98a
    o = subprocess.run([sys.executable, 'tools/scans/families_audit/chk_stars67/compare_67.py', bo, after], capture_output=True, text=True).stdout
    psame = True                                             # 29 Oct 2026 (Eddie, workbook trim): the pairs-workbook part retired
    out36[nm] = (o, psame)
rem = {}
for l in out36['Ascot'][0].splitlines():
    m = re.match(r'(.+?): rows before \d+, after \d+, new-star rows (\d+) (\{.*?\});', l)
    if m:
        rem[m.group(1)] = (int(m.group(2)), eval(m.group(3)))
# 24 Oct 2026 (Eddie, half + 4-digit rule; both builds run on today's families.py): SKY STARS new-star rows 347 -> 326, LONG TIMELINE 59 -> 55
W36 = {'POSITIONS': {'SKY': 21, 'P01': 42}, 'SKY JUPITER': {'SKY': 21, 'SKY STARS': 326, 'T24': 21, 'TIMELINE': 40, 'LONG TIMELINE': 55},
       'P01': {'T1': 21, 'T2': 21, 'T24': 42, 'T24 H / J': 21, 'T24 HORSE RATIOS': 21, 'T24 JOCKEY RATIOS': 21}}
ok_ = (all(line(o, 'sheets with differences other than title rows:') == 'sheets with differences other than title rows: 0' and ps for o, ps in out36.values())
       and {k: v[1] for k, v in rem.items() if v[0]} == W36
       and not any('new-star rows 0 {}' not in l for l in out36['Lingfield'][0].splitlines() if ': rows before' in l))
check('T36 against the build before (commit ad0defaf), both races --full: every row not naming one of the 21 identical, row for row (values, '
      'fills, comments) - the 46 earlier stars\' rows, SKY STARS pairs and timeline events included - only count titles differ; the 21 only in '
      'POSITIONS, T1, T2, T24, T24 H / J, T24 RATIOS, the SKY block, SKY STARS, TIMELINE, LONG TIMELINE (rows as reported); '
      'Lingfield (old batch) gains nothing (builds by ad0defaf / 18aff98a code; pairs part retired 29 Oct 2026)', ok_,
      f"Ascot new-star rows { {k: v for k, v in rem.items() if v[0]} }")
RX21 = re.compile(r'\b(' + '|'.join(NEW21) + r')\b')
tln21 = lambda w, t: sum(1 for x in w['SKY JUPITER'].iter_rows(values_only=True) if x and x[0] == t and x[4] and RX21.search(str(x[4])))
# 24 Oct 2026 (Eddie, half + 4-digit rule): TIMELINE 90 (29 / 7 / 54) -> 90 (29 / 8 / 53), new-star 40; LONG 126 (61 / 65), new-star 59 -> 108 (47 / 61), 55
# 29 Oct 2026 (Eddie, half removed): TIMELINE 90 (29 / 8 / 53) -> 87 (27 / 7 / 53), LONG 108 (47 / 61) -> 104 (45 / 59); new-star 40 / 55 unchanged
check('T36 rebuilt Ascot: TIMELINE 87 events (40 new-star; 27 before the off, 7 in the first 5 min, 53 after), LONG TIMELINE 104 '
      '(55 new-star; 45 separating, 59 applying; built by the code of 18aff98a, LONG >= 80; on the families of 29 Oct 2026 (no half) - before: 90 (29 / 7 / 54), '
      '126 (59 new-star; 61 / 65))', tln21(gw_s, 'TIMELINE') == 40 and tln21(gw_s, 'LONG TIMELINE') == 55
      and tlt(gw_s)[0].endswith('87 events, 27 before the off, 7 in the first 5 min, 53 after') and tlt(gw_s)[1].endswith('104 events, 45 separating, 59 applying'),
      f"{tln21(gw_s, 'TIMELINE')} / {tln21(gw_s, 'LONG TIMELINE')}; {tlt(gw_s)[0][-70:]}; {tlt(gw_s)[1][-40:]}")
#@ quick: T36 by-hand cells from POSITIONS on the current build
# by hand from POSITIONS (rebuilt Ascot): two cells each of T1, T2, T24 (one ratio table cell), the SKY block and SKY STARS; clusters in T1,
# T24 and SKY STARS
cases = [  # (table, chart, star, plane, column, value, target text, offset, score by hand, fill)
    ('T1', 'H', 'Acumen', 'Dec', 10, 12.001827, '12 (whole)', lambda v: v - 12, lambda v: sq(0.02, abs(v - 12)), GREEN),
    ('T1', 'H', 'El Nath', 'RA', 2, 158.570201, '98×φ (golden)', lambda v: v - 98 * PHI, lambda v: sq(0.03236, abs(v - 98 * PHI)), GREEN),
    ('T2', 'J', 'Murzim', 'RA', 2, 179.445573, '10×φ^6 (phi power)', lambda v: v - 10 * PHI ** 6, lambda v: sq(0.005 * 10 * PHI ** 6, abs(v - 10 * PHI ** 6)), GREEN),
]
for tid, C, star, pl, c0, exp, tt, off, sc, fill in cases:
    x = row(gw['P01'], tid, f'{C} Jupiter → {C} {star}')
    v = dist(gp[('P01', C, 'Jupiter')], gp[('P01', C, star)], pl)
    check(f'T36 {tid} {C} Jupiter → {C} {star} {pl} by hand from POSITIONS: {exp} -> {tt} ({sc(exp):.2f})', abs(v - exp) < 1e-6 and x[c0].value == round(v, 6)
          and x[c0 + 2].value == tt and x[c0 + 4].value == f'{off(v):+.5f}' and f'{x[c0 + 5].value:.2f}' == f'{sc(v):.2f}' and fillrgb(x[c0]) == fill,
          f'{v:.6f}; {x[c0 + 2].value}, {x[c0 + 4].value}, {x[c0 + 5].value} (by hand {sc(v):.2f}), {fillrgb(x[c0])}')
# T2's SKY cell S Jupiter → S El Nath Dec (the J El Nath row): ninth 8.88889 (_NINTH_RANGES tolerance 0.013889), +5 min from the POSITIONS rates
x = row(gw['P01'], 'T2', 'J Jupiter → J El Nath')
sj, se = gp[('SKY', 'S', 'Jupiter')], gp[('SKY', 'S', 'El Nath')]
v0, v5 = abs(sj[1] - se[1]), abs(sj[1] + 5 * rates['Jupiter'][1] - se[1] - 5 * rates['El Nath'][1])
check('T36 T2 J El Nath row, SKY cell S Jupiter → S El Nath Dec by hand from POSITIONS: 8.884166, ninth 8.88889 −0.00472 outside the 4-digit '
      'match (8884 vs 8888; was 66.00 before 24 Oct 2026) -> shown outside, no score',
      abs(v0 - 8.884166) < 1e-6 and x[27].value == round(v0, 6) and x[28].value == round(v5, 6) and x[29].value == '8.88889 ninth — not a 4-digit match (-0.00472)'
      and x[31].value == f'{v0 - 80 / 9:+.5f}' and x[32].value is None,
      f'{v0:.6f} -> {v5:.6f}; {x[27].value}, {x[28].value}, {x[29].value}, {x[31].value}, {x[32].value}')
# T24 JOCKEY J Facies – J Neptune Dec (cluster); T24 JOCKEY RATIOS J Mirfak RA d(Mirfak, Jupiter) / d(Mirfak, Uranus)
v = dist(gp[('P01', 'J', 'Facies')], gp[('P01', 'J', 'Neptune')], 'Dec')
hdr, rws = t24_blocks(gw['P01'])['T24 JOCKEY']
c = rws['J Facies'][hdr.index(('Neptune', 'Dec'))]
want = f'nearest: 1×φ golden — not a 4-digit match ({v - PHI:+.5f})'
check('T36 T24 JOCKEY J Facies – J Neptune Dec by hand from POSITIONS: 1.616795, 1×φ −0.00124 outside the 4-digit match (1616 vs 1618; was 92.49 '
      'before 24 Oct 2026) -> nearest, no fill', abs(v - 1.616795) < 1e-6
      and c.value == round(v, 6) and first(c) == want and fillrgb(c) is None, f'{v:.6f}; {first(c)!r} vs {want!r}')
# retired 29 Oct 2026 (Eddie, workbook trim): T24 JOCKEY RATIOS J Mirfak RA d(Mirfak, Jupiter) / d(Mirfak, Uranus) 1.199952 6/5 B (94.78) -
# the T24 RATIOS tables are gone
# SKY block: S Jupiter → S Meissa RA (silver 24.14 = 10 (1 + √2)) and S Jupiter → S Al Rescha RA (18×φ). 24 Oct 2026 (Eddie, 4-digit rule): both
# were hits (90.61, separating 13:53; 51.96, separating 13:33); 24.136328 (2413 vs 2414) and 29.133646 (2913 vs 2912) are outside the digit
# match -> shown outside, no score, no timing
for star, tt, tv, tol in (('Meissa', '24.14 silver — not a 4-digit match (-0.00581)', 10 * (1 + 2 ** 0.5), max(0.005 * 10 * (1 + 2 ** 0.5), 0.025)),
                          ('Al Rescha', '18×φ golden — not a 4-digit match (+0.00903)', 18 * PHI, 0.03236)):
    x = row(gw['SKY JUPITER'], 'SKY', f'S Jupiter → S {star}')
    sx = gp[('SKY', 'S', star)]
    v0 = ra_d(sj[0], sx[0]); v5 = ra_d(sj[0] + 5 * rates['Jupiter'][0], sx[0] + 5 * rates[star][0])
    t = (tv - v0) / ((v5 - v0) / 5)
    tim = f"separating, exact {(A_OFF + datetime.timedelta(minutes=t)).strftime('%H:%M')} ({-t:.1f} min ago)"
    check(f'T36 SKY block S Jupiter → S {star} RA by hand from POSITIONS: {v0:.6f} -> {tt} {v0 - tv:+.5f}, no score, no timing (was '
          f'{sq(tol, abs(v0 - tv)):.2f}, {tim})',
          t < 0 and x[2].value == round(v0, 6) and x[3].value == round(v5, 6) and x[4].value == tt and x[6].value == f'{v0 - tv:+.5f}'
          and x[7].value is None and x[10].value is None,
          f'{x[2].value}, {x[3].value}, {x[4].value}, {x[6].value}, {x[7].value}, {x[10].value}')
# SKY STARS: S Pleiades – S Diadem Dec 2 (area A, absolute) and S Facies – S Deneb Dec 1+1/√2 (φ and √2, 0.15 %; cluster)
ss = ss_rows(gw)
for a, b, pl, tv, tt, area, kind in (('Pleiades', 'Diadem', 'Dec', 2.0, '2 · octave · Perfect', 'AREA A', 'abs'),
                                     ('Facies', 'Deneb', 'Dec', 1 + 1 / 2 ** 0.5, '1+1/√2 · √2 division partner · Dissonant', 'AREA Φ AND √2', 'pct')):
    J0, J5 = gp[('SKY', 'S', 'Jupiter')], (gp[('SKY', 'S', 'Jupiter')][0] + 5 * rates['Jupiter'][0], gp[('SKY', 'S', 'Jupiter')][1] + 5 * rates['Jupiter'][1])
    X0, Y0 = gp[('SKY', 'S', a)], gp[('SKY', 'S', b)]
    X5, Y5 = ((p[0] + 5 * rates[z][0], p[1] + 5 * rates[z][1]) for p, z in ((X0, a), (Y0, b)))
    dx, dy, dxy = pdist(J0, X0, pl), pdist(J0, Y0, pl), pdist(X0, Y0, pl)
    r0 = max(dx, dy) / min(dx, dy)
    r5 = (pdist(J5, X5, pl) / pdist(J5, Y5, pl)) if dx > dy else (pdist(J5, Y5, pl) / pdist(J5, X5, pl))
    t = (tv - r0) / ((r5 - r0) / 5)
    clk = (A_OFF + datetime.timedelta(minutes=t)).strftime('%H:%M')
    tim = f"applying, exact {clk} (in {t:.1f} min)" if t > 0 else f"separating, exact {clk} ({-t:.1f} min ago)"
    pos_ = 'between' if abs(dxy - (dx + dy)) <= 0.01 * (dx + dy) else 'far side'
    offt, sct = (f'{r0 - tv:+.5f}', sq(0.02, abs(r0 - tv))) if kind == 'abs' else (f'{(r0 / tv - 1) * 100:+.5f}%', sq(0.15, abs((r0 / tv - 1) * 100)))
    blk, x = ss[(f'S {a} – S {b}', pl)]
    check(f'T36 SKY STARS S {a} – S {b} {pl} by hand from POSITIONS: {dx:.6f} / {dy:.6f} / {dxy:.6f}, {r0:.6f} {tt.split(" · ")[0]} {offt} ({sct:.2f}), {tim}, {pos_}',
          [round(v, 6) for v in (dx, dy, dxy, r0)] == [x[4], x[6], x[8], x[10]] and x[11] == tt and x[12] == offt
          and f'{x[13]:.2f}' == f'{sct:.2f}' and x[14] == tim and x[15] == pos_ and area in blk.upper(),
          f'sheet {x[10:16]}; {blk[:40]}')
# Orion Nebula Cluster by hand from POSITIONS (rebuilt Ascot): T1 H Jupiter → H ONC RA (outside: nearest 100×φ^1, no score, no fill) and its
# SKY cell S Jupiter → S ONC RA (silver 24.14, applying); T24 HORSE H ONC – H Venus Dec (16 whole) and H ONC – H Mars RA (conjunction, red);
# SKY STARS S ONC – S Phact Equator (unison); its pairs with Meissa, Alnilam, Phact: RA "together" (left out), the rest by the filter
ONC = 'Orion Nebula Cluster'
x = row(gw['P01'], 'T1', f'H Jupiter → H {ONC}')
v = dist(gp[('P01', 'H', 'Jupiter')], gp[('P01', 'H', ONC)], 'RA')
check('T36 T1 H Jupiter → H Orion Nebula Cluster RA by hand from POSITIONS: 160.819052, nearest 100×φ^1 outside ±0.809, −0.98435, no score, no fill',
      abs(v - 160.819052) < 1e-6 and x[2].value == round(v, 6) and x[4].value == '100×φ^1 (phi power) — outside ±0.809'
      and x[6].value == f'{v - 100 * PHI:+.5f}' and x[7].value is None and fillrgb(x[2]) is None, f'{v:.6f}; {x[4].value}, {x[6].value}, {x[7].value}')
so = gp[('SKY', 'S', ONC)]
v0 = ra_d(sj[0], so[0]); v5 = ra_d(sj[0] + 5 * rates['Jupiter'][0], so[0] + 5 * rates[ONC][0])
tv, tol = 10 * (1 + 2 ** 0.5), max(0.005 * 10 * (1 + 2 ** 0.5), 0.025)
t_ = (tv - v0) / ((v5 - v0) / 5)
tim = f"applying, exact {(A_OFF + datetime.timedelta(minutes=t_)).strftime('%H:%M')} (in {t_:.1f} min)"
check(f'T36 T1 H Orion Nebula Cluster row, SKY cell S Jupiter → S ONC RA by hand from POSITIONS: {v0:.6f}, 24.14 silver {v0 - tv:+.5f} outside '
      f'the 4-digit match (2417 vs 2414; was {sq(tol, abs(v0 - tv)):.2f}, {tim} before 24 Oct 2026) -> shown outside, no score, no timing',
      abs(v0 - 24.173991) < 1e-6 and x[18].value == round(v0, 6) and x[19].value == round(v5, 6)
      and x[20].value == '24.14 silver — not a 4-digit match (+0.03185)' and x[22].value == f'{v0 - tv:+.5f}' and x[23].value is None and x[26].value is None, f'{x[18].value}, {x[19].value}, {x[20].value}, {x[22].value}, {x[23].value}, {x[26].value}')
hdr, rws = t24_blocks(gw['P01'])['T24 HORSE']
for body, pl, exp, want, fill in (('Venus', 'Dec', 16.000211, lambda v: f'16 · whole · {v - 16:+.5f} · {sq(0.02, abs(v - 16)):.2f}', GREEN),
                                  ('Mars', 'RA', 0.333693, lambda v: f'0 · conjunction (RA < 1.5) · {v:+.5f} · {sq(1.5, v):.2f}', '00FF8080')):
    v = dist(gp[('P01', 'H', ONC)], gp[('P01', 'H', body)], pl)
    c = rws[f'H {ONC}'][hdr.index((body, pl))]
    check(f'T36 T24 HORSE H Orion Nebula Cluster – H {body} {pl} by hand from POSITIONS: {v:.6f}, {want(v)}, fill {fill}', abs(v - exp) < 1e-6
          and c.value == round(v, 6) and first(c) == want(v) and fillrgb(c) == fill, f'{v:.6f}; {first(c)!r}; {fillrgb(c)}')
J0, J5 = gp[('SKY', 'S', 'Jupiter')], (gp[('SKY', 'S', 'Jupiter')][0] + 5 * rates['Jupiter'][0], gp[('SKY', 'S', 'Jupiter')][1] + 5 * rates['Jupiter'][1])
X0, Y0 = gp[('SKY', 'S', ONC)], gp[('SKY', 'S', 'Phact')]
X5, Y5 = ((p[0] + 5 * rates[z][0], p[1] + 5 * rates[z][1]) for p, z in ((X0, ONC), (Y0, 'Phact')))
dx, dy, dxy = pdist(J0, X0, 'Equator'), pdist(J0, Y0, 'Equator'), pdist(X0, Y0, 'Equator')
r0 = max(dx, dy) / min(dx, dy)
r5 = (pdist(J5, X5, 'Equator') / pdist(J5, Y5, 'Equator')) if dx > dy else (pdist(J5, Y5, 'Equator') / pdist(J5, X5, 'Equator'))
t_ = (1 - r0) / ((r5 - r0) / 5)
tim = (f"applying, exact {(A_OFF + datetime.timedelta(minutes=t_)).strftime('%H:%M')} (in {t_:.1f} min)" if t_ > 0 else
       f"separating, exact {(A_OFF + datetime.timedelta(minutes=t_)).strftime('%H:%M')} ({-t_:.1f} min ago)")
pos_ = 'between' if abs(dxy - (dx + dy)) <= 0.01 * (dx + dy) else 'far side'
blk, x = ss[(f'S {ONC} – S Phact', 'Equator')]
check(f'T36 SKY STARS S Orion Nebula Cluster – S Phact Equator by hand from POSITIONS: {dx:.6f} / {dy:.6f} / {dxy:.6f}, {r0:.6f} unison '
      f'{r0 - 1:+.5f} ({sq(0.02, abs(r0 - 1)):.2f}), {tim}, {pos_}', [round(v, 6) for v in (dx, dy, dxy, r0)] == [x[4], x[6], x[8], x[10]]
      and x[11] == '1 · unison · Perfect' and x[12] == f'{r0 - 1:+.5f}' and f'{x[13]:.2f}' == f'{sq(0.02, abs(r0 - 1)):.2f}' and x[14] == tim
      and x[15] == pos_ and 'AREA A' in blk.upper(), f'sheet {x[10:16]}')
# its pairs with Meissa, Alnilam, Phact (orientation: RA order of the new stars, so S Meissa – S ONC): every RA pair "together" (d < 1.5),
# not listed; listed exactly: ONC – Alnilam Dec (6/5) and ONC – Phact Equator (unison)
cx = pos_ctx(gw, A_OFF)
prs = [(f'S Meissa – S {ONC}'), f'S {ONC} – S Alnilam', f'S {ONC} – S Phact']
allr = {(h['pair'], h['plane']): h for h in JT.pair_rows(cx, {'H': {}, 'J': {}}, [x for x in SS_ORDER if x in cx.S0], tid='SKY STARS', sky=True)
        if h['pair'] in prs}
listed = sorted(k for k in ss if k[0] in prs)
check('T36 SKY STARS pairs of the Orion Nebula Cluster with Meissa, Alnilam, Phact: RA together (d(X, Y) 0.037663 / 0.231326 / 1.089043), not '
      'listed; listed exactly ONC – Alnilam Dec (6/5) and ONC – Phact Equator (unison)',
      all(allr[(p_, 'RA')]['together'] for p_ in prs) and listed == [(f'S {ONC} – S Alnilam', 'Dec'), (f'S {ONC} – S Phact', 'Equator')]
      and [round(pdist(gp[('SKY', 'S', a)], gp[('SKY', 'S', b)], 'RA'), 6) for a, b in ((ONC, 'Meissa'), (ONC, 'Alnilam'), (ONC, 'Phact'))]
      == [0.037663, 0.231326, 1.089043], f'{sorted(allr)}; listed {listed}')
# Test 37 (Eddie, 22 Oct 2026): all 67 stars in every timeline - runner TIMELINE H / J and LONG TIMELINE H / J take star events off-sheet
# (S Jup→star, star→Jup, S Jup on star pair, star apex), STAR TIMELINE (sky) and STAR TIMELINE H / J (the Jupiter–star events), LONG
# TIMELINE threshold >= 90 (was 80). "Before" = commit 18aff98a.
#@ full: T37 comparison 18aff98a -> 632447ff (old against old)
print('test 37: all 67 stars in every timeline, STAR TIMELINE, LONG >= 90')
SSRC = ('S Jup→star', 'star→Jup', 'S Jup on star pair', 'star apex')
out37 = {}
for nm, bo, after in (('Lingfield', lf_s, lf_t), ('Ascot', gf_s, gf_t)):   # 18aff98a -> 632447ff (the builds above)
    out37[nm] = subprocess.run([sys.executable, 'tools/scans/families_audit/chk_startl/compare_tl.py', bo, after], capture_output=True, text=True).stdout
# 24 Oct 2026 (Eddie, half + 4-digit rule; both builds run on today's families.py): sky LONG 126 (>= 90: 68) -> 69 became 108 (60) -> 61, STAR 65 -> 66;
# P01 LONG H 17 (8) -> 50 became 14 (7) -> 49, LONG J 23 (6) -> 49 became 21 (8) -> 51, TIMELINE J 11 -> 63 became 13 -> 65
# 29 Oct 2026 (Eddie, half removed; both builds on today's families): sky LONG 108 (60) -> 61 became 104 (59) -> 60, STAR 66 -> 65, TIMELINE 90 -> 87;
# P01 LONG J 21 (8) -> 51 became 19 (6) -> 49, TIMELINE J 13 -> 65 became 11 -> 63
W37 = {'Ascot': ['SKY JUPITER  LONG TIMELINE    before  104 (>= 90: 59) after   60', 'SKY JUPITER  STAR TIMELINE    before    0 after   65',
                 'SKY JUPITER  TIMELINE         before   87 after   87', 'P01          LONG TIMELINE H  before   14 (>= 90: 7) after   49',
                 'P01          LONG TIMELINE J  before   19 (>= 90: 6) after   49', 'P01          STAR TIMELINE H  before    0 after   57',
                 'P01          STAR TIMELINE J  before    0 after   52', 'P01          TIMELINE H       before   13 after   68',
                 'P01          TIMELINE J       before   11 after   63']}
check('T37 against the build before (commit 18aff98a), both races --full: every earlier TIMELINE / TIMELINE H / J event kept with the same cells '
      '(only "also" may gain a star source), every LONG event scoring >= 90 kept, every row outside the timelines identical; counts as reported',
      all(line(o, 'mismatches:') == 'mismatches: 0' for o in out37.values()) and all(any(l.startswith(w_) for l in out37['Ascot'].splitlines()) for w_ in W37['Ascot']),
      ' | '.join(f"{k}: {line(o, 'mismatches:')}" for k, o in out37.items()))
# retired 29 Oct 2026 (Eddie, workbook trim): every runner star event recomputed from POSITIONS (own code star_expect: S Jup→star, star→Jup,
# S Jup on star pair, star apex) in TIMELINE H / J and LONG TIMELINE H / J, none missing or extra; the rebuilt Ascot LONG TIMELINE H / J star
# events' real times (engine crossing within ±1 min) and "not within 1 year" (daily scan) - the runner timelines are gone
#@ quick: T37 STAR TIMELINE and LONG >= 90 on the current builds
# STAR TIMELINE (sky) and STAR TIMELINE H / J = the TIMELINE rows between Jupiter and stars only (own test: no body name, at least one star)
BODYW = re.compile(r'\b(' + '|'.join(b for b in JT.SKY_BODIES + JT.NATAL_BODIES + ['equator'] if b != 'Jupiter') + r')\b')
STW = re.compile(r'\b(' + '|'.join(re.escape(x) for x in sorted(set(JT.STARS) | set(EXPECT67), key=len, reverse=True)) + r')\b')
isstar = lambda rel: bool(STW.search(rel)) and not BODYW.search(STW.sub('', rel))
bad, det = [], []
for nm, w in (('Ascot', gw), ('Lingfield', lfw)):
    # 29 Oct 2026 (Eddie, workbook trim): STAR TIMELINE H / J (runner sheets) are gone - the sky STAR TIMELINE only
    for sh, pairs in [('SKY JUPITER', [('TIMELINE', 'STAR TIMELINE')])]:
        for src_t, st_t in pairs:
            allr = [x for x in w[sh].iter_rows(values_only=True) if x and x[0] in (src_t, st_t)]
            # 22 Oct 2026 (test 39): the sky STAR TIMELINE keeps S Jupiter-centred events only and has three more columns
            sky_ = sh == 'SKY JUPITER'
            a_ = [x[1:13] if sky_ else x[1:] for x in allr if x[0] == src_t and isstar(str(x[4] or '')) and not str(x[3]).startswith('———')
                  and (not sky_ or str(x[4]).startswith('S Jupiter'))]
            b_ = [x[1:13] if sky_ else x[1:] for x in allr if x[0] == st_t and not str(x[3]).startswith('———')]
            mk = [x[3] for x in allr if x[0] == st_t and str(x[3]).startswith('———')]
            if a_ != b_ or mk != ['——— OFF ———', '——— OFF + 5 MIN ———']:
                bad.append((nm, sh, st_t, len(a_), len(b_), mk))
            det.append(f'{nm} {sh} {st_t} {len(b_)}')
check('T37 STAR TIMELINE (sky; STAR TIMELINE H / J gone 29 Oct 2026) = exactly the TIMELINE rows whose relationship is between Jupiter and stars only (same '
      'cells, same order; sky since 22 Oct 2026: S Jupiter-centred only, first 12 cells), with the OFF / OFF + 5 MIN markers, both races', not bad, f'{det[:3]} ...; {bad[:3]}')
# LONG threshold >= 90 (22 Oct 2026): every LONG row scores >= 90 at the off; the >= 80 events of the build before scoring >= 90 all kept
# (compare_tl above); titles
lows = [(sh, x[0], x[10]) for w in (gw, lfw) for sh in ['SKY JUPITER'] for x in w[sh].iter_rows(values_only=True)   # runner LONG H / J gone 29 Oct 2026
        if x and str(x[0] or '').startswith('LONG TIMELINE') and x[3] and not str(x[3]).startswith('———') and x[1] != 'exact (race-local)'
        and (x[10] is None or x[10] < 90)]
check('T37 every sky LONG TIMELINE event scores >= 90 at the off (was >= 80), both races; titles say >= 90 (LONG TIMELINE H / J gone 29 Oct 2026)', not lows
      and all('scoring >= 90 at the off' in str(x[0]) for x in gw['SKY JUPITER'].iter_rows(values_only=True) if str(x[0] or '').startswith('LONG TIMELINE (')),
      f'{lows[:3]}')
# retired 29 Oct 2026 (Eddie, workbook trim): by hand from POSITIONS (rebuilt Ascot P01), two events per runner star source: TIMELINE J
# S Jupiter → J Meissa RA silver (96.72) and J Al Rescha RA 18×φ (92.63); LONG TIMELINE J S Murzim → J Jupiter RA 10×φ^6 real −68.62 days and
# LONG TIMELINE H S Pollux → H Jupiter RA not within 1 year; TIMELINE J S Jupiter: J Mirach – J Menkar RA 3 and TIMELINE H H Pleiades –
# H Nunki Equator 3/2; TIMELINE H Nunki star apex Equator 15/8 and LONG TIMELINE J Sadalsuud Dec √2 - the runner timelines are gone
# Test 38 (Eddie, 22 Oct 2026): LONG STAR TIMELINE (sky) and LONG STAR TIMELINE H / J = the LONG TIMELINE rows between Jupiter and stars only;
# every other row unchanged. "Before" = commit 632447ff.
#@ full: T38 comparison 632447ff -> 00b6db9e (old against old)
print('test 38: LONG STAR TIMELINE')
out38 = {}
for nm, bo, after in (('Lingfield', lf_t, lf_u), ('Ascot', gf_t, gf_u)):    # 632447ff -> 00b6db9e (the builds above)
    out38[nm] = subprocess.run([sys.executable, 'tools/scans/families_audit/chk_longstar/check_longstar.py', bo, after], capture_output=True, text=True).stdout
# 24 Oct 2026 (Eddie, half + 4-digit rule): Ascot sky 59 (28 / 31) -> 55 (25 / 30), H 45 (21 / 24) -> 44 (21 / 23); Lingfield P01 J 3 (0 / 3) -> 2 (0 / 2),
# P07 J 3 (1 / 2) -> 3 (2 / 1), P09 H 3 (2 / 1) -> 2 (2 / 0), P11 J 2 (1 / 1) -> 3 (1 / 2)
W38 = {'Ascot': {'SKY JUPITER LONG STAR TIMELINE': '55 (25 / 30)', 'P01 LONG STAR TIMELINE H': '44 (21 / 23)', 'P01 LONG STAR TIMELINE J': '44 (21 / 23)'},
       'Lingfield': {'SKY JUPITER LONG STAR TIMELINE': '3 (0 / 3)', 'P01 LONG STAR TIMELINE H': '3 (1 / 2)', 'P01 LONG STAR TIMELINE J': '2 (0 / 2)',
                     'P03 LONG STAR TIMELINE H': '3 (1 / 2)', 'P03 LONG STAR TIMELINE J': '4 (1 / 3)', 'P05 LONG STAR TIMELINE H': '4 (2 / 2)',
                     'P05 LONG STAR TIMELINE J': '2 (1 / 1)', 'P07 LONG STAR TIMELINE H': '4 (3 / 1)', 'P07 LONG STAR TIMELINE J': '2 (1 / 1)',
                     'P09 LONG STAR TIMELINE H': '1 (1 / 0)', 'P09 LONG STAR TIMELINE J': '3 (1 / 2)', 'P11 LONG STAR TIMELINE H': '4 (1 / 3)',
                     'P11 LONG STAR TIMELINE J': '2 (1 / 1)'}}
# 29 Oct 2026 (Eddie, half removed): Lingfield P05 J 3 (1 / 2) -> 2 (1 / 1), P07 J 3 (2 / 1) -> 2 (1 / 1), P09 H 2 (2 / 0) -> 1 (1 / 0),
# P11 J 3 (1 / 2) -> 2 (1 / 1); Ascot unchanged
got38 = {nm: {f'{m.group(1)} {m.group(2)}': m.group(3) for m in re.finditer(r'^(\S+(?: JUPITER)?)\s+(LONG STAR TIMELINE(?: [HJ])?)\s+(\d+ \(\d+ / \d+\))$', o, re.M)}
         for nm, o in out38.items()}
check('T38 LONG STAR TIMELINE (sky) / LONG STAR TIMELINE H / J, both races --full (chk_longstar/check_longstar.py): every row = a row of the '
      'matching LONG TIMELINE (identical cells and fills), every LONG row between Jupiter and stars only in it, none missing or extra, same order, '
      'SEPARATING / APPLYING with their counts, title counts, every row >= 90; every other row of every sheet identical to the build before '
      '(commit 632447ff); counts (separating / applying) as reported', all(line(o, 'mismatches:') == 'mismatches: 0' for o in out38.values())
      and got38 == W38, f"{ {k: line(o, 'mismatches:') for k, o in out38.items()} }; {got38['Ascot']}")
#@ quick: T38 LONG STAR TIMELINE sort on the current builds
# sort (own check): inside each section by score at the off, highest first; ties: separating most recent first, applying soonest first
bad = []
for nm, w in (('Ascot', gw), ('Lingfield', lfw)):
    for sh in ['SKY JUPITER']:                                  # 29 Oct 2026 (Eddie, workbook trim): LONG STAR TIMELINE H / J gone
        for tid in ('LONG STAR TIMELINE',):
            sec, parts = None, {}
            for x in w[sh].iter_rows(values_only=True):
                if x and x[0] == tid:
                    if str(x[3]).startswith('———'):
                        sec = x[3].split()[1].rstrip(':'); parts[sec] = []
                    else:
                        parts[sec].append(x)
            for sec, rs in parts.items():
                nf = lambda x: x[1] == 'not within 1 year'
                # scores on the sheet are rounded to 2 dp: two rows may show the same score with different unrounded scores (e.g. Ascot SKY
                # 95.3332 before 95.3325), so the time tie-break is checked only where the next row's shown score is not lower (unrounded order kept)
                sc = [x[10] for x in rs]
                if sc != sorted(sc, reverse=True) or any((x[11] == 'separating') != (sec == 'SEPARATING') for x in rs):
                    bad.append((nm, sh, tid, sec))
check('T38 sky LONG STAR TIMELINE sections in order (own check; H / J blocks gone 29 Oct 2026): by score at the off, highest first (shown 2-dp scores never rise; the time '
      'tie-break is the LONG TIMELINE order, compared row for row above); separating / applying in the right section', not bad, f'{bad[:3]}')
# Test 39 (Eddie, 22 Oct 2026): sky STAR / LONG STAR TIMELINE trimmed to S Jupiter-centred events, 95 / 100 / 95 times; runner T24 STAR
# SELECT tables. "Before" = commit 00b6db9e.
#@ full: T39 comparison 00b6db9e -> current (an older commit, not the one before the current change)
print('test 39: sky star timelines trimmed, 95 / 100 / 95 (T24 STAR SELECT removed)')
out39 = {}
for nm, inp, after, bfull, bdef in (('Lingfield', 'reference/classic/race_batch_20210402_lingfield_aw_1435.xlsx', lf, lf_u, lf_ud),
                                    ('Ascot', f'{tmp}/race_batch_20120619_ascot_1430.xlsx', gf, gf_u, gf_ud)):
    for fl in ('--full', None):
        # 22 Oct 2026 (T24 repeats): the "after" is the build by 5e10f7bf's code (before the repeats), so this keeps checking its own change
        ao = H5E[(nm.lower(), 'full' if fl else 'def')]
        out39[(nm, fl or 'default')] = subprocess.run([sys.executable, 'tools/scans/families_audit/chk_select/compare_sel.py', bfull if fl else bdef, ao],
                                                      capture_output=True, text=True).stdout
# 24 Oct 2026 (Eddie, half + 4-digit rule): Ascot LONG STAR before 59, after 45, dropped 14 -> 55, 44, 11; STAR 65 -> 66 (both builds on today's families)
DROP39 = {'Ascot': ('SKY JUPITER LONG STAR TIMELINE: before 55, after 44, dropped 11: S Aculeus – S Rukbat RA 16×φ (golden) (99.87); S Antares – '
                    'S Alphard Dec 17.7778 (ninth) (96.89)'), 'Lingfield': 'SKY JUPITER LONG STAR TIMELINE: before 3, after 2, dropped 1: S Sirius – S Rigel RA 14×φ (golden) (92.35)'}
check('T39 against the build before (commit 00b6db9e), both races, default and --full (chk_select/compare_sel.py): sky STAR / LONG STAR '
      'TIMELINE = before minus exactly the star-to-star distance events (Ascot LONG 11 since 24 Oct 2026 (was 14), e.g. S Aculeus – S Rukbat RA 16×φ; STAR 0; Lingfield '
      'LONG 1, S Sirius – S Rigel RA 14×φ), every remaining row S Jupiter-centred, its cells identical, the header + "> 95 since", "100", '
      '"< 95 from"; every other row identical (T24 STAR SELECT removed the same day: none on the runner sheets)',
      all(line(o, 'mismatches:') == 'mismatches: 0' for o in out39.values()) and all(DROP39[k[0]] in o for k, o in out39.items())
      and all('STAR TIMELINE: before 65, after 65, dropped 0' in o for k, o in out39.items() if k[0] == 'Ascot'),   # 66 -> 65 (half removed, 29 Oct 2026)
      f"{ {k: line(o, 'mismatches:') for k, o in out39.items()} }")
#@ quick: T39 95 / 100 / 95 cells re-searched with the engine on the current builds; T40 current builds
# (b) every 95 / 100 / 95 cell re-searched with the engine (own score code from POSITIONS + engine; families.py for the targets)
_s.path.insert(0, 'tools/jup2')
import motion as MO2


def sky_engine(w, off, race, batch):
    jd = EN._to_jd_ut(EN._parse_time(off, 'Europe/London'))
    site = MO2.off_and_site(race, 'TRANS', batch, None, None)[1]
    b0 = EN.bodies_at(jd, *site, False)
    sk = {x[3]: x for x in w['POSITIONS'].iter_rows(min_row=5, values_only=True) if x and x[0] == 'SKY' and x[3]}
    cache = {}

    def at(p, m):
        if m not in cache:
            cache[m] = EN.bodies_at(jd + m / 1440.0, *site, False)
        eb = p + '_B' if sk[p][4] == 'star' else p
        b = cache[m]
        return ((sk[p][5] + (b[eb]['ra'] - b0[eb]['ra'] + 180) % 360 - 180) % 360, sk[p][6] + b[eb]['dec'] - b0[eb]['dec'])
    return at


def dpl(p, q, pl):
    if pl == 'RA':
        x = abs(p[0] - q[0]) % 360
        return min(x, 360 - x)
    return abs(p[1] - q[1]) if pl == 'Dec' else abs(abs(p[1]) - abs(q[1]))


def own_score(at, r, m, big, side):
    rel, pl, kind, tgt = r[4], r[5], r[6], r[7]
    if kind == 'distance':
        a, b = rel.split(' → ')
        lab, fam = tgt.rsplit(' (', 1)
        v = dpl(at(a[2:], m), at(b[2:], m), pl)
        return next((d['score'] for d in families.evaluate(v, pl if pl in ('RA', 'Dec') else None, 'deg') if d['family'] == fam[:-1] and d['target'] == lab), 0.0)
    x, y = rel.split(': ')[1].split(' – ')
    J, X, Y = at('Jupiter', m), at(x[2:], m), at(y[2:], m)
    dx, dy = dpl(J, X, pl), dpl(J, Y, pl)
    if kind == 'ratio':
        v = dx / dy if big else dy / dx
        return next((d['score'] for d in families.evaluate(v, None, 'ratio') if d['target'] == tgt.split(' · ')[0]), 0.0)
    close = 100 * min(dx, dy) / max(dx, dy)
    if 100 - close > 10:
        return 0.0
    if pl == 'RA':
        mid = (X[0] + ((Y[0] - X[0] + 180) % 360 - 180) / 2) % 360
        near, far = (J[0] - mid + 180) % 360 - 180, (J[0] - mid) % 360 - 180
        off, tol = (near if side == 'near' else far), 0.5
    else:
        f = (lambda v: v) if pl == 'Dec' else abs
        off, tol = f(J[1]) - (f(X[1]) + f(Y[1])) / 2, 0.15
    if abs(off) > tol:
        return 0.0
    return 0.5 * (10 * (10 - (100 - close)) / 10) ** 2 + 0.5 * (10 * (tol - abs(off)) / tol) ** 2


def mins_of(cell):
    if cell in ('—', 'not within 1 year'):
        return cell
    t = cell.split('(')[1].rstrip(')').replace('−', '-')
    return float(t.split()[0]) * (1440 if 'days' in t else 1)


bad, n95, nd, nnf = [], 0, 0, 0
for nm, w, off, race, batch in (('Ascot', gw, '2012-06-19 14:30', 'ascot_1430_20120619', f'{tmp}/race_batch_20120619_ascot_1430.xlsx'),
                                ('Lingfield', lfw, '2021-04-02 14:35', 'lingfield_aw_1435_20210402', 'reference/classic/race_batch_20210402_lingfield_aw_1435.xlsx')):
    at = sky_engine(w, off, race, batch)
    for tid, c95 in (('STAR TIMELINE', 13), ('LONG STAR TIMELINE', 15)):
        for r in w['SKY JUPITER'].iter_rows(values_only=True):
            if not (r and r[0] == tid and r[3] and not str(r[3]).startswith('———')):
                continue
            big = side = None
            if r[6] != 'distance':
                x, y = r[4].split(': ')[1].split(' – ')
                J, X, Y = at('Jupiter', 0), at(x[2:], 0), at(y[2:], 0)
                big = dpl(J, X, r[5]) >= dpl(J, Y, r[5])
                if r[5] == 'RA':
                    mid = (X[0] + ((Y[0] - X[0] + 180) % 360 - 180) / 2) % 360
                    side = 'near' if abs((J[0] - mid + 180) % 360 - 180) <= abs((J[0] - mid) % 360 - 180) else 'far'
            sc = lambda m: own_score(at, r, m, big, side)
            since, exact, frm = (mins_of(v) for v in r[c95:c95 + 3])
            if since == '—':
                nd += 1
                m0 = 0.0 if exact == 'not within 1 year' else exact
                if not (frm == '—' and sc(m0) < 95):
                    bad.append((nm, tid, r[4], r[5], 'dash', sc(m0)))
                continue
            for v, sgn in ((since, -1), (frm, 1)):
                if v == 'not within 1 year':
                    nnf += 1
                    m0 = 0.0 if exact == 'not within 1 year' else exact
                    if not all(sc(m0 + sgn * dd * 1440) >= 95 for dd in range(1, 367)):
                        bad.append((nm, tid, r[4], r[5], 'nf', sgn))
                    continue
                n95 += 1
                dm = 8.0 if abs(v) >= 1440 else 1.0               # shown in days beyond a day (0.01 day = 14.4 min): ±8 min
                ins, outs = sc(v - sgn * dm), sc(v + sgn * dm)      # inside the 95 window on the exact-time side, outside beyond
                if not (ins >= 95 > outs):
                    bad.append((nm, tid, r[4], r[5], 'since' if sgn < 0 else 'from', v, round(ins, 3), round(outs, 3)))
check('T39 every "> 95 since" / "< 95 from" cell of the sky STAR / LONG STAR TIMELINE re-searched (own score code: family / interval on its own '
      'tolerance, midpoint total; engine positions): the score is >= 95 one minute inside and < 95 one minute outside (8 min for a cell shown '
      'in days, 0.01 day); "—" only where the score '
      'is under 95 at the exact time; "not within 1 year" by a daily ±1 year scan; both races', not bad and n95 > 0,
      f'{n95} crossings, {nd} dashes, {nnf} not within 1 year; {bad[:3]}')
# (c) T24 STAR SELECT removed (Eddie, 22 Oct 2026, the same day): no T24 STAR SELECT title or row on any sheet (both races, --full)
left = [(nm, sh) for nm, w in (('Ascot', gw), ('Lingfield', lfw)) for sh in w.sheetnames
        if any(x and any('T24 STAR SELECT' in str(v) for v in x[:2] if v) for x in w[sh].iter_rows(max_col=2, values_only=True))]
check('T39 T24 STAR SELECT removed (22 Oct 2026): no T24 STAR SELECT title or row on any sheet, both races (--full)', not left, f'{left[:5]}')
# 95 crossing: S Jupiter → S Castor RA on 54 (whole, ±0.02): score 95 at |v − 54| = 0.02 (1 − √0.95) = 0.000506
at = sky_engine(gw, '2012-06-19 14:30', 'ascot_1430_20120619', f'{tmp}/race_batch_20120619_ascot_1430.xlsx')
r = next(x for x in gw['SKY JUPITER'].iter_rows(values_only=True) if x and x[0] == 'STAR TIMELINE' and x[4] == 'S Jupiter → S Castor' and x[5] == 'RA')
dv = lambda m: dpl(at('Jupiter', m), at('Castor', m), 'RA') - 54
since, frm = mins_of(r[13]), mins_of(r[15])
lim = 0.02 * (1 - 0.95 ** 0.5)
check(f'T39 STAR TIMELINE S Jupiter → S Castor RA (54 whole) by hand with the engine: score 95 where |d − 54| = {lim:.6f}; "> 95 since" {r[13]}, '
      f'"100" {r[14]}, "< 95 from" {r[15]}', abs(dv(since - 1)) > lim > abs(dv(since + 1)) and abs(dv(frm + 1)) > lim > abs(dv(frm - 1)),
      f'd−54 at since ±1: {dv(since - 1):+.6f} / {dv(since + 1):+.6f}; at from ±1: {dv(frm - 1):+.6f} / {dv(frm + 1):+.6f}')
# ---------------------------------------------------------------------------------------------------------------------------------
# Test 40 (Eddie, 22 Oct 2026): T24 STAR SELECT removed from the runner sheets; every other block unchanged. "Before" = commit 0ad7d2e5.
print('test 40: builds by the current code (Lingfield, rebuilt Ascot) for tests 41 - 43')
# retired 29 Oct 2026 (Eddie, workbook trim): T40 against the build before (commit 0ad7d2e5, chk_rmsel/compare_rmsel.py): each runner sheet =
# before minus exactly the T24 STAR SELECT tables, every other row identical, pairs identical; and the code check (write_t24sel, star_select,
# SEL_NOTE gone, t24r_cell without tags) - T24 STAR SELECT, t24r_cell and the pairs workbook are gone. Only the --full builds by the current
# code stay (they feed tests 41 - 43; --full = the default since the trim)
os.makedirs(f'{tmp}/after40', exist_ok=True)
for n in ('lingfield', 'ascot'):
    ao = f'{tmp}/after40/{n}_full_jupiter_tables.xlsx'
    inp = 'reference/classic/race_batch_20210402_lingfield_aw_1435.xlsx' if n == 'lingfield' else f'{tmp}/race_batch_20120619_ascot_1430.xlsx'
    subprocess.run([sys.executable, 'tools/jupiter_tables.py', inp, '--full', '--out', ao], capture_output=True, text=True)
# ---------------------------------------------------------------------------------------------------------------------------------
# Test 41 (Eddie, 22 Oct 2026): T24 repeats flag - T24 SKY / HORSE / JOCKEY: the same best target (family + target number, no x10 / /10
# folding; conjunction cells out) in two or more body cells of a star row; "repeats" column, dotted outline (thick midpoint outline kept),
# comment line, title counts. "Before" = commit 5e10f7bf.
#@ full: T41 builds by 10b18003 and the comparison 5e10f7bf -> 10b18003 (old against old)
print('test 41: T24 repeats flag')
# 24 Oct 2026 (Eddie, "not a 4-digit match" wording): builds by the code of commit 10b18003 (repeats, half, 4-digit rule, old wording),
# default and --full - test 41 compares these with the 5e10f7bf builds (the repeats alone), test 43 with now (the wording alone)
bf = 'tools/_jt_before_10b18003.py'
open(bf, 'w').write(old_code('10b18003'))
os.makedirs(f'{tmp}/h10', exist_ok=True)
for n in ('lingfield', 'ascot'):
    inp = 'reference/classic/race_batch_20210402_lingfield_aw_1435.xlsx' if n == 'lingfield' else f'{tmp}/race_batch_20120619_ascot_1430.xlsx'
    for f in ('full', 'def'):
        run_build([bf, inp] + (['--full'] if f == 'full' else []) + ['--out', f'{tmp}/h10/{n}_{f}_jupiter_tables.xlsx'],
                       capture_output=True, text=True)
os.remove(bf)
out41 = {}
for (n, f), bo in H5E.items():
    ao = f'{tmp}/h10/{n}_{f}_jupiter_tables.xlsx'               # 24 Oct 2026: by 10b18003's code (old wording; the repeats as now)
    out41[(n, f)] = subprocess.run([sys.executable, 'tools/scans/families_audit/chk_t24rep/check_rep.py', bo, ao], capture_output=True, text=True).stdout
check('T41 against the build before (commit 5e10f7bf), both races, default and --full (chk_t24rep/check_rep.py): every repeats cell of T24 SKY / '
      'HORSE / JOCKEY recomputed from POSITIONS (own code, families.evaluate: best target per cell, conjunction cells out, every target in >= 2 '
      'cells with its cells, scores and order; none missing or extra); dotted outline exactly on repeated cells (thick kept on midpoint cells); '
      'comment = before + the repeat lines; title = the counts; header "repeats"; every other cell of every sheet identical (values, fills, '
      'comments, borders, formats)', all(line(o, 'mismatches:') == 'mismatches: 0' for o in out41.values())
      # 24 Oct 2026 (Eddie, half + 4-digit rule): Ascot SKY 24 / 67 rows, 27 targets -> 1 / 67, 1; Lingfield SKY 8 / 21, 8 -> 1 / 21, 1
      and all('SKY JUPITER T24 SKY: 1 / 67 rows, 1 targets' in o for k, o in out41.items() if k[0] == 'ascot')
      and all('SKY JUPITER T24 SKY: 1 / 21 rows, 1 targets' in o for k, o in out41.items() if k[0] == 'lingfield'),
      f"{ {k: line(o, 'mismatches:') for k, o in out41.items()} }")
# by hand from POSITIONS: an RA–Dec repeat, a Dec–Dec repeat, a midpoint body cell that is also a repeat.
# 24 Oct 2026 (Eddie, half + 4-digit rule): the first cases (T24 HORSE H Menkent silver 24.14 Makemake RA 24.137213 / Rahu Dec 24.066157; T24 JOCKEY
# J Orion Nebula Cluster ninth 16.6667 Uranus Dec 16.673056 / Gonggong Dec 16.652082; T24 HORSE H Castor RA ninth 144.444 ×3 Quaoar / Rahu / Vesta)
# are all outside the 4-digit match now (2413 / 2406 vs 2414; 1667 / 1665 vs 1666; 1446 / 1438 / 1435 vs 1444), so no repeat is left there. Moved
# to the repeats there are now (found by the rule): rebuilt Ascot T24 SKY S Rukbat 56 whole (Juno RA, Ceres Dec), T24 JOCKEY J Algol 63 whole
# (Gonggong Dec, Uranus Dec); Lingfield T24 SKY S Betelgeuse 16 whole (Juno Dec, Makemake Dec), both cells also midpoint cells (thick kept).
#@ quick: T41 by-hand repeats cells on the current builds; T42 family rules by hand on the current families


def t24cell(w, sh, tid, lab, body, pl):
    rows_ = list(w[sh].iter_rows())
    i = next(k for k, r in enumerate(rows_) if r[0].value == 'T24' and str(r[1].value).startswith(tid))
    h1, h2 = rows_[i + 1], rows_[i + 2]
    cm = [k for k, c in enumerate(h1) if c.value in ('RA midpoints', 'Dec midpoints')]
    r = next(r for r in rows_[i + 3:i + 80] if r[1].value == lab)
    lo, hi = (2, cm[0]) if pl == 'RA' else (cm[0] + 1, cm[1])
    return r[next(k for k in range(lo, hi) if h2[k].value == body)], r[cm[1] + 1].value


def posmap(w, who):
    return {(x[2], x[3]): (x[5], x[6]) for x in w['POSITIONS'].iter_rows(min_row=5, values_only=True) if x and x[0] == who and x[3]}


w41 = openpyxl.load_workbook(f'{tmp}/after40/ascot_full_jupiter_tables.xlsx')
wl41 = openpyxl.load_workbook(f'{tmp}/after40/lingfield_full_jupiter_tables.xlsx')
for w, sh, who, tid, ch, star, n_, (B1, P1), (B2, P2), thick_ in (
        (w41, 'SKY JUPITER', 'SKY', 'T24 SKY', 'S', 'Rukbat', 56, ('Juno', 'RA'), ('Ceres', 'Dec'), False),
        (w41, 'P01', 'P01', 'T24 JOCKEY', 'J', 'Algol', 63, ('Gonggong', 'Dec'), ('Uranus', 'Dec'), False),
        (wl41, 'SKY JUPITER', 'SKY', 'T24 SKY', 'S', 'Betelgeuse', 16, ('Juno', 'Dec'), ('Makemake', 'Dec'), True)):
    pp = posmap(w, who)
    a, b = dpl(pp[(ch, star)], pp[(ch, B1)], P1), dpl(pp[(ch, star)], pp[(ch, B2)], P2)
    sa_, sb_ = sq(0.02, abs(a - n_)), sq(0.02, abs(b - n_))
    (B1, P1, sa_), (B2, P2, sb_) = sorted(((B1, P1, sa_), (B2, P2, sb_)), key=lambda z: -z[2])
    ca, rp = t24cell(w, sh, tid, f'{ch} {star}', B1, P1)
    cb, _ = t24cell(w, sh, tid, f'{ch} {star}', B2, P2)
    want = f'{n_} whole ×2 ({B1} {P1} ({sa_:.2f}), {B2} {P2} ({sb_:.2f}))'
    st = 'thick' if thick_ else 'dotted'
    check(f'T41 repeat by hand from POSITIONS ({"Ascot" if w is w41 else "Lingfield"} {tid}): {ch} {star} – {ch} {B1} {P1} '
          f'{dpl(pp[(ch, star)], pp[(ch, B1)], P1):.6f}, – {ch} {B2} {P2} {dpl(pp[(ch, star)], pp[(ch, B2)], P2):.6f}, both whole {n_} (±0.02): '
          f'{sa_:.2f} / {sb_:.2f}; repeats cell {rp!r}; outline {st}' + ('; both midpoint cells, both comment lines' if thick_ else ''),
          rp == want and ca.border.left.style == cb.border.left.style == st
          and ca.comment.text.endswith(f'repeat: {n_} whole ×2 (with {B2} {P2} ({sb_:.2f}))')
          and cb.comment.text.endswith(f'repeat: {n_} whole ×2 (with {B1} {P1} ({sa_:.2f}))')
          and (not thick_ or all('midpoint with ' in c.comment.text for c in (ca, cb))), f'{rp!r}; {ca.comment.text!r}; {cb.comment.text!r}')
# ---------------------------------------------------------------------------------------------------------------------------------
# Test 42 (Eddie, 24 Oct 2026; half REMOVED 29 Oct 2026): families.py - HALF family (0.5 ... 199.5, ±0.01 absolute, squared) and the 4-DIGIT RULE for golden, φ powers,
# ninths, √2, silver (first 4 significant digits truncated, same order of magnitude, as the exact target; inside the match the old score).
# "Before" = commit 83850b1c (its own tools/ and rebuild_kit/, i.e. the old families.py).
print('test 42: half family removed (29 Oct 2026) and the 4-digit rule')
import importlib.util as _ilu
os.makedirs(f'{tmp}/h85', exist_ok=True)
extract(f'83850b1c', f'{tmp}/h85')
_sp = _ilu.spec_from_file_location('families_old', f'{tmp}/h85/tools/families.py')
FO = _ilu.module_from_spec(_sp)
_sp.loader.exec_module(FO)
F = families
sqs = lambda tol, dev: (10 * (tol - dev) / tol) ** 2
# (a) 29 Oct 2026 (Eddie): the half family is REMOVED. Retired: 'T42 half by hand' (27.5 -> 27.5 (100.00); 27.505 -> 25.00; 27.515 none;
# 0.495 -> 0.5; 199.505 -> 199.5; 200.5 none; order whole, half, golden). Now: no half family, no half score, never a half nearest target
try:
    F.score_family(27.5, 'half'); _hraise = False
except ValueError:
    _hraise = True
_ev = [d['family'] for v in (27.5, 27.505, 0.5, 1.497, 199.5) for d in F.evaluate(v, 'RA')]
_out = [F.outside(v, 'Dec')['family'] for v in (27.52, 10.487259, 12.499623, 83.5)]
check('T42 half removed (Eddie, 29 Oct 2026): "half" not a family (score_family raises), DEG_FAMILIES whole, golden, phi power, ninth, sqrt2, '
      'silver, conj; evaluate(27.5, 27.505, 0.5, 1.497, 199.5) offers no half; outside() never names half (27.52, 10.487259, 12.499623, 83.5 -> '
      f'{_out}); RA 1.497 -> conjunction first ({F.evaluate(1.497, "RA")[0]["family"]})',
      _hraise and F.DEG_FAMILIES == ('whole', 'golden', 'phi power', 'ninth', 'sqrt2', 'silver', 'conj') and 'half' not in _ev
      and 'half' not in _out and F.evaluate(1.497, 'RA')[0]['family'] == 'conj' and not hasattr(F, 'score_half'), f'{_ev} {_out}')
# (b) 4-digit rule by hand
S24, S2 = 10 * (1 + 2 ** 0.5), 1 + 2 ** 0.5
cases = [('silver 24.1421 / 24.1400', F.digit_match(24.14, S24), True), ('silver 24.1421 / 24.1500', F.digit_match(24.15, S24), False),
         ('silver 24.1421 / 24.1399', F.digit_match(24.1399, S24), False), ('φ^-1 0.618034 / 0.6180', F.digit_match(0.6180, 1 / F.PHI), True),
         ('2.414 vs 24.14 (no scaling)', F.digit_match(2.414, S24) or F.digit_match(24.14, S2), False),
         ('ninth 166.667 / 166.60', F.digit_match(166.60, 166.666667), True), ('ninth 166.667 / 166.70', F.digit_match(166.70, 166.666667), False),
         ('golden 59×φ 95.4640 / 95.46', F.digit_match(95.46, 59 * F.PHI), True), ('golden 59×φ 95.4640 / 95.44 (inside ±0.03236)', F.digit_match(95.44, 59 * F.PHI), False),
         ('√2 14.1421 / 14.141031', F.digit_match(14.141031, 10 * 2 ** 0.5), True), ('√2 14.1421 / 14.13 (inside ±0.0707)', F.digit_match(14.13, 10 * 2 ** 0.5), False),
         ('digits4 0.7071 -> (-1, 7071)', F.digits4(0.7071) == (-1, 7071), True)]
check('T42 digits4 by hand: ' + '; '.join(f'{c[0]}: {"match" if c[2] else "no match"}' for c in cases), all(c[1] == c[2] for c in cases),
      str([c for c in cases if c[1] != c[2]]))
lost = [(v, f) for v, f in ((24.1399, 'silver'), (24.15, 'silver'), (166.70, 'ninth'), (95.44, 'golden'), (14.13, 'sqrt2'))
        if F.score_family(v, f)[0] or not FO.score_family(v, f)[0]]
kept = [(v, f, FO.score_family(v, f)[0], F.score_family(v, f)[0]) for v, f in ((24.14, 'silver'), (24.1499, 'silver'), (0.6180, 'phi power'),
        (166.60, 'ninth'), (95.46, 'golden'), (14.141031, 'sqrt2'), (1.414256, 'sqrt2'), (158.570201, 'golden'))]
check('T42 the score inside a match = the old score (families of 83850b1c): ' + ', '.join(f'{v} {f} {b:.2f}' for v, f, a, b in kept)
      + '; outside the match no score, though inside the old tolerance: ' + ', '.join(f'{v} {f}' for v, f in
      ((24.1399, 'silver'), (24.15, 'silver'), (166.70, 'ninth'), (95.44, 'golden'), (14.13, 'sqrt2'))),
      not lost and all(a > 0 and abs(a - b) < 1e-12 for v, f, a, b in kept), f'{lost} {kept}')
# (c) real cells, rebuilt Ascot P01 (--full, current code), by hand from POSITIONS
w42 = openpyxl.load_workbook(f'{tmp}/after40/ascot_full_jupiter_tables.xlsx')
pp = {(x[2], x[3]): (x[5], x[6]) for x in w42['POSITIONS'].iter_rows(min_row=5, values_only=True) if x and x[0] == 'P01' and x[3]}
r42 = {(r[0], r[1]): r for r in w42['P01'].iter_rows(values_only=True) if r and r[0] in ('T1', 'T2')}
want = []
# 29 Oct 2026 (Eddie, half removed): J Rasalhague RA 12.499623 was '12.5 (half)' (92.61) -> nearest ninth 12.2222, no score; H Rahu Dec
# 10.487259 was '10.5 (half) — outside ±0.01' -> nearest ninth 11.1111
for tid, ch, X, pl, exp in (('T2', 'J', 'Rasalhague', 'RA', ('12.2222 (ninth) — outside ±0.0722', None)),
                            ('T1', 'H', 'Rahu', 'Dec', ('11.1111 (ninth) — outside ±0.0611', None)),
                            ('T2', 'J', 'Kaus Australis', 'RA', ('conj (conj)', lambda v: sqs(1.5, v))),
                            ('T1', 'H', 'El Nath', 'RA', ('98×φ (golden)', lambda v: sqs(0.03236, abs(v - 98 * F.PHI)))),
                            ('T2', 'J', 'Gonggong', 'Dec', ('1.414 (sqrt2)', lambda v: sqs(0.025, abs(v - 2 ** 0.5)))),
                            ('T1', 'H', 'Sun', 'RA', ('41.4 silver — not a 4-digit match (-0.04177)', None))):
    r = r42[(tid, f'{ch} Jupiter → {ch} {X}')]
    c0 = 2 if pl == 'RA' else 10
    v = dpl(pp[(ch, 'Jupiter')], pp[(ch, X)], pl)
    sc = exp[1](v) if exp[1] else None
    ok = abs(r[c0] - v) < 1e-6 and r[c0 + 2] == exp[0] and (r[c0 + 5] is None if sc is None else abs(r[c0 + 5] - sc) < 0.006)
    want.append((ok, f'{tid} {ch} Jupiter → {ch} {X} {pl} {v:.6f} {exp[0]} {"" if sc is None else f"({sc:.2f})"}'))
check('T42 by hand from POSITIONS (rebuilt Ascot P01): ' + '; '.join(w for _, w in want) + ' - a former half hit and a former half nearest target, both now '
      'the nearest ninth (no score), a conjunction that stays conjunction, a golden and a √2 match (old score), a silver outside the 4-digit match (41.379591: 4137 '
      'vs 4142; was 63.73)', all(o for o, _ in want), str([w for o, w in want if not o]))
# (d) retired 29 Oct 2026 (Eddie, half removed): compare_half.py, builds by 83850b1c vs b19fac1f (old families -> half + 4-digit rule), both
# races, default and --full; expected hits per family incl. 'half 0 -> 66' (Ascot T24 SKY) and 'half 0 -> 14' (Lingfield T1). The half part
# no longer holds; the 4-digit rule stays checked in (b), (c) and test 43
# ---------------------------------------------------------------------------------------------------------------------------------
# Test 43 (Eddie, 24 Oct 2026): a value inside its family's old tolerance that fails the 4-digit match reads '<target> <family> — not a
# 4-digit match (<offset>)' (T24 comments: 'nearest: ' + that); beyond the old tolerance '— outside ±<tol>' stays. No score, no fill on either.
# "Before" = builds by the code of commit 10b18003 (half + 4-digit rule, the old wording).
#@ full: T43 comparison 10b18003 -> b19fac1f builds (old against old)
print('test 43: "not a 4-digit match" wording')
out43 = {}
# 27 Oct 2026: "after" = builds by b19fac1f's code on today's families (as h10): the whole-on-target fix (test 45) in both, no SKY STARS tables (test 44)
bf = 'tools/_jt_before_b19fac1f.py'
open(bf, 'w').write(old_code('b19fac1f'))
os.makedirs(f'{tmp}/hf43', exist_ok=True)
for n in ('lingfield', 'ascot'):
    inp = 'reference/classic/race_batch_20210402_lingfield_aw_1435.xlsx' if n == 'lingfield' else f'{tmp}/race_batch_20120619_ascot_1430.xlsx'
    for f in ('full', 'def'):
        run_build([bf, inp] + (['--full'] if f == 'full' else []) + ['--out', f'{tmp}/hf43/{n}_{f}_jupiter_tables.xlsx'],
                       capture_output=True, text=True)
os.remove(bf)
for n in ('lingfield', 'ascot'):
    for f in ('full', 'def'):
        bo = f'{tmp}/h10/{n}_{f}_jupiter_tables.xlsx'
        out43[(n, f)] = subprocess.run([sys.executable, 'tools/scans/families_audit/chk_half/compare_wording.py', f'{tmp}/h85/tools/families.py', bo,
                                        f'{tmp}/hf43/{n}_{f}_jupiter_tables.xlsx'], capture_output=True, text=True).stdout
check('T43 against the build before (commit 10b18003), both races, default and --full, with their pairs workbooks (chk_half/compare_wording.py): '
      'only nearest-target texts and comments differ, each from "— outside ..." to "<target> <family> — not a 4-digit match (<offset>)" with the '
      'same target, family and offset, its value inside the family\'s old tolerance (old score > 0) and scoring nothing now; values, fills, '
      'borders, formats and every other text identical', all(line(o, 'mismatches:') == 'mismatches: 0' for o in out43.values())
      and all(int(line(o, 'changed cells:').split()[2]) > 0 for o in out43.values()),
      f"{ {k: (line(o, 'changed cells:'), line(o, 'mismatches:')) for k, o in out43.items()} }")
#@ quick: T43 by-hand wording on the current build
# by hand from POSITIONS (rebuilt Ascot P01): one of each wording
w43 = openpyxl.load_workbook(f'{tmp}/after40/ascot_full_jupiter_tables.xlsx')
pp = posmap(w43, 'P01')
sk = {x[3]: (x[5], x[6]) for x in w43['POSITIONS'].iter_rows(min_row=5, values_only=True) if x and x[0] == 'SKY' and x[3]}
r43 = {(r[0], r[1]): r for r in w43['P01'].iter_rows(values_only=True) if r and r[0] in ('T1', 'T2')}
v1 = dpl(pp[('J', 'Jupiter')], pp[('J', 'Hamal')], 'Dec')
t1_ = f'29×φ golden — not a 4-digit match ({v1 - 29 * F.PHI:+.5f})'
x1 = r43[('T2', 'J Jupiter → J Hamal')]
v2 = dpl(sk['Jupiter'], sk['Pollux'], 'RA')
t2_ = f'35×φ (golden) — outside ±0.0324'
x2 = r43[('T1', 'H Jupiter → H Pollux')]
check(f'T43 by hand from POSITIONS: T2 J Jupiter → J Hamal Dec {v1:.6f}, 29×φ {29 * F.PHI:.6f}, |offset| {abs(v1 - 29 * F.PHI):.5f} < 0.03236 '
      f'but {F.digits4(v1)[1]} vs {F.digits4(29 * F.PHI)[1]} -> {t1_!r}; T1 H Pollux row SKY S Jupiter → S Pollux RA {v2:.6f}, 35×φ '
      f'{35 * F.PHI:.6f}, |offset| {abs(v2 - 35 * F.PHI):.5f} > 0.03236 -> {t2_!r}; neither scored nor filled',
      x1[12] == t1_ and x1[15] is None and abs(v1 - 29 * F.PHI) < 0.03236 and not F.digit_match(v1, 29 * F.PHI)
      and x2[20] == t2_ and x2[23] is None and abs(v2 - 35 * F.PHI) > 0.03236, f'{x1[12]!r} {x1[15]}; {x2[20]!r} {x2[23]}')
# ---------------------------------------------------------------------------------------------------------------------------------
# Test 44 (Eddie, 27 Oct 2026): T24 HORSE / JOCKEY (SKY STARS) on every runner sheet (stars whose S Jupiter → S star SKY block cell scores,
# the T24 row copied, sky hit, SCORES) and the SKY STARS SUMMARY sheet. "Before" = commit b19fac1f. Kelso 14:50 2022, rebuilt Ascot 14:30 2012,
# Lingfield 14:35 2021, default and --full (chk_t24ss/check_t24ss.py).
#@ full: T44 builds by f1327932 (old), inputs of the T45 comparison
print('test 44: builds by f1327932\'s code for test 45 (the SKY STARS checks retired)')
# retired 29 Oct 2026 (Eddie, workbook trim): T44 against the build before (commit b19fac1f, chk_t24ss/check_t24ss.py; selection, copied T24
# cells, sky hit, SCORES, headers, summary, pairs identical); Kelso selects Eddie's 14 stars; by hand from POSITIONS (Kelso P01) H Rukbat sky hit,
# SCORES H Rukbat – H Gonggong RA / H Sun Dec, conjunction H Antares – H Mars RA red - T24 HORSE / JOCKEY (SKY STARS), the SKY STARS SUMMARY
# and the pairs workbook are gone (the Kelso build by b19fac1f's code removed). The builds by f1327932's code stay: test 45 compares with them.
os.makedirs(f'{tmp}/after44', exist_ok=True)
os.makedirs(f'{tmp}/hf13', exist_ok=True)
extract(f'f1327932', f'{tmp}/hf13')
KELSO = 'reference/classic/race_batch_20220504_kelso_1450.xlsx'
INP44 = {'kelso': KELSO, 'lingfield': 'reference/classic/race_batch_20210402_lingfield_aw_1435.xlsx', 'ascot': f'{tmp}/race_batch_20120619_ascot_1430.xlsx'}
for n, inp in INP44.items():
    for f in ('def', 'full'):
        fl = ['--full'] if f == 'full' else []
        run_build([f'{tmp}/hf13/tools/jupiter_tables.py', inp] + fl + ['--out', f'{tmp}/after44/{n}_{f}_jupiter_tables.xlsx'],
                       capture_output=True, text=True)
# ---------------------------------------------------------------------------------------------------------------------------------
# Test 45 (Eddie, 27 Oct 2026): whole numbers checked on the TARGET (families.score_whole_t: t = round(v), 1 <= t <= 200, |v − t| < 0.02,
# squared); active_vibrations.score_whole needed 1 <= v, so 0.98 < v < 1 never scored on 1. "Before" = commit f1327932 (builds in hf13).
#@ quick: T45 whole-on-target family rule by hand (current families)
print('test 45: whole numbers on the target')
import active_vibrations as AV
cases = [(0.997954, ('1', 80.59)), (0.98, None), (0.981, ('1', 0.25)), (1.002, ('1', 81.00)), (199.99, ('200', 25.00)), (200.01, ('200', 25.00))]
got45 = []
for v, want in cases:
    s_, lab, t = F.score_family(v, 'whole')
    got45.append((v, (lab, round(s_, 2)) if s_ else None))
check('T45 whole on the target by hand (10 (0.02 − dev) / 0.02)²: 0.997954 -> 1 (80.59), 0.98 -> none (dev 0.02), 0.981 -> 1 (0.25), 1.002 -> 1 (81.00) '
      f'= active_vibrations.score_whole ({AV.score_whole(1.002):.2f}), 199.99 / 200.01 -> 200 (25.00); best of 0.997954 Dec = whole 1',
      all(g == (v, w) for g, (v, w) in zip(got45, cases)) and abs(AV.score_whole(1.002) - F.score_family(1.002, 'whole')[0]) < 1e-9
      and F.best(0.997954, 'Dec')['family'] == 'whole', got45)
#@ full: T45 comparison f1327932 -> 35255b16 and the Meissa cell on the 35255b16 build (old)
L67 = f'{tmp}/l67/race_batch_20210402_lingfield_aw_1435.xlsx'          # Eddie's 67-star rebuild of Lingfield 14:35 2021 (27 Oct 2026)
os.makedirs(f'{tmp}/l67', exist_ok=True); os.makedirs(f'{tmp}/after45', exist_ok=True)
shutil.copy(f'{AP}/racecards/20210402_lingfield_aw/race_batch_20210402_lingfield_aw_1435.xlsx', L67)
INP45 = dict(INP44, l67=L67)
# 27 Oct 2026 (S Jupiter column + Jupiter rows, test 46): the "after" builds are by 35255b16's code, so test 45 checks its own change alone
os.makedirs(f'{tmp}/hf35', exist_ok=True)
extract(f'35255b16', f'{tmp}/hf35')
# rows changed / dropped / added / with a 0.98-1 value / follow-on (counts, default and --full)
EXP45 = {('kelso', 'def'): (125, 0, 2, 125, 2), ('kelso', 'full'): (127, 0, 3, 127, 3), ('lingfield', 'def'): (20, 0, 0, 20, 0),
         ('lingfield', 'full'): (20, 0, 0, 20, 0), ('ascot', 'def'): (10, 0, 0, 10, 0), ('ascot', 'full'): (10, 0, 0, 10, 0),
         ('l67', 'def'): (31, 0, 0, 31, 0), ('l67', 'full'): (31, 0, 0, 31, 0)}
out45 = {}
for n, inp in INP45.items():
    for f in ('def', 'full'):
        fl = ['--full'] if f == 'full' else []
        bo, ao = f'{tmp}/after44/{n}_{f}_jupiter_tables.xlsx', f'{tmp}/after45/{n}_{f}_jupiter_tables.xlsx'
        if n == 'l67':
            run_build([f'{tmp}/hf13/tools/jupiter_tables.py', inp] + fl + ['--out', bo], capture_output=True, text=True)
        run_build([f'{tmp}/hf35/tools/jupiter_tables.py', inp] + fl + ['--out', ao], capture_output=True, text=True)
        o = subprocess.run([sys.executable, 'tools/scans/families_audit/chk_whole/diff_whole.py', bo, ao], capture_output=True, text=True).stdout
        out45[(n, f)] = (o, tuple(int(x) for x in re.findall(r': (\d+)', line(o, 'rows changed:'))))
fol = [l for o, _ in out45.values() for l in o.splitlines() if '[FOLLOWS]' in l]
fol_ok = all(re.search(r"\['T24', 'T24 JOCKEY \(natal\)|T19 NO RATIO OR MIDPOINT HIT|THREADS row", l) for l in fol)
check('T45 against the build before (commit f1327932), Kelso, Lingfield (repo batch and Eddie\'s 67-star rebuild), rebuilt Ascot, default and --full, '
      'with the pairs workbooks (chk_whole/diff_whole.py): every changed, added or dropped row holds a value 0.98 < v < 1, except the follow-ons '
      '(Kelso: P01 T24 JOCKEY title repeat count, P03 T19 "no ratio or midpoint hit" count, THREADS P03 Eris "+ 1 weaker"); counts as expected',
      all(out45[k][1] == EXP45[k] for k in EXP45) and fol_ok, f'{ {k: v[1] for k, v in out45.items()} }; follow-ons {[l[:60] for l in fol]}')
# by hand from POSITIONS: Eddie's case on the 67-star Lingfield, P09 T24 HORSE H Meissa – H Transpluto Dec
w45 = openpyxl.load_workbook(f'{tmp}/after45/l67_def_jupiter_tables.xlsx')
p45 = posmap(w45, 'P09')
v45 = dpl(p45[('H', 'Meissa')], p45[('H', 'Transpluto')], 'Dec')
r45 = next(r for r in w45['P09'].iter_rows() if r[0].value == 'T24' and r[1].value == 'H Meissa')
c45 = next(c for c in r45 if c.value == round(v45, 6))
check(f'T45 by hand from POSITIONS (Lingfield 67-star P09): T24 HORSE H Meissa – H Transpluto Dec {v45:.6f} -> 1 · whole · {v45 - 1:+.5f} · '
      f'{sqs(0.02, abs(v45 - 1)):.2f}, filled green (was "nearest: 1 · whole — outside ±0.02")',
      c45.comment.text.split('\n')[0] == f'1 · whole · {v45 - 1:+.5f} · {sqs(0.02, abs(v45 - 1)):.2f}' and c45.fill.start_color.rgb.endswith('C6EFCE'),
      c45.comment.text)
# ---------------------------------------------------------------------------------------------------------------------------------
# Test 46 (Eddie, 27 Oct 2026): T24 HORSE / JOCKEY "S Jupiter" column (d(S Jupiter, S star) = the SKY block cell, after Ketu in RA and Dec)
# and the four Jupiter rows on top of T24 HORSE / JOCKEY (SKY STARS) (S Jupiter → S body, natal Jupiter → natal body, S Jupiter → natal
# body, S body → natal Jupiter), carried into the summary. "Before" = commit 35255b16 (builds in after45). Kelso, Lingfield, rebuilt Ascot,
# default and --full (chk_sjcol/check_sjcol.py).
#@ quick: T46 T24 HORSE five rows and S Jupiter column on the current Kelso build
print('test 46: T24 S Jupiter column and the Jupiter rows of T24 HORSE (current code, Kelso)')
KELSO = 'reference/classic/race_batch_20220504_kelso_1450.xlsx'
# retired 29 Oct 2026 (Eddie, workbook trim): T46 against the build before (commit 35255b16, chk_sjcol/check_sjcol.py: S Jupiter column = SKY block
# cells; the four Jupiter rows of the (SKY STARS) tables = T1 / T2 SKY, T1 / T2, T5 - T8, T9; SCORES; summary; pairs identical) - the (SKY STARS)
# tables, the summary and the pairs workbook are gone, and check_sjcol.py imports T24SS_JNOTE from tools/jupiter_tables.py (gone). The a2d20618
# builds (after46) are removed.
# 29 Oct 2026 (Eddie, workbook trim), ADAPTED: the by-hand cases below now read a build by the current code (Kelso, default), the S Jupiter
# column and the Jupiter rows of T24 HORSE itself (rows 'S Jupiter → S body', 'H Jupiter → H body', 'S Jupiter → H body', 'S body → H Jupiter',
# then 'stars', between the header rows and the star rows) - before, they read the same rows on T24 HORSE (SKY STARS) of the a2d20618 build
w46f = f'{tmp}/kelso_now_jupiter_tables.xlsx'
subprocess.run([sys.executable, 'tools/jupiter_tables.py', KELSO, '--out', w46f], capture_output=True, text=True)
# by hand from POSITIONS (Kelso P01)
w46 = openpyxl.load_workbook(w46f)
p46 = posmap(w46, 'P01')
s46 = {x[3]: (x[5], x[6]) for x in w46['POSITIONS'].iter_rows(min_row=5, values_only=True) if x and x[0] == 'SKY' and x[3]}
ws46 = w46['P01']
t46 = next(r for r in range(1, ws46.max_row + 1) if ws46.cell(row=r, column=1).value == 'T24' and str(ws46.cell(row=r, column=2).value).startswith('T24 HORSE (natal)'))
h46 = [ws46.cell(row=t46 + 2, column=c).value for c in range(1, ws46.max_column + 1)]
sjc = {pl: [c for c in range(1, len(h46) + 1) if h46[c - 1] == 'S Jupiter'][k] for k, pl in enumerate(('RA', 'Dec'))}
star46 = lambda x: next(r for r in range(t46 + 3, t46 + 80) if ws46.cell(row=r, column=2).value == f'H {x}')
vcol = lambda pl, b: [c for c in range(1, len(h46) + 1) if h46[c - 1] == b][0 if pl == 'RA' else 1]
jr46 = {ws46.cell(row=r, column=2).value: r for r in range(t46 + 3, t46 + 8)}
check('T46 T24 HORSE (current code, Kelso P01): the five rows after the header rows are S Jupiter → S body, H Jupiter → H body, S Jupiter → H body, '
      'S body → H Jupiter, stars (29 Oct 2026); the first star row follows',
      [ws46.cell(row=t46 + 3 + k, column=2).value for k in range(5)] == ['S Jupiter → S body', 'H Jupiter → H body', 'S Jupiter → H body',
                                                                       'S body → H Jupiter', 'stars']
      and str(ws46.cell(row=t46 + 8, column=2).value).startswith('H '), [ws46.cell(row=t46 + 3 + k, column=2).value for k in range(6)])
cases46 = [  # (what, cell, value by hand, plane, comment)
    ('T24 HORSE S Jupiter column H Aldebaran RA = d(S Jupiter, S Aldebaran)', ws46.cell(row=star46('Aldebaran'), column=sjc['RA']),
     dpl(s46['Jupiter'], s46['Aldebaran'], 'RA'), 'RA', lambda v: f'70 · whole · {v - 70:+.5f} · {sqs(0.02, abs(v - 70)):.2f}'),
    ('T24 HORSE S Jupiter column H Fomalhaut Dec = d(S Jupiter, S Fomalhaut)', ws46.cell(row=star46('Fomalhaut'), column=sjc['Dec']),
     dpl(s46['Jupiter'], s46['Fomalhaut'], 'Dec'), 'Dec', lambda v: f'28 · whole · {v - 28:+.5f} · {sqs(0.02, abs(v - 28)):.2f}'),
    ('S Jupiter → S body, Venus Dec', ws46.cell(row=jr46['S Jupiter → S body'], column=vcol('Dec', 'Venus')),
     dpl(s46['Jupiter'], s46['Venus'], 'Dec'), 'Dec', lambda v: f'1 · whole · {v - 1:+.5f} · {sqs(0.02, abs(v - 1)):.2f}'),
    ('H Jupiter → H body, Uranus RA', ws46.cell(row=jr46['H Jupiter → H body'], column=vcol('RA', 'Uranus')),
     dpl(p46[('H', 'Jupiter')], p46[('H', 'Uranus')], 'RA'), 'RA', lambda v: f'148 · whole · {v - 148:+.5f} · {sqs(0.02, abs(v - 148)):.2f}'),
    ('H Jupiter → H body, Uranus Dec (conjunction, red)', ws46.cell(row=jr46['H Jupiter → H body'], column=vcol('Dec', 'Uranus')),
     dpl(p46[('H', 'Jupiter')], p46[('H', 'Uranus')], 'Dec'), 'Dec', lambda v: f'0 · conjunction (Dec < 0.3) · {v:+.5f} · {sqs(0.3, v):.2f}'),
    ('S Jupiter → H body, Gonggong RA', ws46.cell(row=jr46['S Jupiter → H body'], column=vcol('RA', 'Gonggong')),
     dpl(s46['Jupiter'], p46[('H', 'Gonggong')], 'RA'), 'RA', lambda v: f'23 · whole · {v - 23:+.5f} · {sqs(0.02, abs(v - 23)):.2f}'),
    ('S Jupiter → H body, Venus RA (conjunction, red)', ws46.cell(row=jr46['S Jupiter → H body'], column=vcol('RA', 'Venus')),
     dpl(s46['Jupiter'], p46[('H', 'Venus')], 'RA'), 'RA', lambda v: f'0 · conjunction (RA < 1.5) · {v:+.5f} · {sqs(1.5, v):.2f}'),
    ('S body → H Jupiter, Rahu Dec', ws46.cell(row=jr46['S body → H Jupiter'], column=vcol('Dec', 'Rahu')),
     dpl(s46['Rahu'], p46[('H', 'Jupiter')], 'Dec'), 'Dec', lambda v: f'11 · whole · {v - 11:+.5f} · {sqs(0.02, abs(v - 11)):.2f}'),
]
for what, c, v, pl, txt in cases46:
    red = v < (1.5 if pl == 'RA' else 0.3)
    ok = c.value == round(v, 6) and c.comment.text.split('\n')[0] == txt(v) and (c.fill.start_color.rgb.endswith('FF8080') == red)
    check(f'T46 by hand from POSITIONS (Kelso P01): {what} {v:.6f} -> {txt(v)!r}' + (' red' if red else ''), ok,
          f'{c.value} {c.comment.text if c.comment else None!r} {c.fill.start_color.rgb}')
jj = ws46.cell(row=jr46['S body → H Jupiter'], column=vcol('RA', 'Jupiter'))
check('T46 Jupiter column: blank on S Jupiter → S body and H Jupiter → H body; d(S Jupiter, H Jupiter) on the other two (by hand '
      f'{dpl(s46["Jupiter"], p46[("H", "Jupiter")], "RA"):.6f} RA)',
      ws46.cell(row=jr46['S Jupiter → S body'], column=vcol('RA', 'Jupiter')).value is None
      and ws46.cell(row=jr46['H Jupiter → H body'], column=vcol('RA', 'Jupiter')).value is None
      and jj.value == ws46.cell(row=jr46['S Jupiter → H body'], column=vcol('RA', 'Jupiter')).value == round(dpl(s46['Jupiter'], p46[('H', 'Jupiter')], 'RA'), 6),
      jj.value)
# ---------------------------------------------------------------------------------------------------------------------------------
# retired 29 Oct 2026 (Eddie, workbook trim): Test 47 (commit a2d20618 -> d2d4651b, chk_display/check_display.py): SKY STARS SUMMARY rows 1-2 =
# the header rows, freeze C3, every 'stars' row of the (SKY STARS) tables grey A to the last SCORES column, no column narrower, pairs identical;
# Kelso summary header / freeze / runner AA width 40 - the summary, the (SKY STARS) tables and the pairs workbook are gone (d2d4651b builds removed)
# retired 29 Oct 2026 (Eddie, workbook trim): Test 48 (commit d2d4651b -> now, chk_skyhit/check_skyhit.py): "sky hit" moved to the end of the
# (SKY STARS) tables and the summary, exact widths (values / S Jupiter / SCORES 11, midpoints 40, repeats 60, sky hit 36; runner AA, AZ, BA 11),
# pairs identical; Kelso P01 (SKY STARS) header C 'Sun', H Rukbat sky hit in DB, widths - those tables are gone (the widths they set with them)
# ---------------------------------------------------------------------------------------------------------------------------------
# Test 49 (Eddie, 29 Oct 2026): workbook trim - runner sheets keep T1-T8, T24 HORSE, T24 JOCKEY only; SKY STARS SUMMARY, every --full-only
# sheet and the _pairs workbook gone (--full accepted, = default); POSITIONS and SKY JUPITER unchanged; the four Jupiter rows of the old
# (SKY STARS) tables + a grey 'stars' row now on top of T24 HORSE / JOCKEY, the Jupiter cell = d(S Jupiter, natal Jupiter) computed directly
# (= the old T9 value). "Before" = builds by acf68341's own tools/, "after" by 2cc2af2b's (30 Oct 2026). Kelso, Lingfield, rebuilt Ascot, default and --full
# (chk_trim/check_trim.py).
#@ full: T49 - T53 comparisons acf68341 -> 2cc2af2b -> 9e5d5ea5 -> 5ea09b4e -> 93b54a43 -> 3da4ae61 and their by-hand cells on those builds (old against old)
print('test 49: workbook trim (Eddie, 29 Oct 2026)')
INP49 = {'kelso': 'reference/classic/race_batch_20220504_kelso_1450.xlsx', 'lingfield': 'reference/classic/race_batch_20210402_lingfield_aw_1435.xlsx',
         'ascot': f'{tmp}/race_batch_20120619_ascot_1430.xlsx'}
# 30 Oct 2026 (MATCH tables, test 51): before = acf68341's own tools/, after = 2cc2af2b's own tools/ (the trim alone, both with their families)
for c_, d_ in (('acf68341', 'hacf'), ('2cc2af2b', 'h2cc'), ('9e5d5ea5', 'h9e5')):
    os.makedirs(f'{tmp}/{d_}', exist_ok=True)
    extract(f'{c_}', f'{tmp}/{d_}')
bf = f'{tmp}/hacf/tools/jupiter_tables.py'
os.makedirs(f'{tmp}/before49', exist_ok=True)
os.makedirs(f'{tmp}/after49', exist_ok=True)
out49 = {}
for n, inp in INP49.items():
    for f in ('def', 'full'):
        fl = ['--full'] if f == 'full' else []
        bo, ao = f'{tmp}/before49/{n}_{f}_jupiter_tables.xlsx', f'{tmp}/after49/{n}_{f}_jupiter_tables.xlsx'
        rb = run_build([bf, inp] + fl + ['--out', bo], capture_output=True, text=True)
        ra = run_build([f'{tmp}/h2cc/tools/jupiter_tables.py', inp] + fl + ['--out', ao], capture_output=True, text=True)
        out49[(n, f, 'size')] = (re.findall(r'\(([\d.]+ MB)\)', rb.stdout), re.findall(r'\(([\d.]+ MB)\)', ra.stdout))
        out49[(n, f, 'pairs')] = os.path.exists(bo.replace('_jupiter_tables', '_pairs')) and not os.path.exists(ao.replace('_jupiter_tables', '_pairs'))
        out49[(n, f)] = subprocess.run([sys.executable, 'tools/scans/families_audit/chk_trim/check_trim.py', bo, ao], capture_output=True, text=True).stdout
z49 = lambda p: subprocess.run(f'unzip -p {p} "xl/worksheets/*.xml" "xl/comments/*.xml" xl/styles.xml xl/workbook.xml | shasum', shell=True, capture_output=True,
                               text=True).stdout
same49 = {n: z49(f'{tmp}/after49/{n}_def_jupiter_tables.xlsx') == z49(f'{tmp}/after49/{n}_full_jupiter_tables.xlsx') for n in INP49}
print('T49 sizes (before tables + pairs, after tables):', {k[:2]: v for k, v in out49.items() if k[2:] == ('size',)})
check('T49 against the build before (commit acf68341), Kelso, Lingfield, rebuilt Ascot, default and --full (chk_trim/check_trim.py): sheets = POSITIONS, '
      'SKY JUPITER, P##; POSITIONS and SKY JUPITER identical cell for cell (values, fills, fonts, borders, number formats, comments, merges, widths, '
      'freeze); runner sheets: rows 1-3, T1-T8, T24 HORSE, T24 JOCKEY identical after the row shift, every other row empty, widths of every used '
      'column and freeze C4 unchanged; the five new rows = the old (SKY STARS) Jupiter rows from column B to the repeats column, Jupiter cell = the '
      'old T9 value, "stars" grey A to the repeats column; no removed table anywhere; the old build wrote a pairs workbook, the new one none',
      all(line(o, 'mismatches:') == 'mismatches: 0' for k, o in out49.items() if len(k) == 2) and all(o for k, o in out49.items() if k[2:] == ('pairs',)),
      f"{ {k: (line(o, 'cells compared:'), line(o, 'mismatches:')) for k, o in out49.items() if len(k) == 2} }")
check('T49 --full is accepted and writes the same workbook as the default (sheet, comment, style and workbook XML identical), three races', all(same49.values()), same49)
w49 = openpyxl.load_workbook(f'{tmp}/after49/kelso_def_jupiter_tables.xlsx')
p49 = w49['P01']
h49 = next(r for r in range(1, p49.max_row + 1) if p49.cell(row=r, column=1).value == 'T24' and str(p49.cell(row=r, column=2).value).startswith('T24 HORSE'))
lab49 = [p49.cell(row=h49 + 3 + i, column=2).value for i in range(6)]
jra = next(k for k in range(3, 40) if p49.cell(row=h49 + 2, column=k).value == 'Jupiter')
pos49 = {(x[0], x[2], x[3]): (x[5], x[6]) for x in w49['POSITIONS'].iter_rows(min_row=5, values_only=True) if x and x[3]}
sj, hj = pos49[('SKY', 'S', 'Jupiter')], pos49[('P01', 'H', 'Jupiter')]
d49 = abs((sj[0] - hj[0] + 180) % 360 - 180)
check('T49 Kelso P01 T24 HORSE: the five rows straight after the headers (S Jupiter → S body, H Jupiter → H body, S Jupiter → H body, S body → H '
      f'Jupiter, stars), then H Alpheratz; Jupiter RA cell on rows 3-4 = d(S Jupiter, H Jupiter) from POSITIONS {d49:.6f} (CLAUDE.md: 167.634941), '
      'blank on rows 1-2; no other sheet than POSITIONS, SKY JUPITER, P##',
      lab49 == ['S Jupiter → S body', 'H Jupiter → H body', 'S Jupiter → H body', 'S body → H Jupiter', 'stars', 'H Alpheratz']
      and p49.cell(row=h49 + 5, column=jra).value == p49.cell(row=h49 + 6, column=jra).value == round(d49, 6) == 167.634941
      and p49.cell(row=h49 + 3, column=jra).value is None and p49.cell(row=h49 + 4, column=jra).value is None
      and w49.sheetnames == ['POSITIONS', 'SKY JUPITER', 'P01', 'P03', 'P05', 'P07', 'P09', 'P11'],
      (lab49, p49.cell(row=h49 + 5, column=jra).value, w49.sheetnames))
# ---------------------------------------------------------------------------------------------------------------------------------
# Test 50 (Eddie, 29 Oct 2026): the half family (0.5 ... 199.5, ±0.01) removed from families.py - no half score, never a half nearest target.
# "Before" = builds by the code of 2cc2af2b (its own tools/ and rebuild_kit/, i.e. families with half). Kelso, Lingfield, rebuilt Ascot,
# default only (= --full) (chk_nohalf/check_nohalf.py).
print('test 50: half family removed')
INP50 = {'kelso': 'reference/classic/race_batch_20220504_kelso_1450.xlsx', 'lingfield': 'reference/classic/race_batch_20210402_lingfield_aw_1435.xlsx',
         'ascot': f'{tmp}/race_batch_20120619_ascot_1430.xlsx'}
os.makedirs(f'{tmp}/after50', exist_ok=True)
out50 = {}
for n, inp in INP50.items():
    # 30 Oct 2026: before = the 2cc2af2b build of test 49, after = 9e5d5ea5's own tools/ (the half removal alone)
    bo, ao = f'{tmp}/after49/{n}_def_jupiter_tables.xlsx', f'{tmp}/after50/{n}_jupiter_tables.xlsx'
    run_build([f'{tmp}/h9e5/tools/jupiter_tables.py', inp, '--out', ao], capture_output=True, text=True)
    out50[n] = subprocess.run([sys.executable, 'tools/scans/families_audit/chk_nohalf/check_nohalf.py', bo, ao], capture_output=True, text=True).stdout
check('T50 against the build before (commit 2cc2af2b, families with half), Kelso, Lingfield, rebuilt Ascot (chk_nohalf/check_nohalf.py): every '
      'changed cell is a cell that scored or was named as half before, or in its value group (target, offset, score, arrow, also, timing), or a '
      'T24 repeat on a half target; dropped rows (SKY STARS / PAIRS rows listed for a half leg, timeline events on half) named half; otherwise '
      'only title counts and timeline follow-ons; no half family text anywhere after (the midpoint "ratio half" / "distance half" are not the family)',
      all(line(o, 'mismatches:') == 'mismatches: 0' for o in out50.values()), {n: line(o, 'mismatches:') for n, o in out50.items()})
# by hand from POSITIONS
w50 = {n: openpyxl.load_workbook(f'{tmp}/after50/{n}_jupiter_tables.xlsx') for n in ('kelso', 'lingfield')}
pk, pl_ = posmap(w50['kelso'], 'P03'), posmap(w50['lingfield'], 'P11')
skk = {x[3]: (x[5], x[6]) for x in w50['kelso']['POSITIONS'].iter_rows(min_row=5, values_only=True) if x and x[0] == 'SKY' and x[3]}
k3 = w50['kelso']['P03']
t24j = next(r for r in range(1, k3.max_row + 1) if k3.cell(row=r, column=1).value == 'T24' and str(k3.cell(row=r, column=2).value).startswith('T24 JOCKEY'))
prow = next(r for r in range(t24j, t24j + 80) if k3.cell(row=r, column=2).value == 'J Pollux')
tcol = next(k for k in range(3, 30) if k3.cell(row=t24j + 2, column=k).value == 'Transpluto')
c1 = k3.cell(row=prow, column=tcol)
v1 = dpl(pk[('J', 'Pollux')], pk[('J', 'Transpluto')], 'RA')
s1 = (10 * (0.03236 - abs(v1 - 17 * F.PHI)) / 0.03236) ** 2
r3 = next(r for r in w50['lingfield']['P11'].iter_rows(values_only=True) if r[0] == 'T3' and r[1] == 'H Jupiter → J Ketu')
v2 = dpl(pl_[('H', 'Jupiter')], pl_[('J', 'Ketu')], 'RA')
s2 = (10 * (1.5 - v2) / 1.5) ** 2
rg = next(r for r in w50['kelso']['SKY JUPITER'].iter_rows(values_only=True) if r[0] == 'SKY' and r[1] == 'S Jupiter → S Gonggong')
v3 = dpl(skk['Jupiter'], skk['Gonggong'], 'RA')
check(f'T50 by hand from POSITIONS: lost half hit - Kelso P03 T24 JOCKEY J Pollux – J Transpluto RA {v1:.6f} (was 27.5 half 91.23) -> 17×φ '
      f'{17 * F.PHI:.6f} golden ({s1:.2f}), yellow; conjunction back - Lingfield P11 T3 H Jupiter → J Ketu RA {v2:.6f} (was 1.5 half 55.30) -> '
      f'conj ({s2:.2f}); nearest off half - Kelso SKY S Jupiter → S Gonggong RA {v3:.6f} (was 21.5 half outside) -> 22.2222 ninth outside ±0.0722',
      abs(c1.value - v1) < 1e-6 and c1.comment.text.split('\n')[0] == f'17×φ · golden · {v1 - 17 * F.PHI:+.5f} · {s1:.2f}'
      and c1.fill.start_color.rgb == YELLOW
      and abs(r3[2] - v2) < 1e-6 and r3[4] == 'conj (conj)' and abs(r3[7] - s2) < 0.006
      and abs(rg[2] - v3) < 1e-6 and rg[4] == '22.2222 (ninth) — outside ±0.0722' and rg[7] is None,
      (c1.value, c1.comment.text.split('\n')[0], r3[2:8], rg[2:8]))
# ---------------------------------------------------------------------------------------------------------------------------------
# Test 51 (Eddie, 30 Oct 2026): MATCH HORSE / MATCH JOCKEY on every runner sheet after T24 JOCKEY - SKY block hits against T24 HORSE / JOCKEY
# hits: EXACT (same family + target, or across families the same first 4 significant digits) and RATIO. RATIO tightened the same day (Eddie,
# 30 Oct 2026): whole only when the circle factor gcd(A, B, 360) >= 5 (+ 'circle factor', 'slice of 360' columns), golden only when p, q <= 6
# (was gcd >= 2 for both). "Before" = builds by a76a3e99's own tools/ (MATCH with the first rule); before that commit's tables the sheets are
# also checked against 9e5d5ea5 through it. Kelso, Lingfield, rebuilt Ascot (chk_match/check_match.py).
# Grouped by sky cell (Eddie, 30 Oct 2026): one list per table under grey sky-cell header rows; before = 8f42f04e's own tools/ (same rows,
# only positions, the header rows and the title change); Ripon 15:25 16 Jun 2022 added (batch copied from ~/astronomy-project/racecards).
# Side by side (Eddie, 30 Oct 2026, test 52): the "after" of test 51 = 5ea09b4e's own tools/ (MATCH HORSE | JOCKEY is test 52's change).
print('test 51: MATCH HORSE / MATCH JOCKEY')
os.makedirs(f'{tmp}/after51', exist_ok=True)
os.makedirs(f'{tmp}/h5ea', exist_ok=True)
extract(f'5ea09b4e', f'{tmp}/h5ea')
os.makedirs(f'{tmp}/h8f4', exist_ok=True)       # 30 Oct 2026 (grouped by sky cell): before = 8f42f04e (sky-score order, two sections)
extract(f'8f42f04e', f'{tmp}/h8f4')
INP51 = dict(INP50, ripon='reference/classic/race_batch_20220616_ripon_1525.xlsx')
out51, out51b = {}, {}
for n, inp in INP51.items():
    bo, ao = f'{tmp}/h8f4/{n}_jupiter_tables.xlsx', f'{tmp}/after51/{n}_jupiter_tables.xlsx'
    run_build([f'{tmp}/h8f4/tools/jupiter_tables.py', inp, '--out', bo], capture_output=True, text=True)
    run_build([f'{tmp}/h5ea/tools/jupiter_tables.py', inp, '--out', ao], capture_output=True, text=True)
    out51[n] = subprocess.run([sys.executable, 'tools/scans/families_audit/chk_match/check_match.py', bo, ao], capture_output=True, text=True).stdout
    if n in INP50:
        out51b[n] = subprocess.run([sys.executable, 'tools/scans/families_audit/chk_match/check_match.py', f'{tmp}/after50/{n}_jupiter_tables.xlsx', ao],
                                   capture_output=True, text=True).stdout
cnt51 = {n: [l for l in o.splitlines() if 'MATCH' in l and ': SKY ' in l] for n, o in out51.items()}
print('T51 counts:', cnt51)
check('T51 against the builds before (commits 8f42f04e, two sections in the sky-score order: the same rows cell for cell, only positions, '
      'header rows and title differ; and 9e5d5ea5, no MATCH), Kelso, Lingfield, rebuilt Ascot, Ripon 15:25 2022 (8f42f04e only) '
      '(chk_match/check_match.py): every other cell, merge, width and freeze identical; MATCH HORSE then MATCH JOCKEY after T24 JOCKEY; rebuilt '
      'from the SKY block cells and the T24 cells on the sheet (scoring, no conjunction; T24 star rows x body columns without S Jupiter, Jupiter '
      'rows without S Jupiter → S body) on the tightened rules (whole circle factor gcd(A, B, 360) >= 5 with circle factor and slice of 360; '
      'golden p, q <= 6): every pair, none missing or extra, every cell, score fills, title counts incl. k sky cells; GROUPED (30 Oct 2026): '
      'one list, a grey header row (A to U) above each sky cell with its relationship, plane, value, target, family, score and EXACT / RATIO '
      'counts; groups by the sky score, highest first (ties SKY block order); inside a group EXACT (T24 score highest first, T24 order), then '
      'RATIO by q, p, T24 score, T24 order; no blank rows between groups, two between MATCH HORSE and MATCH JOCKEY',
      all(line(o, 'mismatches:') == 'mismatches: 0' for o in list(out51.values()) + list(out51b.values()))
      and all(len(v) == 2 * {'kelso': 6, 'lingfield': 6, 'ascot': 1, 'ripon': 4}[n] for n, v in cnt51.items())
      and all(sum(int(x.split()[4]) for x in o.splitlines() if 'in a new position' in x) == {'kelso': 1080, 'lingfield': 151, 'ascot': 204,
                                                                                          'ripon': 1419}[n] for n, o in out51.items()),
      {n: (line(o, 'mismatches:'), line(out51b.get(n, ''), 'mismatches:')) for n, o in out51.items()})
# by hand from POSITIONS (Kelso P01 MATCH HORSE)
w51 = openpyxl.load_workbook(f'{tmp}/after51/kelso_jupiter_tables.xlsx', read_only=True)
p51 = posmap(w51, 'P01')
s51 = {x[3]: (x[5], x[6]) for x in w51['POSITIONS'].iter_rows(min_row=5, values_only=True) if x and x[0] == 'SKY' and x[3]}
m51 = [r for r in w51['P01'].iter_rows(values_only=True) if r and r[0] == 'MATCH HORSE' and r[1] in ('EXACT', 'RATIO')]
h51 = [r[1] for r in w51['P01'].iter_rows(values_only=True) if r and r[0] == 'MATCH HORSE' and r[1] not in ('EXACT', 'RATIO')]
mk = lambda r: (r[1], r[7], r[8], r[13], r[14], r[15])     # 30 Oct 2026: two columns after 'shared factor'
M51 = {mk(r): r for r in m51}
sd = lambda x, pl: round(dpl(s51['Jupiter'], s51[x], pl), 6)
hd = lambda x, b, pl: round(dpl(p51[('H', x)], p51[('H', b)], pl), 6)
sq = lambda tol, dev: (10 * (tol - dev) / tol) ** 2
cases = []
v1, v2 = sd('Rukbat', 'Dec'), hd('Hamal', 'Quaoar', 'Dec')                 # exact, same family: 39 whole / 39 whole
r1 = M51.get(('EXACT', 'S Jupiter → S Rukbat', 'Dec', 'H Hamal', 'Quaoar', 'Dec'))
cases.append(('exact whole 39: S Rukbat Dec %.6f (%.2f) / H Hamal – H Quaoar Dec %.6f (%.2f)' % (v1, sq(.02, abs(v1 - 39)), v2, sq(.02, abs(v2 - 39))),
              r1 is not None and r1[9] == v1 and r1[16] == v2 and r1[11] == r1[18] == '39' and abs(r1[20] - min(sq(.02, abs(v1 - 39)), sq(.02, abs(v2 - 39)))) < 0.006))
v1, v2 = sd('Praesepe', 'RA'), hd('Praesepe', 'Venus', 'RA')               # exact across families: 81×φ golden / 131 whole, 4 digits 1310
r2 = M51.get(('EXACT', 'S Jupiter → S Praesepe', 'RA', 'H Praesepe', 'Venus', 'RA'))
cases.append((f'exact across families: S Praesepe RA {v1:.6f} 81×φ golden / H Praesepe – H Venus RA {v2:.6f} 131 whole, digits '
              f'{F.digits4(v1)} = {F.digits4(v2)}', r2 is not None and (r2[10], r2[11], r2[17], r2[18]) == ('golden', '81×φ', 'whole', '131')
              and F.digits4(v1) == F.digits4(v2) and abs(v1 - 81 * F.PHI) < 0.03236 and abs(v2 - 131) < 0.02))
# 30 Oct 2026 (tightened): whole 90 / 70 kept (circle factor 10, 1/36); whole 30 / 28 dropped (circle factor 2; was 15/14 under gcd >= 2);
# golden 93φ / 62φ kept (3/2); golden 88φ / 86φ dropped by the cap (44/43; was kept)
v1, v2 = sd('Aldebaran', 'RA'), hd('Bellatrix', 'Chiron', 'RA')
r3 = M51.get(('RATIO', 'S Jupiter → S Aldebaran', 'RA', 'H Bellatrix', 'Chiron', 'RA'))
cases.append((f'whole kept: S Aldebaran RA {v1:.6f} 70 / H Bellatrix – H Chiron RA {v2:.6f} 90 -> 9/7, factor 10, circle factor gcd(70, 90, 360) = 10, 1/36',
              r3 is not None and r3[2:7] == ('9/7', round(90 / 70, 6), '10', 10, '1/36') and abs(v1 - 70) < .02 and abs(v2 - 90) < .02))
v1, v2 = sd('Fomalhaut', 'Dec'), hd('Kaus Australis', 'Mars', 'RA')
cases.append((f'whole dropped: S Fomalhaut Dec {v1:.6f} 28 / H Kaus Australis – H Mars RA {v2:.6f} 30, circle factor gcd(28, 30, 360) = 2 < 5',
              ('RATIO', 'S Jupiter → S Fomalhaut', 'Dec', 'H Kaus Australis', 'Mars', 'RA') not in M51 and abs(v1 - 28) < .02 and abs(v2 - 30) < .02
              and math.gcd(math.gcd(28, 30), 360) == 2))
v1, v2 = sd('Ras Algethi', 'RA'), hd('Menkar', 'Saturn', 'RA')
r4 = M51.get(('RATIO', 'S Jupiter → S Ras Algethi', 'RA', 'H Menkar', 'Saturn', 'RA'))
cases.append((f'golden kept: S Ras Algethi RA {v1:.6f} 62×φ / H Menkar – H Saturn RA {v2:.6f} 93×φ -> 3/2, factor 31×φ, circle columns blank',
              r4 is not None and r4[2:7] == ('3/2', 1.5, '31×φ', None, None) and abs(v1 - 62 * F.PHI) < .03236 and abs(v2 - 93 * F.PHI) < .03236))
v1, v2 = sd('Haumea', 'RA'), hd('Praesepe', 'Quaoar', 'RA')
cases.append((f'golden dropped by the cap: S Haumea RA {v1:.6f} 88×φ / H Praesepe – H Quaoar RA {v2:.6f} 86×φ -> 44/43 (44 > 6)',
              ('RATIO', 'S Jupiter → S Haumea', 'RA', 'H Praesepe', 'Quaoar', 'RA') not in M51 and abs(v1 - 88 * F.PHI) < .03236
              and abs(v2 - 86 * F.PHI) < .03236))
# a pair left out (gcd 1): S Rukbat RA 68 whole against every scoring whole T24 HORSE star cell on an odd target from POSITIONS
odd = [(x, b, pl, round(v)) for x in JT.T1_T2_T24_STARS if ('H', x) in p51 for b in JT.T24_BODIES if ('H', b) in p51 and b != 'Moon'
       for pl in ('RA', 'Dec') for v in [dpl(p51[('H', x)], p51[('H', b)], pl)] if abs(v - round(v)) < 0.02 and round(v) % 2 == 1 and round(v) % 17
       and v >= 1.5]
left = [o for o in odd if ('RATIO', 'S Jupiter → S Rukbat', 'RA', f'H {o[0]}', o[1], o[2]) in M51 or ('EXACT', 'S Jupiter → S Rukbat', 'RA', f'H {o[0]}', o[1], o[2]) in M51]
cases.append((f'left out (gcd 1): S Rukbat RA 68 whole against {len(odd)} whole T24 HORSE cells on targets coprime to 68 (e.g. {odd[0][:3]} {odd[0][3]}): '
              'none listed', odd and not left))
# order (30 Oct 2026, revised): the top-scoring sky cell S Jupiter → S Aldebaran RA (70 whole) opens RATIO: 2/1 against 35 (H Capella – H Mercury
# Dec, then H Algorab – H Transpluto RA, by T24 score), then 7/1 against 10; EXACT opens with the best sky cell that has one, S Fomalhaut Dec 28
rr51 = [r for r in m51 if r[1] == 'RATIO']
ex51 = [r for r in m51 if r[1] == 'EXACT']
va, vc, vt = sd('Aldebaran', 'RA'), hd('Capella', 'Mercury', 'Dec'), hd('Algorab', 'Transpluto', 'RA')
sca, scc, sct = (sq(.02, abs(v - t)) for v, t in ((va, 70), (vc, 35), (vt, 35)))
cases.append((f'order: RATIO opens with S Aldebaran RA {va:.6f} 70 ({sca:.2f}, the top sky score) - 2/1 H Capella – H Mercury Dec {vc:.6f} '
              f'({scc:.2f}), 2/1 H Algorab – H Transpluto RA {vt:.6f} ({sct:.2f}), then 7/1; EXACT opens with S Fomalhaut Dec 28',
              [(r[2], r[13], r[14]) for r in rr51[:3]] == [('2/1', 'H Capella', 'Mercury'), ('2/1', 'H Algorab', 'Transpluto'), ('7/1', 'H Aculeus', 'Saturn')]
              and all(r[7] == 'S Jupiter → S Aldebaran' and r[8] == 'RA' for r in rr51[:3]) and rr51[0][12] == max(r[12] for r in m51)
              and scc > sct and ex51[0][7:9] == ('S Jupiter → S Fomalhaut', 'Dec')
              and m51[0] == rr51[0] and h51[1].startswith(f'S Jupiter → S Aldebaran RA {va:.6f} 70 whole ({sca:.2f}) — EXACT 0, RATIO ')))
# Ripon 15:25 16 Jun 2022 P01 MATCH HORSE (Eddie's example): the Murzim RA 90 group = 1 EXACT (H Alnilam – H Mars RA) then its RATIO rows;
# the Facies Dec 25 group = 3 EXACT then 17 RATIO (counts as built)
w51r = openpyxl.load_workbook(f'{tmp}/after51/ripon_jupiter_tables.xlsx', read_only=True)
p51r = posmap(w51r, 'P01')
s51r = {x[3]: (x[5], x[6]) for x in w51r['POSITIONS'].iter_rows(min_row=5, values_only=True) if x and x[0] == 'SKY' and x[3]}
mr = [r for r in w51r['P01'].iter_rows(values_only=True) if r and r[0] == 'MATCH HORSE'][1:]
grp51 = {}
for r in mr:
    if r[1] in ('EXACT', 'RATIO'):
        grp51[cur].append(r)
    else:
        cur = r[1]; grp51[cur] = []
vm, va2, vf = round(dpl(s51r['Jupiter'], s51r['Murzim'], 'RA'), 6), round(dpl(p51r[('H', 'Alnilam')], p51r[('H', 'Mars')], 'RA'), 6), \
    round(dpl(s51r['Jupiter'], s51r['Facies'], 'Dec'), 6)
hm = next((h for h in grp51 if h.startswith(f'S Jupiter → S Murzim RA {vm:.6f} 90 whole ({sq(.02, abs(vm - 90)):.2f}) — ')), None)
hf = next((h for h in grp51 if h.startswith(f'S Jupiter → S Facies Dec {vf:.6f} 25 whole ({sq(.02, abs(vf - 25)):.2f}) — ')), None)
secs = lambda h: [r[1] for r in grp51.get(h, [])]
cases.append((f'Ripon P01 MATCH HORSE: S Murzim RA {vm:.6f} 90 group - EXACT 1 (H Alnilam – H Mars RA {va2:.6f}), RATIO 42; S Facies Dec {vf:.6f} 25 '
              'group - EXACT 3, RATIO 17; header counts = the rows under it, EXACT before RATIO',
              hm is not None and hf is not None and hm.endswith('— EXACT 1, RATIO 42') and hf.endswith('— EXACT 3, RATIO 17')
              and secs(hm) == ['EXACT'] + ['RATIO'] * 42 and secs(hf) == ['EXACT'] * 3 + ['RATIO'] * 17
              and grp51[hm][0][13:17] == ('H Alnilam', 'Mars', 'RA', va2)
              and all(h.endswith(f"— EXACT {secs(h).count('EXACT')}, RATIO {secs(h).count('RATIO')}") for h in grp51)))
check('T51 by hand from POSITIONS (Kelso P01 MATCH HORSE; Ripon P01): ' + '; '.join(c for c, _ in cases), all(ok for _, ok in cases), [c for c, ok in cases if not ok])

# ---------------------------------------------------------------------------------------------------------------------------------
# Test 52 (Eddie, 30 Oct 2026): MATCH HORSE | JOCKEY on every runner sheet, after MATCH JOCKEY (two blank rows): the two MATCH tables side
# by side, cells copied from the sheet. Key = sky cell, section, ratio p/q, T24 family, T24 target; several rows of a key paired by T24
# score; grouped by sky cell (union), grey header rows 'EXACT h | j, RATIO h | j'. Before = 5ea09b4e's own tools/ (builds of test 51);
# Kelso, Lingfield, rebuilt Ascot, Ripon (chk_matchhj/check_hj.py).
print('test 52: MATCH HORSE | JOCKEY')
os.makedirs(f'{tmp}/after52', exist_ok=True)
os.makedirs(f'{tmp}/h93b', exist_ok=True)              # 31 Oct 2026 (test 53): after = 93b54a43's own tools/ (MATCH HORSE | JOCKEY alone)
extract(f'93b54a43', f'{tmp}/h93b')
out52 = {}
for n, inp in INP51.items():
    ao = f'{tmp}/after52/{n}_jupiter_tables.xlsx'
    run_build([f'{tmp}/h93b/tools/jupiter_tables.py', inp, '--out', ao], capture_output=True, text=True)
    out52[n] = subprocess.run([sys.executable, 'tools/scans/families_audit/chk_matchhj/check_hj.py', f'{tmp}/after51/{n}_jupiter_tables.xlsx', ao],
                              capture_output=True, text=True).stdout
print('T52 keys:', {n: [l for l in o.splitlines() if l.startswith('P')] for n, o in out52.items()})
check('T52 against the build before (commit 5ea09b4e), Kelso, Lingfield, rebuilt Ascot, Ripon (chk_matchhj/check_hj.py): every existing cell '
      'of every sheet identical (value, fill, font, border, number format, alignment, comment), merges, widths, freeze; on each runner sheet two '
      'blank rows after MATCH JOCKEY, then MATCH HORSE | JOCKEY and nothing after; every MATCH HORSE / MATCH JOCKEY row on its side exactly once, '
      'cell for cell (value and style), shared A-M = the source B-M on both sides; key, pairing by T24 score, the shorter side empty; groups '
      'by sky score (union), EXACT then RATIO by q, p, T24 target; grey header rows A-AC with EXACT h | j, RATIO h | j = the source headers; '
      'HORSE / JOCKEY merged over N-U / V-AC; title totals',
      all(line(o, 'mismatches:') == 'mismatches: 0' for o in out52.values())
      and all(len([l for l in o.splitlines() if l.startswith('P')]) == {'kelso': 6, 'lingfield': 6, 'ascot': 1, 'ripon': 4}[n] for n, o in out52.items()),
      {n: (line(o, 'cells compared:'), line(o, 'mismatches:')) for n, o in out52.items()})
# by hand (Ripon P01, Eddie's example): the S Wezen Dec 17×φ group - the EXACT line pairs H Menkent – H Orcus Dec with J Alhena – J Eris Dec
w52 = openpyxl.load_workbook(f'{tmp}/after52/ripon_jupiter_tables.xlsx', read_only=True)
p52 = posmap(w52, 'P01')
s52 = {x[3]: (x[5], x[6]) for x in w52['POSITIONS'].iter_rows(min_row=5, values_only=True) if x and x[0] == 'SKY' and x[3]}
m52 = [r for r in w52['P01'].iter_rows(values_only=True) if r and r[0] == 'MATCH HORSE | JOCKEY']
vw = round(dpl(s52['Jupiter'], s52['Wezen'], 'Dec'), 6)
uh, uj, uw = dpl(p52[('H', 'Menkent')], p52[('H', 'Orcus')], 'Dec'), dpl(p52[('J', 'Alhena')], p52[('J', 'Eris')], 'Dec'), \
    dpl(s52['Jupiter'], s52['Wezen'], 'Dec')
vh, vj = round(uh, 6), round(uj, 6)
g17 = 17 * F.PHI
U52 = {vw: uw, vh: uh, vj: uj}                         # score from the unrounded distance (the 6-dp value gives 69.40 for 69.39)
sc = lambda v: (10 * (0.03236 - abs(U52[v] - g17)) / 0.03236) ** 2
iw = next((i for i, r in enumerate(m52) if str(r[1]).startswith('S Jupiter → S Wezen Dec ')), None)
ln = m52[iw + 1] if iw is not None else ()
cases52 = [(f'Ripon P01 header S Jupiter → S Wezen Dec {vw:.6f} 17×φ golden ({sc(vw):.2f}) — EXACT 1 | 1, RATIO 0 | 0; EXACT line H Menkent – Orcus Dec '
            f'{vh:.6f} ({sc(vh):.2f}) | J Alhena – Eris Dec {vj:.6f} ({sc(vj):.2f})',
            iw is not None and m52[iw][1] == f'S Jupiter → S Wezen Dec {vw:.6f} 17×φ golden ({sc(vw):.2f}) — EXACT 1 | 1, RATIO 0 | 0'
            and ln[1:3] == ('EXACT', '1') and ln[7:13] == ('S Jupiter → S Wezen', 'Dec', vw, 'golden', '17×φ', round(sc(vw), 2))
            and ln[13:21] == ('H Menkent', 'Orcus', 'Dec', vh, 'golden', '17×φ', round(sc(vh), 2), round(sc(vh), 2))
            and ln[21:29] == ('J Alhena', 'Eris', 'Dec', vj, 'golden', '17×φ', round(sc(vj), 2), round(sc(vj), 2))
            and (vh, round(sc(vh), 2), vj, round(sc(vj), 2)) == (27.501175, 69.39, 27.503545, 82.13))]
t52 = m52[0][1] if m52 else ''
cases52.append(('Ripon P01 keys (Eddie\'s rough count 116 horse, 101 jockey, 56 shared): title "161 keys: 56 with both sides, 60 horse only, 45 jockey '
                'only"', '161 keys: 56 with both sides, 60 horse only, 45 jockey only; 15 sky cells' in t52))
check('T52 by hand from POSITIONS (Ripon P01): ' + '; '.join(c for c, _ in cases52), all(ok for _, ok in cases52), [c for c, ok in cases52 if not ok])

# Test 53 (Eddie, 31 Oct 2026): runner sheets freeze D4 (was C4); MATCH HORSE / MATCH JOCKEY gain BODY rows (SKY block S Jupiter → S X > 50 with
# the T24 row '{C} Jupiter → {C} body' cell of X > 50) and CONJ STAR rows (SKY block S Jupiter → S star conjunction > 50 with the star's T24
# row cells > 50 and the T24 row 'S Jupiter → {C} body' cells > 50); conjunction sky cells become groups; MATCH HORSE | JOCKEY takes them in;
# new SUMMARY sheet. Before = 93b54a43's own tools/ (builds of test 52); Kelso, Lingfield, rebuilt Ascot, Ripon, and Newcastle AW 16:33 2021
# (none of the four has a conjunction sky star cell over 50; Newcastle has S Sadalmelik RA). chk_match31/check_m31.py.
print('test 53: BODY, CONJ STAR, SUMMARY, freeze D4')
os.makedirs(f'{tmp}/after53', exist_ok=True)
os.makedirs(f'{tmp}/h3da', exist_ok=True)              # 31 Oct 2026 (test 54): after = 3da4ae61's own tools/ (this change alone)
extract(f'3da4ae61', f'{tmp}/h3da')
INP53 = dict(INP51, newcastle='reference/classic/race_batch_20210802_newcastle_aw_1633.xlsx')
out53 = {}
for n, inp in INP53.items():
    bo, ao = f'{tmp}/after52/{n}_jupiter_tables.xlsx', f'{tmp}/after53/{n}_jupiter_tables.xlsx'
    if not os.path.exists(bo):
        run_build([f'{tmp}/h93b/tools/jupiter_tables.py', inp, '--out', bo], capture_output=True, text=True)
    run_build([f'{tmp}/h3da/tools/jupiter_tables.py', inp, '--out', ao], capture_output=True, text=True)
    out53[n] = subprocess.run([sys.executable, 'tools/scans/families_audit/chk_match31/check_m31.py', bo, ao], capture_output=True, text=True).stdout
print('T53 rows:', {n: [l for l in o.splitlines() if re.match(r'P\d\d ', l)] for n, o in out53.items()})
check('T53 against the build before (commit 93b54a43), Kelso, Lingfield, rebuilt Ascot, Ripon, Newcastle AW 16:33 2021 (chk_match31/check_m31.py): '
      'POSITIONS and SKY JUPITER identical (freeze unchanged), runner sheets freeze D4, every cell before the MATCH tables and every width '
      'identical; each MATCH table minus the BODY / CONJ STAR rows and the conjunction-only groups = before cell for cell (titles and header '
      'counts excepted, header counts extended by BODY / CONJ STAR); BODY / CONJ STAR rows rebuilt from the SKY block and T24 cells (set, cells, '
      'order, blanks, fills, header counts, group order by sky score); MATCH HORSE | JOCKEY: each new row once on its side, keys, order, counts; '
      'SUMMARY rebuilt from MATCH HORSE | JOCKEY (filter, groups, counts, cells and styles, merges, blank rows)',
      all(line(o, 'mismatches:') == 'mismatches: 0' for o in out53.values())
      and all(len([l for l in o.splitlines() if re.match(r'P\d\d ', l)]) == {'kelso': 6, 'lingfield': 6, 'ascot': 1, 'ripon': 4, 'newcastle': 6}[n]
              for n, o in out53.items()),
      {n: line(o, 'mismatches:') for n, o in out53.items()})
# by hand from POSITIONS: Ripon P01 BODY (S Juno Dec conjunction / H Jupiter → H Juno Dec 11) and its SUMMARY line; Newcastle P01 CONJ STAR
# (S Sadalmelik RA conjunction / J Sadalmelik – J Gonggong RA conjunction, S Jupiter → J Makemake Dec 45)
cases53 = []
sq = lambda d, tol: (10 * (tol - d) / tol) ** 2
w53 = openpyxl.load_workbook(f'{tmp}/after53/ripon_jupiter_tables.xlsx', read_only=True)
p53 = posmap(w53, 'P01')
s53 = {x[3]: (x[5], x[6]) for x in w53['POSITIONS'].iter_rows(min_row=5, values_only=True) if x and x[0] == 'SKY' and x[3]}
ds, dh = dpl(s53['Jupiter'], s53['Juno'], 'Dec'), dpl(p53[('H', 'Jupiter')], p53[('H', 'Juno')], 'Dec')
m53 = [r for r in w53['P01'].iter_rows(values_only=True) if r and r[0] == 'MATCH HORSE' and r[1] == 'BODY']
row53 = ('MATCH HORSE', 'BODY', None, None, None, None, None, 'S Jupiter → S Juno', 'Dec', round(ds, 6), 'conj', 'conj', round(sq(ds, 0.3), 2),
         'H Jupiter → H body', 'Juno', 'Dec', round(dh, 6), 'whole', '11', round(sq(abs(dh - 11), 0.02), 2), round(sq(ds, 0.3), 2))
cases53.append((f'Ripon P01 MATCH HORSE BODY S Jupiter → S Juno Dec {ds:.6f} conjunction ({sq(ds, 0.3):.2f}) / H Jupiter → H Juno Dec {dh:.6f} 11 '
                f'whole ({sq(abs(dh - 11), 0.02):.2f}), the only BODY row', [tuple(r[:21]) for r in m53] == [row53]
                and (round(ds, 6), round(sq(ds, 0.3), 2), round(dh, 6), round(sq(abs(dh - 11), 0.02), 2)) == (0.078975, 54.28, 10.997519, 76.73)))
hd53 = f'S Jupiter → S Juno Dec {ds:.6f} conjunction ({sq(ds, 0.3):.2f}) — EXACT 0 | 0, RATIO 0 | 0, BODY 1 | 0, CONJ STAR 0 | 0'
sm = [r for r in w53['SUMMARY'].iter_rows(values_only=True)]
i53 = next((i for i, r in enumerate(sm) if r[1] == hd53), None)
hj53 = [r for r in w53['P01'].iter_rows(values_only=True) if r and r[0] == 'MATCH HORSE | JOCKEY' and r[1] in ('EXACT', 'RATIO', 'BODY', 'CONJ STAR')]
big = lambda x: isinstance(x, (int, float)) and x > 50
out_ = next(r for r in hj53 if big(r[12]) and not big(r[19]) and not big(r[27]))
cases53.append((f'Ripon SUMMARY P01: header {hd53!r}, then the line = the MATCH HORSE | JOCKEY line (jockey side empty); a line with sky > 50 and both '
                f'T24 scores <= 50 ({out_[7]} {out_[8]} / {out_[13] or out_[21]} {out_[14] or out_[22]} {out_[15] or out_[23]}, '
                f'{out_[19]} | {out_[27]}) left out',
                sm[0][0] == 'P01 1 Bollin Joan / David Allan' and i53 is not None and tuple(sm[i53 + 1][:21]) == ('MATCH HORSE | JOCKEY',) + row53[1:]
                and tuple(sm[i53 + 1][21:29]) == (None,) * 8 and tuple(out_[:29]) not in [tuple(r[:29]) for r in sm]))
w53n = openpyxl.load_workbook(f'{tmp}/after53/newcastle_jupiter_tables.xlsx', read_only=True)
pn = posmap(w53n, 'P01')
sn = {x[3]: (x[5], x[6]) for x in w53n['POSITIONS'].iter_rows(min_row=5, values_only=True) if x and x[0] == 'SKY' and x[3]}
dsn = dpl(sn['Jupiter'], sn['Sadalmelik'], 'RA')
dg, dm = dpl(pn[('J', 'Sadalmelik')], pn[('J', 'Gonggong')], 'RA'), dpl(sn['Jupiter'], pn[('J', 'Makemake')], 'Dec')
mn = [r[:21] for r in w53n['P01'].iter_rows(values_only=True) if r and r[0] == 'MATCH JOCKEY' and r[1] == 'CONJ STAR']
skyn = ('S Jupiter → S Sadalmelik', 'RA', round(dsn, 6), 'conj', 'conj', round(sq(dsn, 1.5), 2))
cases53.append((f'Newcastle AW 16:33 2021 P01 MATCH JOCKEY CONJ STAR: S Jupiter → S Sadalmelik RA {dsn:.6f} conjunction ({sq(dsn, 1.5):.2f}) / '
                f'J Sadalmelik – J Gonggong RA {dg:.6f} conjunction ({sq(dg, 1.5):.2f}), S Jupiter → J Makemake Dec {dm:.6f} 45 whole '
                f'({sq(abs(dm - 45), 0.02):.2f}); 5 rows, star row first',
                len(mn) == 5 and mn[0][13] == 'J Sadalmelik' and mn[2][13] == 'S Jupiter → J body'
                and ('MATCH JOCKEY', 'CONJ STAR', None, None, None, None, None) + skyn + ('J Sadalmelik', 'Gonggong', 'RA', round(dg, 6), 'conj', 'conj',
                                                                                        round(sq(dg, 1.5), 2), round(min(sq(dg, 1.5), sq(dsn, 1.5)), 2)) in mn
                and ('MATCH JOCKEY', 'CONJ STAR', None, None, None, None, None) + skyn + ('S Jupiter → J body', 'Makemake', 'Dec', round(dm, 6), 'whole', '45',
                                                                                        round(sq(abs(dm - 45), 0.02), 2), round(sq(dsn, 1.5), 2)) in mn
                and (round(dsn, 6), round(sq(dsn, 1.5), 2), round(dg, 6), round(dm, 6)) == (0.211451, 73.79, 0.179281, 44.999747)))
fr = {n: {s_: openpyxl.load_workbook(f'{tmp}/after53/{n}_jupiter_tables.xlsx')[s_].freeze_panes for s_ in ('POSITIONS', 'SKY JUPITER', 'P01')} for n in ('ascot',)}
cases53.append(('freeze: runner sheets D4, POSITIONS and SKY JUPITER none (as before)', fr == {'ascot': {'POSITIONS': None, 'SKY JUPITER': None, 'P01': 'D4'}}))
check('T53 by hand from POSITIONS: ' + '; '.join(c for c, _ in cases53), all(ok for _, ok in cases53), [c for c, ok in cases53 if not ok])

# Test 54 (Eddie, 31 Oct 2026): MATCH groups in three blocks under darker grey band rows - FIXED STARS (star or equator sky cell, not
# conjunction), CONJUNCTIONS (star / equator conjunction), OTHER BODIES (body cells, then body conjunctions at the end); inside by sky score;
# SUMMARY filter on unrounded scores. Before = 3da4ae61's own tools/ (builds of test 53); five races + Goodwood 18:53 27 Aug 2021 (the only
# one with an equator group: S Jupiter → S equator Dec 14 whole). chk_match32/check_m32.py.
#@ full: T54 comparison 3da4ae61 -> e6267326 builds (old against old; e6267326's tools/ = ef4298c6's, the block-order change)
print('test 54: MATCH blocks, SUMMARY unrounded')
os.makedirs(f'{tmp}/after54', exist_ok=True)
INP54 = {'kelso': 'reference/classic/race_batch_20220504_kelso_1450.xlsx', 'lingfield': 'reference/classic/race_batch_20210402_lingfield_aw_1435.xlsx',
         'ascot': f'{tmp}/race_batch_20120619_ascot_1430.xlsx', 'ripon': 'reference/classic/race_batch_20220616_ripon_1525.xlsx',
         'newcastle': 'reference/classic/race_batch_20210802_newcastle_aw_1633.xlsx', 'goodwood': 'reference/classic/race_batch_20210827_goodwood_1853.xlsx'}
extract('3da4ae61', f'{tmp}/h3da')                      # the build before (also built in test 53 with --full; cached)
extract('e6267326', f'{tmp}/he62')                      # the build after (Part 1 commit; also the 'before' of test 55)
os.makedirs(f'{tmp}/after53', exist_ok=True)
sqd = lambda d, tol: (10 * (tol - d) / tol) ** 2
out54 = {}
for n, inp in INP54.items():
    bo, ao = f'{tmp}/after53/{n}_jupiter_tables.xlsx', f'{tmp}/after54/{n}_jupiter_tables.xlsx'
    if not os.path.exists(bo):
        run_build([f'{tmp}/h3da/tools/jupiter_tables.py', inp, '--out', bo], capture_output=True, text=True)
    run_build([f'{tmp}/he62/tools/jupiter_tables.py', inp, '--out', ao], capture_output=True, text=True)
    out54[n] = subprocess.run([sys.executable, 'tools/scans/families_audit/chk_match32/check_m32.py', bo, ao], capture_output=True, text=True).stdout
print('T54 groups per block:', {n: [l for l in o.splitlines() if re.match(r'P\d\d ', l)] for n, o in out54.items()})
check('T54 builds by e6267326 (= ef4298c6 tools/) against the build before (commit 3da4ae61), Kelso, Lingfield, rebuilt Ascot, Ripon, Newcastle AW 16:33 2021, Goodwood 18:53 2021 '
      '(chk_match32/check_m32.py): every sheet and runner-sheet cell before the MATCH tables identical; each MATCH table = before with its groups '
      'moved only (every group cell for cell, titles excepted); order FIXED STARS (incl. the equator), CONJUNCTIONS, OTHER BODIES (body '
      'conjunctions last), inside by the sky score recomputed from POSITIONS (unrounded), ties SKY block order; band rows (A6A6A6, A to the last '
      'column) exactly before each block present; title counts; SUMMARY rebuilt on unrounded scores (bands only for shown blocks); no line '
      'the 2-dp filter would add or drop',
      all(line(o, 'mismatches:') == 'mismatches: 0' for o in out54.values())
      and all(len([l for l in o.splitlines() if re.match(r'P\d\d ', l)]) == {'kelso': 6, 'lingfield': 6, 'ascot': 1, 'ripon': 4, 'newcastle': 6,
                                                                              'goodwood': 5}[n] for n, o in out54.items())
      and all(l.endswith('2-dp edge lines []') for o in out54.values() for l in o.splitlines() if re.match(r'P\d\d ', l)),
      {n: line(o, 'mismatches:') for n, o in out54.items()})

#@ quick: T55 comparison with e6267326 (the commit before the current change); T54 by-hand cells on the current builds
print('test 55: SUMMARY freeze D4, ratio 1 light blue in column C')
INP55 = {'kelso': 'reference/classic/race_batch_20220504_kelso_1450.xlsx', 'lingfield': 'reference/classic/race_batch_20210402_lingfield_aw_1435.xlsx',
         'ascot': f'{tmp}/race_batch_20120619_ascot_1430.xlsx', 'ripon': 'reference/classic/race_batch_20220616_ripon_1525.xlsx',
         'newcastle': 'reference/classic/race_batch_20210802_newcastle_aw_1633.xlsx', 'goodwood': 'reference/classic/race_batch_20210827_goodwood_1853.xlsx'}
extract('e6267326', f'{tmp}/he62')
for d in ('before55', 'after55'):
    os.makedirs(f'{tmp}/{d}', exist_ok=True)
sqd = lambda d, tol: (10 * (tol - d) / tol) ** 2
out55 = {}
for n, inp in INP55.items():
    bo, ao = f'{tmp}/before55/{n}_jupiter_tables.xlsx', f'{tmp}/after55/{n}_jupiter_tables.xlsx'
    run_build([f'{tmp}/he62/tools/jupiter_tables.py', inp, '--out', bo], capture_output=True, text=True)
    subprocess.run([sys.executable, 'tools/jupiter_tables.py', inp, '--out', ao], capture_output=True, text=True)
    out55[n] = subprocess.run([sys.executable, 'tools/scans/families_audit/chk_ratio1/check_r1.py', bo, ao], capture_output=True, text=True).stdout
print('T55 filled C cells:', {n: line(o, 'totals:') for n, o in out55.items()})
check('T55 against the build before (commit e6267326), Kelso, Lingfield, rebuilt Ascot, Ripon, Newcastle AW 16:33 2021, Goodwood 18:53 2021 '
      '(chk_ratio1/check_r1.py): every sheet, value, font, border, alignment, number format, comment, merge and width identical; freeze '
      'SUMMARY None -> D4, every other sheet unchanged (runner sheets D4, POSITIONS / SKY JUPITER none); the only fill change is column C on '
      'MATCH HORSE / MATCH JOCKEY / MATCH HORSE | JOCKEY data rows (runner sheets and SUMMARY), none -> BDD7EE exactly where column C is '
      "'1' (EXACT rows, every one of them); no BDD7EE anywhere else",
      all(line(o, 'mismatches:') == 'mismatches: 0' for o in out55.values())
      and all(line(o, 'freeze:').startswith("freeze: {'POSITIONS': None, 'SKY JUPITER': None, 'P01': 'D4', 'SUMMARY': 'D4'}") for o in out55.values())
      and {n: line(o, 'totals:') for n, o in out55.items()} == {n: f'totals: MATCH HORSE {h}, MATCH JOCKEY {j}, MATCH HORSE | JOCKEY {hj} | SUMMARY {sm}'
           for n, (h, j, hj, sm) in {'kelso': (90, 94, 136, 38), 'lingfield': (17, 13, 26, 9), 'ascot': (9, 14, 16, 3), 'ripon': (53, 46, 73, 9),
                                     'newcastle': (50, 41, 71, 1), 'goodwood': (84, 86, 128, 4)}.items()},
      {n: line(o, 'mismatches:') for n, o in out55.items()})
os.makedirs(f'{tmp}/after54', exist_ok=True)               # T54's by-hand cases (below) read the current builds
for n in INP55:
    shutil.copy(f'{tmp}/after55/{n}_jupiter_tables.xlsx', f'{tmp}/after54/{n}_jupiter_tables.xlsx')
cases54 = []
def blocks54(path, P, tid):
    """-> [(band, [group header texts])] of one MATCH table"""
    out = []
    for r in openpyxl.load_workbook(path, read_only=True)[P].iter_rows(values_only=True):
        if r and r[0] == tid and r[1] in ('FIXED STARS', 'CONJUNCTIONS', 'OTHER BODIES'):
            out.append((r[1], []))
        elif r and r[0] == tid and isinstance(r[1], str) and ' — EXACT ' in r[1] and len(r[1]) < 300:
            out[-1][1].append(r[1])
    return out
bn = blocks54(f'{tmp}/after54/newcastle_jupiter_tables.xlsx', 'P01', 'MATCH JOCKEY')
cases54.append(('Newcastle P01 MATCH JOCKEY: bands FIXED STARS, CONJUNCTIONS, OTHER BODIES; CONJUNCTIONS = the S Sadalmelik RA group only',
                [b for b, _ in bn] == ['FIXED STARS', 'CONJUNCTIONS', 'OTHER BODIES']
                and [g.split(' — ')[0] for g in bn[1][1]] == ['S Jupiter → S Sadalmelik RA 0.211451 conjunction (73.79)']))
br = blocks54(f'{tmp}/after54/ripon_jupiter_tables.xlsx', 'P01', 'MATCH HORSE')
ob = dict(br).get('OTHER BODIES', [])
cases54.append(('Ripon P01 MATCH HORSE: S Murzim RA 90.000720 90 whole (92.93) under FIXED STARS; no CONJUNCTIONS band; S Juno Dec 0.078975 '
                'conjunction (54.28) the last group of OTHER BODIES, after every non-conjunction body group',
                [b for b, _ in br] == ['FIXED STARS', 'OTHER BODIES']
                and any(g.startswith('S Jupiter → S Murzim RA 90.000720 90 whole (92.93) — ') for g in dict(br)['FIXED STARS'])
                and ob[-1].startswith('S Jupiter → S Juno Dec 0.078975 conjunction (54.28) — ')
                and all(' conjunction (' not in g for g in ob[:-1]) and len(ob) == 5))
wg = openpyxl.load_workbook(f'{tmp}/after54/goodwood_jupiter_tables.xlsx', read_only=True)
sg = {x[3]: (x[5], x[6]) for x in wg['POSITIONS'].iter_rows(min_row=5, values_only=True) if x and x[0] == 'SKY' and x[3]}
ve = abs(sg['Jupiter'][1])
he = f'S Jupiter → S equator Dec {ve:.6f} 14 whole ({sqd(abs(ve - 14), 0.02):.2f})'
bg = blocks54(f'{tmp}/after54/goodwood_jupiter_tables.xlsx', 'P01', 'MATCH HORSE')
cases54.append((f'Goodwood P01 MATCH HORSE: the equator group {he!r} under FIXED STARS (|Dec S Jupiter| from POSITIONS)',
                (round(ve, 6), round(sqd(abs(ve - 14), 0.02), 2)) == (14.000933, 90.89)
                and any(g.startswith(he + ' — ') for g in dict(bg).get('FIXED STARS', []))))
check('T54 by hand: ' + '; '.join(c for c, _ in cases54), all(ok for _, ok in cases54), [c for c, ok in cases54 if not ok])
shutil.rmtree(tmp, ignore_errors=True)
print(f"old builds: {NBUILT_OLD[0]} built, {NBUILT_OLD[1]} from tests/.cache ({'--full' if FULL else 'quick'} run)")
print(f'\n{sum(RES)}/{len(RES)} passed')
sys.exit(0 if all(RES) else 1)
