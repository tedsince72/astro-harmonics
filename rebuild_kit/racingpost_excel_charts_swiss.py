#!/usr/bin/env python3
"""
racingpost_excel_charts_swiss.py  v2.3
--------------------------------------
Takes a race batch CSV (one race) and produces one Excel file with:
  - TRANS tab   : celestial chart at race time / racecourse coordinates
  - P01, P02... : natal charts, horse then its jockey, fixed by cloth order
  - NATAL_HOURLY: every natal body at each hour of its birth day in
                  standard time (00:00 to 24:00, no summer time), one block
                  per P tab
  - META tab    : versions, race details and one row per P tab (last sheet)

Calls celestial_bodies_swisseph.py for every chart (Swiss Ephemeris engine).

Changes in v2.3 (engine v2.2, decision NT-1):
  - Natal charts use STANDARD time, never summer time: noon and the birth
    day are taken at the birth zone's standard offset (Eddie 30/09/2026:
    clock changes are artificial). UK and Irish births are therefore always
    12:00 UTC, as in the original v2.5 engine. The engine does the work;
    this script only records the new rule in META (natal_time).
  - The new engine version changes the fingerprint, so every cached jockey
    chart is recalculated automatically on first use.

Changes in v2.2 (engine v2.1, decision H3-A):
  - The engine's hourly birth-day positions are written to a NATAL_HOURLY
    sheet (tab, name, body, hour, local_time, utc_time, ra, dec), so natal-
    to-natal distances can be followed hour by hour through the same day.
  - Jockey cache stores the hourly file beside the chart; a cached chart is
    only used when both files are present.

Changes in v2.1 (engine v2.0, plan item ENG-2):
  - Natal charts use the birth country's LOCAL calendar day: the engine is
    called in --natal-day mode at local noon in the birth time zone, and
    each natal row carries the day's range in ra_lo, ra_hi, dec_lo, dec_hi.
  - META records the engine version as well as its fingerprint.

Changes in v2.0 (plan items CH-1 to CH-8):
  - Fixed tab numbers: the k-th horse by cloth number is always P(2k-1) and
    its jockey always P(2k). A subject that cannot be charted (no DOB, no
    jockey row) gets a placeholder tab in its fixed place, named
    "<name> [NO DATA: reason]", so later tabs never shift.
  - Birth dates checked against the race date: a 2-digit year that would put
    a birth after the race is moved back 100 years (e.g. '20-May-67' = 1967);
    a birth date still on or after the race date stops the race.
  - Jockey chart cache keyed by name, DOB, birth country and a fingerprint of
    the engine file, so it can never serve charts from an older engine or mix
    up two riders with the same name.
  - Birth-country time zone recorded for every subject (racing_common
    COUNTRY_TZ).
  - Racecourse matched exactly (no partial matching, no London fallback).
    Birth country matched exactly (no online geocoding, no London fallback).
  - Fail loudly: an unknown course or country, an unreadable date, or any
    engine failure stops the race with exit code 1 and no workbook is written.
  - Non-runners (status NR) are skipped.

Usage:
  python racingpost_excel_charts_swiss.py --csv race_batch_20260926_haydock_1635.csv \\
      --ephemeris-path ephe --engine celestial_bodies_swisseph.py

Requires: openpyxl, racing_common.py, celestial_bodies_swisseph.py, Swiss Ephemeris files
"""

import argparse
import csv
import hashlib
import os
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

import racing_common as rc

CHARTS_VERSION = "2.3"
NATAL_TIME = "12:00 standard time (no summer time); range 00:00-24:00 standard"


class ChartError(Exception):
    """The race cannot be charted correctly. The message says why."""


# ---------------------------------------------------------------------------
# UK Racecourse coordinates (lat, lon, elevation_m) - matched EXACTLY
# ---------------------------------------------------------------------------

RACECOURSE_COORDS = {
    "aintree":          (53.4728, -2.9383,   19),
    "ascot":            (51.4082, -0.6187,   43),
    "ayr":              (55.4621, -4.6359,    6),
    "bangor on dee":    (52.9131, -3.0084,   43),
    "bath":             (51.3981, -2.3308,  175),
    "beverley":         (53.8471, -0.4104,   42),
    "brighton":         (50.8321, -0.1325,  150),
    "carlisle":         (54.8827, -2.9364,   92),
    "cartmel":          (54.2018, -2.9535,   10),
    "catterick":        (54.3727, -1.7069,  128),
    "chelmsford":             (51.7367,  0.4754,   38),
    "chelmsford city":         (51.7367,  0.4754,   38),
    "cheltenham":       (51.9006, -2.0622,   78),
    "chepstow":         (51.6421, -2.6749,   71),
    "chester":          (53.1866, -2.9024,    8),
    "doncaster":        (53.5244, -1.0875,   16),
    "epsom":            (51.3357, -0.2565,  155),
    "exeter":           (50.7336, -3.4619,  109),
    "fakenham":         (52.8320,  0.8425,   52),
    "ffos las":         (51.7472, -4.2103,   70),
    "folkestone":       (51.1032,  1.1482,   85),
    "fontwell park":    (50.8515, -0.6538,   20),
    "goodwood":         (50.8998, -0.7576,  107),
    "hamilton":         (55.7668, -4.0480,   58),
    "haydock":          (53.4714, -2.6316,   32),
    "hereford":         (52.0600, -2.7160,   50),
    "hexham":           (54.9613, -2.1097,  296),
    "huntingdon":       (52.3396, -0.1807,   15),
    "kelso":            (55.6136, -2.4311,   50),
    "kempton":                (51.4119, -0.3712,   20),
    "kempton aw":              (51.4119, -0.3712,   20),
    "leicester":        (52.6113, -1.1307,   90),
    "lichfield":        (52.7100, -1.8200,   98),
    "lingfield":             (51.1826, -0.0028,  132),
    "lingfield aw":            (51.1826, -0.0028,  132),
    "ludlow":           (52.3686, -2.7153,  132),
    "market rasen":     (53.3979, -0.3429,   23),
    "musselburgh":      (55.9406, -3.0564,    9),
    "newbury":          (51.3957, -1.3132,   91),
    "newcastle":              (54.9832, -1.6247,   35),
    "newcastle aw":            (54.9832, -1.6247,   35),
    "newmarket":              (52.2405,  0.4108,   67),
    "newmarket july":          (52.2357,  0.4158,   67),
    "newmarket rowley":        (52.2405,  0.4108,   67),
    "newton abbot":     (50.5365, -3.5920,   10),
    "nottingham":       (52.9419, -1.1228,   34),
    "perth":            (56.3892, -3.4265,   23),
    "plumpton":         (50.9360, -0.0663,  113),
    "pontefract":       (53.6915, -1.3025,   30),
    "redcar":           (54.5968, -1.0766,    5),
    "ripon":            (54.1359, -1.5206,   55),
    "salisbury":        (51.0832, -1.8219,   96),
    "sandown":          (51.3684, -0.3473,   34),
    "sedgefield":       (54.6561, -1.4514,  105),
    "southwell":              (53.0699, -0.9499,   37),
    "southwell aw":            (53.0699, -0.9499,   37),
    "stratford":        (52.1920, -1.7108,   50),
    "taunton":          (51.0224, -3.0872,   20),
    "thirsk":           (54.2338, -1.3488,   46),
    "thurles":          (52.6833, -7.8833,   65),
    "uttoxeter":        (52.8897, -1.8586,  108),
    "warwick":          (52.2914, -1.5905,   74),
    "wetherby":         (53.9335, -1.3834,   50),
    "wincanton":        (51.0647, -2.4226,  130),
    "windsor":          (51.4813, -0.6185,   28),
    "wolverhampton":          (52.5668, -2.1143,  130),
    "wolverhampton aw":        (52.5668, -2.1143,  130),
    "worcester":        (52.2019, -2.2197,   25),
    "yarmouth":         (52.6083,  1.7196,    4),
    "york":             (53.9948, -1.0859,   19),
    # Irish courses
    "curragh":          (53.1553, -6.7991,  100),
    "leopardstown":     (53.2800, -6.2100,   60),
    "tipperary":        (52.4713, -8.1557,   71),
    "galway":           (53.3026, -8.9699,   21),
    "cork":             (51.8716, -8.4933,   30),
    # Malaysian courses
    "selangor":         ( 3.0456, 101.7144,   35),
    "perak":            ( 4.5784, 101.0833,   40),
    # Japan JRA courses
    "sapporo":          (43.0522, 141.3380,   15),
    "hakodate":         (41.7989, 140.7808,   25),
    "fukushima":        (37.7543, 140.4440,   70),
    "niigata":          (37.8108, 138.9872,    5),
    "tokyo":            (35.6972, 139.4794,   37),
    "nakayama":         (35.7753, 139.9253,   18),
    "chukyo":           (35.0653, 136.9833,   52),
    "kyoto":            (34.9828, 135.7011,   65),
    "hanshin":          (34.8072, 135.3594,   26),
    "kokura":           (33.8694, 130.8383,   18),
    # added v2.0: Racing Post course keys not previously listed
    "towcester":        (52.1386, -0.9906,  110),
    "fontwell":         (50.8515, -0.6538,   20),
    "chelmsford aw":    (51.7367,  0.4754,   38),
}


# ---------------------------------------------------------------------------
# Country centroid coordinates (lat, lon) - matched EXACTLY
# Used for natal charts when only birth_country is known
# ---------------------------------------------------------------------------

COUNTRY_CENTROIDS = {
    "england":          (52.3555, -1.1743),
    "great britain":    (54.0000, -2.5000),
    "scotland":         (56.4907, -4.2026),
    "wales":            (52.1307, -3.7837),
    "northern ireland": (54.7877, -6.4923),
    "ireland":          (53.1424, -7.6921),
    "france":           (46.2276,  2.2137),
    "germany":          (51.1657, 10.4515),
    "italy":            (41.8719, 12.5674),
    "spain":            (40.4637, -3.7492),
    "portugal":         (39.3999, -8.2245),
    "usa":              (37.0902, -95.7129),
    "australia":        (-25.2744, 133.7751),
    "new zealand":      (-40.9006, 174.8860),
    "south africa":     (-30.5595, 22.9375),
    "canada":           (56.1304, -106.3468),
    "japan":            (36.2048, 138.2529),
    "brazil":           (-14.2350, -51.9253),
    "argentina":        (-38.4161, -63.6167),
    "belgium":          (50.5039,  4.4699),
    "denmark":          (56.2639,  9.5018),
    "sweden":           (60.1282, 18.6435),
    "norway":           (60.4720,  8.4689),
    "qatar":            (25.3548, 51.1839),
    "uae":              (23.4241, 53.8478),
    "romania":          (45.9432, 24.9668),
    "swaziland":        (-26.5225, 31.4659),
    "serbia":           (44.0165, 21.0059),
    "czech republic":   (49.8175, 15.4730),
    "poland":           (51.9194, 19.1451),
    "hungary":          (47.1625, 19.5033),
    "chile":            (-35.6751, -71.5430),
    "peru":             (-9.1900, -75.0152),
    "uruguay":          (-32.5228, -55.7658),
    # added v2.0 so every country in racing_common.COUNTRY_TZ has a centroid
    "netherlands":      (52.1326,  5.2913),
    "switzerland":      (46.8182,  8.2275),
    "austria":          (47.5162, 14.5501),
    "slovakia":         (48.6690, 19.6990),
    "greece":           (39.0742, 21.8243),
    "turkey":           (38.9637, 35.2433),
    "russia":           (61.5240, 105.3188),
    "mexico":           (23.6345, -102.5528),
    "zimbabwe":         (-19.0154, 29.1549),
    "morocco":          (31.7917, -7.0926),
    "saudi arabia":     (23.8859, 45.0792),
    "india":            (20.5937, 78.9629),
    "hong kong":        (22.3193, 114.1694),
    "south korea":      (35.9078, 127.7669),
    "china":            (35.8617, 104.1954),
    "kazakhstan":       (48.0196, 66.9237),
}


# ---------------------------------------------------------------------------
# Excel styling
# ---------------------------------------------------------------------------

HEADER_FILL  = PatternFill("solid", fgColor="1F4E79")   # dark blue
HEADER_FONT  = Font(name="Arial", size=9, bold=True, color="FFFFFF")
NAME_FILL    = PatternFill("solid", fgColor="1F4E79")
NAME_FONT    = Font(name="Arial", size=9, bold=True, color="FFFFFF")
ROW_FILL_A   = PatternFill("solid", fgColor="D6E4F0")   # light blue
ROW_FILL_B   = PatternFill("solid", fgColor="FFFFFF")
DATA_FONT    = Font(name="Arial", size=9)
COL_WIDTHS   = {
    "A": 22,   # name
    "B": 14,   # body_A
    "C": 14,   # body_B
    "D": 22,   # measurement_type
    "E": 10, "F": 10, "G": 10, "H": 10,   # ra fields
    "I": 10, "J": 10, "K": 10, "L": 10,   # dec fields
    "M": 12, "N": 12,                       # topo dist
    "O": 12,                                # dist_AU
    "P": 10, "Q": 10, "R": 10,             # mid xyz
    "S": 6,                                 # moon_involved
}

FIELDNAMES = [
    "name",
    "body_A", "body_B", "measurement_type",
    "ra_A", "ra_B", "ra_diff", "ra_midpoint",
    "dec_A", "dec_B", "dec_diff", "dec_midpoint",
    "topo_dist_AU_A", "topo_dist_AU_B", "dist_AU",
    "mid_x_AU", "mid_y_AU", "mid_z_AU",
    "moon_involved",
    "ra_lo", "ra_hi", "dec_lo", "dec_hi",   # natal birth-day ranges (engine v2.0)
]


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def parse_args():
    p = argparse.ArgumentParser(description=f"Race charts v{CHARTS_VERSION}")
    p.add_argument("--csv", required=True, help="Race batch CSV file")
    p.add_argument("--outdir", default=None,
                   help="Output directory (default: same folder as CSV)")
    p.add_argument("--ephemeris-path", default=".",
                   help="Path to Swiss Ephemeris .se1 files directory")
    p.add_argument("--hip-path", default=".",
                   help="Kept for CLI compatibility - unused by Swiss engine")
    p.add_argument("--engine", default="celestial_bodies_swisseph.py",
                   help="Path to celestial_bodies_swisseph.py")
    return p.parse_args()


def engine_fingerprint(engine_path: str) -> str:
    """First 10 hex digits of the engine file's SHA-1: any change to the
    engine changes the fingerprint, so cached charts are never reused."""
    with open(engine_path, "rb") as f:
        return hashlib.sha1(f.read()).hexdigest()[:10]


def engine_version(engine_path: str) -> str:
    """ENGINE_VERSION as written in the engine file ('unknown' if absent)."""
    m = re.search(r'^ENGINE_VERSION\s*=\s*"([^"]+)"', Path(engine_path).read_text(), re.M)
    return m.group(1) if m else "unknown"


def get_racecourse_coords(racecourse: str) -> Tuple[float, float, int]:
    key = rc.venue_name(racecourse)
    if key not in RACECOURSE_COORDS:
        raise ChartError(f"Unknown racecourse '{racecourse}' - add it to "
                         f"RACECOURSE_COORDS (exact Racing Post course name)")
    return RACECOURSE_COORDS[key]


def get_birth_coords(birth_country: str) -> Tuple[float, float]:
    key = (birth_country or "").strip().lower()
    if key not in COUNTRY_CENTROIDS:
        raise ChartError(f"No centroid for birth country '{birth_country}' - add it "
                         f"to COUNTRY_CENTROIDS")
    return COUNTRY_CENTROIDS[key]


def _minus_100_years(d: date) -> date:
    try:
        return d.replace(year=d.year - 100)
    except ValueError:            # 29 Feb in a non-leap target year
        return d.replace(year=d.year - 100, day=28)


def parse_dob(dob_str: str, race_day: date) -> Tuple[date, bool]:
    """
    'DD-Mon-YYYY' or 'DD-Mon-YY' -> (date, century_corrected).
    A 2-digit year is first read as 2000-2068 / 1969-1999 (Python's rule);
    if that puts the birth on or after the race date it is moved back 100
    years. Any birth still on or after the race date is refused.
    """
    s = (dob_str or "").strip()
    corrected = False
    try:
        d = datetime.strptime(s, "%d-%b-%Y").date()
    except ValueError:
        try:
            d = datetime.strptime(s, "%d-%b-%y").date()
        except ValueError:
            raise ChartError(f"Unreadable date of birth '{dob_str}'")
        if d >= race_day:
            d = _minus_100_years(d)
            corrected = True
    if d >= race_day:
        raise ChartError(f"Date of birth {d} is not before the race date {race_day}")
    return d, corrected


def parse_race_datetime(race_date: str, race_time: str) -> datetime:
    """'15-Jun-26' + '13:50' -> datetime (UK local time)"""
    try:
        return datetime.strptime(f"{race_date.strip()} {race_time.strip()}", "%d-%b-%y %H:%M")
    except ValueError:
        raise ChartError(f"Could not read race date/time '{race_date}' '{race_time}'")


def cloth_of(row: Dict) -> Optional[int]:
    """Cloth number from the 'cloth' column, or the name's leading number."""
    c = (row.get("cloth") or "").strip()
    if c.isdigit():
        return int(c)
    m = re.match(r"^(\d+)\s+", row.get("name", ""))
    return int(m.group(1)) if m else None


def plan_tabs(subjects: List[Dict]) -> List[Tuple[str, str, int, Optional[Dict]]]:
    """
    Fixed tab plan: [(tab, role, cloth, csv_row or None)].
    k-th horse by cloth -> P(2k-1); the jockey with the same cloth -> P(2k).
    """
    horses = [s for s in subjects if s["subject_type"] == "horse"]
    jockeys = [s for s in subjects if s["subject_type"] == "jockey"]
    for s in horses:
        if cloth_of(s) is None:
            raise ChartError(f"Horse row without a cloth number: '{s['name']}'")
    cloths = [cloth_of(h) for h in horses]
    if len(set(cloths)) != len(cloths):
        raise ChartError(f"Duplicate horse cloth numbers: {sorted(cloths)}")
    jockey_by_cloth: Dict[int, Dict] = {}
    for j in jockeys:
        c = cloth_of(j)
        if c is None or c not in cloths:
            print(f"  [WARN] Jockey row not matched to a runner, ignored: '{j['name']}'")
            continue
        if c in jockey_by_cloth:
            raise ChartError(f"Two jockey rows for cloth {c}")
        jockey_by_cloth[c] = j

    plan = []
    for k, h in enumerate(sorted(horses, key=cloth_of), 1):
        c = cloth_of(h)
        plan.append((f"P{2 * k - 1:02d}", "horse", c, h))
        plan.append((f"P{2 * k:02d}", "jockey", c, jockey_by_cloth.get(c)))
    return plan


def jockey_cache_file(cache_dir: Path, name: str, dob: date, country: str, fp: str) -> Path:
    clean = re.sub(r"^\d+\s+", "", name.strip()).lower()
    slug = lambda s: re.sub(r"_+", "_", re.sub(r"[^\w\-]", "_", s)).strip("_")
    return cache_dir / f"{slug(clean)}__{dob.isoformat()}__{slug(country)}__eng{fp}.csv"


def run_engine(engine: str, lat: float, lon: float, elev: int,
               time_str: str, tz: str,
               ephemeris_path: str, hip_path: str,
               out_prefix: str, natal_day: bool = False) -> Optional[str]:
    """
    Call celestial_bodies_swisseph.py and return path to output CSV,
    or None if failed. natal_day=True: local-noon values + whole-day ranges.
    """
    cmd = [
        sys.executable, engine,
        "--lat1", str(lat), "--lon1", str(lon), "--elevation1", str(elev),
        "--time1", time_str, "--tz1", tz,
        "--lat2", str(lat), "--lon2", str(lon), "--elevation2", str(elev),
        "--time2", time_str, "--tz2", tz,
        "--ephemeris-path", ephemeris_path,
        "--hip-path", hip_path,
        "--csv", out_prefix,
        "--quiet",
    ] + (["--natal-day"] if natal_day else [])
    try:
        result = subprocess.run(cmd, capture_output=True, text=True, timeout=120)
        if result.returncode != 0:
            print(f"    [ERROR] Engine failed: {result.stderr[:200]}")
            return None
        # Engine writes <prefix>_<t1>_vs_<t2>.csv — find it
        parent = Path(out_prefix).parent
        stem   = Path(out_prefix).stem
        matches = sorted(m for m in parent.glob(f"{stem}*.csv")
                         if not m.name.endswith("_hourly.csv"))
        if matches:
            return str(matches[0])
        print(f"    [WARN] Engine ran but output CSV not found at {out_prefix}*.csv")
        return None
    except subprocess.TimeoutExpired:
        print(f"    [ERROR] Engine timed out")
        return None
    except Exception as e:
        print(f"    [ERROR] {e}")
        return None


def read_engine_csv(csv_path: str) -> List[Dict]:
    """Read engine output CSV into list of dicts."""
    rows = []
    try:
        with open(csv_path, newline="", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for row in reader:
                rows.append(row)
    except Exception as e:
        print(f"    [WARN] Could not read {csv_path}: {e}")
    return rows


def write_tab(ws, tab_label: str, subject_name: str, rows: List[Dict]):
    """Write one tab (TRANS or P01 etc) to the worksheet."""
    # Header row
    for col_idx, field in enumerate(FIELDNAMES, 1):
        cell = ws.cell(row=1, column=col_idx, value=field)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    # Data rows
    for row_idx, row in enumerate(rows, 2):
        fill = ROW_FILL_A if row_idx % 2 == 0 else ROW_FILL_B
        for col_idx, field in enumerate(FIELDNAMES, 1):
            if field == "name":
                val = subject_name
            else:
                val = row.get(field, "")
                # Try numeric conversion
                try:
                    val = float(val)
                except (ValueError, TypeError):
                    pass
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = NAME_FONT if field == "name" else DATA_FONT
            cell.fill = NAME_FILL if field == "name" else fill

    # Column widths
    for col_letter, width in COL_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width

    # Freeze header
    ws.freeze_panes = "A2"


def hourly_twin(chart_csv: str) -> Path:
    """The engine's hourly file for a natal chart CSV (<chart>_hourly.csv)."""
    return Path(chart_csv[:-4] + "_hourly.csv")


HOURLY_FIELDS = ["tab", "name", "body", "hour", "local_time", "utc_time", "ra", "dec"]


def write_hourly(wb, blocks: List[Tuple[str, str, List[Dict]]]):
    """NATAL_HOURLY sheet: each natal body at every hour of its birth day."""
    ws = wb.create_sheet("NATAL_HOURLY")
    for c, h in enumerate(HOURLY_FIELDS, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    r = 1
    for tab, name, rows in blocks:
        for row in rows:
            r += 1
            vals = [tab, name, row["body"], int(row["hour"]), row["local_time"], row["utc_time"],
                    float(row["ra"]) if row["ra"] != "" else None, float(row["dec"])]
            for c, v in enumerate(vals, 1):
                ws.cell(row=r, column=c, value=v)
    for col, w in zip("ABCDEFGH", (7, 26, 14, 6, 22, 18, 20, 20)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"


def write_meta(wb, info: Dict, tabs: List[Dict]):
    """META sheet: versions and race details, then one row per P tab."""
    ws = wb.create_sheet("META")
    r = 1
    for k, v in info.items():
        ws.cell(row=r, column=1, value=k).font = HEADER_FONT
        ws.cell(row=r, column=1).fill = HEADER_FILL
        ws.cell(row=r, column=2, value=v).font = DATA_FONT
        r += 1
    r += 1
    cols = ["tab", "role", "cloth", "name", "dob_csv", "dob_used", "dob_century_corrected",
            "birth_country", "birth_tz", "natal_time", "status", "cache"]
    for c, h in enumerate(cols, 1):
        cell = ws.cell(row=r, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for t in tabs:
        r += 1
        for c, h in enumerate(cols, 1):
            ws.cell(row=r, column=c, value=t.get(h, "")).font = DATA_FONT
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["D"].width = 28
    ws.column_dimensions["K"].width = 30


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def build(args) -> Path:
    csv_path = Path(args.csv)
    if not csv_path.exists():
        raise ChartError(f"CSV not found: {csv_path}")
    if not Path(args.engine).exists():
        raise ChartError(f"Engine not found: {args.engine}")

    outdir = Path(args.outdir) if args.outdir else csv_path.parent
    outdir.mkdir(parents=True, exist_ok=True)
    xl_path = outdir / f"{csv_path.stem}.xlsx"
    if xl_path.exists():
        xl_path.unlink()     # never leave a stale workbook behind a failed run

    with open(csv_path, newline="", encoding="utf-8") as f:
        subjects = list(csv.DictReader(f))
    if not subjects:
        raise ChartError(f"CSV is empty: {csv_path}")
    race_date = subjects[0]["race_date"]
    race_time = subjects[0]["race_time"]
    racecourse = subjects[0]["racecourse"]

    nr = [s["name"] for s in subjects if s.get("status", "runner") == "NR"]
    if nr:
        print(f"[INFO] Skipping non-runner rows: {', '.join(nr)}")
    subjects = [s for s in subjects if s.get("status", "runner") != "NR"]

    if not race_time:
        raise ChartError("race_time is blank in the CSV - re-scrape the race")
    race_dt = parse_race_datetime(race_date, race_time)
    race_day = race_dt.date()
    race_time_str = race_dt.strftime("%Y-%m-%d %H:%M")
    rc_lat, rc_lon, rc_elev = get_racecourse_coords(racecourse)
    fp = engine_fingerprint(args.engine)
    ev = engine_version(args.engine)
    try:
        ev_ok = tuple(int(x) for x in ev.split(".")) >= (2, 2)
    except ValueError:
        ev_ok = False
    if not ev_ok:
        raise ChartError(f"Engine v{ev} found; v2.2 or later is needed "
                         f"(natal charts in standard time, decision NT-1)")
    plan = plan_tabs(subjects)

    print(f"[INFO] racingpost_excel_charts_swiss.py v{CHARTS_VERSION}  engine v{engine_version(args.engine)} ({fp})")
    print(f"[INFO] Race: {racecourse} {race_time} on {race_date}  "
          f"({len(plan) // 2} runners, {len(plan)} tabs)")
    print(f"[INFO] Racecourse coords: {rc_lat}, {rc_lon}, elev={rc_elev}m")
    print(f"[INFO] Race time (local): {race_time_str} Europe/London")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    tab_info: List[Dict] = []
    hourly_blocks: List[Tuple[str, str, List[Dict]]] = []
    cache_dir = Path(args.ephemeris_path) / "jockeys_natal_swiss"
    cache_dir.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory() as tmpdir:
        # ── TRANS: race time at the racecourse ─────────────────────────────
        print(f"\n[STEP 1] Computing TRANS chart (race time at {racecourse})...")
        trans_csv = run_engine(args.engine, rc_lat, rc_lon, rc_elev,
                               race_time_str, "Europe/London",
                               args.ephemeris_path, args.hip_path,
                               os.path.join(tmpdir, "trans"))
        if not trans_csv:
            raise ChartError("Engine failed for the race-time (TRANS) chart")
        trans_rows = read_engine_csv(trans_csv)
        if not trans_rows:
            raise ChartError("Engine produced an empty TRANS chart")
        write_tab(wb.create_sheet("TRANS"), "TRANS", f"{racecourse} {race_time}", trans_rows)
        print(f"  → TRANS: {len(trans_rows)} rows")

        # ── Natal tabs in the fixed plan ──────────────────────────────────
        print(f"\n[STEP 2] Computing natal charts ({len(plan)} tabs)...")
        for tab, role, cloth, row in plan:
            info = {"tab": tab, "role": role, "cloth": cloth,
                    "natal_time": NATAL_TIME}
            ws = wb.create_sheet(tab)

            reason = None
            if row is None:
                name = f"{cloth} (no {role} row)"
                reason = f"no {role} row in CSV"
            else:
                name = row["name"]
                info["dob_csv"] = row.get("dob", "")
                info["birth_country"] = row.get("birth_country", "")
                if not row.get("dob", "").strip():
                    reason = "no date of birth"
            info["name"] = name

            if reason:
                write_tab(ws, tab, f"{name} [NO DATA: {reason}]", [{}])
                info["status"] = f"NO DATA: {reason}"
                tab_info.append(info)
                print(f"  [{tab}] {name} — NO DATA ({reason}): placeholder tab")
                continue

            try:
                dob, corrected = parse_dob(row["dob"], race_day)
            except ChartError as e:
                raise ChartError(f"{tab} {name}: {e}")
            country = row.get("birth_country", "")
            try:
                tz = rc.country_tz(country)
            except rc.ScrapeError as e:
                raise ChartError(f"{tab} {name}: {e}")
            b_lat, b_lon = get_birth_coords(country)
            info.update({"dob_used": dob.isoformat(), "dob_century_corrected":
                         "yes" if corrected else "", "birth_tz": tz})
            if corrected:
                print(f"  [{tab}] {name}: DOB '{row['dob']}' read as {dob.isoformat()}")

            natal_csv = None
            cache_file = None
            if role == "jockey":
                cache_file = jockey_cache_file(cache_dir, name, dob, country, fp)
                if cache_file.exists() and hourly_twin(str(cache_file)).exists():
                    natal_csv = str(cache_file)
                    info["cache"] = "hit"
            if not natal_csv:
                natal_csv = run_engine(args.engine, b_lat, b_lon, 0,
                                       f"{dob.isoformat()} 12:00", tz,
                                       args.ephemeris_path, args.hip_path,
                                       os.path.join(tmpdir, f"natal_{tab}"),
                                       natal_day=True)
                if not natal_csv:
                    raise ChartError(f"Engine failed for {tab} {name}")
                if not hourly_twin(natal_csv).exists():
                    raise ChartError(f"Engine wrote no hourly file for {tab} {name} "
                                     f"(engine v2.2 or later needed)")
                if cache_file is not None:
                    shutil.copy(natal_csv, cache_file)
                    shutil.copy(hourly_twin(natal_csv), hourly_twin(str(cache_file)))
                    info["cache"] = "new"
            natal_rows = read_engine_csv(natal_csv)
            if not natal_rows:
                raise ChartError(f"Empty natal chart for {tab} {name}")
            write_tab(ws, tab, name, natal_rows)
            hourly_blocks.append((tab, name, read_engine_csv(str(hourly_twin(natal_csv)))))
            info["status"] = "OK"
            tab_info.append(info)
            print(f"  [{tab}] {name} ({role}) — DOB {dob.isoformat()}, {country}"
                  f"{' (cached)' if info.get('cache') == 'hit' else ''}")

    write_hourly(wb, hourly_blocks)
    write_meta(wb, {
        "charts_version": CHARTS_VERSION,
        "racing_common_version": rc.COMMON_VERSION,
        "engine_file": Path(args.engine).name,
        "engine_version": engine_version(args.engine),
        "engine_fingerprint": fp,
        "created_utc": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S"),
        "source_csv": csv_path.name,
        "racecourse": racecourse,
        "race_local_time": race_time_str,
        "racecourse_lat_lon_elev": f"{rc_lat}, {rc_lon}, {rc_elev}",
    }, tab_info)

    wb.save(xl_path)
    placeholders = [t["tab"] for t in tab_info if t["status"] != "OK"]
    print(f"\n[RESULT] Saved: {xl_path}")
    print(f"[RESULT] Tabs: TRANS + {len(plan)} natal (P01–P{len(plan):02d}) + NATAL_HOURLY + META"
          + (f"  |  placeholders: {', '.join(placeholders)}" if placeholders else ""))
    return xl_path


def main() -> int:
    args = parse_args()
    try:
        build(args)
    except ChartError as e:
        print(f"\n[FAILED] {e}")
        print("[FAILED] No workbook written for this race.")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
